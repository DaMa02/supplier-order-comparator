"""Guided configuration of unrecognised document schemas."""

from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
for cartella in (ROOT / "app", ROOT / "scripts"):
    if str(cartella) not in sys.path:
        sys.path.insert(0, str(cartella))

import pipeline_jobs  # noqa: E402
import schema_mapping  # noqa: E402
from inspect_sources import profile_file  # noqa: E402
from pipeline_jobs import ConfigurazionePipeline, PipelineJobManager  # noqa: E402


def profilo_caricato(percorso: Path) -> dict:
    """Build the profile `POST /api/upload` writes for an uploaded document."""

    profilo = profile_file(percorso)
    profilo["upload_role"] = "supplier"
    profilo["sha256"] = hashlib.sha256(percorso.read_bytes()).hexdigest()
    profilo.setdefault("profile_id", percorso.stem)
    return profilo


class SchemaMappingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.adapters = schema_mapping.carica_adattatori(ROOT / "references" / "adapters.json")

    def workbook(self, name: str, rows: list[list[object]], sheet: str = "Listino") -> dict:
        path = self.root / name
        book = Workbook()
        ws = book.active
        ws.title = sheet
        for row in rows:
            ws.append(row)
        book.save(path)
        book.close()
        return profile_file(path)

    def test_propone_un_fornitore_dal_contenuto_non_dal_nome_del_file(self) -> None:
        profilo = self.workbook("nome-casuale.xlsx", [
            ["COD.ART.", "DES.ARTICOLO", "UM", "QT", "LISTINO", "COD.EAN"],
            ["0001", "Prodotto uno", "PZ", 6, 2.5, "8000000000001"],
        ])

        risposta = schema_mapping.prepara_pendenti(
            [profilo], ["nome-casuale.xlsx"], self.adapters, "run-1"
        )
        proposta = risposta["documents"][0]["suggestion"]

        self.assertEqual(proposta["supplierId"], "cipresso")
        self.assertEqual(proposta["columns"]["description"], 2)
        self.assertEqual(proposta["columns"]["unit_price_net"], 5)
        self.assertEqual(proposta["orderColumn"], 7)

    def test_il_box_gestionale_prevale_sulla_somiglianza_con_un_listino(self) -> None:
        profilo = self.workbook("export.xlsx", [
            ["COD.ART.", "DES.ARTICOLO", "UM", "QT", "LISTINO", "COD.EAN"],
            ["0001", "Prodotto uno", "PZ", 6, 2.5, "8000000000001"],
        ])
        profilo["upload_role"] = "master"
        profilo["ai_preflight"]["role"] = "master"

        risposta = schema_mapping.prepara_pendenti(
            [profilo], ["export.xlsx"], self.adapters, "run-1"
        )

        self.assertEqual(risposta["documents"][0]["suggestion"]["role"], "master")

    def test_la_validazione_usa_il_lettore_vero_e_mostra_un_campione(self) -> None:
        profilo = self.workbook("qualsiasi.xlsx", [
            ["COD.ART.", "DES.ARTICOLO", "UM", "QT", "LISTINO", "COD.EAN"],
            ["0001", "Prodotto uno", "PZ", 6, 2.5, "8000000000001"],
            ["0002", "Prodotto due", "PZ", 12, 1.1, "8000000000002"],
        ])
        payload = {
            "runId": "run-1",
            "mappings": [{
                "profileId": profilo["profile_id"],
                "role": "supplier",
                "supplierId": "cipresso",
                "sheet": "Listino",
                "headerRow": 1,
                "dataStartRow": 2,
                "columns": {
                    "supplier_code": 1,
                    "description": 2,
                    "unit": 3,
                    "pieces_per_carton": 4,
                    "unit_price_net": 5,
                    "ean": 6,
                },
                "orderColumn": 7,
            }],
        }

        esito, decisioni = schema_mapping.valida_mappature(
            [profilo], [profilo["file_name"]], self.adapters, payload, "run-1"
        )

        self.assertEqual(esito["documents"][0]["rowsUsable"], 2)
        self.assertEqual(esito["documents"][0]["sample"][0]["description"], "Prodotto uno")
        self.assertEqual(decisioni[0]["field_mapping"]["columns"]["unit_price_net"], "LISTINO")
        self.assertTrue(decisioni[0]["field_mapping"]["order_header_blank_confirmed"])
        self.assertEqual(decisioni[0]["user_confirmation"]["status"], "CONFIRMED")

    def test_colonne_essenziali_sbagliate_non_passano_per_buone(self) -> None:
        profilo = self.workbook("sbagliato.xlsx", [
            ["Codice", "Nome", "Pezzi", "Prezzo"],
            ["1", "Prodotto", 6, 2.5],
        ])
        payload = {"runId": "run-1", "mappings": [{
            "profileId": profilo["profile_id"], "role": "supplier",
            "supplierName": "Nuovo", "sheet": "Listino", "headerRow": 1,
            "dataStartRow": 2,
            # Price and name columns swapped: the reader will produce no usable rows.
            "columns": {"supplier_code": 1, "description": 4,
                        "pieces_per_carton": 3, "unit_price_net": 2},
            "orderColumn": 5,
        }]}

        with self.assertRaisesRegex(ValueError, "nessun prodotto ordinabile"):
            schema_mapping.valida_mappature(
                [profilo], [profilo["file_name"]], self.adapters, payload, "run-1"
            )

    def test_un_file_cambiato_dopo_l_anteprima_viene_rifiutato(self) -> None:
        profilo = self.workbook("cambiato.xlsx", [
            ["Codice", "Nome", "Pezzi", "Prezzo"],
            ["1", "Prodotto", 6, 2.5],
        ])
        Path(profilo["path"]).write_bytes(b"non e' piu' lo stesso file")
        payload = {"runId": "run-1", "mappings": [{"profileId": profilo["profile_id"]}]}

        with self.assertRaisesRegex(ValueError, "cambiato dopo l'anteprima"):
            schema_mapping.valida_mappature(
                [profilo], [profilo["file_name"]], self.adapters, payload, "run-1"
            )

    def gestionale(self, name: str, rows: list[list[object]]) -> dict:
        return self.workbook(name, rows, sheet="Gestionale")

    def mappatura_gestionale(self, profilo: dict) -> dict:
        return {"runId": "run-1", "mappings": [{
            "profileId": profilo["profile_id"], "role": "master",
            "sheet": "Gestionale", "headerRow": 1, "dataStartRow": 2,
            "columns": {"ean": 1, "description": 2, "unit": 3,
                        "suggested_colli": 4, "last_unit_price": 5, "vat": 6},
        }]}

    def test_un_gestionale_senza_codici_non_passa_per_buono(self) -> None:
        # The code column maps to empty cells: recomputing would build a comparison
        # where no product can be matched to any supplier offer.
        profilo = self.gestionale("senza-codici.xlsx", [
            ["EAN", "DESCRIZIONE", "UM", "COLLI", "PREZZO", "IVA"],
            ["", "Prodotto uno", "PZ", 1, 2.5, 22],
            ["", "Prodotto due", "PZ", 2, 1.1, 22],
            ["", "Prodotto tre", "PZ", 1, 3.0, 22],
            ["", "Prodotto quattro", "PZ", 4, 0.9, 22],
            ["", "Prodotto cinque", "PZ", 1, 7.4, 22],
        ])

        with self.assertRaisesRegex(ValueError, "Codice EAN"):
            schema_mapping.valida_mappature(
                [profilo], [profilo["file_name"]], self.adapters,
                self.mappatura_gestionale(profilo), "run-1",
            )

    def test_un_gestionale_mappato_bene_resta_valido(self) -> None:
        profilo = self.gestionale("gestionale.xlsx", [
            ["EAN", "DESCRIZIONE", "UM", "COLLI", "PREZZO", "IVA"],
            ["8000000000001", "Prodotto uno", "PZ", 1, 2.5, 22],
            ["8000000000002", "Prodotto due", "PZ", 2, 1.1, 22],
            ["8000000000003", "Prodotto tre", "PZ", 1, 3.0, 22],
        ])

        esito, _decisioni = schema_mapping.valida_mappature(
            [profilo], [profilo["file_name"]], self.adapters,
            self.mappatura_gestionale(profilo), "run-1",
        )
        documento = esito["documents"][0]

        self.assertEqual(documento["rowsRead"], 3)
        self.assertEqual(documento["rowsUsable"], 3)
        self.assertEqual(documento["rowsDiscarded"], 0)
        self.assertEqual(documento["sample"][0]["ean"], "8000000000001")

    def test_qualche_riga_sporca_non_ferma_un_gestionale_buono(self) -> None:
        # Department subtotals and separator rows make the document dirty, not
        # badly mapped; the test asserts how many rows get discarded without
        # blocking the recompute.
        profilo = self.gestionale("con-totali.xlsx", [
            ["EAN", "DESCRIZIONE", "UM", "COLLI", "PREZZO", "IVA"],
            ["8000000000001", "Prodotto uno", "PZ", 1, 2.5, 22],
            ["", "TOTALE REPARTO", "", "", 999.0, ""],
            ["8000000000003", "Prodotto tre", "PZ", 1, 3.0, 22],
            ["8000000000004", "Prodotto quattro", "PZ", 4, 0.9, 22],
            ["", "--- SEZIONE DUE ---", "", "", "", ""],
            ["8000000000006", "Prodotto sei", "PZ", 1, 1.5, 22],
        ])

        esito, _decisioni = schema_mapping.valida_mappature(
            [profilo], [profilo["file_name"]], self.adapters,
            self.mappatura_gestionale(profilo), "run-1",
        )
        documento = esito["documents"][0]

        self.assertEqual(documento["rowsRead"], 6)
        self.assertEqual(documento["rowsUsable"], 4)
        self.assertEqual(documento["rowsDiscarded"], 2)
        self.assertNotIn("TOTALE REPARTO", [voce["description"] for voce in documento["sample"]])

    def test_due_ruoli_sulla_stessa_colonna_sono_un_errore_leggibile(self) -> None:
        profilo = self.workbook("duplicata.xlsx", [
            ["Codice", "Nome", "Pezzi", "Prezzo"],
            ["1", "Prodotto", 6, 2.5],
        ])
        scelta = {
            "role": "supplier", "supplierName": "Nuovo", "sheet": "Listino",
            "headerRow": 1, "dataStartRow": 2,
            "columns": {"description": 2, "unit_price_net": 2, "pieces_per_carton": 3},
            "orderColumn": 5,
        }

        with self.assertRaisesRegex(ValueError, "assegnata sia"):
            schema_mapping.decisione_da_mappatura(profilo, scelta, self.adapters)

    def test_la_decisione_porta_l_impronta_del_documento_confermato(self) -> None:
        """The mapping is keyed to the file's content hash, not to its name.

        Without the fingerprint, deleting a badly mapped price list and
        re-uploading a different one under the same name would silently
        reapply the old column mapping to it.
        """

        profilo = self.workbook("betulla.xlsx", [
            ["Codice", "Nome", "Pezzi", "Prezzo"],
            ["1", "Prodotto", 6, 2.5],
        ])
        scelta = {
            "role": "supplier", "supplierName": "Nuovo", "sheet": "Listino",
            "headerRow": 1, "dataStartRow": 2,
            "columns": {"description": 2, "unit_price_net": 4, "pieces_per_carton": 3},
            "orderColumn": 5,
        }

        decisione = schema_mapping.decisione_da_mappatura(profilo, scelta, self.adapters)

        self.assertEqual(decisione["file_sha256"], profilo["sha256"])
        self.assertTrue(decisione["file_sha256"], "il profilo deve portare un'impronta vera")


class LaProposta(unittest.TestCase):
    """When the wizard proposes an already-known supplier, and when it must not.

    A weak match (one header out of five, coincidentally "ORDINE") once got
    proposed as a specific supplier with a low similarity score. Confirming
    that suggestion overwrites that supplier's real price list, so a proposal
    below a sane similarity threshold must come back empty instead.
    """

    LISTINI = ROOT / "listini-storici"

    def setUp(self) -> None:
        self.adattatori = schema_mapping.carica_adattatori(ROOT / "references" / "adapters.json")

    def proposta(self, nome: str) -> dict:
        percorso = self.LISTINI / nome
        if not percorso.is_file():
            self.skipTest(f"Manca il listino {nome}")
        profilo = profile_file(percorso)
        profilo["upload_role"] = "supplier"
        pendenti = schema_mapping.prepara_pendenti(
            [profilo], [profilo["file_name"]], self.adattatori, "run-prova",
        )
        return pendenti["documents"][0]["suggestion"]

    def test_una_intestazione_su_cinque_non_e_un_riconoscimento(self) -> None:
        for nome in ("OFFERTE AGOSTO 4.xlsx",
                     "LISTINO QUERCIA AGGIORNATO DEL 06-08-2026.xlsx",
                     "ACERO LISTINO SETTIMANA 26.xlsx",
                     "28.1 06-10lug.xlsx"):
            with self.subTest(listino=nome):
                self.assertEqual(self.proposta(nome)["supplierId"], "")

    def test_i_listini_veri_continuano_a_essere_proposti(self) -> None:
        """The similarity threshold must not suppress a correct proposal."""

        for nome, atteso in (("LISTINO BETULLA VALIDO FINO AL 28-07-26.xlsx", "betulla"),
                             ("3listino_Cipresso.xlsx", "cipresso"),
                             ("Listino3_33.xlsx", "cipresso"),
                             ("formattato_104233.xls", "noce")):
            with self.subTest(listino=nome):
                self.assertEqual(self.proposta(nome)["supplierId"], atteso)

    def test_senza_fornitore_scelto_lo_dice_invece_di_chiedere_un_nome(self) -> None:
        percorso = self.LISTINI / "OFFERTE AGOSTO 4.xlsx"
        if not percorso.is_file():
            self.skipTest("Manca il listino delle offerte")
        profilo = profile_file(percorso)

        with self.assertRaises(ValueError) as errore:
            schema_mapping.decisione_da_mappatura(profilo, {
                "role": "supplier", "supplierId": "", "supplierName": "",
                "sheet": "Foglio1", "headerRow": 3, "dataStartRow": 6,
                "columns": {"description": 2, "unit_price_net": 5, "pieces_per_carton": 4},
                "orderColumn": 8,
            }, self.adattatori)

        self.assertIn("scegli il fornitore", str(errore.exception))

    def test_senza_colonne_dice_quali_colonne_non_come_e_fatto_il_json(self) -> None:
        """The error must name the missing fields, not describe the JSON shape."""

        percorso = self.LISTINI / "OFFERTE AGOSTO 4.xlsx"
        if not percorso.is_file():
            self.skipTest("Manca il listino delle offerte")
        profilo = profile_file(percorso)

        with self.assertRaises(ValueError) as errore:
            schema_mapping.decisione_da_mappatura(profilo, {
                "role": "supplier", "supplierId": "", "supplierName": "Fornitore Di Prova",
                "sheet": "Foglio1", "headerRow": 3, "dataStartRow": 6,
                "columns": {}, "orderColumn": 8,
            }, self.adattatori)

        self.assertIn("completa nome prodotto, prezzo, pezzi per collo.", str(errore.exception))
        self.assertNotIn("oggetto", str(errore.exception))

    def test_chi_ha_gia_un_listino_in_questa_run_si_sa_prima_di_confermare(self) -> None:
        betulla = self.LISTINI / "LISTINO BETULLA VALIDO FINO AL 28-07-26.xlsx"
        offerte = self.LISTINI / "OFFERTE AGOSTO 4.xlsx"
        for percorso in (betulla, offerte):
            if not percorso.is_file():
                self.skipTest(f"Manca {percorso.name}")
        profili = [profile_file(betulla), profile_file(offerte)]
        for profilo in profili:
            profilo["upload_role"] = "supplier"

        pendenti = schema_mapping.prepara_pendenti(
            profili, [offerte.name], self.adattatori, "run-prova",
        )

        occupati = pendenti["occupied"]
        self.assertEqual([voce["supplierId"] for voce in occupati], ["betulla"])
        self.assertEqual(occupati[0]["fileName"], betulla.name)
        self.assertTrue(pendenti["documents"][0]["modifiedAt"])


class InizioDeiProdottiTests(unittest.TestCase):
    """The data-start rule is "after the row that says LISTINO", not a fixed row number.

    A hardcoded row doesn't survive a week: on one supplier's price list the rows
    before the product table are a promotional block whose length changes between
    price lists, so the real data starts after the single marker row in column A.
    """

    maxDiff = None

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.adapters = schema_mapping.carica_adattatori(ROOT / "references" / "adapters.json")

    def listino(self, promozionali: int = 25, separatore: str = "LISTINO") -> dict:
        righe: list[list[object]] = [["Articolo", "EAN", "Descrizione", "Imballo", "Prezzo", "Ordine"]]
        for numero in range(promozionali):
            righe.append([f"OMA{numero}", f"80000000009{numero:02d}", f"OMAGGIO {numero}", 12, 1.05, None])
        righe.append([separatore, None, None, None, None, None])
        for numero in range(20):
            righe.append([f"FAT{numero}", f"80014807193{numero:02d}", f"AXO {numero}", 18, 0.94, None])
        percorso = self.root / "listino.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "Listino"
        for riga in righe:
            ws.append(riga)
        book.save(percorso)
        book.close()
        return profile_file(percorso)

    def scelta(self, profilo: dict, **cambiamenti: object) -> dict:
        base: dict[str, object] = {
            "role": "supplier", "supplierName": "Va. Pa.", "sheet": "Listino",
            "headerRow": 1, "dataStartRow": 28,
            "dataStartMarker": {"column": 1, "match": "equals", "text": "LISTINO", "offset": 1},
            "columns": {"supplier_code": 1, "ean": 2, "description": 3,
                        "pieces_per_carton": 4, "unit_price_net": 5},
            "orderColumn": 6,
        }
        base.update(cambiamenti)
        return base

    def test_l_anteprima_mostra_il_taglio_invece_di_farlo_indovinare(self) -> None:
        """The preview must show the actual cut point, not leave the operator guessing it."""

        profilo = self.listino()
        foglio = schema_mapping.serializza_foglio(schema_mapping.fogli_del_profilo(profilo)[0])

        self.assertEqual([voce["row"] for voce in foglio["sectionBreaks"]], [27])
        self.assertEqual(foglio["sectionBreaks"][0]["text"], "LISTINO")
        self.assertEqual(foglio["sectionBreaks"][0]["data_from"], 28)
        # The window the page renders: from headerRow-1 to dataStartRow+7.
        visibili = {voce["row"] for voce in foglio["rows"]}
        self.assertLessEqual({26, 27, 28, 29, 30}, visibili)

    def test_le_righe_del_taglio_non_cacciano_la_fine_del_documento(self) -> None:
        """Adding rows around the section break must not push out the document's tail rows.

        A fixed cap on visible rows keeps the lowest-numbered ones; with the
        separator rows inserted in the middle, that would silently drop the
        last rows of the price list — the only place to see whether the
        document ends with products or with totals.
        """

        profilo = self.listino()
        foglio = schema_mapping.serializza_foglio(schema_mapping.fogli_del_profilo(profilo)[0])

        visibili = {voce["row"] for voce in foglio["rows"]}
        self.assertIn(foglio["maxRow"], visibili)
        # The section-break rows are still there too: the cap doesn't trade one for the other.
        self.assertLessEqual({26, 27, 28}, visibili)

    def test_la_decisione_porta_la_regola_e_il_numero_a_cui_si_e_risolta(self) -> None:
        profilo = self.listino()

        decisione = schema_mapping.decisione_da_mappatura(profilo, self.scelta(profilo), self.adapters)

        mappatura = decisione["field_mapping"]
        self.assertEqual(mappatura["data_start_marker"],
                         {"column": "A", "equals": "LISTINO", "offset": 1})
        # The resolved row number is kept too, for the order-file writer;
        # the reader itself never looks at it once the marker exists.
        self.assertEqual(mappatura["data_start_row"], 28)

    def test_la_regola_si_verifica_sul_documento_prima_di_entrare(self) -> None:
        """A data-start marker that isn't checked against the document would be
        worse than the fixed row number it's meant to replace."""

        profilo = self.listino()
        scelta = self.scelta(profilo, dataStartMarker={
            "column": 1, "match": "equals", "text": "PREZZI", "offset": 1})

        with self.assertRaises(ValueError) as errore:
            schema_mapping.decisione_da_mappatura(profilo, scelta, self.adapters)

        messaggio = str(errore.exception)
        self.assertIn("A27", messaggio)
        self.assertIn("LISTINO", messaggio)

    def test_la_riga_dichiarata_deve_stare_sotto_il_separatore(self) -> None:
        """If the declared data-start row and the marker's offset disagree, fail rather than guess."""

        profilo = self.listino()
        scelta = self.scelta(profilo, dataStartRow=30)

        with self.assertRaises(ValueError) as errore:
            schema_mapping.decisione_da_mappatura(profilo, scelta, self.adapters)

        self.assertIn("A29", str(errore.exception))

    def test_la_lettura_di_prova_taglia_dove_dice_la_regola(self) -> None:
        profilo = self.listino()
        payload = {"runId": "run-1", "mappings": [{
            "profileId": profilo["profile_id"], **self.scelta(profilo)}]}

        esito, decisioni = schema_mapping.valida_mappature(
            [profilo], [profilo["file_name"]], self.adapters, payload, "run-1")

        documento = esito["documents"][0]
        self.assertEqual(documento["rowsUsable"], 20)
        self.assertEqual(documento["rowsDiscarded"], 0)
        self.assertEqual(documento["sample"][0]["description"], "AXO 0")
        self.assertEqual(decisioni[0]["field_mapping"]["data_start_marker"]["equals"], "LISTINO")

    def test_senza_regola_la_mappatura_resta_quella_di_prima(self) -> None:
        """The marker key is optional: documents that don't need it never see it."""

        profilo = self.listino()
        scelta = self.scelta(profilo)
        scelta.pop("dataStartMarker")

        decisione = schema_mapping.decisione_da_mappatura(profilo, scelta, self.adapters)

        self.assertNotIn("data_start_marker", decisione["field_mapping"])
        self.assertEqual(decisione["field_mapping"]["data_start_row"], 28)


class LeColonneSiRivedonoQuandoSiVuole(unittest.TestCase):
    """The manual column selector for a price list the pipeline already recognises.

    The schema-mapping wizard only opens for a document the adapter registry
    doesn't recognise; a recognised supplier never triggers it, which left no
    way to revise the columns of an already-known price list — including the
    order column, which decides where quantities land in the file sent back
    to the supplier. The reader still exists and works; it was just unreachable.
    """

    NOME = "listino-betulla.xlsx"

    def setUp(self) -> None:
        self.temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporanea.cleanup)
        self.radice = Path(self.temporanea.name)
        self.uploads = self.radice / "uploads"
        self.uploads.mkdir(parents=True)
        # A realistic price list: an Excel file with an empty order column,
        # which is the column this test suite exercises overriding.
        percorso = self.uploads / self.NOME
        libro = Workbook()
        foglio = libro.active
        foglio.title = "Listino"
        foglio.append(["EAN", "CodArt", "ORDINE", "Descrizione", "PzCt", "Cessione"])
        foglio.append(["8000000000001", "A1", None, "PRIMO PRODOTTO", 6, 1.5])
        foglio.append(["8000000000002", "A2", None, "SECONDO PRODOTTO", 12, 2.5])
        libro.save(percorso)
        libro.close()
        (self.uploads / "upload_profiles.json").write_text(json.dumps({
            "schema_version": 1,
            "profiles": [profilo_caricato(percorso)],
            "errors": [],
        }), encoding="utf-8")
        self.gestore = PipelineJobManager(ConfigurazionePipeline(
            data_dir=self.radice,
            uploads_dir=self.uploads,
            review_path=self.radice / "review_data.json",
            state_path=self.radice / "state.json",
        ))

    def test_l_anteprima_si_apre_su_un_documento_qualunque(self) -> None:
        esito = self.gestore.colonne_del_documento(self.NOME)

        self.assertTrue(esito["ok"])
        self.assertEqual(len(esito["documents"]), 1)
        self.assertEqual(esito["documents"][0]["fileName"], self.NOME)
        # The preview exists to let the operator choose; without sample rows there's nothing to choose from.
        self.assertTrue(esito["documents"][0]["sheets"][0]["rows"])

    def test_un_nome_che_non_e_fra_i_caricamenti_non_si_apre(self) -> None:
        for nome in ("", "altro.xlsx", "../secrets.json", "C:/Windows/win.ini"):
            with self.subTest(nome=nome):
                with self.assertRaises(ValueError):
                    self.gestore.colonne_del_documento(nome)

    def test_un_documento_cancellato_da_fuori_lo_dice_subito(self) -> None:
        """A separate guard is needed for a file removed from disk after upload.

        A made-up file name is already rejected by the check against the
        profile registry; that check alone can't catch a file that's still
        listed in the registry but has since been deleted outside the app.
        Without this guard, the error would surface later, somewhere else.
        """

        (self.uploads / self.NOME).unlink()

        with self.assertRaises(ValueError) as fermata:
            self.gestore.colonne_del_documento(self.NOME)
        self.assertIn("non è fra i caricamenti", str(fermata.exception))

    def mappatura(self, colonna_ordine: str) -> dict:
        documento = self.gestore.colonne_del_documento(self.NOME)["documents"][0]
        return {
            "runId": PipelineJobManager.CHIAVE_COLONNE_A_MANO,
            "fileName": self.NOME,
            "mappings": [{
                "profileId": documento["profileId"],
                "role": "supplier",
                "supplierChoice": "__new__",
                "supplierName": "Nuovo Fornitore",
                "sheet": "Listino",
                "headerRow": 1,
                "dataStartRow": 2,
                "columns": {"ean": 1, "supplier_code": 2, "description": 4,
                            "pieces_per_carton": 5, "unit_price_net": 6},
                "orderColumn": colonna_ordine,
            }],
        }

    def test_la_prova_legge_i_dati_veri_e_non_scrive_niente(self) -> None:
        esito = self.gestore.prova_colonne_del_documento(self.mappatura("3"))

        self.assertEqual(esito["documents"][0]["rowsUsable"], 2)
        self.assertFalse((self.radice / "decisioni_schemi.json").exists())

    def test_salvare_scrive_la_colonna_d_ordine_scelta(self) -> None:
        esito = self.gestore.salva_colonne_del_documento(self.mappatura("3"))

        decisioni = json.loads((self.radice / "decisioni_schemi.json").read_text(encoding="utf-8"))
        salvata = decisioni["decisions"][0]
        self.assertEqual(salvata["file_name"], self.NOME)
        self.assertEqual(salvata["field_mapping"]["order_column"], "C")
        # Saving a column correction must not trigger a recompute on its own: an
        # AI-assisted run costs real time and shouldn't start without the operator
        # asking for it. The page must still know the shown prices are stale.
        self.assertEqual(esito["pipeline"]["stato"], pipeline_jobs.IN_ATTESA)
        cambiamento = esito["pipeline"]["cambiamento"]
        self.assertEqual(cambiamento["tipo"], "colonne")
        self.assertIn(self.NOME, cambiamento["documenti"])

    def test_e_si_puo_cambiare_da_una_settimana_all_altra(self) -> None:
        """This is the actual use case: the order column shifts from one week's file to the next."""

        self.gestore.salva_colonne_del_documento(self.mappatura("3"))
        self.gestore.salva_colonne_del_documento(self.mappatura("7"))

        decisioni = json.loads((self.radice / "decisioni_schemi.json").read_text(encoding="utf-8"))
        self.assertEqual(len(decisioni["decisions"]), 1, "una decisione per documento, non una pila")
        self.assertEqual(decisioni["decisions"][0]["field_mapping"]["order_column"], "G")


if __name__ == "__main__":
    unittest.main()
