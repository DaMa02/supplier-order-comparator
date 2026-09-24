#!/usr/bin/env python3
"""The order-quantity column can be moved by the user, not fixed at setup.

Besides the guided mapping screen — which only opens when the app can't
recognize a document's columns — a known supplier's order column also lives
in the adapter registry and can be moved from the page, instead of requiring
a hand edit of a JSON file.

Three supplier shapes this must handle, each genuinely different — measured
on real price lists:

* A supplier with a dedicated reader and a header row: the order column has a
  title (e.g. "ORDINE") that the writer verifies before writing.
* A supplier declaring `from_field_mapping`: the registry and the confirmed
  field mapping must agree on the same column, and its column can have no
  title at all (an empty cell across every row).
* A supplier with no header row whatsoever: there's no cell above the column
  to look at, so declaring one "empty" instead of "absent" would block it.

And the two rejections that matter: a column the app already reads (the order
would overwrite it) and a column full of formulas (writing over it destroys
them).
"""

from __future__ import annotations

import dataclasses
import json
import sys
import tempfile
import unittest
from pathlib import Path
from typing import Any

from openpyxl import Workbook

RADICE = Path(__file__).resolve().parents[1]
for cartella in (RADICE / "app", RADICE / "scripts"):
    if str(cartella) not in sys.path:
        sys.path.insert(0, str(cartella))

import colonna_ordine  # noqa: E402
import launcher  # noqa: E402
import registro  # noqa: E402
import schema_mapping  # noqa: E402
from inspect_sources import profile_file  # noqa: E402
from pipeline_jobs import LavoroGiaInCorso  # noqa: E402
from server import ReviewStore  # noqa: E402


# --------------------------------------------------------------------------
# What gets declared, without touching disk or the registry.
# --------------------------------------------------------------------------
class CheCosaSiDichiara(unittest.TestCase):
    def test_una_colonna_con_titolo_diventa_l_intestazione_attesa(self) -> None:
        """`expected_header` reflects what the document has, not a preference."""

        nuova = colonna_ordine.dichiarazione_aggiornata(
            {"order_column": "C", "expected_header": "ORDINE"}, "K", "QUANTITA"
        )

        self.assertEqual(nuova["order_column"], "K")
        self.assertEqual(nuova["expected_header"], "QUANTITA")
        self.assertNotIn("allow_blank_header_if_confirmed", nuova)

    def test_una_colonna_senza_titolo_si_dichiara_vuota_e_confermata(self) -> None:
        """Without this declaration the column would be written with no
        check at all: a missing `expected_header` means "nothing to verify",
        so picking a titleless column would disable the safeguard instead of
        enabling it."""

        nuova = colonna_ordine.dichiarazione_aggiornata(
            {"order_column": "C", "expected_header": "ORDINE"}, "J", ""
        )

        self.assertNotIn("expected_header", nuova)
        self.assertIs(nuova["allow_blank_header_if_confirmed"], True)
        self.assertIs(nuova["order_header_blank_confirmed"], True)

    def test_senza_riga_di_intestazione_non_si_dichiara_nessuna_cella_vuota(self) -> None:
        """A supplier with no header row: there's no cell above the column to
        look at. Declaring it "empty" wouldn't make the check stricter, it
        would make it impossible — `source_rule` requires a header row to run
        it and would fail, blaming the registry for something the document
        simply doesn't have."""

        nuova = colonna_ordine.dichiarazione_aggiornata(
            {"order_column": "D"}, "S", "", c_e_intestazione=False
        )

        self.assertEqual(nuova["order_column"], "S")
        self.assertNotIn("expected_header", nuova)
        self.assertNotIn("allow_blank_header_if_confirmed", nuova)
        self.assertNotIn("order_header_blank_confirmed", nuova)

    def test_la_procedura_di_scrittura_resta_dov_era(self) -> None:
        """A supplier's `.xls` written in-place: moving the column must not
        change how that document is patched."""

        nuova = colonna_ordine.dichiarazione_aggiornata(
            {
                "order_column": "I",
                "from_field_mapping": True,
                "mode": "patch_xls_in_posizione",
                "required_columns": ["ean"],
            },
            "N",
            "",
        )

        self.assertEqual(nuova["order_column"], "N")
        self.assertIs(nuova["from_field_mapping"], True)
        self.assertEqual(nuova["mode"], "patch_xls_in_posizione")
        self.assertEqual(nuova["required_columns"], ["ean"])

    def test_la_mappatura_dice_la_stessa_cosa_del_registro(self) -> None:
        registro_nuovo = colonna_ordine.dichiarazione_aggiornata(
            {"order_column": "G", "from_field_mapping": True}, "H", "ORDINI"
        )
        mappatura = colonna_ordine.mappatura_aggiornata(
            {"order_column": "G", "order_header_blank_confirmed": True}, "H", "ORDINI"
        )

        self.assertEqual(registro_nuovo["order_column"], mappatura["order_column"])
        self.assertEqual(mappatura["order_header_expected"], "ORDINI")
        self.assertNotIn("order_header_blank_confirmed", mappatura)

    def test_la_colonna_si_indica_per_lettera_o_per_numero(self) -> None:
        """The page sends a column number, the registry stores a letter; the
        conversion must live in one place, or the two forms could disagree,
        e.g. "AA" vs 27."""

        self.assertEqual(colonna_ordine.indice_scelto("C"), 3)
        self.assertEqual(colonna_ordine.indice_scelto("c"), 3)
        self.assertEqual(colonna_ordine.indice_scelto(3), 3)
        self.assertEqual(colonna_ordine.indice_scelto("3"), 3)
        self.assertEqual(colonna_ordine.indice_scelto("AA"), 27)
        self.assertEqual(colonna_ordine.lettera_di_indice(27), "AA")
        for scarto in (None, "", 0, -2, "3C", True):
            self.assertIsNone(colonna_ordine.indice_scelto(scarto), scarto)


# --------------------------------------------------------------------------
# The columns offered as candidates.
# --------------------------------------------------------------------------
class ColonneFraCuiScegliere(unittest.TestCase):
    def setUp(self) -> None:
        temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(temporanea.cleanup)
        self.radice = Path(temporanea.name)

    def profilo(self, righe: list[list[Any]], nome: str = "listino.xlsx") -> dict[str, Any]:
        percorso = self.radice / nome
        libro = Workbook()
        pagina = libro.active
        pagina.title = "Sheet1"
        for riga in righe:
            pagina.append(riga)
        libro.save(percorso)
        libro.close()
        return profile_file(percorso)

    def test_offre_anche_la_prima_colonna_libera_in_fondo(self) -> None:
        """A price list with no order column yet must still be able to get
        one: offering only already-written columns would leave nowhere to
        put it."""

        profilo = self.profilo([["EAN", "Descrizione"], ["8000000000001", "SAPONE"]])

        colonne = schema_mapping.colonne_del_foglio(profilo, "Sheet1", 1, 2)

        self.assertEqual([voce["lettera"] for voce in colonne], ["A", "B", "C"])
        self.assertEqual(colonne[2]["valori"], 0)

    def test_conta_le_formule_di_ogni_colonna(self) -> None:
        """The count isn't a sample: `inspect_sources` scans every row. It's
        the number that decides, since writing the order over a formula
        column would erase every formula."""

        profilo = self.profilo([
            ["EAN", "PzCt", "TOTALI"],
            ["8000000000001", 6, "=B2*2"],
            ["8000000000002", 4, "=B3*2"],
        ])

        colonne = {voce["lettera"]: voce for voce in
                   schema_mapping.colonne_del_foglio(profilo, "Sheet1", 1, 2)}

        self.assertEqual(colonne["C"]["formule"], 2)
        self.assertEqual(colonne["B"]["formule"], 0)
        self.assertEqual(colonne["B"]["esempio"], "6")

    def test_una_colonna_d_ordine_oltre_l_ultima_scritta_resta_raggiungibile(self) -> None:
        """A column that's empty across every row doesn't show up in the
        profile at all. Without `fino_a`, the only new column offered would
        be the one right next to the one already in use."""

        profilo = self.profilo([["COD", "DES"], ["019654", "SPAZZOLA"]])

        senza = schema_mapping.colonne_del_foglio(profilo, "Sheet1", 1, 2)
        con = schema_mapping.colonne_del_foglio(profilo, "Sheet1", 1, 2, fino_a=4)

        self.assertEqual([voce["lettera"] for voce in senza], ["A", "B", "C"])
        self.assertEqual([voce["lettera"] for voce in con], ["A", "B", "C", "D", "E"])

    def test_le_colonne_occupate_si_vedono_marcate_non_tolte(self) -> None:
        """Hiding occupied columns would make the document look like it has
        fewer columns than it does, and a user counting "the one after the
        price" counts what's visible."""

        effettiva = {
            "columns": [{"campo": "unit_price_net", "etichetta": "Prezzo", "colonna": 2}],
            "orderColumn": {"colonna": 3, "lettera": "C"},
        }
        colonne = colonna_ordine.colonne_per_la_scelta(effettiva, [
            {"colonna": 1, "lettera": "A"},
            {"colonna": 2, "lettera": "B"},
            {"colonna": 3, "lettera": "C"},
        ])

        self.assertEqual([voce["lettera"] for voce in colonne], ["A", "B", "C"])
        self.assertEqual(colonne[1]["occupataDa"], "Prezzo")
        self.assertIs(colonne[1]["scegliibile"], False)
        self.assertIs(colonne[2]["attuale"], True)


# --------------------------------------------------------------------------
# A price list openpyxl can read but can't write into: the legacy .xls format.
# --------------------------------------------------------------------------
class UnListinoChePerScriverciDentroNonVaBene(unittest.TestCase):
    """A legacy `.xls` file must fail with a clear, actionable message.

    `source_rule` opens the document with openpyxl, which doesn't support
    Excel 97-2003 at all: without a translated message, the user would see
    openpyxl's raw English error — "openpyxl does not support the old .xls
    file format, please use xlrd to read this file" — inside a generic
    "can't read this price list" message, easy to misread as "it reads but
    can't write quantities". The fix is a one-line resave in Excel, and the
    message must say so.
    """

    # First eight bytes of an Excel 97-2003 (OLE2) file. The format is
    # identified by content, not extension — the same signature `inspect_sources` reads.
    FIRMA_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

    def setUp(self) -> None:
        temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(temporanea.cleanup)
        self.radice = Path(temporanea.name)

    def test_un_xls_senza_procedura_dichiarata_dice_che_cosa_fare(self) -> None:
        sorgente = self.radice / "New Larice.xls"
        sorgente.write_bytes(self.FIRMA_OLE2 + b"\x00" * 512)

        regola, motivo = launcher.source_rule(
            "larice", sorgente, {},
            {"id": "prova_v1", "supplier_id": "larice", "order_write": {
                "sheet": "FIRST", "header_row": 11, "data_start_row": 12, "order_column": "D",
            }},
        )

        self.assertIsNone(regola)
        self.assertIn("Excel 97-2003", motivo)
        self.assertIn(".xlsx", motivo)
        self.assertNotIn("openpyxl", motivo)

    def test_l_xls_che_si_compila_in_posizione_non_incontra_la_guardia(self) -> None:
        """A supplier whose `.xls` is patched in place must be unaffected:
        that branch returns before this guard, which never even sees it."""

        sorgente = self.radice / "formattato.xls"
        sorgente.write_bytes(self.FIRMA_OLE2 + b"\x00" * 512)

        regola, motivo = launcher.source_rule(
            "noce", sorgente,
            {"fieldMapping": {
                "sheet": "Foglio1", "header_row": 1, "data_start_row": 2,
                "order_column": "I", "columns": {"ean": "cat"},
            }},
            {"id": "prova_v1", "supplier_id": "noce", "order_write": {
                "mode": "patch_xls_in_posizione", "from_field_mapping": True,
                "sheet": "Foglio1", "header_row": 1, "data_start_row": 2, "order_column": "I",
            }},
        )

        self.assertIsNone(motivo)
        self.assertEqual(regola["compilazione"], "patch_xls_in_posizione")


# --------------------------------------------------------------------------
# The actual check that gates writing to the document.
# --------------------------------------------------------------------------
class LaVerificaDelDocumento(unittest.TestCase):
    """`source_rule` decides whether a price list is eligible to be written
    into. This covers the blank-header confirmation stored inside
    `order_write` — the only place it can live for a supplier with a
    dedicated reader and no field mapping."""

    def setUp(self) -> None:
        temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(temporanea.cleanup)
        self.radice = Path(temporanea.name)

    def listino(self, intestazioni: list[Any], righe: list[list[Any]]) -> Path:
        percorso = self.radice / "listino.xlsx"
        libro = Workbook()
        pagina = libro.active
        pagina.title = "Sheet1"
        pagina.append(intestazioni)
        for riga in righe:
            pagina.append(riga)
        libro.save(percorso)
        libro.close()
        return percorso

    @staticmethod
    def adattatore(**scrittura: Any) -> dict[str, Any]:
        base = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "order_column": "C",
        }
        base.update(scrittura)
        return {"id": "prova_v1", "supplier_id": "betulla", "order_write": base}

    def test_la_conferma_della_cella_vuota_puo_stare_nel_registro(self) -> None:
        sorgente = self.listino(["EAN", "Descrizione", None], [["8000000000001", "SAPONE", None]])

        regola, motivo = launcher.source_rule(
            "betulla", sorgente, {},
            self.adattatore(allow_blank_header_if_confirmed=True, order_header_blank_confirmed=True),
        )

        self.assertIsNone(motivo)
        self.assertIsNotNone(regola)
        self.assertEqual(regola["order_column"], "C")
        self.assertIs(regola["blank_header_confirmed"], True)

    def test_una_cella_che_non_e_piu_vuota_ferma_la_compilazione(self) -> None:
        """The stored declaration isn't enough on its own: today's document is
        checked too."""

        sorgente = self.listino(["EAN", "Descrizione", "PREZZO"], [["8000000000001", "SAPONE", 2]])

        regola, motivo = launcher.source_rule(
            "betulla", sorgente, {},
            self.adattatore(allow_blank_header_if_confirmed=True, order_header_blank_confirmed=True),
        )

        self.assertIsNone(regola)
        self.assertIn("non è più vuota", motivo)

    def test_del_testo_sotto_la_colonna_ferma_la_compilazione(self) -> None:
        sorgente = self.listino(
            ["EAN", "Descrizione", None],
            [["8000000000001", "SAPONE", None], ["8000000000002", "OLIO", "OMAGGIO"]],
        )

        regola, motivo = launcher.source_rule(
            "betulla", sorgente, {},
            self.adattatore(allow_blank_header_if_confirmed=True, order_header_blank_confirmed=True),
        )

        self.assertIsNone(regola)
        self.assertIn("riga 3", motivo)


class ConCheCosaSiScriveDentroQuestoDocumento(unittest.TestCase):
    """The run's own decision determines what's written, not the supplier's name.

    An `if supplier == "..."` shortcut is a recurring class of bug in this
    project, and here it's especially costly since it's on the write side: an
    order written into the wrong column leaves the app and goes to the
    supplier.

    Two cases where the supplier name alone isn't enough: a supplier can have
    two schemas with two different `order_write` blocks, and an order column
    moved from the page is stored as a `__locale` registry entry.
    """

    def setUp(self) -> None:
        temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(temporanea.cleanup)
        self.radice = Path(temporanea.name)
        self.registro = self.radice / "adapters.json"
        self.registro.write_text(json.dumps({"schema_version": 1, "adapters": [
            {"id": "tizio_v1", "supplier_id": "tizio", "kind": "supplier",
             "display_name": "TIZIO",
             "order_write": {"sheet": "FIRST", "header_row": 1, "data_start_row": 2,
                             "order_column": "C", "allow_blank_header_if_confirmed": True}},
            {"id": "tizio_secondo_v1", "supplier_id": "tizio", "kind": "supplier",
             "display_name": "TIZIO",
             "order_write": {"sheet": "FIRST", "header_row": 1, "data_start_row": 2,
                             "order_column": "D", "expected_header": "ORDINE"}},
        ]}, ensure_ascii=False), encoding="utf-8")
        self.vecchio = launcher.ADAPTERS_PATH
        launcher.ADAPTERS_PATH = self.registro
        self.addCleanup(setattr, launcher, "ADAPTERS_PATH", self.vecchio)
        self.registro_di_registro = registro.REGISTRO
        registro.REGISTRO = self.registro
        self.addCleanup(setattr, registro, "REGISTRO", self.registro_di_registro)

    def per_fornitore(self) -> dict[str, Any]:
        return launcher.adattatori_compilabili(self.registro)

    def test_fra_due_schemi_dello_stesso_fornitore_vale_quello_della_decisione(self) -> None:
        """Picking the registry's first matching entry instead of the one the
        run actually chose would lose `expected_header` — the check that the
        column is titled "ORDINE" before writing to it.
        """

        scelto = launcher.adattatore_del_documento(
            {"adapterId": "tizio_secondo_v1"}, "tizio", self.per_fornitore(), self.registro,
        )

        self.assertEqual(scelto["id"], "tizio_secondo_v1")
        self.assertEqual(scelto["order_write"]["order_column"], "D")
        self.assertEqual(scelto["order_write"]["expected_header"], "ORDINE")

    def test_la_colonna_spostata_a_mano_vince_su_quella_spedita(self) -> None:
        """The run's decision still names the shipped adapter's id; looking it
        up literally would write the order into the old column."""

        imparato = registro.percorso_imparato(self.registro)
        # An entry that overrides a shipped one is only valid when it carries
        # the shipped entry's own stamp; without it, it's a stale entry to
        # set aside, not today's active choice.
        spedita = registro.adattatore("tizio_v1", self.registro)
        imparato.write_text(json.dumps({"schema_version": 1, "adapters": [
            {"id": "tizio_v1__locale", "supplier_id": "tizio", "kind": "supplier",
             "display_name": "TIZIO", "derivato_da": "tizio_v1",
             "sopra_spedito": registro.sopra_spedito_di(spedita),
             "order_write": {"sheet": "FIRST", "header_row": 1, "data_start_row": 2,
                             "order_column": "H", "allow_blank_header_if_confirmed": True}},
        ]}, ensure_ascii=False), encoding="utf-8")

        scelto = launcher.adattatore_del_documento(
            {"adapterId": "tizio_v1"}, "tizio", self.per_fornitore(), self.registro,
        )

        self.assertEqual(scelto["id"], "tizio_v1__locale")
        self.assertEqual(scelto["order_write"]["order_column"], "H")

    def test_una_scheda_senza_adattatore_ripiega_sul_fornitore(self) -> None:
        """Comparisons run before this feature existed declare no adapter on
        the file entry: falling back to the supplier's own entry beats no
        entry at all."""

        scelto = launcher.adattatore_del_documento({}, "tizio", self.per_fornitore(), self.registro)

        self.assertEqual(scelto["id"], "tizio_v1")

    def test_un_adattatore_che_il_registro_non_conosce_ripiega_sul_fornitore(self) -> None:
        scelto = launcher.adattatore_del_documento(
            {"adapterId": "sparito_v1"}, "tizio", self.per_fornitore(), self.registro,
        )

        self.assertEqual(scelto["id"], "tizio_v1")


# --------------------------------------------------------------------------
# The full command, against a real (fixture) registry and comparison.
# --------------------------------------------------------------------------
class IlComandoCompleto(unittest.TestCase):
    def setUp(self) -> None:
        temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(temporanea.cleanup)
        self.radice = Path(temporanea.name)
        self.dati = self.radice / "data"
        self.corrente = self.dati / "current"
        self.esecuzioni = self.corrente / "esecuzioni"
        self.uploads = self.corrente / "uploads"
        self.uploads.mkdir(parents=True)
        (self.esecuzioni / "run-prova").mkdir(parents=True)
        self.registro = self.radice / "adapters.json"
        # The fixture registry must also be reachable where path-less callers
        # look for it (`fornitori_compilabili`); otherwise the test would
        # partly run against the project's real registry.
        self.registro_vero = launcher.ADAPTERS_PATH
        launcher.ADAPTERS_PATH = self.registro
        self.addCleanup(setattr, launcher, "ADAPTERS_PATH", self.registro_vero)
        self.registro_di_registro = registro.REGISTRO
        registro.REGISTRO = self.registro
        self.addCleanup(setattr, registro, "REGISTRO", self.registro_di_registro)

    # -- gli ingredienti ---------------------------------------------------
    def scrivi_listino(self) -> Path:
        percorso = self.uploads / "listino.xlsx"
        libro = Workbook()
        pagina = libro.active
        pagina.title = "Sheet1"
        pagina.append(["EAN", "Descrizione", "ORDINE", "PzCt", "Prezzo", None, "TOTALI"])
        pagina.append(["8000000000001", "SAPONE", None, 6, 1.5, None, "=D2*E2"])
        pagina.append(["8000000000002", "OLIO", None, 4, 2.5, None, "=D3*E3"])
        libro.save(percorso)
        libro.close()
        return percorso

    def scrivi_registro(self, **extra: Any) -> None:
        adattatore = {
            "id": "prova_v1",
            "supplier_id": "betulla",
            "display_name": "BETULLA",
            "kind": "supplier",
            "file_types": [".xlsx"],
            "header_signature": {
                "kind": "headers",
                "sheet": "FIRST",
                "header_row": 1,
                "data_start_row": 2,
                # Positions are keyed by the HEADER tokens, not the field
                # names: `_verifica_posizioni` looks for "descrizione" in the
                # document's header row, and "description" would never match,
                # forcing this document to always read as SCHEMA_VARIATO and
                # defeating the recognition test.
                "columns": {"ean": 1, "descrizione": 2, "pzct": 4, "prezzo": 5},
                "required": ["EAN", "Descrizione", "PzCt", "Prezzo"],
            },
            "header_aliases": {
                "ean": ["EAN"],
                "description": ["Descrizione"],
                "pieces_per_carton": ["PzCt"],
                "unit_price_net": ["Prezzo"],
            },
            "order_write": {
                "sheet": "FIRST",
                "header_row": 1,
                "data_start_row": 2,
                "order_column": "C",
                "expected_header": "ORDINE",
            },
        }
        adattatore.update(extra)
        self.registro.write_text(
            json.dumps({"schema_version": 1, "adapters": [adattatore]}, ensure_ascii=False),
            encoding="utf-8",
        )

    def scrivi_run(self, sorgente: Path) -> None:
        cartella = self.esecuzioni / "run-prova"
        profilo = profile_file(sorgente)
        profilo["profile_id"] = "prova"
        (cartella / "input_profiles.json").write_text(
            json.dumps({"profiles": [profilo]}, ensure_ascii=False), encoding="utf-8",
        )
        (cartella / "input_manifest.json").write_text(
            json.dumps({"files": [{
                "file_name": sorgente.name,
                "profile_id": "prova",
                "ai_preflight": {
                    "role": "supplier",
                    "supplier_id": "betulla",
                    "adapter_id": "prova_v1",
                },
            }]}, ensure_ascii=False),
            encoding="utf-8",
        )

    def scrivi_confronto(self, sorgente: Path) -> None:
        (self.corrente).mkdir(parents=True, exist_ok=True)
        (self.corrente / "review_data.json").write_text(json.dumps({
            "run": {"id": "run-prova", "pipelineRunId": "run-prova", "status": "ready"},
            "files": [{
                "id": "file:1",
                "name": sorgente.name,
                "role": "supplier",
                "supplier": "BETULLA",
                "supplierId": "betulla",
                "sourcePath": str(sorgente),
                "adapterId": "prova_v1",
            }],
            "suppliers": [{"id": "betulla", "name": "BETULLA", "minimumOrder": 0}],
            "products": [],
            "warnings": [],
        }, ensure_ascii=False), encoding="utf-8")

    def negozio(self) -> ReviewStore:
        sorgente = self.scrivi_listino()
        self.scrivi_registro()
        self.scrivi_run(sorgente)
        self.scrivi_confronto(sorgente)
        store = ReviewStore(
            self.corrente / "review_data.json",
            self.corrente / "state.json",
            self.uploads,
            self.corrente / "outputs",
            self.corrente / "writer_config.json",
            self.dati / "history",
            self.corrente / "ordini",
            conferme_path=self.dati / "history" / "conferme.db",
        )
        store.pipeline_jobs.configurazione = dataclasses.replace(
            store.pipeline_jobs.configurazione, adapters_path=self.registro,
        )
        self.addCleanup(self.chiudi, store)
        return store

    @staticmethod
    def chiudi(store: ReviewStore) -> None:
        magazzino = getattr(store, "_conferme", None)
        if magazzino is not None:
            magazzino.chiudi()

    def voce_del_registro(self, identificativo: str = "prova_v1") -> dict[str, Any]:
        """The entry as the app sees it: shipped and learned merged together.

        A column moved from the page never lands in `references/adapters.json`
        — that file is under git, and starting the app resets it — but in the
        separate learned registry, outside git.

        Nor does it land under the same `id`: the new entry is
        `prova_v1__locale`, and the shipped one stays untouched. This asks
        `voce_in_uso`, not `adattatore` — the same question the app asks when
        it needs to know what writes today's order.
        """

        letta = registro.voce_in_uso(identificativo, registro.adattatori(self.registro))
        if not letta:
            raise AssertionError(f"«{identificativo}» non è nel registro effettivo")
        return letta

    def scrittura_nel_registro(self) -> dict[str, Any]:
        return self.voce_del_registro()["order_write"]

    # -- le prove ----------------------------------------------------------
    def test_le_colonne_offerte_dicono_che_cosa_contengono(self) -> None:
        store = self.negozio()

        esito = store.colonne_d_ordine("betulla")
        per_lettera = {voce["lettera"]: voce for voce in esito["colonne"]}

        self.assertEqual(esito["attuale"]["lettera"], "C")
        self.assertEqual(per_lettera["B"]["occupataDa"], "Nome del prodotto")
        self.assertIs(per_lettera["B"]["scegliibile"], False)
        self.assertEqual(per_lettera["G"]["formule"], 2)
        self.assertIs(per_lettera["F"]["scegliibile"], True)

    def test_la_colonna_nuova_entra_nel_registro_e_nella_compilazione(self) -> None:
        store = self.negozio()

        esito = store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        self.assertIs(esito["ok"], True)
        self.assertEqual(esito["colonna"], "F")
        self.assertEqual(self.scrittura_nel_registro()["order_column"], "F")
        # No title above column F: it's declared blank, and the writer will
        # recheck that against the document before writing.
        self.assertIs(self.scrittura_nel_registro()["allow_blank_header_if_confirmed"], True)
        configurazione = json.loads((self.corrente / "writer_config.json").read_text(encoding="utf-8"))
        self.assertEqual(configurazione["supplier_write_rules"]["betulla"]["order_column"], "F")

    def test_se_la_scrittura_non_si_rifa_lo_dice_invece_di_tacere(self) -> None:
        """By this point the column has already changed: the declaration is
        written and has passed the document check. Letting the exception
        propagate unhandled would claim nothing happened, when almost
        everything did.
        """

        store = self.negozio()

        def non_si_rifa(_review):
            raise RuntimeError("Writer XLSX non attivato; manca: Node 18+")

        store.riconfigura_compilazione = non_si_rifa

        esito = store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        self.assertIs(esito["ok"], True)
        self.assertEqual(self.scrittura_nel_registro()["order_column"], "F")
        self.assertIn("Node 18+", esito["avviso"])
        self.assertIn("Node 18+", esito["message"])
        self.assertNotIn("Vale da subito", esito["message"])

    def test_il_registro_tiene_la_versione_di_prima(self) -> None:
        """If a later run regresses, the only way to see what changed is
        still having the previous version on record."""

        store = self.negozio()

        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        voce = self.voce_del_registro()
        self.assertEqual(voce["schema_version"], 2)
        self.assertEqual(voce["previous_versions"][0]["order_write"]["order_column"], "C")

    def test_la_colonna_spostata_non_entra_nel_registro_spedito(self) -> None:
        """Moving the order column is a decision made at the store; writing
        it into the shipped file would lose it on the next app restart,
        since that file gets reset to the shipped version.
        """

        store = self.negozio()
        spedito_prima = self.registro.read_bytes()

        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        self.assertEqual(self.registro.read_bytes(), spedito_prima)
        imparato = registro.percorso_imparato(self.registro)
        self.assertTrue(imparato.is_file())
        self.assertEqual(self.voce_del_registro()["order_write"]["order_column"], "F")

    def test_la_voce_spedita_resta_intera_e_quella_nuova_prende_un_id_suo(self) -> None:
        """Moving the order column must not write a full copy of the shipped
        entry into the learned registry under the same `id`: that would bury
        the shipped entry and stop any future update to it from ever
        reaching the app — commercial terms, header aliases, column
        positions, all of it.
        """

        store = self.negozio()
        spedita_prima = registro.adattatore("prova_v1", self.registro)

        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        # The new entry has its own id and records where it came from.
        imparate = json.loads(
            registro.percorso_imparato(self.registro).read_text(encoding="utf-8")
        )["adapters"]
        self.assertEqual([voce["id"] for voce in imparate], ["prova_v1__locale"])
        self.assertEqual(imparate[0]["derivato_da"], "prova_v1")
        self.assertEqual(imparate[0]["order_write"]["order_column"], "F")

        # And the shipped entry is still there, unchanged: nothing shadowed it.
        self.assertEqual(registro.adattatore("prova_v1", self.registro), spedita_prima)
        self.assertEqual(spedita_prima["order_write"]["order_column"], "C")

    def test_la_pagina_mostra_subito_la_colonna_nuova(self) -> None:
        """The run's decision still names the shipped adapter's id; looking
        it up literally would find the shipped entry with the old column,
        undoing a move that already happened.
        """

        store = self.negozio()
        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        documento = store.pipeline_jobs.documento_del_fornitore("betulla")

        self.assertEqual(documento["adattatore"]["id"], "prova_v1__locale")
        self.assertEqual(documento["effettiva"]["orderColumn"]["lettera"], "F")
        self.assertEqual(store.colonne_d_ordine("betulla")["attuale"]["lettera"], "F")

    def test_al_prossimo_ricalcolo_la_colonna_resta_quella_nuova(self) -> None:
        """The test that proves the move actually stuck.

        A later recompute ignores any past decision and reruns `riconosci` on
        the document, with two candidates: the shipped entry and the local
        one. Both read the document identically — same confidence, same
        checks, same required fields — since only the order column differs
        between them. If ties were broken in favor of the shipped entry, the
        column would silently revert on the next recompute.
        """

        sorgente = self.uploads / "listino.xlsx"
        store = self.negozio()
        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        esito = registro.riconosci(profile_file(sorgente)["details"], self.registro)

        self.assertEqual(esito["state"], "SCHEMA_NOTO")
        self.assertEqual(esito["adapter_id"], "prova_v1__locale")
        vincente = registro.adattatore("prova_v1__locale", self.registro)
        self.assertEqual(vincente["order_write"]["order_column"], "F")

    def test_spostarla_due_volte_non_moltiplica_le_voci(self) -> None:
        """Each move is a new version of the same local entry, not a new
        entry: a registry that grows on every change would become unreadable
        exactly when it needs to be understood."""

        store = self.negozio()
        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})
        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "H"})

        imparate = json.loads(
            registro.percorso_imparato(self.registro).read_text(encoding="utf-8")
        )["adapters"]
        self.assertEqual([voce["id"] for voce in imparate], ["prova_v1__locale"])
        self.assertEqual(imparate[0]["order_write"]["order_column"], "H")
        self.assertEqual(imparate[0]["schema_version"], 3)
        # History starts from the shipped column and passes through F.
        storia = [voce["order_write"]["order_column"] for voce in imparate[0]["previous_versions"]]
        self.assertEqual(storia, ["C", "F"])

    def test_una_colonna_di_formule_si_rifiuta(self) -> None:
        store = self.negozio()

        with self.assertRaises(ValueError) as errore:
            store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "G"})

        self.assertIn("formule", str(errore.exception))
        self.assertEqual(self.scrittura_nel_registro()["order_column"], "C")

    def test_una_colonna_che_il_programma_legge_si_rifiuta(self) -> None:
        store = self.negozio()

        with self.assertRaises(ValueError) as errore:
            store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "E"})

        self.assertIn("Prezzo", str(errore.exception))
        self.assertEqual(self.scrittura_nel_registro()["order_column"], "C")

    def test_una_colonna_fuori_dal_foglio_si_rifiuta(self) -> None:
        store = self.negozio()

        with self.assertRaises(ValueError):
            store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "ZZ"})

        self.assertEqual(self.scrittura_nel_registro()["order_column"], "C")

    def test_se_la_verifica_del_documento_non_passa_non_si_scrive_niente(self) -> None:
        """The real gate isn't in this module: it's `source_rule`, the same
        function that decides whether a price list is written to. This checks
        that when it says no, the registry stays untouched — otherwise a
        column would end up recorded that can never actually be used."""

        store = self.negozio()
        # A data row past the end of the sheet: the check must stop here.
        self.scrivi_registro(order_write={
            "sheet": "FIRST", "header_row": 1, "data_start_row": 99,
            "order_column": "C", "expected_header": "ORDINE",
        })

        with self.assertRaises(ValueError) as errore:
            store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        self.assertIn("riga 99", str(errore.exception))
        self.assertEqual(self.scrittura_nel_registro()["order_column"], "C")
        self.assertFalse((self.corrente / "writer_config.json").is_file())

    def test_un_fornitore_che_non_e_nel_confronto_lo_dice(self) -> None:
        store = self.negozio()

        with self.assertRaises(ValueError) as errore:
            store.cambia_colonna_d_ordine({"supplierId": "larice", "colonna": "F"})

        self.assertIn("nessun listino di questo fornitore", str(errore.exception))

    def test_durante_un_ricalcolo_non_si_cambia(self) -> None:
        """The write configuration is rebuilt when a recompute finishes;
        changing it while one is running would get silently overwritten a
        moment later."""

        store = self.negozio()
        store.pipeline_jobs.in_corso = lambda: True

        with self.assertRaises(LavoroGiaInCorso):
            store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        self.assertEqual(self.scrittura_nel_registro()["order_column"], "C")

    def test_per_i_fornitori_con_mappatura_si_cambiano_tutte_e_due(self) -> None:
        """`source_rule` refuses when the registry and the confirmed field
        mapping name different columns; updating only one would leave the
        supplier unwritable, with a message that blames the mapping."""

        store = self.negozio()
        self.scrivi_registro(
            field_mapping={
                "sheet": "Sheet1",
                "header_row": 1,
                "data_start_row": 2,
                "columns": {"ean": "EAN", "description": "Descrizione",
                            "pieces_per_carton": "PzCt", "unit_price_net": "Prezzo"},
                "order_column": "C",
                "order_header_expected": "ORDINE",
            },
            order_write={
                "from_field_mapping": True,
                "order_column": "C",
                "expected_header": "ORDINE",
                "required_columns": ["ean", "description"],
            },
        )
        review = json.loads((self.corrente / "review_data.json").read_text(encoding="utf-8"))
        review["files"][0]["fieldMapping"] = {
            "sheet": "Sheet1", "header_row": 1, "data_start_row": 2,
            "columns": {"ean": "EAN", "description": "Descrizione",
                        "pieces_per_carton": "PzCt", "unit_price_net": "Prezzo"},
            "order_column": "C", "order_header_expected": "ORDINE",
        }
        (self.corrente / "review_data.json").write_text(
            json.dumps(review, ensure_ascii=False), encoding="utf-8")

        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        voce = self.voce_del_registro()
        self.assertEqual(voce["order_write"]["order_column"], "F")
        self.assertEqual(voce["field_mapping"]["order_column"], "F")
        riletto = json.loads((self.corrente / "review_data.json").read_text(encoding="utf-8"))
        self.assertEqual(riletto["files"][0]["fieldMapping"]["order_column"], "F")

    def test_la_decisione_scritta_a_mano_segue(self) -> None:
        """A manual per-file decision overrides the registry; leaving it
        stale would show the new column today but silently revert to the
        old one on the next recompute."""

        store = self.negozio()
        decisioni = store.pipeline_jobs.configurazione.decisioni_manuali_path
        decisioni.write_text(json.dumps({"decisions": [{
            "file_name": "listino.xlsx",
            "role": "supplier",
            "supplier_id": "betulla",
            "adapter_id": "prova_v1",
            "field_mapping": {"order_column": "C", "order_header_expected": "ORDINE"},
        }]}, ensure_ascii=False), encoding="utf-8")

        store.cambia_colonna_d_ordine({"supplierId": "betulla", "colonna": "F"})

        riletto = json.loads(decisioni.read_text(encoding="utf-8"))
        self.assertEqual(riletto["decisions"][0]["field_mapping"]["order_column"], "F")
        self.assertNotIn("order_header_expected", riletto["decisions"][0]["field_mapping"])


if __name__ == "__main__":
    unittest.main()
