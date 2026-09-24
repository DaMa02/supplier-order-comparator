#!/usr/bin/env python3
"""The bridge between supplier price lists and the promotion engine.

Covers what the engine alone can't: what a price-list reader recognizes as a
commercial condition, what it fails to parse, and whether that failure leaves
a trace the user can see instead of disappearing silently.

The reader must work for any supplier the adapter registry declares, not one
hardcoded name: the tests run both against the real larice price list (13
conditions out of 13) and against a synthetic supplier the code never
references by name.
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

RADICE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RADICE / "scripts"))
sys.path.insert(0, str(RADICE / "app"))

import promotion_bridge  # noqa: E402
from promotions import KIND_AMBIGUOUS, KIND_INCLUDED_PACK, KIND_THRESHOLD_GIFT  # noqa: E402


# larice price-list columns used by the reader: G description, J reward name,
# P discount code, R barcode.
COLONNA_G = 7
COLONNA_J = 10
COLONNA_P = 16
COLONNA_R = 18


def scrivi_listino(
    percorso: Path,
    righe: list[dict[str, object]],
    colonne: dict[str, int] | None = None,
) -> Path:
    """Build a fake larice price list with only the columns that matter.

    `colonne` lets a test move them, to prove the reader follows the adapter
    registry rather than a hardcoded column position.
    """

    posizioni = colonne or {"g": COLONNA_G, "j": COLONNA_J, "p": COLONNA_P, "r": COLONNA_R}
    libro = Workbook()
    foglio = libro.active
    foglio.title = "Canvass di prova"
    for numero, contenuto in enumerate(righe, start=1):
        for chiave, colonna in posizioni.items():
            valore = contenuto.get(chiave)
            if valore not in (None, ""):
                foglio.cell(row=numero, column=colonna, value=valore)
    libro.save(percorso)
    libro.close()
    return percorso


def riga_merce(nome: str, ean: str) -> dict[str, object]:
    return {"g": nome, "p": "TP", "r": ean}


class LetturaDeiBlocchiLarice(unittest.TestCase):
    def setUp(self) -> None:
        self.cartella = Path(tempfile.mkdtemp(prefix="ponte-promozioni-"))
        self.servizio = promotion_bridge.PromotionService()

    def leggi(self, righe: list[dict[str, object]]) -> list[dict]:
        percorso = scrivi_listino(self.cartella / "larice.xlsx", righe)
        promozioni, errore = self.servizio.condizioni_di("larice", percorso)
        self.assertIsNone(errore)
        return promozioni

    def test_l_ean_della_riga_premio_entra_nel_contratto(self):
        """The reward row's EAN must reach `reward.ean`, not just get parsed
        and discarded: without it there's no way to identify the free item
        when the goods arrive."""

        promozioni = self.leggi([
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("BIOPUNTO LATTE CORPO 75ML", "8059651466354"),
            riga_merce("BIOPUNTO LATTE CORPO 200ML", "8059373360282"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "BIOPUNTO SET MARE", "p": "SM", "r": "8050507999999"},
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(promozioni[0]["reward"]["ean"], "8050507999999")
        self.assertEqual(promozioni[0]["reward"]["description"], "BIOPUNTO SET MARE")
        # The reward row doesn't count toward the eligible quantity.
        self.assertEqual(promozioni[0]["eligible"]["source_rows"], [2, 3])

    def test_una_grafia_reale_del_verbo_apre_comunque_il_blocco(self):
        """The header pattern must accept "ACQUISTANO" (missing a D) as well
        as the canonical form: real price lists use it, and a stricter regex
        would drop the condition silently."""

        promozioni = self.leggi([
            {"g": "ACQUISTANO 5 CT TRA"},
            riga_merce("DENT. SENSODENT 80+20ML FRESH CLEAN", "5059833580611"),
            riga_merce("DENT. SENSODENT 80+20ML COMPLEX", "5059905473735"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "DENT. SENSODENT 15 ML", "p": "SM", "r": "5059187155220"},
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(promozioni[0]["threshold"], {"qty": 5, "unit": "cartoni"})

    def test_un_blocco_che_non_si_calcola_resta_da_verificare(self):
        """A block with header and eligible rows but no computable threshold
        must surface as ambiguous, not disappear: otherwise a commercial
        condition would go unreported."""

        promozioni = self.leggi([
            {"g": "ACQUISTANDO 3 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO", "j": "", "p": "SM"},
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_AMBIGUOUS)
        self.assertFalse(promozioni[0]["confirmed"])
        self.assertFalse(promozioni[0]["economic_effect"]["affects_total"])
        self.assertIn("non ricomposta", promozioni[0]["source_text"])
        self.assertEqual(promozioni[0]["eligible"]["source_rows"], [2])

    def test_una_unita_mai_vista_apre_comunque_il_blocco(self):
        """Header recognition is deliberately looser than threshold
        calculation: requiring a known unit too would make an unfamiliar one
        (e.g. "SCATOLE") drop the condition entirely instead of flagging it
        for review."""

        promozioni = self.leggi([
            {"g": "ACQUISTANDO 10 SCATOLE TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_AMBIGUOUS)
        self.assertIn("ACQUISTANDO 10 SCATOLE", promozioni[0]["source_text"])
        self.assertEqual(promozioni[0]["eligible"]["source_rows"], [2])

    def test_un_blocco_troppo_lungo_lascia_una_traccia(self):
        """A block is abandoned past 500 rows, but that must be visible, not
        silent."""

        righe: list[dict[str, object]] = [{"g": "ACQUISTANDO 4 CT TRA"}]
        righe.extend(riga_merce(f"PRODOTTO {n}", f"800000000{n:04d}") for n in range(3))
        # Rows with no EAN: these are what triggers the distance-from-header
        # count.
        righe.extend({"g": ""} for _ in range(520))
        righe.append({"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"})

        promozioni = self.leggi(righe)

        tracce = [p for p in promozioni if p["kind"] == KIND_AMBIGUOUS]
        # Two distinct losses: the abandoned header, and later the reward
        # row left without one. Both must be reported.
        self.assertEqual(len(tracce), 2)
        abbandono = next(p for p in tracce if "500 righe" in p["source_text"])
        self.assertEqual(abbandono["eligible"]["source_rows"], [2, 3, 4])
        self.assertTrue(any("manca l'intestazione" in p["source_text"] for p in tracce))
        self.assertEqual([p["kind"] for p in promozioni], [KIND_AMBIGUOUS, KIND_AMBIGUOUS])

    def test_una_seconda_intestazione_non_cancella_la_prima_in_silenzio(self):
        promozioni = self.leggi([
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "ACQUISTANDO 6 CT TRA"},
            riga_merce("PRODOTTO DUE", "8000000000002"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])

        tipi = [p["kind"] for p in promozioni]
        self.assertEqual(tipi.count(KIND_THRESHOLD_GIFT), 1)
        self.assertEqual(tipi.count(KIND_AMBIGUOUS), 1)
        traccia = next(p for p in promozioni if p["kind"] == KIND_AMBIGUOUS)
        self.assertIn("ACQUISTANDO 2 CT TRA", traccia["source_text"])
        self.assertIn("intestazione successiva", traccia["source_text"])

    def test_un_blocco_aperto_a_fine_foglio_lascia_una_traccia(self):
        promozioni = self.leggi([
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_AMBIGUOUS)
        self.assertIn("finisce prima", promozioni[0]["source_text"])

    def test_una_riga_premio_senza_intestazione_lascia_una_traccia(self):
        promozioni = self.leggi([
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])

        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_AMBIGUOUS)
        self.assertIn("manca l'intestazione", promozioni[0]["source_text"])

    def test_le_condizioni_non_ricomposte_sono_contate(self):
        """The count of unparsed conditions must be exposed, not something an
        operator has to tally by hand."""

        percorso = scrivi_listino(self.cartella / "larice.xlsx", [
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
        ])
        review = {"files": [{"supplierId": "larice", "sourcePath": str(percorso)}], "products": []}

        promozioni = self.servizio.detect(review)

        self.assertEqual(self.servizio.condizioni_da_verificare["larice"], 1)
        self.assertEqual(len(promozioni), 1)
        # The count must survive a second read, which hits the cache.
        self.servizio.detect(review)
        self.assertEqual(self.servizio.condizioni_da_verificare["larice"], 1)

    def test_gli_sconti_gia_nel_prezzo_non_entrano_nell_elenco(self):
        """A flat per-row discount that's already reflected in the offer's
        `unitPriceNet` must be dropped from the reported list, not just
        counted: keeping it would bury threshold-gift conditions (the ones
        that actually need review) under a large number of no-op discounts.
        """

        percorso = scrivi_listino(self.cartella / "larice.xlsx", [
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "SALE LAVASTOVIGLIE KG1", "p": "SM", "r": "8000000009999"},
        ])
        review = {
            "files": [{"supplierId": "larice", "sourcePath": str(percorso)}],
            "products": [
                {
                    "id": f"product:{numero}",
                    "ean": f"800000000{numero:04d}",
                    "offers": [{
                        "supplierId": "larice",
                        "sourceRow": numero,
                        "available": True,
                        "discountRate": 0.1,
                    }],
                }
                for numero in range(1, 6)
            ],
        }

        promozioni = self.servizio.detect(review)

        self.assertEqual([p["kind"] for p in promozioni], [KIND_THRESHOLD_GIFT])
        self.assertEqual(self.servizio.sconti_gia_nel_prezzo, 5)

    def test_uno_sconto_che_il_prezzo_non_contiene_resta(self):
        """A discount not already reflected in the price changes the total,
        so it must stay in the reported list."""

        promozione = promotion_bridge.detect_numeric_discount(
            supplier="larice",
            source_reference="larice:riga:1",
            source_text="Sconto numerico 10%",
            discount_value=0.1,
            eligible={"products": ["product:1"]},
            already_applied=False,
        )

        tenute = self.servizio._solo_quelle_che_cambiano_qualcosa([promozione])

        self.assertEqual(len(tenute), 1)
        self.assertEqual(self.servizio.sconti_gia_nel_prezzo, 0)

    def test_un_listino_senza_perdite_non_inventa_tracce(self):
        promozioni = self.leggi([
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])

        self.assertEqual([p["kind"] for p in promozioni], [KIND_THRESHOLD_GIFT])


class ConUnRegistroFinto(unittest.TestCase):
    """Swaps in a fake adapter registry for the duration of a test.

    Always starts from the real registry and changes one declaration at a
    time: an adapter also carries other facts (e.g. larice's row codes,
    which mark `SM` as the reward rather than orderable goods), and a
    stripped-down fake would exercise a shape that never occurs in practice.
    """

    def setUp(self) -> None:
        self.cartella = Path(tempfile.mkdtemp(prefix="registro-finto-"))
        self.registro_vero = promotion_bridge.registro.REGISTRO

    def tearDown(self) -> None:
        promotion_bridge.registro.REGISTRO = self.registro_vero

    def registro_finto(self, modifica) -> None:
        documento = json.loads(self.registro_vero.read_text(encoding="utf-8"))
        modifica(documento)
        percorso = self.cartella / "adapters.json"
        percorso.write_text(json.dumps(documento, ensure_ascii=False), encoding="utf-8", newline="")
        promotion_bridge.registro.REGISTRO = percorso

    @staticmethod
    def adattatore(documento: dict, identificativo: str) -> dict:
        return next(voce for voce in documento["adapters"] if voce["id"] == identificativo)


class ColonneDalRegistro(ConUnRegistroFinto):
    """The reader must get larice's column positions from the adapter
    registry, not a hardcoded layout: a supplier can add a column, or the
    user can confirm a schema change that the registry then learns, and the
    reader must follow along instead of silently losing threshold-gift
    conditions from the summary."""

    SOGLIA = [
        {"g": "ACQUISTANDO 5 CT TRA"},
        riga_merce("DENT. SENSODENT 80+20ML FRESH CLEAN", "5059833580611"),
        {"g": "IN OMAGGIO 1 CT DI", "j": "DENT. SENSODENT 15 ML", "p": "SM", "r": "5059187155220"},
    ]

    # Today's column positions, for tests that move or drop just one.
    COLONNE_DI_OGGI = {
        "description": "G", "reward_description": "J", "discount": "P", "ean": "R",
    }

    def setUp(self) -> None:
        super().setUp()
        self.servizio = promotion_bridge.PromotionService()

    def registro_con(self, colonne_di_larice: dict[str, object]) -> None:
        """A registry identical to the real one except for larice's columns."""

        def sostituisci(documento: dict) -> None:
            self.adattatore(documento, "larice_v1")["column_map"] = colonne_di_larice

        self.registro_finto(sostituisci)

    def senza(self, *colonne: str) -> dict[str, object]:
        """Today's columns minus the ones the test wants to drop."""

        return {campo: dove for campo, dove in self.COLONNE_DI_OGGI.items() if campo not in colonne}

    def test_una_colonna_spostata_nel_registro_sposta_anche_il_lettore(self):
        """The registry already reflects an inserted column; the reader must
        follow it. Reading the wrong (adjacent) column raises nothing in
        openpyxl, so a stale hardcoded position would drop the condition
        without any error.
        """

        self.registro_con({
            "description": "H", "reward_description": "K", "discount": "Q", "ean": "S",
        })
        percorso = scrivi_listino(
            self.cartella / "larice.xlsx", self.SOGLIA,
            colonne={"g": 8, "j": 11, "p": 17, "r": 19},
        )

        promozioni, errore = self.servizio.condizioni_di("larice", percorso)

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(promozioni[0]["threshold"], {"qty": 5, "unit": "cartoni"})
        self.assertEqual(promozioni[0]["reward"]["description"], "DENT. SENSODENT 15 ML")
        self.assertEqual(promozioni[0]["reward"]["ean"], "5059187155220")
        self.assertTrue(promozioni[0]["confirmed"])
        # The reward row still doesn't count toward the eligible quantity:
        # the `SM` code marks it regardless of where the columns moved to.
        self.assertEqual(promozioni[0]["eligible"]["source_rows"], [2])
        # The cell reference must point at today's columns, not the old ones.
        self.assertEqual(promozioni[0]["source_reference"], "Canvass di prova!H1:K3")

    def test_una_colonna_che_il_registro_non_dichiara_ferma_la_lettura(self):
        """A missing column must stop the read with an error, not guess:
        reading an adjacent column would silently produce zero thresholds."""

        self.registro_con(self.senza("ean"))
        percorso = scrivi_listino(self.cartella / "larice.xlsx", self.SOGLIA)
        review = {"files": [{"supplierId": "larice", "sourcePath": str(percorso)}], "products": []}

        self.assertEqual(self.servizio.detect(review), [])
        self.assertEqual(len(self.servizio.load_errors), 1)
        self.assertEqual(self.servizio.load_errors[0]["supplier"], "larice")
        self.assertEqual(self.servizio.load_errors[0]["supplierName"], "LARICE")
        self.assertEqual(
            self.servizio.load_errors[0]["message"],
            "Non riesco a leggere le condizioni commerciali di LARICE: non so più dove il "
            "listino tiene il codice a barre.",
        )

    def test_il_nome_dell_articolo_in_omaggio_vale_come_le_altre_colonne(self):
        """The reward-name column is required like the others, not optional.

        Without it the threshold still parses, but as unconfirmed rather
        than confirmed, which would silently downgrade the condition's
        status instead of raising a clear error.
        """

        self.registro_con(self.senza("reward_description"))
        percorso = scrivi_listino(self.cartella / "larice.xlsx", self.SOGLIA)

        promozioni, errore = self.servizio.condizioni_di("larice", percorso)

        self.assertEqual(promozioni, [])
        self.assertEqual(
            errore["message"],
            "Non riesco a leggere le condizioni commerciali di LARICE: non so più dove il "
            "listino tiene il nome dell'articolo in omaggio.",
        )

    def test_due_colonne_che_mancano_si_dicono_come_le_direbbe_una_persona(self):
        self.registro_con(self.senza("description", "ean"))
        percorso = scrivi_listino(self.cartella / "larice.xlsx", self.SOGLIA)

        _promozioni, errore = self.servizio.condizioni_di("larice", percorso)

        self.assertEqual(
            errore["message"],
            "Non riesco a leggere le condizioni commerciali di LARICE: non so più dove il "
            "listino tiene le descrizioni e il codice a barre.",
        )

    def test_una_colonna_dichiarata_in_un_modo_che_non_si_legge_non_si_tira_a_indovinare(self):
        """A column declared as a header name ("9") must not be resolved by
        position: larice's price list has no header row, so treating "9" as
        column 9 would pick a column no other reader agrees on."""

        self.registro_con({**self.COLONNE_DI_OGGI, "ean": "9"})
        percorso = scrivi_listino(self.cartella / "larice.xlsx", self.SOGLIA)

        promozioni, errore = self.servizio.condizioni_di("larice", percorso)

        self.assertEqual(promozioni, [])
        self.assertIn("il codice a barre", errore["message"])


class IlListinoCheNonSiRisolve(unittest.TestCase):
    """A price list the review references but fails to open (moved, deleted
    after a recompute, a dropped network path) must surface as a load error,
    not silently read as "no promotions"."""

    def setUp(self) -> None:
        self.cartella = Path(tempfile.mkdtemp(prefix="listino-sparito-"))
        self.servizio = promotion_bridge.PromotionService()

    def test_un_listino_che_non_si_apre_piu_diventa_un_avviso(self):
        sparito = self.cartella / "larice.xlsx"
        review = {"files": [{"supplierId": "larice", "sourcePath": str(sparito)}], "products": []}

        self.assertEqual(self.servizio.detect(review), [])

        self.assertEqual(len(self.servizio.load_errors), 1)
        self.assertEqual(self.servizio.load_errors[0]["supplier"], "larice")
        self.assertEqual(self.servizio.load_errors[0]["supplierName"], "LARICE")
        self.assertEqual(
            self.servizio.load_errors[0]["message"],
            "Non riesco a leggere le condizioni commerciali di LARICE: il listino non è fra "
            "i documenti caricati.",
        )

    def test_un_fornitore_che_la_review_non_dichiara_non_diventa_un_avviso(self):
        """A warning nobody could ever silence stops being read: no larice
        file among the documents means nothing was lost to report."""

        self.assertEqual(self.servizio.detect({"files": [], "products": []}), [])
        self.assertEqual(self.servizio.load_errors, [])

    def test_un_listino_che_si_apre_non_segnala_niente(self):
        """The same review with the file back in place stays silent: proves
        the warning is keyed on the file, not on the supplier."""

        percorso = scrivi_listino(self.cartella / "larice.xlsx", [
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])
        review = {"files": [{"supplierId": "larice", "sourcePath": str(percorso)}], "products": []}

        promozioni = self.servizio.detect(review)

        self.assertEqual([p["kind"] for p in promozioni], [KIND_THRESHOLD_GIFT])
        self.assertEqual(self.servizio.load_errors, [])


class LePromozioniDiOggiNonCambiano(unittest.TestCase):
    """Pins the exact field-by-field contract the summary shows today.

    Reading column positions from the registry instead of a hardcoded layout
    must not change anything the user sees with today's price list and
    mapping — the condition identifier included, since the user's existing
    confirmations key off it.
    """

    def test_il_contratto_di_una_soglia_e_di_una_condizione_ambigua_e_lo_stesso(self):
        cartella = Path(tempfile.mkdtemp(prefix="promozioni-di-oggi-"))
        percorso = scrivi_listino(cartella / "larice.xlsx", [
            {"g": "ACQUISTANDO 5 CT TRA"},
            riga_merce("DENT. SENSODENT 80+20ML FRESH CLEAN", "5059833580611"),
            riga_merce("DENT. SENSODENT 80+20ML COMPLEX", "5059905473735"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "DENT. SENSODENT 15 ML", "p": "SM", "r": "5059187155220"},
            {"g": "ACQUISTANDO 10 SCATOLE TRA"},
            riga_merce("BIOPUNTO LATTE CORPO 75ML", "8059651466354"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "BIOPUNTO SET MARE", "p": "SM", "r": "8050507999999"},
        ])

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di("larice", percorso)

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 2)
        soglia, ambigua = promozioni
        self.assertEqual(soglia["id"], "promo:larice:f43d3a2ac7fd207f")
        self.assertEqual(soglia["source_reference"], "Canvass di prova!G1:J4")
        self.assertEqual(
            soglia["source_text"],
            "ACQUISTANDO 5 CT TRA IN OMAGGIO 1 CT DI DENT. SENSODENT 15 ML",
        )
        self.assertEqual(soglia["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(soglia["threshold"], {"qty": 5, "unit": "cartoni"})
        self.assertEqual(soglia["reward"], {
            "kind": "prodotto",
            "description": "DENT. SENSODENT 15 ML",
            "qty": 1,
            "unit": "cartoni",
            "pieces_per_unit": None,
            "ean": "5059187155220",
            "supplier_code": None,
        })
        self.assertEqual(soglia["eligible"], {
            "products": [], "eans": [], "source_rows": [2, 3], "mix_allowed": True, "group": None,
        })
        self.assertEqual(soglia["certainty"], "alta")
        self.assertTrue(soglia["confirmed"])
        self.assertTrue(soglia["repeatable"])
        self.assertEqual(soglia["economic_effect"], {
            "type": "informational_reward",
            "discount_rate": None,
            "deterministic": True,
            "active": True,
            "affects_total": False,
            "affects_supplier_choice": False,
            "base_price_field": "unitPricePreDiscount",
            "already_applied": False,
        })
        self.assertEqual(ambigua["id"], "promo:larice:4c32daca2df97376")
        self.assertEqual(ambigua["source_reference"], "Canvass di prova!G5:J7")
        self.assertEqual(ambigua["kind"], KIND_AMBIGUOUS)
        self.assertEqual(
            ambigua["source_text"],
            "ACQUISTANDO 10 SCATOLE TRA IN OMAGGIO 1 CT DI BIOPUNTO SET MARE "
            "[condizione non ricomposta: il testo del blocco non dice una soglia calcolabile]",
        )
        self.assertEqual(ambigua["eligible"]["source_rows"], [6])
        self.assertFalse(ambigua["confirmed"])


class IlDiPiuGiaCompresoNelPrezzo(ConUnRegistroFinto):
    """Whether a promotional "N+1 free" text is already included in the
    listed price must be a per-supplier declaration in the adapter registry,
    not a hardcoded `supplier == "betulla"` check: otherwise declaring the
    same deal for another supplier would have no effect, and their offer
    would stay flagged for manual review instead of being recognized.
    """

    # betulla's real price-list text: "11+1 Gratis" is not a discount to
    # apply, the listed price already includes it.
    OFFERTA = "LINDA SETA Assorbenti Ultra Con Ali Lunghi 11+1 Gratis Pz"

    def tipi_delle_promozioni(self, fornitore: str) -> list[str]:
        review = {
            "files": [],
            "products": [{
                "id": "prodotto-1",
                "ean": "8000000000010",
                "offers": [{
                    "supplierId": fornitore,
                    "available": True,
                    "sourceRow": 12,
                    "promotionText": self.OFFERTA,
                }],
            }],
        }
        return [p["kind"] for p in promotion_bridge.PromotionService().detect(review)]

    def test_il_fornitore_che_lo_dichiara_oggi_lo_ottiene_ancora(self):
        """With the real registry, betulla's result is unchanged."""

        self.assertEqual(self.tipi_delle_promozioni("betulla"), [KIND_INCLUDED_PACK])

    def test_senza_la_dichiarazione_la_regola_non_vale_piu(self):
        """Proves the rule actually comes from the registry: with the
        declaration removed, the "N+1" text goes back to needing manual
        review."""

        def togli(documento: dict) -> None:
            self.adattatore(documento, "betulla_v1")["commercial_rules"].pop(
                "promotion_included_in_product", None
            )

        self.registro_finto(togli)

        self.assertEqual(self.tipi_delle_promozioni("betulla"), [KIND_AMBIGUOUS])

    def test_vale_per_qualunque_fornitore_lo_dichiari(self):
        """Not the same check wearing a different name: the declaration is
        added to cipresso and only affects cipresso."""

        self.assertEqual(self.tipi_delle_promozioni("cipresso"), [KIND_AMBIGUOUS])

        def dichiara(documento: dict) -> None:
            self.adattatore(documento, "cipresso_v1")["commercial_rules"][
                "promotion_included_in_product"
            ] = True

        self.registro_finto(dichiara)

        self.assertEqual(self.tipi_delle_promozioni("cipresso"), [KIND_INCLUDED_PACK])
        # betulla, which already declared it, is unaffected.
        self.assertEqual(self.tipi_delle_promozioni("betulla"), [KIND_INCLUDED_PACK])

    def test_una_dichiarazione_che_non_e_un_si_non_vale(self):
        """The string `"true"` is not a valid declaration: a badly learned
        adapter must not be able to silently change a supplier's price
        handling."""

        def quasi(documento: dict) -> None:
            self.adattatore(documento, "betulla_v1")["commercial_rules"][
                "promotion_included_in_product"
            ] = "true"

        self.registro_finto(quasi)

        self.assertEqual(self.tipi_delle_promozioni("betulla"), [KIND_AMBIGUOUS])


class ColonnaOffertaNoce(unittest.TestCase):
    """noce's `descrizione_offerta` column is currently always empty in real
    files, so this path can only be exercised with a synthetic case: it
    guards that, whenever the supplier does fill it in, the text reaches the
    promotion detectors.
    """

    INTESTAZIONI = [
        "codice_a_barre", "codice", "descrizione_articolo", "pezzi_x_cartone",
        "cartoni_x_stra", "strati_x_pal", "prezzo", "quantita", "offerta",
        "Importo", "descrizione_reparto", "cat", "ragione_sociale", "variato",
        "descrizione_offerta", "Iva",
    ]

    def setUp(self) -> None:
        self.cartella = Path(tempfile.mkdtemp(prefix="noce-offerta-"))
        adattatori = json.loads((RADICE / "references" / "adapters.json").read_text(encoding="utf-8"))
        self.adattatore = next(a for a in adattatori["adapters"] if a["id"] == "noce_xls_v1")

    def _listino(self, testo_offerta: str) -> Path:
        libro = Workbook()
        foglio = libro.active
        foglio.title = "Foglio1"
        foglio.cell(row=4, column=4, value="i prezzi offerta sono in grassetto")
        for indice, nome in enumerate(self.INTESTAZIONI, start=2):
            foglio.cell(row=5, column=indice, value=nome)
        valori = {
            "codice_a_barre": "8000000000001",
            "codice": "0000000449070",
            "descrizione_articolo": "NEVAL DOCCIA MEN BOOST ML.250",
            "pezzi_x_cartone": 12,
            "prezzo": 0.98,
            "quantita": 0,
            "offerta": "SI",
            "cat": "NOFOOD",
            "ragione_sociale": "BEIERSDORF",
            "descrizione_offerta": testo_offerta,
            "Iva": "22",
        }
        for indice, nome in enumerate(self.INTESTAZIONI, start=2):
            if nome in valori:
                foglio.cell(row=6, column=indice, value=valori[nome])
        percorso = self.cartella / "noce.xlsx"
        libro.save(percorso)
        libro.close()
        return percorso

    def test_la_colonna_dell_offerta_arriva_fino_ai_rilevatori(self):
        import prepare_manifest_sources as pms
        from catalog_search import _promotion_text
        from promotions import detect_promotions

        testo = "ACQUISTA 15 COLLI IN OMAGGIO 1 COLLO (12 PEZZI) DI NEVAL DOCCIA MEN BOOST ML.250"
        percorso = self._listino(testo)

        record, _avvisi = pms.read_mapped_xlsx_supplier(
            percorso, "noce", dict(self.adattatore["field_mapping"])
        )

        self.assertEqual(len(record), 1)
        self.assertEqual(record[0]["availability"], testo)
        self.assertTrue(record[0]["usable"], "la colonna dell'offerta non deve rendere la riga inutilizzabile")

        promozioni = detect_promotions(
            supplier="noce",
            source_reference="noce:riga:6",
            source_text=_promotion_text(record[0]),
            eligible={"eans": [record[0]["ean"]]},
        )

        self.assertEqual([p["kind"] for p in promozioni], [KIND_THRESHOLD_GIFT])
        self.assertEqual(promozioni[0]["threshold"], {"qty": 15, "unit": "colli"})
        self.assertEqual(promozioni[0]["reward"]["unit"], "colli")

    def test_una_colonna_offerta_vuota_non_cambia_niente(self):
        import prepare_manifest_sources as pms

        percorso = self._listino("")
        record, _avvisi = pms.read_mapped_xlsx_supplier(
            percorso, "noce", dict(self.adattatore["field_mapping"])
        )

        self.assertEqual(len(record), 1)
        self.assertEqual(record[0]["availability"], "")
        self.assertTrue(record[0]["usable"])


class UnFornitoreInventatoDichiaraDoveTieneLeSueCondizioni(ConUnRegistroFinto):
    """The engine must read any supplier the registry declares, not only
    ones the code names explicitly.

    A synthetic supplier that appears nowhere else in the program should
    become readable purely by adding an adapter registry entry, with no code
    change.
    """

    FORNITORE = "bianchi"

    def adattatore_inventato(self, condizioni: dict | None) -> dict:
        """A supplier that only exists in the registry, with its own columns."""

        voce = {
            "id": "bianchi_v1",
            "schema_version": 1,
            "kind": "supplier",
            "supplier_id": self.FORNITORE,
            "display_name": "BIANCHI & FIGLI",
            "file_types": [".xlsx"],
            "column_map": {
                "description": "C",
                "reward_description": "E",
                "discount": "F",
                "ean": "H",
            },
            "row_markers": {
                "field": "discount_raw",
                "codes": {
                    "RG": {
                        "means": "Riga regalo: non è merce acquistabile.",
                        "orderable": False,
                        "row_type": "OMAGGIO",
                    }
                },
            },
        }
        if condizioni is not None:
            voce["commercial_conditions"] = condizioni
        return voce

    def registro_con_bianchi(self, condizioni: dict | None) -> None:
        def aggiungi(documento: dict) -> None:
            documento["adapters"].append(self.adattatore_inventato(condizioni))

        self.registro_finto(aggiungi)

    def scrivi(self, foglio_chiamato: str, righe: list[dict[str, object]]) -> Path:
        """A price list shaped by bianchi's columns, not larice's."""

        libro = Workbook()
        foglio = libro.active
        foglio.title = foglio_chiamato
        for numero, contenuto in enumerate(righe, start=1):
            for colonna, valore in contenuto.items():
                if valore not in (None, ""):
                    foglio.cell(row=numero, column=colonna, value=valore)
        percorso = self.cartella / "bianchi.xlsx"
        libro.save(percorso)
        libro.close()
        return percorso

    # -- block layout (larice's), applied to a different supplier ----------

    A_BLOCCHI = {
        "layout": "blocchi",
        "sheet": "Condizioni",
        "data_start_row": 3,
        "fields": {
            "text": "description",
            "reward": "reward_description",
            "ean": "ean",
            "row_code": "discount",
        },
    }

    LISTINO_A_BLOCCHI = [
        {3: "LISTINO SETTIMANA 34"},
        {3: "codice a barre", 5: "omaggio"},
        {3: "ACQUISTANDO 4 CT TRA"},
        {3: "SAPONE MANI 300 ML", 6: "TP", 8: "8011111111111"},
        {3: "SAPONE MANI 500 ML", 6: "TP", 8: "8011111111112"},
        {3: "IN OMAGGIO 1 CT DI", 5: "TOSTAPANE 750W", 6: "RG", 8: "8011111119999"},
    ]

    def test_una_soglia_di_un_fornitore_che_il_codice_non_nomina(self):
        self.registro_con_bianchi(self.A_BLOCCHI)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        soglia = promozioni[0]
        self.assertEqual(soglia["supplier"], "bianchi")
        self.assertEqual(soglia["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(soglia["threshold"], {"qty": 4, "unit": "cartoni"})
        self.assertEqual(soglia["reward"]["description"], "TOSTAPANE 750W")
        self.assertEqual(soglia["reward"]["ean"], "8011111119999")
        self.assertTrue(soglia["confirmed"])
        # The two goods rows, not the reward row: this supplier's row-marker
        # code, its own ("RG" rather than larice's "SM").
        self.assertEqual(soglia["eligible"]["source_rows"], [4, 5])
        # Reference columns and sheet name are this supplier's own, not larice's.
        self.assertEqual(soglia["source_reference"], "Condizioni!C3:E6")

    # -- reward written inline in the text, no dedicated column ------------

    A_BLOCCHI_IN_UNA_COLONNA_SOLA = {
        "layout": "blocchi",
        "sheet": "Condizioni",
        "data_start_row": 3,
        "fields": {
            "text": "description",
            "reward": "description",
            "ean": "ean",
            "row_code": "discount",
        },
    }

    LISTINO_COL_PREMIO_NEL_TESTO = [
        {3: "LISTINO SETTIMANA 34"},
        {3: "codice a barre"},
        {3: "ACQUISTANDO 4 CT TRA"},
        {3: "SAPONE MANI 300 ML", 6: "TP", 8: "8011111111111"},
        {3: "SAPONE MANI 500 ML", 6: "TP", 8: "8011111111112"},
        {3: "IN OMAGGIO 1 CT DI TOSTAPANE 750W", 6: "RG", 8: "8011111119999"},
    ]

    def test_il_premio_nella_stessa_colonna_del_testo_non_esce_scritto_due_volte(self):
        """The registry must support text and reward name in the same column
        (larice's newer canvass layout keeps both in column E).

        Reading that column for both roles must not duplicate the reward
        name in the rendered sentence (e.g. "TOSTAPANE 750W IN OMAGGIO 1 CT
        DI TOSTAPANE 750W"): the threshold itself would be correct, but the
        text shown to the user would not.
        """

        self.registro_con_bianchi(self.A_BLOCCHI_IN_UNA_COLONNA_SOLA)
        percorso = self.scrivi("Condizioni", self.LISTINO_COL_PREMIO_NEL_TESTO)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        soglia = promozioni[0]
        self.assertEqual(soglia["threshold"], {"qty": 4, "unit": "cartoni"})
        self.assertEqual(soglia["reward"]["description"], "TOSTAPANE 750W")
        self.assertEqual(soglia["reward"]["ean"], "8011111119999")
        self.assertEqual(soglia["eligible"]["source_rows"], [4, 5])

    def test_due_colonne_diverse_continuano_a_leggersi_tutt_e_due(self):
        """Converse check: a dedicated reward column still works as before."""

        self.registro_con_bianchi(self.A_BLOCCHI)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)

        promozioni, _errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertEqual(promozioni[0]["reward"]["description"], "TOSTAPANE 750W")

    def test_la_stessa_soglia_arriva_fino_al_riepilogo(self):
        """Checks the whole chain, not just the reader in isolation:
        `detect` must pick up any supplier present in the run's files, not
        only ones it names internally.
        """

        self.registro_con_bianchi(self.A_BLOCCHI)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)
        servizio = promotion_bridge.PromotionService()

        promozioni = servizio.detect({
            "files": [{"supplierId": self.FORNITORE, "sourcePath": str(percorso)}],
            "products": [],
        })

        self.assertEqual([p["supplier"] for p in promozioni], ["bianchi"])
        self.assertEqual(servizio.load_errors, [])
        # Counted per supplier: larice isn't among this run's files, so it
        # gets no count at all rather than a zero, which would read as
        # "read, nothing to verify".
        self.assertEqual(servizio.condizioni_da_verificare, {"bianchi": 0})

    # -- row layout: one condition written entirely in a single column -----

    A_RIGA = {
        "layout": "riga",
        "sheet": "Condizioni",
        "data_start_row": 2,
        "fields": {"text": "description", "ean": "ean", "row_code": "discount"},
    }

    def test_una_condizione_scritta_per_intero_dentro_una_riga(self):
        """The layout betulla and noce would use if they started declaring
        promotions: the whole condition sits in one cell, not spread across
        rows."""

        self.registro_con_bianchi(self.A_RIGA)
        percorso = self.scrivi("Condizioni", [
            # The header row sits above `data_start_row` and deliberately
            # carries a promotional word, so that if the data start row were
            # ignored this non-condition would get picked up too.
            {3: "TABELLA PROMOZIONE SETTIMANALE", 8: "codice a barre"},
            {3: "SAPONE MANI 300 ML", 6: "TP", 8: "8011111111111"},
            {
                3: "ACQUISTA 5 CT IN OMAGGIO 1 CT (12 PEZZI) DI SALE LAVASTOVIGLIE KG1",
                6: "TP",
                8: "8011111111112",
            },
        ])

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        soglia = promozioni[0]
        self.assertEqual(soglia["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(soglia["threshold"], {"qty": 5, "unit": "cartoni"})
        self.assertEqual(soglia["reward"]["description"], "SALE LAVASTOVIGLIE KG1")
        self.assertEqual(soglia["reward"]["pieces_per_unit"], 12)
        # The eligible product is the one on the same row, identified both by
        # row and by EAN.
        self.assertEqual(soglia["eligible"]["source_rows"], [3])
        self.assertEqual(soglia["eligible"]["eans"], ["8011111111112"])
        self.assertEqual(soglia["source_reference"], "Condizioni!C3")

    def test_a_riga_una_descrizione_qualunque_non_diventa_un_offerta(self):
        """The risk with row layout: the text column is often just the
        product description, and an ordinary product name must not be
        mistaken for a promotion."""

        self.registro_con_bianchi(self.A_RIGA)
        percorso = self.scrivi("Condizioni", [
            {3: "descrizione"},
            {3: "SAPONE MANI 300 ML", 6: "TP", 8: "8011111111111"},
            {3: "DENTIFRICIO MENTA 75 ML", 6: "TP", 8: "8011111111112"},
        ])

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(promozioni, [])

    def test_a_riga_la_riga_non_acquistabile_resta_fuori(self):
        """The row-marker code applies in both layouts: a row the supplier
        marks non-orderable never becomes a condition."""

        self.registro_con_bianchi(self.A_RIGA)
        percorso = self.scrivi("Condizioni", [
            {3: "descrizione"},
            {3: "ACQUISTA 5 CT IN OMAGGIO 1 CT DI TOSTAPANE", 6: "RG", 8: "8011111119999"},
        ])

        promozioni, _errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertEqual(promozioni, [])

    # -- behavior when the registry declares nothing -----------------------

    def test_senza_la_dichiarazione_non_si_legge_e_non_si_avvisa_nessuno(self):
        """Proves the rule really comes from the registry, and that silence
        is the correct outcome here: a supplier with no declared conditions
        has none, and warning about it on every recompute would just be
        noise nobody can act on.
        """

        self.registro_con_bianchi(None)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)
        servizio = promotion_bridge.PromotionService()

        promozioni = servizio.detect({
            "files": [{"supplierId": self.FORNITORE, "sourcePath": str(percorso)}],
            "products": [],
        })

        self.assertEqual(promozioni, [])
        self.assertEqual(servizio.load_errors, [])

    def test_una_forma_di_scrittura_che_il_motore_non_conosce_si_ferma_e_lo_dice(self):
        """An unknown layout value in a badly learned registry entry must
        fail with an error, not fall back to reading "something anyway"."""

        self.registro_con_bianchi({**self.A_BLOCCHI, "layout": "a fisarmonica"})
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertEqual(promozioni, [])
        self.assertEqual(
            errore["message"],
            "Non riesco a leggere le condizioni commerciali di BIANCHI & FIGLI: non so "
            "più dove il listino tiene in che modo scrive le sue condizioni.",
        )

    def test_il_nome_nei_messaggi_e_quello_dichiarato_dal_registro(self):
        """Warning messages must use the display name declared in the
        registry, not the internal supplier id or another supplier's name."""

        senza_ean = {
            **self.A_BLOCCHI,
            "fields": {k: v for k, v in self.A_BLOCCHI["fields"].items() if k != "ean"},
        }
        self.registro_con_bianchi(senza_ean)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)
        servizio = promotion_bridge.PromotionService()

        servizio.detect({
            "files": [{"supplierId": self.FORNITORE, "sourcePath": str(percorso)}],
            "products": [],
        })

        self.assertEqual(len(servizio.load_errors), 1)
        self.assertEqual(servizio.load_errors[0]["supplier"], "bianchi")
        self.assertEqual(servizio.load_errors[0]["supplierName"], "BIANCHI & FIGLI")
        self.assertEqual(
            servizio.load_errors[0]["message"],
            "Non riesco a leggere le condizioni commerciali di BIANCHI & FIGLI: non so "
            "più dove il listino tiene il codice a barre.",
        )

    def test_il_listino_sparito_di_un_fornitore_inventato_e_un_avviso_suo(self):
        self.registro_con_bianchi(self.A_BLOCCHI)
        servizio = promotion_bridge.PromotionService()

        servizio.detect({
            "files": [{"supplierId": self.FORNITORE, "sourcePath": str(self.cartella / "via.xlsx")}],
            "products": [],
        })

        self.assertEqual(len(servizio.load_errors), 1)
        self.assertEqual(
            servizio.load_errors[0]["message"],
            "Non riesco a leggere le condizioni commerciali di BIANCHI & FIGLI: il "
            "listino non è fra i documenti caricati.",
        )

    def test_due_fornitori_che_dichiarano_si_leggono_sempre_nello_stesso_ordine(self):
        """Detected promotions must be returned in a stable order across
        multiple suppliers: the page renders this list directly, so two
        reads of the same run producing a different order would make an
        unchanged comparison look like it had changed.
        """

        self.registro_con_bianchi(self.A_BLOCCHI)
        bianchi = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)
        larice = scrivi_listino(self.cartella / "larice.xlsx", [
            {"g": "ACQUISTANDO 2 CT TRA"},
            riga_merce("PRODOTTO UNO", "8000000000001"),
            {"g": "IN OMAGGIO 1 CT DI", "j": "PREMIO", "p": "SM", "r": "8000000009999"},
        ])
        review = {
            "files": [
                {"supplierId": "larice", "sourcePath": str(larice)},
                {"supplierId": self.FORNITORE, "sourcePath": str(bianchi)},
            ],
            "products": [],
        }

        # bianchi comes after larice in the registry; it comes first here
        # because the output order is alphabetical, not declaration order.
        letti = [p["supplier"] for p in promotion_bridge.PromotionService().detect(review)]

        self.assertEqual(letti, ["bianchi", "larice"])

    def test_di_due_adattatori_si_usa_quello_del_documento_che_si_ha_in_mano(self):
        """When a supplier has more than one adapter (e.g. noce's separate
        `.xls` and CSV entries), the reader must pick the one matching the
        file actually being read, not just the first match: reading with the
        wrong format's declaration means looking for a column that isn't
        there.
        """

        def due_adattatori(documento: dict) -> None:
            # The first entry has a different format with different columns:
            # if the format check were missing, this one would be picked.
            csv = self.adattatore_inventato(dict(self.A_RIGA))
            csv["id"] = "bianchi_csv_v1"
            csv["file_types"] = [".csv"]
            csv["column_map"] = {"description": "A", "ean": "B", "discount": "C"}
            documento["adapters"].append(csv)
            documento["adapters"].append(self.adattatore_inventato(self.A_BLOCCHI))

        self.registro_finto(due_adattatori)
        percorso = self.scrivi("Condizioni", self.LISTINO_A_BLOCCHI)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["kind"], KIND_THRESHOLD_GIFT)
        self.assertEqual(promozioni[0]["source_reference"], "Condizioni!C3:E6")

    def test_il_foglio_dichiarato_e_quello_che_si_legge(self):
        """The reader must use the sheet name declared in the registry, not
        always the workbook's first sheet."""

        self.registro_con_bianchi(self.A_BLOCCHI)
        libro = Workbook()
        libro.active.title = "Copertina"
        foglio = libro.create_sheet("Condizioni")
        for numero, contenuto in enumerate(self.LISTINO_A_BLOCCHI, start=1):
            for colonna, valore in contenuto.items():
                if valore not in (None, ""):
                    foglio.cell(row=numero, column=colonna, value=valore)
        percorso = self.cartella / "bianchi.xlsx"
        libro.save(percorso)
        libro.close()

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            self.FORNITORE, percorso
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 1)
        self.assertEqual(promozioni[0]["source_reference"], "Condizioni!C3:E6")


class SoloChiLoDichiaraVieneLetto(unittest.TestCase):
    """Today, only larice declares commercial conditions in the registry.

    Declaring a plain description column as the condition text for another
    supplier would surface plausible-looking product names as promotions
    that aren't: inventing offers for a supplier that has none is worse than
    missing real ones.
    """

    def test_nel_registro_di_oggi_lo_dichiara_solo_larice(self):
        self.assertEqual(
            sorted(promotion_bridge.fornitori_con_condizioni_dichiarate()), ["larice"]
        )

    def test_gli_altri_fornitori_del_confronto_non_producono_ne_offerte_ne_avvisi(self):
        cartella = Path(tempfile.mkdtemp(prefix="senza-condizioni-"))
        listini = []
        for fornitore in ("betulla", "cipresso", "noce"):
            percorso = cartella / f"{fornitore}.xlsx"
            libro = Workbook()
            libro.active.cell(row=1, column=1, value="EAN")
            libro.save(percorso)
            libro.close()
            listini.append({"supplierId": fornitore, "sourcePath": str(percorso)})
        servizio = promotion_bridge.PromotionService()

        promozioni = servizio.detect({"files": listini, "products": []})

        self.assertEqual(promozioni, [])
        self.assertEqual(servizio.load_errors, [])


class IlListinoLariceVero(unittest.TestCase):
    """A regression check synthetic price lists can't cover: reading the
    real larice file must still yield exactly 13 threshold-gift conditions,
    all complete (threshold with unit, reward with name and EAN, resolved
    eligible rows, repeatable, confirmed) and none left "to verify".
    """

    LISTINO = RADICE / "app" / "data" / "current" / "uploads" / "33-34.1 07-21 ago.xlsx"

    def test_le_tredici_condizioni_restano_tredici_e_restano_complete(self):
        if not self.LISTINO.is_file():
            self.skipTest(f"il listino LARICE vero non c'è: {self.LISTINO.name}")

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            "larice", self.LISTINO
        )

        self.assertIsNone(errore)
        self.assertEqual([p["kind"] for p in promozioni], [KIND_THRESHOLD_GIFT] * 13)
        for promozione in promozioni:
            with self.subTest(riferimento=promozione["source_reference"]):
                self.assertEqual(promozione["supplier"], "larice")
                self.assertEqual(promozione["certainty"], "alta")
                self.assertTrue(promozione["confirmed"])
                self.assertTrue(promozione["repeatable"])
                self.assertEqual(promozione["threshold"]["unit"], "cartoni")
                self.assertGreater(promozione["threshold"]["qty"], 0)
                self.assertTrue(promozione["reward"]["description"])
                self.assertTrue(promozione["reward"]["ean"])
                self.assertTrue(promozione["eligible"]["source_rows"])
        # The reward names the operator picked out when deciding what to keep.
        premi = {p["reward"]["description"] for p in promozioni}
        self.assertIn("RESALINA SALE LAVASTOVIGLIE KG1", premi)
        self.assertIn("BISTECCHIERA 1000W", premi)
        self.assertIn("TOSTIERA ELETTRICA 750W", premi)


class UnDocumentoExcel97(ConUnRegistroFinto):
    """noce sends only legacy Excel 97-2003 (`.xls`) files, and the reader
    must actually open that format: accepting a declaration but never
    reading the document would be the same hardcoded-supplier problem in a
    less visible form. The real price list also acts as a measurement: its
    offer-text column is empty on every row, so zero is the correct count.
    """

    LISTINO = RADICE / "app" / "data" / "current" / "uploads" / "formattato_104233.xls"

    def test_il_listino_xls_si_legge_e_non_ha_nessuna_condizione(self):
        if not self.LISTINO.is_file():
            self.skipTest(f"il listino NOCE vero non c'è: {self.LISTINO.name}")

        def dichiara(documento: dict) -> None:
            self.adattatore(documento, "noce_xls_v1")["commercial_conditions"] = {
                "layout": "riga",
                "sheet": "Foglio1",
                "data_start_row": 6,
                "fields": {"text": "descrizione_offerta", "ean": "codice_a_barre"},
            }

        self.registro_finto(dichiara)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            "noce", self.LISTINO
        )

        self.assertIsNone(errore, "il .xls dev'essere leggibile, non un errore")
        self.assertEqual(promozioni, [])

    def test_la_colonna_sbagliata_del_listino_xls_si_legge_lo_stesso(self):
        """Proves the `.xls` file is actually read, not just opened: its
        description column has five "N+M GRATIS" texts, and the detector
        must find exactly those when pointed at that column."""

        if not self.LISTINO.is_file():
            self.skipTest(f"il listino NOCE vero non c'è: {self.LISTINO.name}")

        def dichiara(documento: dict) -> None:
            self.adattatore(documento, "noce_xls_v1")["commercial_conditions"] = {
                "layout": "riga",
                "sheet": "Foglio1",
                "data_start_row": 6,
                "fields": {"text": "descrizione_articolo", "ean": "codice_a_barre"},
            }

        self.registro_finto(dichiara)

        promozioni, errore = promotion_bridge.PromotionService().condizioni_di(
            "noce", self.LISTINO
        )

        self.assertIsNone(errore)
        testi = sorted(p["source_text"] for p in promozioni)
        self.assertEqual(len(testi), 5, testi)
        self.assertIn("DANZATRIX PANNO MULTIUSO PZ.3+1 GRATIS", testi)
        self.assertIn("CUKO ALLUMINIO MT.16+4 GRATIS", testi)


class QualeDeiDueAdattatoriDelloStessoFornitore(unittest.TestCase):
    """A supplier with two adapters for the same file format.

    File extension alone is enough to disambiguate noce's `.xls` and `.csv`
    adapters, but not larice's two `.xlsx` adapters: picking the first match
    would silently read the newer canvass layout with the older adapter's
    columns, producing zero conditions with no error.
    """

    VOCI = [
        {"id": "larice_v1", "file_types": [".xlsx"]},
        {"id": "larice_canvass_v1", "file_types": [".xlsx"]},
    ]
    DOCUMENTO = Path("New Larice N°37.xlsx")

    def test_decide_l_identificativo_che_ha_riconosciuto_il_documento(self) -> None:
        scelto = promotion_bridge._adattatore_per_il_documento(
            self.VOCI, self.DOCUMENTO, "larice_canvass_v1"
        )
        self.assertEqual(scelto["id"], "larice_canvass_v1")

    def test_lo_schema_imparato_porta_a_quello_da_cui_deriva(self) -> None:
        """An id-based decision must resolve a learned schema back to its
        base adapter via `registro.adattatore_base`."""

        scelto = promotion_bridge._adattatore_per_il_documento(
            self.VOCI, self.DOCUMENTO, "larice_canvass_v1__locale"
        )
        self.assertEqual(scelto["id"], "larice_canvass_v1")

    def test_senza_identificativo_resta_la_regola_del_formato(self) -> None:
        """The older extension-based rule must still apply, since noce's two
        adapters still rely on it."""

        voci = [
            {"id": "noce_csv_v1", "file_types": [".csv"]},
            {"id": "noce_xls_v1", "file_types": [".xls", ".xlsx"]},
        ]
        scelto = promotion_bridge._adattatore_per_il_documento(
            voci, Path("formattato_104233.xls"), None
        )
        self.assertEqual(scelto["id"], "noce_xls_v1")

    def test_un_identificativo_che_non_e_fra_questi_non_si_indovina(self) -> None:
        """An unrecognized adapter id must resolve to nothing, rather than
        falling back to another document's conditions."""

        scelto = promotion_bridge._adattatore_per_il_documento(
            self.VOCI, self.DOCUMENTO, "qualcosa_di_imparato_v1"
        )
        self.assertEqual(scelto, {})

    def test_con_un_candidato_solo_l_identificativo_non_toglie_niente(self) -> None:
        """A supplier with only one schema must keep reading as before."""

        scelto = promotion_bridge._adattatore_per_il_documento(
            [{"id": "betulla_v1", "file_types": [".xlsx"]}], Path("betulla.xlsx"), "betulla_v1__imparato"
        )
        self.assertEqual(scelto["id"], "betulla_v1")


if __name__ == "__main__":
    unittest.main()
