"""Tests for the stdlib-only `.xls` reader.

Two kinds of test, both needed.

1. Hand-built fixtures. A small `.xls` writer below packs a real OLE2
   container with real BIFF8 records, exercising cases a single real price
   list may not cover: bold formatting, the four ways to pack a compressed
   number, formulas that resolve to text or an error, and a shared-string
   table split mid-word. These run anywhere, need no external files, and are
   fast.

2. Comparison against a real price list. The real `.xls` is read twice: once
   with this reader, once via openpyxl from an Excel-converted `.xlsx` copy.
   Every cell of every sheet is compared, value and bold flag. This is the
   test that matters: a subtly wrong BIFF reader doesn't fail outright, it
   corrupts data far enough into the file that no one notices. The converted
   copy lives outside the repo, so the test skips with a message pointing to
   it when missing; openpyxl is only a comparison oracle here, not a
   dependency of the reader itself.
"""

from __future__ import annotations

import importlib.util
import os
import shutil
import stat
import struct
import sys
import tempfile
import unittest
from pathlib import Path

RADICE = Path(__file__).resolve().parents[1]
PERCORSO_LETTORE = RADICE / "app" / "xls_reader.py"

_SPEC = importlib.util.spec_from_file_location("xls_reader_in_collaudo", PERCORSO_LETTORE)
if _SPEC is None or _SPEC.loader is None:  # pragma: no cover - only if the file goes missing
    raise RuntimeError(f"Impossibile importare {PERCORSO_LETTORE}")
LETTORE = importlib.util.module_from_spec(_SPEC)
# Register the module before executing it: its dataclasses look themselves
# up in sys.modules while being constructed.
sys.modules[_SPEC.name] = LETTORE
_SPEC.loader.exec_module(LETTORE)

XlsError = LETTORE.XlsError
read_workbook = LETTORE.read_workbook
read_sheet_values = LETTORE.read_sheet_values


# The two real files. They live outside the repo; their location can be
# overridden with these two environment variables without touching the test.
LISTINO_XLS = Path(
    os.environ.get(
        "LISTINO_XLS_DI_PROVA",
        str(RADICE / "listini-storici" / "formattato_104233.xls"),
    )
)
LISTINO_XLSX = Path(
    os.environ.get(
        "LISTINO_XLSX_DI_PARAGONE",
        r"C:\Users\HP\AppData\Local\Temp\claude\C--Users-HP-Desktop-ChatGPT"
        r"\41e85a3c-f814-42e2-94b2-a314e1cd0bcf\scratchpad\noce\formattato_104233.xlsx",
    )
)

# Numeric comparison tolerance: the two readers reach the same number by
# different paths (packed RK record vs. decimal text), so allow drift in the
# last digit of a double, not a cent of currency.
TOLLERANZA_ASSOLUTA = 1e-9
TOLLERANZA_RELATIVA = 1e-9


# ---------------------------------------------------------------------------
# Minimal .xls writer for the fixtures below
# ---------------------------------------------------------------------------

DIMENSIONE_SETTORE = 512
DIMENSIONE_MINI = 64
SOGLIA_MINI = 4096
LIBERO = 0xFFFFFFFF
FINE_CATENA = 0xFFFFFFFE
SETTORE_FAT = 0xFFFFFFFD
FIRMA_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _voce_direttorio(
    nome: str,
    tipo: int,
    inizio: int,
    dimensione: int,
    *,
    figlio: int = LIBERO,
    destro: int = LIBERO,
) -> bytes:
    voce = bytearray(128)
    grezzo = nome.encode("utf-16-le") + b"\x00\x00"
    voce[0 : len(grezzo)] = grezzo
    voce[64:66] = struct.pack("<H", len(grezzo))
    voce[66] = tipo
    voce[67] = 1
    voce[68:72] = struct.pack("<I", LIBERO)
    voce[72:76] = struct.pack("<I", destro)
    voce[76:80] = struct.pack("<I", figlio)
    voce[116:120] = struct.pack("<I", inizio)
    voce[120:128] = struct.pack("<Q", dimensione)
    return bytes(voce)


def costruisci_ole2(flussi: list[tuple[str, bytes]]) -> bytes:
    """Pack streams into a minimal but real OLE2 container.

    Streams under 4096 bytes go into the mini-stream, matching Excel's own
    behavior: a small workbook exercises the mini-FAT, a large one the normal
    sector chain.
    """
    mini_dati = bytearray()
    mini_fat: list[int] = []
    inizio_mini: dict[str, int] = {}
    for nome, contenuto in flussi:
        if not 0 < len(contenuto) < SOGLIA_MINI:
            continue
        inizio_mini[nome] = len(mini_fat)
        blocchi = -(-len(contenuto) // DIMENSIONE_MINI)
        for _ in range(blocchi):
            mini_fat.append(len(mini_fat) + 1)
        mini_fat[-1] = FINE_CATENA
        mini_dati += contenuto.ljust(blocchi * DIMENSIONE_MINI, b"\x00")

    settori: list[bytes] = []
    fat: list[int] = []

    def aggiungi(contenuto: bytes) -> tuple[int, int]:
        if not contenuto:
            return FINE_CATENA, 0
        inizio = len(settori)
        blocchi = -(-len(contenuto) // DIMENSIONE_SETTORE)
        for indice in range(blocchi):
            pezzo = contenuto[indice * DIMENSIONE_SETTORE : (indice + 1) * DIMENSIONE_SETTORE]
            settori.append(pezzo.ljust(DIMENSIONE_SETTORE, b"\x00"))
            fat.append(len(fat) + 1)
        fat[-1] = FINE_CATENA
        return inizio, blocchi

    inizio_mini_stream, _ = aggiungi(bytes(mini_dati))
    inizio_grandi: dict[str, int] = {}
    for nome, contenuto in flussi:
        if len(contenuto) >= SOGLIA_MINI:
            inizio_grandi[nome], _ = aggiungi(contenuto)

    voci = [
        _voce_direttorio(
            "Root Entry",
            5,
            inizio_mini_stream,
            len(mini_dati),
            figlio=1 if flussi else LIBERO,
        )
    ]
    for indice, (nome, contenuto) in enumerate(flussi):
        if not contenuto:
            inizio = FINE_CATENA
        elif len(contenuto) < SOGLIA_MINI:
            inizio = inizio_mini[nome]
        else:
            inizio = inizio_grandi[nome]
        destro = indice + 2 if indice + 1 < len(flussi) else LIBERO
        voci.append(_voce_direttorio(nome, 2, inizio, len(contenuto), destro=destro))
    inizio_direttorio, _ = aggiungi(b"".join(voci))

    if mini_fat:
        grezzo_mini = b"".join(struct.pack("<I", valore) for valore in mini_fat)
        avanzo = (-len(grezzo_mini)) % DIMENSIONE_SETTORE
        grezzo_mini += struct.pack("<I", LIBERO) * (avanzo // 4)
        inizio_mini_fat, numero_mini_fat = aggiungi(grezzo_mini)
    else:
        inizio_mini_fat, numero_mini_fat = FINE_CATENA, 0

    per_settore = DIMENSIONE_SETTORE // 4
    numero_fat = 1
    while True:
        necessari = -(-(len(settori) + numero_fat) // per_settore)
        if necessari <= numero_fat:
            break
        numero_fat = necessari

    settori_fat: list[int] = []
    for _ in range(numero_fat):
        settori_fat.append(len(settori))
        settori.append(b"\x00" * DIMENSIONE_SETTORE)
        fat.append(SETTORE_FAT)
    voci_fat = fat + [LIBERO] * (numero_fat * per_settore - len(fat))
    grezzo_fat = b"".join(struct.pack("<I", valore) for valore in voci_fat)
    for indice, numero in enumerate(settori_fat):
        settori[numero] = grezzo_fat[indice * DIMENSIONE_SETTORE : (indice + 1) * DIMENSIONE_SETTORE]

    intestazione = bytearray(512)
    intestazione[0:8] = FIRMA_OLE2
    intestazione[24:26] = struct.pack("<H", 0x003E)
    intestazione[26:28] = struct.pack("<H", 3)
    intestazione[28:30] = struct.pack("<H", 0xFFFE)
    intestazione[30:32] = struct.pack("<H", 9)
    intestazione[32:34] = struct.pack("<H", 6)
    intestazione[44:48] = struct.pack("<I", numero_fat)
    intestazione[48:52] = struct.pack("<I", inizio_direttorio)
    intestazione[56:60] = struct.pack("<I", SOGLIA_MINI)
    intestazione[60:64] = struct.pack("<I", inizio_mini_fat)
    intestazione[64:68] = struct.pack("<I", numero_mini_fat)
    intestazione[68:72] = struct.pack("<I", FINE_CATENA)
    intestazione[72:76] = struct.pack("<I", 0)
    difat = settori_fat + [LIBERO] * (109 - len(settori_fat))
    intestazione[76:512] = b"".join(struct.pack("<I", valore) for valore in difat)
    return bytes(intestazione) + b"".join(settori)


# -- BIFF8 records -----------------------------------------------------------


def rec(codice: int, dati: bytes) -> bytes:
    return struct.pack("<HH", codice, len(dati)) + dati


def bof(tipo: int) -> bytes:
    return rec(0x0809, struct.pack("<HHHHII", 0x0600, tipo, 0x0DBB, 0x07CC, 0, 0))


def eof() -> bytes:
    return rec(0x000A, b"")


def font(peso: int, nome: str = "Calibri") -> bytes:
    # Weight is at offset 6: 400 is normal, 700 is bold.
    dati = struct.pack("<HHHHHBBBB", 220, 0, 0x7FFF, peso, 0, 0, 0, 0, 0)
    dati += bytes([len(nome), 0x01]) + nome.encode("utf-16-le")
    return rec(0x0031, dati)


def xf(indice_font: int) -> bytes:
    return rec(0x00E0, struct.pack("<HH", indice_font, 0) + b"\x00" * 16)


def boundsheet(nome: str, posizione: int, tipo: int = 0, compresso: bool = False) -> bytes:
    if compresso:
        coda = bytes([len(nome), 0x00]) + nome.encode("latin-1")
    else:
        coda = bytes([len(nome), 0x01]) + nome.encode("utf-16-le")
    return rec(0x0085, struct.pack("<IBB", posizione, 0, tipo) + coda)


def labelsst(riga: int, colonna: int, indice: int, indice_xf: int = 0) -> bytes:
    return rec(0x00FD, struct.pack("<HHHI", riga, colonna, indice_xf, indice))


def label(riga: int, colonna: int, testo: str, indice_xf: int = 0) -> bytes:
    coda = struct.pack("<HB", len(testo), 0x01) + testo.encode("utf-16-le")
    return rec(0x0204, struct.pack("<HHH", riga, colonna, indice_xf) + coda)


def numero(riga: int, colonna: int, valore: float, indice_xf: int = 0) -> bytes:
    return rec(0x0203, struct.pack("<HHHd", riga, colonna, indice_xf, valore))


def rk_intero(valore: int, centesimi: bool = False) -> int:
    grezzo = (valore << 2) & 0xFFFFFFFC
    return grezzo | 0x02 | (0x01 if centesimi else 0)


def rk_reale(valore: float, centesimi: bool = False) -> int:
    (bit,) = struct.unpack("<Q", struct.pack("<d", valore))
    if bit & 0x3FFFFFFFF:
        raise AssertionError(f"il numero {valore} non entra in un RK a virgola mobile")
    return ((bit >> 32) & 0xFFFFFFFC) | (0x01 if centesimi else 0)


def rk(riga: int, colonna: int, grezzo: int, indice_xf: int = 0) -> bytes:
    return rec(0x027E, struct.pack("<HHHI", riga, colonna, indice_xf, grezzo))


def mulrk(riga: int, prima_colonna: int, coppie: list[tuple[int, int]]) -> bytes:
    dati = struct.pack("<HH", riga, prima_colonna)
    for indice_xf, grezzo in coppie:
        dati += struct.pack("<HI", indice_xf, grezzo)
    dati += struct.pack("<H", prima_colonna + len(coppie) - 1)
    return rec(0x00BD, dati)


def blank(riga: int, colonna: int, indice_xf: int = 0) -> bytes:
    return rec(0x0201, struct.pack("<HHH", riga, colonna, indice_xf))


def mulblank(riga: int, prima_colonna: int, indici_xf: list[int]) -> bytes:
    dati = struct.pack("<HH", riga, prima_colonna)
    for indice_xf in indici_xf:
        dati += struct.pack("<H", indice_xf)
    dati += struct.pack("<H", prima_colonna + len(indici_xf) - 1)
    return rec(0x00BE, dati)


def boolerr(riga: int, colonna: int, valore: int, e_errore: bool, indice_xf: int = 0) -> bytes:
    return rec(0x0205, struct.pack("<HHHBB", riga, colonna, indice_xf, valore, 1 if e_errore else 0))


def formula(riga: int, colonna: int, memorizzato: bytes, indice_xf: int = 0) -> bytes:
    assert len(memorizzato) == 8
    rgce = b"\x1e\x01\x00"  # arbitrary constant: the reader ignores the formula bytes
    corpo = struct.pack("<HHH", riga, colonna, indice_xf) + memorizzato
    corpo += struct.pack("<HIH", 0, 0, len(rgce)) + rgce
    return rec(0x0006, corpo)


def formula_numero(valore: float) -> bytes:
    return struct.pack("<d", valore)


def formula_testo() -> bytes:
    return bytes([0, 0, 0, 0, 0, 0, 0xFF, 0xFF])


def formula_booleano(vero: bool) -> bytes:
    return bytes([1, 0, 1 if vero else 0, 0, 0, 0, 0xFF, 0xFF])


def formula_errore(codice: int) -> bytes:
    return bytes([2, 0, codice, 0, 0, 0, 0xFF, 0xFF])


def formula_stringa_vuota() -> bytes:
    return bytes([3, 0, 0, 0, 0, 0, 0xFF, 0xFF])


def record_string(testo: str) -> bytes:
    return rec(0x0207, struct.pack("<HB", len(testo), 0x01) + testo.encode("utf-16-le"))


def costruisci_xls(
    fogli: list[tuple[str, bytes]],
    *,
    pesi_font: tuple[int, ...] = (400,),
    font_di_xf: tuple[int, ...] = (0,),
    blocchi_sst: list[bytes] | None = None,
    riempimento: int = 0,
    tipi_foglio: tuple[int, ...] | None = None,
    nomi_compressi: bool = False,
    protetto: bool = False,
    extra_globali: bytes = b"",
) -> bytes:
    """Build a complete `.xls`: container, globals, and sheet substreams.

    `extra_globali` lets `test_xls_writer` inject a `RECALCID` record into the
    globals, the record compilation zeroes to force Excel to recompute on open.
    """
    tipi = tipi_foglio if tipi_foglio is not None else tuple(0 for _ in fogli)

    testa = bof(0x0005) + extra_globali
    if protetto:
        testa += rec(0x002F, struct.pack("<H", 1))
    for peso in pesi_font:
        testa += font(peso)
    for indice in font_di_xf:
        testa += xf(indice)

    coda = b""
    if blocchi_sst:
        coda += rec(0x00FC, blocchi_sst[0])
        for blocco in blocchi_sst[1:]:
            coda += rec(0x003C, blocco)
    coda += eof()

    # BOUNDSHEET records have a fixed layout: compute the globals' total
    # length first, then each sheet's start offset follows from it.
    segnaposto = b"".join(
        boundsheet(nome, 0, tipo, nomi_compressi) for (nome, _), tipo in zip(fogli, tipi)
    )
    lunghezza_globali = len(testa) + len(segnaposto) + len(coda)

    posizioni: list[int] = []
    sottoflussi: list[bytes] = []
    corrente = lunghezza_globali
    for (_nome, celle), tipo in zip(fogli, tipi):
        sottoflusso = bof(0x0010 if tipo == 0 else 0x0020) + celle + eof()
        posizioni.append(corrente)
        corrente += len(sottoflusso)
        sottoflussi.append(sottoflusso)

    globali = testa
    for (nome, _celle), tipo, posizione in zip(fogli, tipi, posizioni):
        globali += boundsheet(nome, posizione, tipo, nomi_compressi)
    globali += coda
    assert len(globali) == lunghezza_globali

    flusso = globali + b"".join(sottoflussi) + b"\x00" * riempimento
    # A small unrelated stream precedes the workbook, as in real files (macro
    # leftovers). This offsets the workbook from the start of the
    # mini-stream, so following the wrong chain would still land on
    # plausible data instead of failing the test outright.
    return costruisci_ole2([("Ctls", b"resti di un modulo" * 6), ("Workbook", flusso)])


class BaseXls(unittest.TestCase):
    """Shared helpers: write a temporary `.xls` and read it back."""

    def scrivi(self, contenuto: bytes, nome: str = "prova.xls") -> Path:
        cartella = Path(tempfile.mkdtemp(prefix="collaudo_xls_"))
        self.addCleanup(shutil.rmtree, cartella, True)
        percorso = cartella / nome
        percorso.write_bytes(contenuto)
        return percorso

    def leggi(self, contenuto: bytes) -> list:
        return read_workbook(self.scrivi(contenuto))

    def valori(self, foglio) -> list[list[object]]:
        return [[valore for valore, _grassetto in riga] for riga in foglio.rows]


# ---------------------------------------------------------------------------
# The OLE2 container
# ---------------------------------------------------------------------------


class ContenitoreOle2Test(BaseXls):
    def test_libro_piccolo_passa_dal_mini_stream(self):
        """A workbook under 4096 bytes lives in the mini sectors, not the normal ones."""
        contenuto = costruisci_xls([("Foglio1", label(0, 0, "ciao"))])
        # The stream is genuinely under the threshold here, so reading goes
        # through the mini-FAT; a bug in that path would yield no cells at all.
        contenitore = LETTORE._ContenitoreOle2(contenuto)
        self.assertLess(len(contenitore.flusso("Workbook")), 4096)
        self.assertEqual(contenitore.flusso("Ctls"), b"resti di un modulo" * 6)
        fogli = self.leggi(contenuto)
        self.assertEqual(self.valori(fogli[0]), [["ciao"]])

    def test_libro_grande_passa_dai_settori_normali(self):
        """Above the threshold, the stream lives in the normal sector chain."""
        contenuto = costruisci_xls([("Foglio1", label(0, 0, "ciao"))], riempimento=9000)
        contenitore = LETTORE._ContenitoreOle2(contenuto)
        self.assertGreater(len(contenitore.flusso("Workbook")), 4096)
        fogli = self.leggi(contenuto)
        self.assertEqual(self.valori(fogli[0]), [["ciao"]])

    def test_file_grande_usa_piu_settori_di_fat(self):
        """With enough sectors, the FAT spans more than one sector."""
        # One FAT sector describes 128 sector entries; beyond that, more FAT
        # pages are needed. Reading only the first would break the chain
        # partway through, silently dropping the tail of the price list.
        celle = b"".join(label(riga, 0, f"riga {riga}") for riga in range(6000))
        contenuto = costruisci_xls([("Foglio1", celle)])
        contenitore = LETTORE._ContenitoreOle2(contenuto)
        self.assertGreater(len(contenitore._fat) // 128, 1, "il file di prova non è abbastanza grande")
        fogli = self.leggi(contenuto)
        self.assertEqual(len(fogli[0].rows), 6000)
        self.assertEqual(fogli[0].rows[5999][0], ("riga 5999", False))

    def test_file_non_ole2(self):
        percorso = self.scrivi(b"questo e' un file di testo qualsiasi" * 40)
        with self.assertRaises(XlsError) as errore:
            read_workbook(percorso)
        self.assertIn("Excel 97-2003", str(errore.exception))

    def test_xlsx_scambiato_per_xls(self):
        """Renaming a `.xlsx` to `.xls` should raise a clear error, not crash."""
        percorso = self.scrivi(b"PK\x03\x04" + b"\x00" * 600)
        with self.assertRaises(XlsError) as errore:
            read_workbook(percorso)
        self.assertIn(".xlsx", str(errore.exception))

    def test_contenitore_senza_libro_di_lavoro(self):
        """An OLE2 file from another program is not "corrupted": it's a different file."""
        contenuto = costruisci_ole2([("Ctls", b"quattro salti")])
        with self.assertRaises(XlsError) as errore:
            read_workbook(self.scrivi(contenuto))
        self.assertIn("un altro programma", str(errore.exception))

    def test_file_protetto_da_password(self):
        contenuto = costruisci_xls([("Foglio1", label(0, 0, "ciao"))], protetto=True)
        with self.assertRaises(XlsError) as errore:
            read_workbook(self.scrivi(contenuto))
        self.assertIn("password", str(errore.exception))

    def test_record_tagliato_a_meta(self):
        """A corrupted file should raise a clear error, not a raw struct exception."""
        celle = label(0, 0, "buona") + rec(0x00FD, b"\x00\x00")
        with self.assertRaises(XlsError) as errore:
            read_workbook(self.scrivi(costruisci_xls([("Foglio1", celle)])))
        self.assertIn("rovinato", str(errore.exception))

    def test_l_originale_non_viene_toccato(self):
        """The supplier's file must never be modified by the reader.

        The file is made read-only before opening: rather than checking its
        mtime afterward, which might not change even if written, this asks
        the OS to refuse the write outright.
        """
        percorso = self.scrivi(costruisci_xls([("Foglio1", label(0, 0, "ciao"))]))
        prima = percorso.read_bytes()
        os.chmod(percorso, stat.S_IREAD)
        self.addCleanup(os.chmod, percorso, stat.S_IWRITE | stat.S_IREAD)
        fogli = read_workbook(percorso)
        self.assertEqual(self.valori(fogli[0]), [["ciao"]])
        self.assertEqual(percorso.read_bytes(), prima)


# ---------------------------------------------------------------------------
# Cell records
# ---------------------------------------------------------------------------


class RecordCelleTest(BaseXls):
    def test_quattro_modi_del_record_rk(self):
        """The four ways an RK record packs a number into four bytes.

        Getting the "x100" flag wrong doesn't fail the read: it silently
        produces prices a hundred times too large.
        """
        celle = (
            rk(0, 0, rk_reale(1.5))
            + rk(0, 1, rk_reale(1.5, centesimi=True))
            + rk(0, 2, rk_intero(1234))
            + rk(0, 3, rk_intero(1234, centesimi=True))
            + rk(0, 4, rk_intero(-7))
            + rk(0, 5, rk_intero(-750, centesimi=True))
        )
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [[1.5, 0.015, 1234, 12.34, -7, -7.5]])

    def test_mulrk_descrive_piu_celle_in_un_record(self):
        celle = mulrk(3, 2, [(0, rk_intero(10)), (0, rk_intero(20)), (0, rk_intero(3050, centesimi=True))])
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(fogli[0].rows[3][2][0], 10)
        self.assertEqual(fogli[0].rows[3][3][0], 20)
        self.assertEqual(fogli[0].rows[3][4][0], 30.5)

    def test_mulblank_occupa_le_colonne_senza_valore(self):
        """Formatted empty cells still exist and count toward the sheet's row width."""
        celle = label(0, 0, "intestazione") + mulblank(1, 1, [0, 0, 0])
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(len(fogli[0].rows), 2)
        self.assertEqual(len(fogli[0].rows[1]), 4)
        self.assertEqual(self.valori(fogli[0])[1], [None, None, None, None])

    def test_numero_label_e_blank(self):
        celle = numero(0, 0, 12.25) + label(0, 1, "Città") + blank(0, 2)
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [[12.25, "Città", None]])

    def test_numero_intero_resta_intero(self):
        """Carton sizes and quantities are integers; they must not turn into 24.0."""
        celle = numero(0, 0, 24.0) + rk(0, 1, rk_intero(6))
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        valori = self.valori(fogli[0])[0]
        self.assertEqual(valori, [24, 6])
        self.assertIsInstance(valori[0], int)
        self.assertIsInstance(valori[1], int)

    def test_booleani_ed_errori(self):
        celle = boolerr(0, 0, 1, False) + boolerr(0, 1, 0, False) + boolerr(0, 2, 0x07, True)
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [[True, False, "#DIV/0!"]])

    def test_formula_restituisce_il_valore_memorizzato(self):
        """A formula cell exposes Excel's cached result, not the formula itself."""
        celle = formula(0, 0, formula_numero(21.0)) + formula(0, 1, formula_numero(3.75))
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [[21, 3.75]])

    def test_formula_con_risultato_di_testo(self):
        """When a formula's result is text, the text is in the STRING record that follows."""
        celle = formula(0, 0, formula_testo()) + record_string("NOCE") + numero(0, 1, 1.0)
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [["NOCE", 1]])

    def test_formula_booleana_errore_e_testo_vuoto(self):
        celle = (
            formula(0, 0, formula_booleano(True))
            + formula(0, 1, formula_errore(0x2A))
            + formula(0, 2, formula_stringa_vuota())
        )
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(self.valori(fogli[0]), [[True, "#N/A", None]])

    def test_righe_dense_e_celle_mancanti(self):
        """rows[r][c] must never raise: missing cells read as (None, False)."""
        celle = label(0, 0, "a") + label(3, 5, "b")
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        righe = fogli[0].rows
        self.assertEqual(len(righe), 4)
        self.assertTrue(all(len(riga) == 6 for riga in righe))
        self.assertEqual(righe[0][0], ("a", False))
        self.assertEqual(righe[3][5], ("b", False))
        self.assertEqual(righe[2][4], (None, False))

    def test_foglio_senza_celle(self):
        fogli = self.leggi(costruisci_xls([("Foglio1", b"")]))
        self.assertEqual(fogli[0].rows, [])


# ---------------------------------------------------------------------------
# The shared string table
# ---------------------------------------------------------------------------


def blocchi_sst_di_prova() -> tuple[list[bytes], list[str]]:
    """Build an SST deliberately split at the four points that trip up a naive parser.

    1. a string that resumes with a different encoding (1-byte chars to 2-byte);
    2. one split the other way (2-byte chars to 1-byte);
    3. one whose header ends exactly at the record boundary, with all its
       characters in the next chunk;
    4. one that ends exactly at the record boundary, so the next chunk starts
       with a fresh header rather than a continuation flag byte.

    A string with rich formatting and one with a phonetic run follow: that
    extra data must be skipped in full, or it eats into the next string's
    header, which is why another plain string comes after them.
    """
    blocchi = [
        # SST: two counters, then the first string, cut halfway through.
        struct.pack("<II", 12, 7) + struct.pack("<HB", 8, 0x00) + b"ALFA",
        # resumes as UTF-16, then the second string starts and is cut again
        b"\x01" + "BETA".encode("utf-16-le") + struct.pack("<HB", 4, 0x01) + "WX".encode("utf-16-le"),
        # the second string resumes compressed; the third's header lands here
        b"\x00" + b"YZ" + struct.pack("<HB", 3, 0x00),
        # the third string's characters, ending exactly at the chunk boundary
        b"\x00" + b"TRE",
        # no continuation flag: a fresh header starts right here
        struct.pack("<HB", 6, 0x00)
        + b"QUINTA"
        # rich formatting: two 4-byte runs to skip
        + struct.pack("<HBH", 4, 0x08, 2)
        + b"RICH"
        + b"\x00" * 8
        # phonetic extension: six bytes to skip
        + struct.pack("<HBI", 3, 0x04, 6)
        + b"SEI"
        + b"\x00" * 6
        + struct.pack("<HB", 6, 0x00)
        + b"ULTIMA",
    ]
    return blocchi, ["ALFABETA", "WXYZ", "TRE", "QUINTA", "RICH", "SEI", "ULTIMA"]


class TestiCondivisiTest(BaseXls):
    def test_sst_spezzata_sui_record_continue(self):
        blocchi, attesi = blocchi_sst_di_prova()
        celle = b"".join(labelsst(indice, 0, indice) for indice in range(len(attesi)))
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)], blocchi_sst=blocchi))
        self.assertEqual(self.valori(fogli[0]), [[testo] for testo in attesi])

    def test_sst_in_un_record_solo(self):
        blocchi = [
            struct.pack("<II", 2, 2)
            + struct.pack("<HB", 4, 0x00)
            + b"UNO "
            + struct.pack("<HB", 3, 0x01)
            + "DUE".encode("utf-16-le")
        ]
        celle = labelsst(0, 0, 0) + labelsst(0, 1, 1)
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)], blocchi_sst=blocchi))
        self.assertEqual(self.valori(fogli[0]), [["UNO ", "DUE"]])

    def test_testo_compresso_con_accenti(self):
        """"Compressed" text is latin-1, not ASCII: accented characters must decode."""
        blocchi = [struct.pack("<II", 1, 1) + struct.pack("<HB", 6, 0x00) + "però!".encode("latin-1") + b"?"]
        fogli = self.leggi(costruisci_xls([("Foglio1", labelsst(0, 0, 0))], blocchi_sst=blocchi))
        self.assertEqual(self.valori(fogli[0]), [["però!?"]])


# ---------------------------------------------------------------------------
# Bold formatting
# ---------------------------------------------------------------------------


class GrassettoTest(BaseXls):
    def test_grassetto_dalla_cella_al_font(self):
        """Resolves bold via cell -> XF -> FONT, the chain this supplier uses to flag a promo price.

        The mapping is deliberately crossed: XF 0 points at font 1 and vice
        versa, so following the format's own number instead of its XF record
        would resolve bold on the wrong cell. Bold is read from the font
        weight, the field that has been authoritative since BIFF5, not the
        legacy options bit.
        """
        celle = label(0, 0, "normale", indice_xf=0) + label(0, 1, "offerta", indice_xf=1)
        contenuto = costruisci_xls(
            [("Foglio1", celle)],
            pesi_font=(700, 400),
            font_di_xf=(1, 0),
        )
        fogli = self.leggi(contenuto)
        self.assertEqual(fogli[0].rows[0][0], ("normale", False))
        self.assertEqual(fogli[0].rows[0][1], ("offerta", True))

    def test_il_font_numero_quattro_non_esiste(self):
        """Excel skips font index 4; ignoring that offsets bold by one font.

        Six fonts are written here, so the last one has index 6, not 5. If the
        skip isn't accounted for, bold resolves against the wrong font.
        """
        celle = label(0, 0, "normale", indice_xf=0) + label(0, 1, "offerta", indice_xf=1)
        contenuto = costruisci_xls(
            [("Foglio1", celle)],
            pesi_font=(400, 400, 400, 400, 400, 700),
            font_di_xf=(0, 6),
        )
        fogli = self.leggi(contenuto)
        self.assertEqual(fogli[0].rows[0][0][1], False)
        self.assertEqual(fogli[0].rows[0][1][1], True)

    def test_grassetto_anche_sulle_celle_numeriche_e_vuote(self):
        celle = rk(0, 0, rk_intero(150, centesimi=True), indice_xf=1) + blank(0, 1, indice_xf=1)
        contenuto = costruisci_xls(
            [("Foglio1", celle)],
            pesi_font=(700, 400),
            font_di_xf=(1, 0),
        )
        fogli = self.leggi(contenuto)
        self.assertEqual(fogli[0].rows[0][0], (1.5, True))
        self.assertEqual(fogli[0].rows[0][1], (None, True))

    def test_indice_xf_fuori_elenco_non_fa_saltare_la_lettura(self):
        """A missing format index means "not bold", not a crash."""
        celle = label(0, 0, "ciao", indice_xf=99)
        fogli = self.leggi(costruisci_xls([("Foglio1", celle)]))
        self.assertEqual(fogli[0].rows[0][0], ("ciao", False))


# ---------------------------------------------------------------------------
# Multiple sheets
# ---------------------------------------------------------------------------


class PiuFogliTest(BaseXls):
    def test_nomi_e_ordine_dei_fogli(self):
        fogli = self.leggi(
            costruisci_xls(
                [
                    ("Listino", label(0, 0, "primo")),
                    ("Città", label(0, 0, "secondo")),
                    ("Vuoto", b""),
                ]
            )
        )
        self.assertEqual([foglio.name for foglio in fogli], ["Listino", "Città", "Vuoto"])
        self.assertEqual(self.valori(fogli[0]), [["primo"]])
        self.assertEqual(self.valori(fogli[1]), [["secondo"]])
        self.assertEqual(fogli[2].rows, [])

    def test_nomi_dei_fogli_scritti_compressi(self):
        fogli = self.leggi(
            costruisci_xls([("Prezzi", label(0, 0, "x"))], nomi_compressi=True)
        )
        self.assertEqual([foglio.name for foglio in fogli], ["Prezzi"])

    def test_i_grafici_non_sono_fogli_di_celle(self):
        fogli = self.leggi(
            costruisci_xls(
                [("Dati", label(0, 0, "x")), ("Grafico", b"")],
                tipi_foglio=(0, 2),
            )
        )
        self.assertEqual([foglio.name for foglio in fogli], ["Dati"])

    def test_read_sheet_values_restituisce_i_soli_valori(self):
        percorso = self.scrivi(
            costruisci_xls(
                [
                    ("Uno", label(0, 0, "primo")),
                    ("Due", numero(0, 0, 7.0) + numero(0, 1, 0.5)),
                ]
            )
        )
        self.assertEqual(read_sheet_values(percorso), [["primo"]])
        self.assertEqual(read_sheet_values(percorso, 1), [[7, 0.5]])

    def test_read_sheet_values_foglio_inesistente(self):
        percorso = self.scrivi(costruisci_xls([("Uno", label(0, 0, "primo"))]))
        with self.assertRaises(XlsError) as errore:
            read_sheet_values(percorso, 5)
        self.assertIn("non esiste", str(errore.exception))


# ---------------------------------------------------------------------------
# Comparison against a real price list
# ---------------------------------------------------------------------------


def _numeri_uguali(mio: object, suo: object) -> bool:
    if isinstance(mio, bool) or isinstance(suo, bool):
        return False
    if not isinstance(mio, (int, float)) or not isinstance(suo, (int, float)):
        return False
    return abs(mio - suo) <= TOLLERANZA_ASSOLUTA + TOLLERANZA_RELATIVA * abs(suo)


class ListinoVeroTest(unittest.TestCase):
    """Cell-by-cell comparison against the Excel-converted copy."""

    @classmethod
    def setUpClass(cls):
        if not LISTINO_XLS.is_file():
            raise unittest.SkipTest(
                f"Manca il listino .xls di prova in {LISTINO_XLS}. "
                "Il percorso si cambia con la variabile d'ambiente LISTINO_XLS_DI_PROVA."
            )
        if not LISTINO_XLSX.is_file():
            raise unittest.SkipTest(
                f"Manca la copia convertita da Excel in {LISTINO_XLSX}. "
                "Serve solo come termine di paragone e sta fuori dal progetto: "
                "si ottiene aprendo il .xls con Excel e salvandolo come .xlsx. "
                "Il percorso si cambia con la variabile d'ambiente LISTINO_XLSX_DI_PARAGONE."
            )
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - depends on the environment
            raise unittest.SkipTest(
                "openpyxl non è installato: serve solo a questo confronto, "
                "non al codice che legge i listini."
            )
        cls.load_workbook = staticmethod(load_workbook)
        cls.fogli = read_workbook(LISTINO_XLS)

    def test_nomi_dei_fogli(self):
        libro = self.load_workbook(LISTINO_XLSX, data_only=True, read_only=True)
        try:
            self.assertEqual([foglio.name for foglio in self.fogli], libro.sheetnames)
        finally:
            libro.close()

    def test_tutte_le_celle_di_tutti_i_fogli(self):
        libro = self.load_workbook(LISTINO_XLSX, data_only=True, read_only=True)
        confrontate = 0
        differenze: list[str] = []
        non_vuote = 0
        try:
            self.assertEqual(len(self.fogli), len(libro.sheetnames))
            for foglio, nome in zip(self.fogli, libro.sheetnames):
                paragone = libro[nome]
                attese = [
                    [
                        (cella.value, bool(getattr(cella, "font", None) and cella.font.bold))
                        for cella in riga
                    ]
                    for riga in paragone.iter_rows()
                ]
                numero_righe = max(len(foglio.rows), len(attese))
                numero_colonne = max(
                    max((len(riga) for riga in foglio.rows), default=0),
                    max((len(riga) for riga in attese), default=0),
                )
                for indice_riga in range(numero_righe):
                    mia_riga = foglio.rows[indice_riga] if indice_riga < len(foglio.rows) else ()
                    sua_riga = attese[indice_riga] if indice_riga < len(attese) else ()
                    for indice_colonna in range(numero_colonne):
                        confrontate += 1
                        mio, mio_grassetto = (
                            mia_riga[indice_colonna]
                            if indice_colonna < len(mia_riga)
                            else (None, False)
                        )
                        suo, suo_grassetto = (
                            sua_riga[indice_colonna]
                            if indice_colonna < len(sua_riga)
                            else (None, False)
                        )
                        if mio is not None:
                            non_vuote += 1
                        uguale = mio == suo if type(mio) is type(suo) else _numeri_uguali(mio, suo)
                        if not uguale or mio_grassetto != suo_grassetto:
                            if len(differenze) < 20:
                                differenze.append(
                                    f"{nome} riga {indice_riga + 1} colonna {indice_colonna + 1}: "
                                    f"letto {mio!r}/{mio_grassetto} "
                                    f"invece di {suo!r}/{suo_grassetto}"
                                )
        finally:
            libro.close()

        print(
            f"\n[collaudo] celle confrontate con la copia convertita da Excel: "
            f"{confrontate} (di cui {non_vuote} con un contenuto)",
            file=sys.stderr,
        )
        self.assertFalse(differenze, "\n".join(differenze))
        # A comparison over too few cells would prove nothing: the real price
        # list has well over 300,000 and all of them must be present.
        self.assertGreater(confrontate, 300_000)
        self.assertGreater(non_vuote, 250_000)

    def test_le_offerte_in_grassetto_vengono_lette(self):
        """Bold is the promo signal: it must be read exactly, never guessed."""
        libro = self.load_workbook(LISTINO_XLSX, data_only=True, read_only=True)
        try:
            attesi = sum(
                1
                for riga in libro[libro.sheetnames[0]].iter_rows()
                for cella in riga
                if getattr(cella, "font", None) and cella.font.bold
            )
        finally:
            libro.close()
        letti = sum(1 for riga in self.fogli[0].rows for _valore, grassetto in riga if grassetto)
        self.assertEqual(letti, attesi)
        # This price list has no active promos, but its header row is bold:
        # a zero count here would say nothing about the signal that matters.
        self.assertGreater(letti, 0)

    def test_il_listino_vero_non_viene_modificato(self):
        stato = LISTINO_XLS.stat()
        read_workbook(LISTINO_XLS)
        nuovo = LISTINO_XLS.stat()
        self.assertEqual(nuovo.st_mtime_ns, stato.st_mtime_ns)
        self.assertEqual(nuovo.st_size, stato.st_size)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
