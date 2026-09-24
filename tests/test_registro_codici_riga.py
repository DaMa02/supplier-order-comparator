"""Supplier row codes live in the adapter registry, not in the code.

Larice marks a threshold's reward row with `SM`: it has a real code, EAN and
description, but it isn't merchandise the store can order. A price-only check
would flag it by accident and stop working the day the supplier fills in a
price for the free item. The rule must change by editing
`references/adapters.json`, without touching the code.
"""

from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
APP = SKILL_ROOT / "app"
ADAPTERS = SKILL_ROOT / "references" / "adapters.json"
for cartella in (SCRIPTS, APP):
    if str(cartella) not in sys.path:
        sys.path.insert(0, str(cartella))

import registro  # noqa: E402
from prepare_sources import larice_discount, read_larice  # noqa: E402
from prepare_manifest_sources import read_mapped_csv_supplier, read_mapped_xlsx_supplier  # noqa: E402
from promotion_bridge import PromotionService  # noqa: E402


# Larice price-list columns (A to R), as declared by the adapter.
COLONNA = {"supplier_code": 3, "pieces_per_carton": 5, "pallet": 6, "description": 7,
           "unit_price_pre_discount": 15, "discount": 16, "vat": 17, "ean": 18}


def scrivi_listino_larice(percorso: Path, righe: list[dict[str, object]]) -> Path:
    """Build a sheet shaped like the Larice price list: header then data rows."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LISTINO"
    for numero, riga in enumerate(righe, start=1):
        for campo, colonna in COLONNA.items():
            if campo in riga:
                sheet.cell(row=numero, column=colonna, value=riga[campo])
    workbook.save(percorso)
    workbook.close()
    return percorso


PRODOTTO = {
    "supplier_code": "12345", "pieces_per_carton": 6, "description": "BAGNO VIDOR 500 ML TALCO",
    "unit_price_pre_discount": 2.50, "discount": "TP", "vat": 22, "ean": "8000000000011",
}
PREMIO = {
    "supplier_code": "27606", "pieces_per_carton": 6, "description": "IN OMAGGIO 1 CT DI",
    "discount": "SM", "vat": 22, "ean": "8059147044172",
}


class CodiciDiRigaTests(unittest.TestCase):
    maxDiff = None

    def setUp(self) -> None:
        self.radice = Path(tempfile.mkdtemp(prefix="collaudo_codici_"))
        self.addCleanup(shutil.rmtree, self.radice, True)

    # -- the rule lives in the registry --------------------------------

    def test_il_registro_dichiara_che_cosa_significa_ogni_codice(self) -> None:
        codici = registro.codici_di_riga(registro.adattatore("larice_v1"))

        self.assertEqual(codici["field"], "discount_raw")
        self.assertEqual(registro.codici_ammessi(codici), {"TP", "SM"})
        self.assertTrue(codici["codes"]["TP"]["orderable"])
        self.assertFalse(codici["codes"]["SM"]["orderable"])
        self.assertEqual(codici["codes"]["SM"]["row_type"], "OMAGGIO")
        self.assertIn("omaggio", codici["codes"]["SM"]["means"].casefold())

    def test_i_codici_ammessi_li_dice_il_registro_e_non_il_codice(self) -> None:
        """When Larice defines what "**" means, only the registry needs updating,
        not a function."""
        _sconto, _tipo, avviso = larice_discount("**", {"TP", "SM"})
        self.assertEqual(avviso, "Codice sconto testuale inatteso: **")

        _sconto, _tipo, ammesso = larice_discount("**", {"TP", "SM", "**"})
        self.assertIsNone(ammesso)

        _sconto, _tipo, vuoto = larice_discount("", set())
        self.assertIsNone(vuoto)

    # -- la riga premio non è merce ----------------------------------------

    def test_la_riga_premio_non_e_ordinabile(self) -> None:
        listino = scrivi_listino_larice(self.radice / "larice.xlsx", [PRODOTTO, PREMIO])

        records, avvisi = read_larice(listino)

        self.assertEqual(avvisi, [])
        prodotto, premio = records
        self.assertTrue(prodotto["usable"])
        self.assertIsNone(prodotto.get("row_type"))
        self.assertFalse(premio["usable"])
        self.assertEqual(premio["row_type"], "OMAGGIO")
        self.assertIn("non è merce acquistabile", premio["not_orderable_reason"])

    def test_una_riga_premio_con_un_prezzo_resta_non_ordinabile(self) -> None:
        """A reward row must stay non-orderable even once it carries a price.

        Some suppliers, like Quercia in its promotional section, do fill in a
        price for the free item; a price-only check would then treat it as a
        purchasable offer.
        """
        premio_con_prezzo = {**PREMIO, "unit_price_pre_discount": 4.90}
        listino = scrivi_listino_larice(self.radice / "larice.xlsx", [PRODOTTO, premio_con_prezzo])

        records, _avvisi = read_larice(listino)

        premio = records[1]
        # The price is still read: no data from the file is hidden.
        self.assertEqual(premio["unit_price_net"], "4.9000")
        # But the row never becomes orderable.
        self.assertFalse(premio["usable"])
        self.assertEqual(premio["row_type"], "OMAGGIO")

    def test_senza_la_regola_nel_registro_la_riga_tornerebbe_acquistabile(self) -> None:
        """Counter-check: the registry decides, not the code."""
        premio_con_prezzo = {**PREMIO, "unit_price_pre_discount": 4.90}
        listino = scrivi_listino_larice(self.radice / "larice.xlsx", [PRODOTTO, premio_con_prezzo])
        registro_senza = self.radice / "adapters_senza_regola.json"
        documento = json.loads(ADAPTERS.read_text(encoding="utf-8"))
        for voce in documento["adapters"]:
            voce.pop("row_markers", None)
        registro_senza.write_text(json.dumps(documento), encoding="utf-8")

        originale = registro.REGISTRO
        registro.REGISTRO = registro_senza
        try:
            records, _avvisi = read_larice(listino)
        finally:
            registro.REGISTRO = originale

        self.assertTrue(records[1]["usable"])
        self.assertIsNone(records[1].get("row_type"))

    # -- the same rule for an arbitrary mapped supplier ------------------

    def test_la_stessa_regola_vale_per_un_fornitore_mappato(self) -> None:
        """The rule belongs to the engine, not to Larice: any supplier can
        declare it in its own mapping and get the same behavior."""
        percorso = self.radice / "nuovo_fornitore.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo;codice\n"
            "8000000000011;PRODOTTO NORMALE;6;2,50;\n"
            "8059147044172;IN REGALO 1 CT DI;6;4,90;OM\n",
            encoding="utf-8",
        )
        mappatura = {
            "header_row": 1,
            "delimiter": ";",
            "italian_numbers": True,
            "columns": {
                "ean": "ean", "description": "descrizione",
                "pieces_per_carton": "pezzi", "unit_price_net": "prezzo", "discount": "codice",
            },
            # `field` names the normalized record field (`discount_raw`), not the
            # column header in the file, so the rule doesn't depend on how the
            # supplier named its column.
            "row_markers": {
                "field": "discount_raw",
                "codes": {"OM": {"means": "Riga regalo, non acquistabile.", "orderable": False, "row_type": "OMAGGIO"}},
            },
        }
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertTrue(records[0]["usable"])
        self.assertFalse(records[1]["usable"])
        self.assertEqual(records[1]["row_type"], "OMAGGIO")
        self.assertEqual(records[1]["unit_price_net"], "4.9000")
        self.assertEqual(lettura["rows_not_orderable"], {"OMAGGIO": 1})

    def test_la_stessa_regola_vale_anche_su_un_foglio_di_calcolo(self) -> None:
        """The format doesn't matter: the rule belongs to the engine, not the reader."""
        percorso = self.radice / "nuovo_fornitore.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "pezzi", "prezzo", "codice"])
        sheet.append(["8000000000011", "PRODOTTO NORMALE", 6, 2.50, None])
        sheet.append(["8059147044172", "REGALO DELLA SOGLIA", 6, 4.90, "OM"])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {
                "ean": "ean", "description": "descrizione",
                "pieces_per_carton": "pezzi", "unit_price_net": "prezzo", "discount": "codice",
            },
            "row_markers": {
                "field": "discount_raw",
                "codes": {"OM": {"means": "Riga regalo, non acquistabile.", "orderable": False, "row_type": "OMAGGIO"}},
            },
        }
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_xlsx_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertTrue(records[0]["usable"])
        self.assertFalse(records[1]["usable"])
        self.assertEqual(records[1]["row_type"], "OMAGGIO")
        self.assertEqual(lettura["rows_not_orderable"], {"OMAGGIO": 1})

    # -- a supplier that doesn't mark the reward row at all --------------

    def test_una_riga_premio_senza_codice_si_riconosce_dal_testo(self) -> None:
        """Larice's newer canvass marks all promoted rows as PROMO, reward rows
        included; the only signal left is the description text, which the
        promotion engine already parses."""
        percorso = self.radice / "canvass.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "pezzi", "prezzo", "codice"])
        sheet.append(["8000000000011", "SH. A/ERBAR. 250 LAVANDA", 12, 0.68, "PROMO"])
        sheet.append(["8000000000011", "IN OMAGGIO 1CT SH. A/ERBAR. 250ML LAVANDA", 12, 0.65, "PROMO"])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {
                "ean": "ean", "description": "descrizione",
                "pieces_per_carton": "pezzi", "unit_price_net": "prezzo", "discount": "codice",
            },
            "row_markers": {
                "field": "discount_raw",
                "reward_rows": {
                    "means": "Riga premio di una soglia con omaggio.",
                    "orderable": False,
                    "row_type": "OMAGGIO",
                },
            },
        }
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_xlsx_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertTrue(records[0]["usable"])
        self.assertFalse(records[1]["usable"])
        self.assertEqual(records[1]["row_type"], "OMAGGIO")
        self.assertEqual(lettura["rows_not_orderable"], {"OMAGGIO": 1})

    def test_senza_reward_rows_il_premio_col_prezzo_batte_la_merce(self) -> None:
        """Counter-check: both rows share an EAN and the reward row is cheaper.
        Without `reward_rows`, the comparison sees two offers for the same
        product and picks the free one as the best price."""
        percorso = self.radice / "canvass_senza_regola.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "pezzi", "prezzo", "codice"])
        sheet.append(["8000000000011", "SH. A/ERBAR. 250 LAVANDA", 12, 0.68, "PROMO"])
        sheet.append(["8000000000011", "IN OMAGGIO 1CT SH. A/ERBAR. 250ML LAVANDA", 12, 0.65, "PROMO"])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {
                "ean": "ean", "description": "descrizione",
                "pieces_per_carton": "pezzi", "unit_price_net": "prezzo", "discount": "codice",
            },
        }

        records, _avvisi = read_mapped_xlsx_supplier(percorso, "nuovo", mappatura)

        ordinabili = [record for record in records if record["usable"]]
        self.assertEqual(len(ordinabili), 2)
        self.assertEqual(min(record["unit_price_net"] for record in ordinabili), "0.6500")

    def test_una_riga_premio_marcata_ordinabile_resta_fuori_lo_stesso(self) -> None:
        """When a code and `reward_rows` disagree, "not orderable" wins.

        Mistaking a product for a reward is visible immediately (it's missing
        from the comparison); mistaking a reward for a product puts a price
        that doesn't exist into the order, and only the supplier catches it.
        """
        percorso = self.radice / "canvass_marcato.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "pezzi", "prezzo", "codice"])
        sheet.append(["8000000000011", "IN OMAGGIO 1CT SH. A/ERBAR. 250ML LAVANDA", 12, 0.65, "TP"])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {
                "ean": "ean", "description": "descrizione",
                "pieces_per_carton": "pezzi", "unit_price_net": "prezzo", "discount": "codice",
            },
            "row_markers": {
                "field": "discount_raw",
                "codes": {"TP": {"means": "Prezzo di transito: merce normale.", "orderable": True}},
                "reward_rows": {
                    "means": "Riga premio di una soglia con omaggio.",
                    "orderable": False,
                    "row_type": "OMAGGIO",
                },
            },
        }

        records, _avvisi = read_mapped_xlsx_supplier(percorso, "nuovo", mappatura)

        self.assertFalse(records[0]["usable"])
        self.assertEqual(records[0]["row_type"], "OMAGGIO")

    # -- the reward row doesn't count toward its own threshold -----------

    def test_il_premio_non_conta_fra_i_prodotti_che_fanno_soglia(self) -> None:
        """A reward row whose description doesn't say "in omaggio" is only
        recognizable by its code; without the rule it would be counted among
        the purchased items that reach the threshold."""
        listino = scrivi_listino_larice(self.radice / "larice.xlsx", [
            {"description": "ACQUISTANDO 3 CT TRA"},
            {**PRODOTTO, "ean": "8000000000011"},
            {**PRODOTTO, "ean": "8000000000012", "description": "BAGNO VIDOR 500 ML VITAMINA C"},
            # The reward row: marked SM but described with the gifted product's name.
            {**PREMIO, "description": "RESALINA SALE LAVASTOVIGLIE KG1"},
            {**PREMIO, "description": "IN OMAGGIO 1 CT DI"},
        ])
        servizio = PromotionService()

        promozioni = servizio.detect({
            "files": [{"supplierId": "larice", "sourcePath": str(listino)}],
            "products": [],
        })

        self.assertEqual(len(promozioni), 1)
        ammessi = (promozioni[0].get("eligible") or {}).get("source_rows") or []
        self.assertEqual(sorted(ammessi), [2, 3])


class IlCanvassNuovoDiLariceLettoDavvero(unittest.TestCase):
    """Runs against the real Larice price list and its shipped adapter.

    The tests above use fixtures built for the purpose; this one needs the
    real file because the defect the rule closes only shows up there: the six
    "IN OMAGGIO ..." rows carry a real price and no code, and five of the six
    repeat the EAN of an item already listed at full price.
    """

    LISTINO = SKILL_ROOT / "listini-storici" / "New Larice N°37(v.0).xls"

    def setUp(self) -> None:
        if not self.LISTINO.is_file():
            self.skipTest(f"Manca il listino {self.LISTINO}")
        self.adattatore = registro.adattatore("larice_canvass_v1", ADAPTERS)

    def test_le_sei_righe_premio_restano_fuori_dall_ordine(self) -> None:
        lettura: dict[str, object] = {}

        records, avvisi = read_mapped_xlsx_supplier(
            self.LISTINO, "larice", self.adattatore["field_mapping"], report=lettura
        )

        self.assertEqual(avvisi, [])
        omaggi = [record for record in records if record.get("row_type") == "OMAGGIO"]
        self.assertEqual(len(omaggi), 6)
        self.assertFalse(any(record["usable"] for record in omaggi))
        self.assertEqual(lettura["rows_not_orderable"]["OMAGGIO"], 6)

    def test_il_prezzo_del_regalo_non_diventa_un_offerta(self) -> None:
        """EAN 8019580330416 is 0.68 as merchandise and 0.65 as the reward."""

        records, _avvisi = read_mapped_xlsx_supplier(
            self.LISTINO, "larice", self.adattatore["field_mapping"]
        )

        offerte = [
            record for record in records
            if record["usable"] and record["ean"] == "8019580330416"
        ]
        self.assertEqual(len(offerte), 1)
        self.assertEqual(offerte[0]["unit_price_net"], "0.6800")

    def test_le_sei_soglie_con_omaggio_si_ricompongono_tutte(self) -> None:
        """Six threshold headers, six reward rows, six resulting conditions.

        The reward name comes out clean: the promo text and the reward
        description share column E, and reading it twice would duplicate it.
        """

        promozioni, errore = PromotionService().condizioni_di(
            "larice", self.LISTINO, "larice_canvass_v1"
        )

        self.assertIsNone(errore)
        self.assertEqual(len(promozioni), 6)
        self.assertEqual(
            sorted(promozione["threshold"]["qty"] for promozione in promozioni),
            [2, 2, 2, 3, 5, 10],
        )
        premi = {promozione["reward"]["description"] for promozione in promozioni}
        self.assertIn("SH. A/ERBAR. 250ML LAVANDA", premi)
        self.assertTrue(
            all("IN OMAGGIO" not in nome for nome in premi),
            f"il nome del premio si è portato dietro il testo della riga: {sorted(premi)}",
        )


if __name__ == "__main__":
    unittest.main()
