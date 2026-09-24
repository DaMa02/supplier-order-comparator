"""Every row dropped from the comparison must be counted, with its reason.

`rows_not_orderable` only counted rows carrying a `row_type`, which only the
adapter registry assigns. A row dropped for an unreadable price, or missing
pieces-per-carton, has no `row_type`, so the count read zero with half a
price list out of the comparison. `row_filter` discarded rows the same way,
with no counter at all, and the four fixed-schema readers (BETULLA, Larice,
Noce, the management-software export) produced no read report at all.
"""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
ADAPTERS = SKILL_ROOT / "references" / "adapters.json"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import registro  # noqa: E402
from prepare_manifest_sources import (  # noqa: E402
    mapped_standalone_displays,
    read_mapped_csv_supplier,
    read_mapped_xlsx_supplier,
)
from prepare_sources import (  # noqa: E402
    decimal_value,
    integrate_larice_displays,
    read_betulla,
    read_gestionale,
    read_larice,
    read_noce,
)


class CasoConCartella(unittest.TestCase):
    def setUp(self) -> None:
        self.temporanea = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporanea.cleanup)
        self.radice = Path(self.temporanea.name)


class RigheNonOrdinabiliContateTests(CasoConCartella):
    """Rows demoted for price, factor or description are counted."""

    MAPPATURA = {
        "header_row": 1,
        "delimiter": ";",
        "italian_numbers": True,
        "columns": {
            "ean": "ean",
            "description": "descrizione",
            "pieces_per_carton": "pezzi",
            "unit_price_net": "prezzo",
        },
    }

    def listino(self, righe: str) -> Path:
        percorso = self.radice / "listino.csv"
        percorso.write_text("ean;descrizione;pezzi;prezzo\n" + righe, encoding="utf-8")
        return percorso

    def test_una_riga_senza_prezzo_si_conta_col_suo_motivo(self) -> None:
        percorso = self.listino(
            "8000000000011;PRODOTTO NORMALE;6;2,50\n"
            "8000000000012;PREZZO ILLEGGIBILE;6;\n"
        )
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertTrue(records[0]["usable"])
        self.assertFalse(records[1]["usable"])
        self.assertEqual(records[1]["unusable_reason"], "senza_prezzo")
        self.assertEqual(lettura["rows_not_orderable"], {"senza_prezzo": 1})

    def test_una_riga_senza_pezzi_per_collo_si_conta_col_suo_motivo(self) -> None:
        percorso = self.listino(
            "8000000000011;PRODOTTO NORMALE;6;2,50\n"
            "8000000000012;PEZZI PER COLLO VUOTI;;2,50\n"
            "8000000000013;PEZZI PER COLLO A ZERO;0;2,50\n"
        )
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual(lettura["rows_not_orderable"], {"senza_pezzi_per_collo": 2})
        self.assertEqual([record["usable"] for record in records], [True, False, False])

    def test_una_riga_senza_descrizione_si_conta_col_suo_motivo(self) -> None:
        percorso = self.listino(
            "8000000000011;PRODOTTO NORMALE;6;2,50\n"
            "8000000000012;;6;2,50\n"
        )
        lettura: dict[str, object] = {}

        read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual(lettura["rows_not_orderable"], {"senza_descrizione": 1})

    def test_il_conteggio_non_e_zero_quando_meta_listino_e_fuori(self) -> None:
        """Regression guard: the count must not read zero when half the rows are dropped."""

        percorso = self.listino(
            "8000000000011;PRIMO;6;2,50\n"
            "8000000000012;SECONDO;6;3,50\n"
            "8000000000013;TERZO;6;4,50\n"
            "8000000000014;QUARTO SENZA PREZZO;6;\n"
            "8000000000015;QUINTO SENZA PEZZI;;5,50\n"
            "8000000000016;;6;6,50\n"
        )
        lettura: dict[str, object] = {}

        read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual(lettura["rows_kept"], 6, "le righe restano lette: non si buttano, si declassano")
        self.assertEqual(
            lettura["rows_not_orderable"],
            {"senza_prezzo": 1, "senza_pezzi_per_collo": 1, "senza_descrizione": 1},
        )
        self.assertEqual(sum(lettura["rows_not_orderable"].values()), 3)

    def test_lo_stesso_conteggio_su_un_foglio_di_calcolo(self) -> None:
        """Format is irrelevant: the count comes from the shared engine, not the reader."""

        percorso = self.radice / "listino.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "pezzi", "prezzo"])
        sheet.append(["8000000000011", "PRODOTTO NORMALE", 6, 2.50])
        sheet.append(["8000000000012", "PREZZO ILLEGGIBILE", 6, None])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {
                "ean": "ean",
                "description": "descrizione",
                "pieces_per_carton": "pezzi",
                "unit_price_net": "prezzo",
            },
        }
        lettura: dict[str, object] = {}

        read_mapped_xlsx_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertEqual(lettura["rows_not_orderable"], {"senza_prezzo": 1})


class FiltroDiRigaContatoTests(CasoConCartella):
    """`row_filter` discarded rows without counting them, unlike the other paths."""

    def test_le_righe_fuori_dal_filtro_si_contano(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo;reparto\n"
            "8000000000011;PRODOTTO CASA;6;2,50;CASA\n"
            "8000000000012;PRODOTTO FOOD;6;2,50;FOOD\n"
            "8000000000013;ALTRO FOOD;6;2,50;FOOD\n",
            encoding="utf-8",
        )
        mappatura = {
            "header_row": 1,
            "delimiter": ";",
            "italian_numbers": True,
            "columns": {
                "ean": "ean",
                "description": "descrizione",
                "pieces_per_carton": "pezzi",
                "unit_price_net": "prezzo",
                "category": "reparto",
            },
            "row_filter": {"field": "category", "equals": "CASA"},
        }
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertEqual(len(records), 1)
        self.assertEqual(lettura["rows_excluded"], {"riga_fuori_dal_filtro": 2})
        self.assertEqual(lettura["rows_kept"], 1)


class LeRigheSeparatoreNonSonoProdottiTests(CasoConCartella):
    """A section-title row can slip into the read products.

    On QUERCIA, separator rows (e.g. "OGNI € 1500,00 OMAGGIO …", "LISTINO")
    have an empty description, empty EAN and no price, but carry their text
    in the "Articolo" column, i.e. in `supplier_code`. A guard requiring all
    three fields empty misses them, since the code column isn't. The
    management-software reader has always dropped such rows under the same
    label; the supplier readers didn't: a separator above the cut inflates
    the row count, one that lands inside the product rows becomes a catalog
    entry with the promo text in place of the item code.
    """

    MAPPATURA = {
        "header_row": 1,
        "delimiter": ";",
        "italian_numbers": True,
        "columns": {
            "supplier_code": "articolo",
            "ean": "ean",
            "description": "descrizione",
            "pieces_per_carton": "pezzi",
            "unit_price_net": "prezzo",
        },
    }

    def listino(self, righe: str) -> Path:
        percorso = self.radice / "listino.csv"
        percorso.write_text("articolo;ean;descrizione;pezzi;prezzo\n" + righe, encoding="utf-8")
        return percorso

    def test_un_separatore_fra_i_prodotti_non_diventa_un_prodotto(self) -> None:
        percorso = self.listino(
            "FAT01;8000000000011;PRIMO;6;2,50\n"
            "OGNI 1500 EURO OMAGGIO A SCELTA;;;;\n"
            "FAT02;8000000000012;SECONDO;6;3,50\n"
        )
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual([record["description"] for record in records], ["PRIMO", "SECONDO"])
        self.assertEqual(lettura["rows_excluded"], {"non_e_una_riga_prodotto": 1})
        self.assertEqual(lettura["rows_kept"], 2)
        # The separator text must not end up as anyone's item code: that's
        # the value that would land in the catalog.
        self.assertNotIn(
            "OGNI 1500 EURO OMAGGIO A SCELTA",
            [str(record.get("supplier_code") or "") for record in records],
        )

    def test_lo_stesso_su_un_foglio_di_calcolo(self) -> None:
        percorso = self.radice / "listino.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["articolo", "ean", "descrizione", "pezzi", "prezzo"])
        sheet.append(["FAT01", "8000000000011", "PRIMO", 6, 2.50])
        sheet.append(["LISTINO", None, None, None, None])
        sheet.append(["FAT02", "8000000000012", "SECONDO", 6, 3.50])
        workbook.save(percorso)
        workbook.close()
        mappatura = {**self.MAPPATURA, "sheet": "FIRST", "data_start_row": 2}
        mappatura.pop("delimiter")
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_xlsx_supplier(percorso, "nuovo", mappatura, report=lettura)

        self.assertEqual([record["description"] for record in records], ["PRIMO", "SECONDO"])
        self.assertEqual(lettura["rows_excluded"], {"non_e_una_riga_prodotto": 1})

    def test_una_riga_col_prezzo_resta_un_prodotto_anche_senza_nome(self) -> None:
        """The cutoff errs toward keeping a row.

        A price or a barcode is enough to say there's an item there: the row
        stays and is demoted with its reason, which is what tells the user
        to check the description column. Only a row with no name, no
        barcode and no price is not merchandise.
        """

        percorso = self.listino(
            "FAT01;8000000000011;PRIMO;6;2,50\n"
            "FAT02;;;6;3,50\n"
            "FAT03;8000000000013;;6;\n"
        )
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual(len(records), 3)
        self.assertEqual(lettura["rows_excluded"], {})
        self.assertEqual(lettura["rows_not_orderable"], {"senza_descrizione": 2})


class RapportoDeiListiniACablatiTests(CasoConCartella):
    """The four fixed-schema readers must report what their read kept, excluded and demoted.

    They were the only ones producing no report at all, and the largest
    price lists at that: the "what got dropped" panel stayed silent where
    it mattered most.
    """

    def test_betulla_dichiara_righe_lette_tenute_e_saltate(self) -> None:
        percorso = self.radice / "betulla.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["EAN", "CodArt", "ORDINE", "Descrizione", "PzCt", "Cessione", "Pedana", "IVA"])
        sheet.append(["8000000000011", "C-1", None, "PRODOTTO BUONO", 6, 2.50, None, 22])
        sheet.append(["8000000000012", "C-2", None, "PREZZO ILLEGGIBILE", 6, None, None, 22])
        sheet.append([None, None, None, "RIGA DI TOTALE SENZA EAN", None, None, None, None])
        workbook.save(percorso)
        workbook.close()
        lettura: dict[str, object] = {}

        records = read_betulla(percorso, report=lettura)

        self.assertEqual(len(records), 2)
        self.assertEqual(lettura["rows_kept"], 2)
        self.assertEqual(lettura["rows_excluded"], {"senza_ean": 1})
        self.assertEqual(lettura["rows_not_orderable"], {"senza_prezzo": 1})

    def test_larice_dichiara_le_righe_senza_prezzo(self) -> None:
        percorso = self.radice / "larice.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        intestazione = [None] * 18
        intestazione[17] = "EAN"
        sheet.append(intestazione)
        for ean, descrizione, prezzo in (
            ("8000000000011", "PRODOTTO BUONO", 2.50),
            ("8000000000012", "PREZZO ILLEGGIBILE", None),
        ):
            riga: list[object] = [None] * 18
            riga[2] = "G-1"
            riga[4] = 6
            riga[6] = descrizione
            riga[14] = prezzo
            riga[17] = ean
            sheet.append(riga)
        workbook.save(percorso)
        workbook.close()
        lettura: dict[str, object] = {}

        records, _avvisi = read_larice(percorso, report=lettura)

        self.assertEqual(len(records), 2)
        self.assertEqual(lettura["rows_kept"], 2)
        # The header row has the word "EAN" in the EAN column, so it's skipped.
        self.assertEqual(lettura["rows_excluded"], {"senza_ean": 1})
        self.assertEqual(lettura["rows_not_orderable"], {"senza_prezzo": 1})

    def test_noce_distingue_prezzo_fattore_e_disponibilita(self) -> None:
        percorso = self.radice / "noce.csv"
        percorso.write_text(
            # The price is quoted because the decimal separator is a comma,
            # as in every price list: "2.50" unquoted would parse as two
            # hundred fifty euros, and now fails to parse at all.
            "catalog_page,ean,product,packaging,availability,price,unit\n"
            '1,8000000000011,PRODOTTO BUONO,CT,Disponibile,"2,50",x 6\n'
            "1,8000000000012,SENZA PREZZO,CT,Disponibile,,x 6\n"
            '1,8000000000013,SENZA MOLTIPLICATORE,CT,Disponibile,"2,50",CT\n'
            '1,8000000000014,ESAURITO,CT,Non disponibile,"2,50",x 6\n',
            encoding="utf-8",
        )
        lettura: dict[str, object] = {}

        records = read_noce(percorso, report=lettura)

        self.assertEqual(len(records), 4)
        self.assertEqual(lettura["rows_kept"], 4)
        self.assertEqual(
            lettura["rows_not_orderable"],
            {"senza_prezzo": 1, "senza_pezzi_per_collo": 1, "non_disponibile": 1},
        )

    def test_il_gestionale_dichiara_le_righe_che_non_sono_prodotti(self) -> None:
        percorso = self.radice / "gestionale.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Foglio1"
        sheet.append(["Tipo", "EAN", "Cod", "Descrizione", "UM", "Colli", "Qta", "Prezzo", "Sconto", "IVA"])
        sheet.append(["C", "8000000000011", "A1", "PRIMO PRODOTTO", "PZ", 2, None, "1,50", None, 22])
        sheet.append(["T", None, None, "TOTALE DI PAGINA", None, None, None, None, None, None])
        sheet.append(["C", "8000000000012", "A2", "SECONDO PRODOTTO", "PZ", 1, None, "2,50", None, 22])
        workbook.save(percorso)
        workbook.close()
        lettura: dict[str, object] = {}

        records = read_gestionale(percorso, report=lettura)

        self.assertEqual(len(records), 2)
        self.assertEqual(lettura["rows_kept"], 2)
        self.assertEqual(lettura["sheet_rows"], 4)
        self.assertEqual(lettura["rows_excluded"], {"non_e_una_riga_prodotto": 2})


class FattoreDOrdineMancanteTests(CasoConCartella):
    """An order factor that is missing or zero must not be treated as one.

    The comparison ranks offers by carton price, which is unit price times
    pieces per carton. If a missing factor defaulted to one, the carton
    would price as a single unit and wrongly win the comparison (the lowest
    price ranks first). If a zero factor were accepted as valid, the order
    total would come out zero, understating the supplier total and letting
    it dodge the minimum-order threshold, since zero is below any threshold.
    """

    MAPPATURA = {
        "header_row": 1,
        "delimiter": ";",
        "italian_numbers": True,
        "columns": {
            "ean": "ean",
            "description": "descrizione",
            "pieces_per_carton": "pezzi",
            "unit_price_net": "prezzo",
        },
    }

    def betulla(self, righe: list[list[object]]) -> Path:
        percorso = self.radice / "betulla.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["EAN", "CodArt", "ORDINE", "Descrizione", "PzCt", "Cessione", "Pedana", "IVA"])
        for riga in righe:
            sheet.append(riga)
        workbook.save(percorso)
        workbook.close()
        return percorso

    def espositore(self, pezzi_del_collo: object) -> dict[str, object]:
        """A minimal display-detection result, with just the fields under test."""

        return {
            "row_classifications": [],
            "rejected_bundle_candidates": [],
            "display_offers": [
                {
                    "offer_id": "LARICE:D-1",
                    "description": "ESPOSITORE MISTO 24 PZ",
                    "supplier_code": "D-1",
                    "pieces_per_carton": pezzi_del_collo,
                    "declared_quantity": 24,
                    "parent_price_pre_discount": "2.0000",
                    "parent_price_post_discount": "1.0000",
                    "source_rows": {"parent_row": 10},
                    "components": [],
                    "confidence": "HIGH",
                }
            ],
        }

    # -- BETULLA --

    def test_betulla_senza_pezzi_per_collo_non_si_puo_ordinare(self) -> None:
        percorso = self.betulla([
            ["8000000000011", "C-1", None, "PRODOTTO BUONO", 6, 2.50, None, 22],
            ["8000000000012", "C-2", None, "PEZZI PER COLLO VUOTI", None, 2.50, None, 22],
            ["8000000000013", "C-3", None, "PEZZI PER COLLO A ZERO", 0, 2.50, None, 22],
        ])
        lettura: dict[str, object] = {}

        records = read_betulla(percorso, report=lettura)

        self.assertEqual([record["usable"] for record in records], [True, False, False])
        self.assertEqual(records[1]["unusable_reason"], "senza_pezzi_per_collo")
        self.assertEqual(records[2]["unusable_reason"], "senza_pezzi_per_collo")
        self.assertEqual(lettura["rows_not_orderable"], {"senza_pezzi_per_collo": 2})

    def test_betulla_con_i_pezzi_per_collo_resta_identico(self) -> None:
        """Regression guard: the fix must not drop rows that already have valid data."""

        percorso = self.betulla([
            ["8000000000011", "C-1", None, "PRODOTTO BUONO", 6, 2.50, None, 22],
            ["8000000000012", "C-2", None, "ALTRO PRODOTTO", 12, "3.75", None, 22],
        ])
        lettura: dict[str, object] = {}

        records = read_betulla(percorso, report=lettura)

        self.assertEqual([record["usable"] for record in records], [True, True])
        self.assertNotIn("unusable_reason", records[0])
        self.assertEqual([record["unit_price_net"] for record in records], ["2.5000", "3.7500"])
        self.assertEqual([record["pieces_per_carton"] for record in records], [6, 12])
        self.assertEqual(lettura["rows_not_orderable"], {})

    # -- Larice display offers --

    def test_l_espositore_col_collo_leggibile_costa_il_collo_intero(self) -> None:
        _righe, offerte, riepilogo = integrate_larice_displays([], self.espositore(24))

        self.assertEqual(offerte[0]["list_price_per_display"], "48.0000")
        self.assertEqual(offerte[0]["net_price_per_display"], "24.0000")
        self.assertTrue(offerte[0]["usable"])
        self.assertEqual(riepilogo["display_offers_not_orderable"], {})

    def test_l_espositore_senza_i_pezzi_del_collo_non_si_puo_ordinare(self) -> None:
        """The display price is never made up: without that number, there is none."""

        for pezzi in (None, 0, "", "a richiesta"):
            with self.subTest(pezzi=pezzi):
                _righe, offerte, riepilogo = integrate_larice_displays([], self.espositore(pezzi))
                offerta = offerte[0]

                self.assertFalse(offerta["usable"])
                self.assertEqual(offerta["unusable_reason"], "senza_pezzi_per_collo")
                self.assertIsNone(offerta["list_price_per_display"])
                self.assertIsNone(offerta["net_price_per_display"])
                # The parent (single-piece) price must not linger: the
                # comparison would fall back to it when the display price is
                # missing, making the offer look unbeatably cheap.
                self.assertIsNone(offerta["parent_price_pre_discount"])
                self.assertIsNone(offerta["parent_price_post_discount"])
                self.assertEqual(riepilogo["display_offers_not_orderable"], {"senza_pezzi_per_collo": 1})

    # -- the fixed value declared in the supplier profile --

    def test_un_valore_fisso_scritto_male_non_diventa_un_fattore(self) -> None:
        """A malformed config value must not crash the read, nor silently
        produce an order factor nobody declared."""

        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO BUONO;6;2,50\n"
            "8000000000012;PEZZI PER COLLO VUOTI;;2,50\n",
            encoding="utf-8",
        )

        for fisso in (0, "0", "a richiesta", ""):
            with self.subTest(valore_fisso=fisso):
                mappatura = {**self.MAPPATURA, "pieces_per_carton_default": fisso}
                lettura: dict[str, object] = {}

                records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura, report=lettura)

                self.assertEqual([record["usable"] for record in records], [True, False])
                self.assertEqual(records[1]["unusable_reason"], "senza_pezzi_per_collo")
                self.assertEqual(lettura["rows_not_orderable"], {"senza_pezzi_per_collo": 1})
                # The row must not report a carton size it doesn't have either:
                # a zero written here reads as a valid number to any later
                # consumer (e.g. `mapped_standalone_displays`), which would make
                # the row orderable through a different path.
                self.assertIsNone(records[1]["pieces_per_carton"])

    def test_un_valore_fisso_buono_vale_ancora_per_le_righe_senza_quel_dato(self) -> None:
        """The other half of the rule: a real number in the right place works."""

        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO BUONO;6;2,50\n"
            "8000000000012;PEZZI PER COLLO VUOTI;;2,50\n",
            encoding="utf-8",
        )
        mappatura = {**self.MAPPATURA, "pieces_per_carton_default": 12}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura)

        self.assertEqual([record["usable"] for record in records], [True, True])
        self.assertEqual([record["pieces_per_carton"] for record in records], ["6.0000", "12.0000"])

    def test_un_moltiplicatore_fisso_scritto_male_non_porta_via_le_righe_buone(self) -> None:
        """A bad fixed value must never override a real value read from the row.

        When present, `order_multiplier` takes priority over pieces per
        carton. If a `0` declared in the profile were accepted as a real
        multiplier, it would override the carton size read from every row,
        including ones where that value is correct, dropping all of them
        from the order.
        """

        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO BUONO;6;2,50\n",
            encoding="utf-8",
        )
        mappatura = {**self.MAPPATURA, "order_multiplier_default": 0}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura)

        self.assertTrue(records[0]["usable"])
        self.assertIsNone(records[0]["order_multiplier"])
        self.assertEqual(records[0]["pieces_per_carton"], "6.0000")

    def test_senza_colonna_e_col_valore_fisso_scritto_male_ci_si_ferma(self) -> None:
        """Silently dropping every row from the order is worse than raising."""

        percorso = self.radice / "listino.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["ean", "descrizione", "prezzo"])
        sheet.append(["8000000000011", "PRODOTTO BUONO", 2.50])
        workbook.save(percorso)
        workbook.close()
        mappatura = {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": {"ean": "ean", "description": "descrizione", "unit_price_net": "prezzo"},
            "pieces_per_carton_default": 0,
        }

        with self.assertRaises(ValueError) as errore:
            read_mapped_xlsx_supplier(percorso, "nuovo", mappatura)

        self.assertIn("non è un numero maggiore di zero", str(errore.exception))


class NumeroDiRigaFisicoTests(CasoConCartella):
    """`source_row` is the physical file row, even when a field wraps onto two lines.

    Whoever checks an order opens the supplier's price list and goes to that
    row. Counting products instead of physical rows breaks that: a single
    quoted, multi-line description shifts every row number after it by one,
    pointing at the wrong product. No CSV price list currently fills written
    orders, so the effect stays inside the comparison; it would become an
    order on the wrong row if one did.
    """

    MAPPATURA = {
        "header_row": 1,
        "delimiter": ";",
        "italian_numbers": True,
        "columns": {
            "ean": "ean",
            "description": "descrizione",
            "pieces_per_carton": "pezzi",
            "unit_price_net": "prezzo",
        },
    }

    def test_un_a_capo_dentro_una_descrizione_non_sposta_le_righe_dopo(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            '8000000000011;"PRODOTTO CON\nIL NOME SU DUE RIGHE";6;2,50\n'
            "8000000000012;PRODOTTO DOPO;6;3,50\n"
            "8000000000013;ULTIMO PRODOTTO;6;4,50\n",
            encoding="utf-8",
        )

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA)

        # Row 1 is the header, rows 2-3 are the first product, then 4 and 5.
        self.assertEqual([record["source_row"] for record in records], [2, 4, 5])
        self.assertEqual(records[1]["description"], "PRODOTTO DOPO")
        self.assertEqual(records[2]["description"], "ULTIMO PRODOTTO")

    def test_senza_a_capo_i_numeri_di_riga_sono_quelli_di_sempre(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRIMO;6;2,50\n"
            "8000000000012;SECONDO;6;3,50\n",
            encoding="utf-8",
        )

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA)

        self.assertEqual([record["source_row"] for record in records], [2, 3])

    def test_lo_stesso_conto_vale_sul_listino_noce(self) -> None:
        percorso = self.radice / "noce.csv"
        percorso.write_text(
            "catalog_page,ean,product,packaging,availability,price,unit\n"
            '1,8000000000011,"PRODOTTO CON\nIL NOME SU DUE RIGHE",CT,Disponibile,2.50,x 6\n'
            "1,8000000000012,PRODOTTO DOPO,CT,Disponibile,3.50,x 6\n"
            "1,8000000000013,ULTIMO PRODOTTO,CT,Disponibile,4.50,x 6\n",
            encoding="utf-8",
        )

        records = read_noce(percorso)

        self.assertEqual([record["source_row"] for record in records], [2, 4, 5])
        self.assertEqual(records[1]["description"], "PRODOTTO DOPO")

    def test_su_noce_un_moltiplicatore_a_zero_non_e_un_moltiplicatore(self) -> None:
        percorso = self.radice / "noce.csv"
        percorso.write_text(
            "catalog_page,ean,product,packaging,availability,price,unit\n"
            '1,8000000000011,PRODOTTO BUONO,CT,Disponibile,"2,50",x 6\n'
            '1,8000000000012,MOLTIPLICATORE A ZERO,CT,Disponibile,"2,50",x 0\n',
            encoding="utf-8",
        )
        lettura: dict[str, object] = {}

        records = read_noce(percorso, report=lettura)

        self.assertEqual([record["usable"] for record in records], [True, False])
        self.assertEqual(records[1]["unusable_reason"], "senza_pezzi_per_collo")
        self.assertEqual(lettura["rows_not_orderable"], {"senza_pezzi_per_collo": 1})


class ListinoLettoAVuotoTests(CasoConCartella):
    """A price list opened with the wrong delimiter is not a supplier with no products.

    When not declared, the CSV delimiter is guessed from the file text; on
    an Italian price list, where every price carries a decimal comma, a
    comma looks like the right guess even when it's the field separator
    instead. Read that way, the file comes out empty, and the supplier shows
    up as "doesn't carry this item" for every item: nothing distinguishes an
    unreadable price list from a supplier that genuinely stocks nothing.
    """

    COLONNE = {
        "ean": "ean",
        "description": "descrizione",
        "pieces_per_carton": "pezzi",
        "unit_price_net": "prezzo",
    }

    def test_il_separatore_indovinato_male_non_produce_un_listino_vuoto(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "8000000000011;PRODOTTO UNO;6;2,50\n"
            "8000000000012;PRODOTTO DUE;6;3,50\n"
            "8000000000013;PRODOTTO TRE;6;4,50\n",
            encoding="utf-8",
        )
        # No header to check and no delimiter declared: without this guard,
        # the file would parse as empty with no warning at all.
        mappatura = {
            "header_row": 0,
            "italian_numbers": True,
            "columns": {"ean": "A", "description": "B", "pieces_per_carton": "C", "unit_price_net": "D"},
        }

        with self.assertRaises(ValueError) as errore:
            read_mapped_csv_supplier(percorso, "nuovo", mappatura)

        messaggio = str(errore.exception)
        self.assertIn("nessuna è risultata ordinabile", messaggio)
        self.assertIn("separatore", messaggio)

    def test_un_separatore_sbagliato_si_riconosce_gia_dall_intestazione(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO UNO;6;2.50\n",
            encoding="utf-8",
        )
        mappatura = {"header_row": 1, "delimiter": ",", "columns": self.COLONNE}

        with self.assertRaises(ValueError) as errore:
            read_mapped_csv_supplier(percorso, "nuovo", mappatura)

        messaggio = str(errore.exception)
        self.assertIn("un campo solo", messaggio)
        self.assertIn("separatore", messaggio)

    def test_col_separatore_giusto_il_listino_si_legge(self) -> None:
        """Control case: the guard must not block a valid read."""

        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO UNO;6;2,50\n"
            "8000000000012;PRODOTTO DUE;6;3,50\n",
            encoding="utf-8",
        )
        mappatura = {
            "header_row": 1,
            "delimiter": ";",
            "italian_numbers": True,
            "columns": self.COLONNE,
        }

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", mappatura)

        self.assertEqual([record["description"] for record in records], ["PRODOTTO UNO", "PRODOTTO DUE"])
        self.assertEqual([record["unit_price_net"] for record in records], ["2.5000", "3.5000"])


class ColonneDiLariceDalRegistroTests(CasoConCartella):
    """The Larice column layout comes from the adapter registry, not hardcoded positions.

    A hardcoded layout keeps reading the old column after Larice moves one,
    or after the user confirms a schema change and the registry learns the
    new mapping. Since this reads prices, the failure isn't missing rows: it's
    a whole price list of wrong prices that looks correct.
    """

    # The columns the registry currently declares.
    COLONNE_DI_OGGI = {
        "supplier_code": "C",
        "pieces_per_carton": "E",
        "pallet": "F",
        "description": "G",
        "unit_price_pre_discount": "O",
        "discount": "P",
        "vat": "Q",
        "ean": "R",
    }

    PRODOTTO = {
        "supplier_code": "12345",
        "pieces_per_carton": 6,
        "pallet": 60,
        "description": "BAGNO VIDOR 500 ML TALCO",
        "unit_price_pre_discount": 2.50,
        "discount": 10,
        "vat": 22,
        "ean": "8000000000011",
    }

    def listino(self, colonne: dict[str, str], riga: dict[str, object]) -> Path:
        """Build a Larice sheet with each value in the column `colonne` names.

        The real file has no header row: the product data starts on row 1.
        """

        percorso = self.radice / "larice.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        for campo, valore in riga.items():
            sheet.cell(row=1, column=column_index_from_string(colonne[campo]), value=valore)
        workbook.save(percorso)
        workbook.close()
        return percorso

    def registro_con(self, colonne: dict[str, str]) -> None:
        """Build a registry copy identical to the real one except for Larice's columns.

        The real registry is reviewed by a person before every commit; a test
        must not be able to touch it, even by accident.
        """

        documento = json.loads(ADAPTERS.read_text(encoding="utf-8"))
        for voce in documento["adapters"]:
            if voce.get("id") == "larice_v1":
                voce["column_map"] = colonne
        percorso = self.radice / "adapters.json"
        percorso.write_text(json.dumps(documento, ensure_ascii=False), encoding="utf-8", newline="")
        originale = registro.REGISTRO
        registro.REGISTRO = percorso
        self.addCleanup(setattr, registro, "REGISTRO", originale)

    def test_sul_listino_di_oggi_i_record_escono_identici(self) -> None:
        """Regression guard: reading from the registry must not change the numbers for a working layout."""

        percorso = self.listino(self.COLONNE_DI_OGGI, self.PRODOTTO)

        records, avvisi = read_larice(percorso)

        self.assertEqual(avvisi, [])
        self.assertEqual(records, [{
            "source": "larice",
            "source_row": 1,
            "ean": "8000000000011",
            "supplier_code": "12345",
            "description": "BAGNO VIDOR 500 ML TALCO",
            "pieces_per_carton": 6,
            "pallet": 60,
            "unit_price_pre_discount": "2.5000",
            "discount_raw": 10,
            "discount_type": "percentuale",
            "discount_rate": "0.1000",
            "unit_price_net": "2.2500",
            "vat": 22,
            "order_column": "D",
            "usable": True,
        }])

    def test_le_colonne_spostate_nel_registro_spostano_anche_il_lettore(self) -> None:
        spostate = {
            **self.COLONNE_DI_OGGI,
            "description": "H",
            "unit_price_pre_discount": "S",
            "ean": "T",
        }
        percorso = self.listino(spostate, self.PRODOTTO)
        self.registro_con(spostate)

        records, _avvisi = read_larice(percorso)

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["ean"], "8000000000011")
        self.assertEqual(records[0]["description"], "BAGNO VIDOR 500 ML TALCO")
        self.assertEqual(records[0]["unit_price_net"], "2.2500")
        self.assertTrue(records[0]["usable"])

    def test_una_colonna_che_il_registro_non_dichiara_ferma_la_lettura(self) -> None:
        """Raising costs an error message; guessing costs a wrong price list."""

        senza = {
            campo: dove for campo, dove in self.COLONNE_DI_OGGI.items()
            if campo not in {"unit_price_pre_discount", "ean"}
        }
        percorso = self.listino(self.COLONNE_DI_OGGI, self.PRODOTTO)
        self.registro_con(senza)

        with self.assertRaises(ValueError) as errore:
            read_larice(percorso)

        messaggio = str(errore.exception)
        self.assertIn("i prezzi e il codice a barre", messaggio)
        self.assertIn("larice.xlsx", messaggio)


class NumeriCheNonSonoNumeriTests(CasoConCartella):
    """NaN and infinity are not prices.

    A spreadsheet can produce them on its own, e.g. from a division by zero
    carried over from a formula. Left unguarded, they multiply and sum like
    any other number without raising, leaving an order total that is not a
    number and that nobody notices.
    """

    MAPPATURA = {
        "header_row": 1,
        "delimiter": ";",
        "italian_numbers": True,
        "columns": {
            "ean": "ean",
            "description": "descrizione",
            "pieces_per_carton": "pezzi",
            "unit_price_net": "prezzo",
        },
    }

    def test_i_valori_con_cui_non_si_fa_aritmetica_valgono_come_non_letti(self) -> None:
        for valore in ("NaN", "nan", "Infinity", "-Infinity", float("nan"), float("inf"), Decimal("NaN")):
            with self.subTest(valore=valore):
                self.assertIsNone(decimal_value(valore))
                self.assertIsNone(decimal_value(valore, italian=True))

    def test_i_numeri_veri_si_leggono_ancora(self) -> None:
        self.assertEqual(decimal_value("2,50", italian=True), Decimal("2.50"))
        self.assertEqual(decimal_value("1.234,50", italian=True), Decimal("1234.50"))
        self.assertEqual(decimal_value(6), Decimal("6"))
        self.assertEqual(decimal_value(Decimal("1.25")), Decimal("1.25"))
        self.assertEqual(decimal_value(0), Decimal("0"))

    def test_un_prezzo_che_non_e_un_numero_non_diventa_una_riga_ordinabile(self) -> None:
        percorso = self.radice / "listino.csv"
        percorso.write_text(
            "ean;descrizione;pezzi;prezzo\n"
            "8000000000011;PRODOTTO BUONO;6;2,50\n"
            "8000000000012;PREZZO CHE NON E UN NUMERO;6;NaN\n",
            encoding="utf-8",
        )
        lettura: dict[str, object] = {}

        records, _avvisi = read_mapped_csv_supplier(percorso, "nuovo", self.MAPPATURA, report=lettura)

        self.assertEqual([record["usable"] for record in records], [True, False])
        self.assertEqual(records[1]["unusable_reason"], "senza_prezzo")
        self.assertIsNone(records[1]["unit_price_net"])
        self.assertEqual(lettura["rows_not_orderable"], {"senza_prezzo": 1})


class EspositoriIsolatiTests(CasoConCartella):
    """A single-row display offer inherits the usability of its own row, nothing more.

    A supplier can sell a pre-assembled display on a single row, with no
    component rows behind it. Without this check, such a row would carry no
    `usable` verdict at all, and its price would be taken as written even
    when the order factor doesn't parse.
    """

    def mappatura(self, colonne: dict[str, str], **regole: object) -> dict[str, object]:
        return {
            "sheet": "FIRST",
            "header_row": 1,
            "data_start_row": 2,
            "columns": colonne,
            "assume_available": True,
            "display_detection": {
                "description_regex": r"\bEXPO\b",
                "supplier_code_regex": r"^E\d{6}$",
                "confidence": "MEDIA",
                **regole,
            },
        }

    def listino(self, righe: list[list[object]]) -> Path:
        percorso = self.radice / "cipresso.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["COD.ART.", "DES.ARTICOLO", "QT", "LISTINO", "COD.EAN"])
        for riga in righe:
            sheet.append(riga)
        workbook.save(percorso)
        workbook.close()
        return percorso

    def espositori(self, righe: list[list[object]], colonne: dict[str, str], **regole: object):
        mappatura = self.mappatura(colonne, **regole)
        percorso = self.listino(righe)
        records, _avvisi = read_mapped_xlsx_supplier(percorso, "cipresso", mappatura)
        return mapped_standalone_displays(records, "cipresso", mappatura)

    COLONNE_A_COLLO = {
        "supplier_code": "COD.ART.",
        "description": "DES.ARTICOLO",
        "pieces_per_carton": "QT",
        "unit_price_net": "LISTINO",
        "ean": "COD.EAN",
    }
    COLONNE_A_MOLTIPLICATORE = {
        "supplier_code": "COD.ART.",
        "description": "DES.ARTICOLO",
        "order_multiplier": "QT",
        "unit_price_net": "LISTINO",
        "ean": "COD.EAN",
    }
    NORMALE = ["A1", "PRODOTTO NORMALE", 6, 2.50, "8000000000012"]

    def test_un_espositore_sano_resta_ordinabile(self) -> None:
        _standard, espositori, audit = self.espositori(
            [["E000001", "EXPO MISTO 120 PZ + 24 PZ", 1, 12.50, "8000000000011"], self.NORMALE],
            self.COLONNE_A_COLLO,
            order_factor_equals=1,
        )

        self.assertEqual(len(espositori), 1)
        self.assertTrue(espositori[0]["usable"])
        self.assertNotIn("unusable_reason", espositori[0])
        self.assertEqual(espositori[0]["net_price_per_display"], "12.5000")
        self.assertEqual(audit["display_offers_not_orderable"], {})

    def test_un_espositore_senza_fattore_d_ordine_non_si_puo_ordinare(self) -> None:
        _standard, espositori, audit = self.espositori(
            [["E000001", "EXPO MISTO", None, 12.50, "8000000000011"], self.NORMALE],
            self.COLONNE_A_COLLO,
        )

        self.assertEqual(len(espositori), 1)
        self.assertFalse(espositori[0]["usable"])
        self.assertEqual(espositori[0]["unusable_reason"], "senza_pezzi_per_collo")
        self.assertEqual(audit["display_offers_not_orderable"], {"senza_pezzi_per_collo": 1})

    def test_un_espositore_senza_prezzo_porta_il_motivo_della_sua_riga(self) -> None:
        _standard, espositori, audit = self.espositori(
            [["E000001", "EXPO MISTO", 1, None, "8000000000011"], self.NORMALE],
            self.COLONNE_A_COLLO,
            order_factor_equals=1,
        )

        self.assertFalse(espositori[0]["usable"])
        self.assertEqual(espositori[0]["unusable_reason"], "senza_prezzo")
        self.assertEqual(audit["display_offers_not_orderable"], {"senza_prezzo": 1})

    def test_il_fattore_puo_venire_dal_moltiplicatore_come_per_le_altre_righe(self) -> None:
        """A row that declares `order_multiplier` instead of pieces-per-carton
        must not lose its display offer: it's the same row, and it's orderable."""

        _standard, espositori, _audit = self.espositori(
            [["E000001", "EXPO MISTO", 1, 12.50, "8000000000011"], self.NORMALE],
            self.COLONNE_A_MOLTIPLICATORE,
            order_factor_equals=1,
        )

        self.assertEqual(len(espositori), 1)
        self.assertTrue(espositori[0]["usable"])

    def test_una_riga_che_si_dichiara_sana_senza_fattore_non_diventa_ordinabile(self) -> None:
        """A defense-in-depth case a spreadsheet-driven test can't reach directly.

        The readers already set `usable: False` themselves when the order
        factor is missing, so no spreadsheet input can exercise this guard on
        its own; it protects `mapped_standalone_displays` as a public
        function, also called by `catalog_search`. A record that claims to be
        usable without saying how many pieces come per carton must still not
        produce an orderable display: the display price is the per-piece
        price times that count, and without it the offer would look as cheap
        as a single piece and win the comparison.
        """

        mappatura = self.mappatura(self.COLONNE_A_COLLO)
        record = {
            "source": "cipresso",
            "source_row": 2,
            "supplier_code": "E000001",
            "description": "EXPO MISTO 120 PZ",
            "ean": "8000000000011",
            "unit_price_net": "12.5000",
            "pieces_per_carton": None,
            "order_multiplier": None,
            "usable": True,
        }

        _standard, espositori, audit = mapped_standalone_displays([record], "cipresso", mappatura)

        self.assertEqual(len(espositori), 1)
        self.assertFalse(espositori[0]["usable"])
        self.assertEqual(espositori[0]["unusable_reason"], "senza_pezzi_per_collo")
        self.assertEqual(audit["display_offers_not_orderable"], {"senza_pezzi_per_collo": 1})


if __name__ == "__main__":
    unittest.main()
