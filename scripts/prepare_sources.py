#!/usr/bin/env python3
"""Read every active source row and produce deterministic EAN matching artifacts."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

import registro
from detect_displays import analyse_workbook


# ---------------------------------------------------------------------------
# Why a read row is excluded from the comparison, declared explicitly rather
# than inferred. `usable` alone is a plain conjunction: it says the row is
# out but never what it was missing, and a drop that can't be counted isn't a
# choice — it can otherwise look like zero unorderable rows while half a
# price list was actually dropped. Defined here because both the dedicated
# readers below and the mapped readers in `prepare_manifest_sources` (which
# imports from this module) use them.
# ---------------------------------------------------------------------------
SENZA_DESCRIZIONE = "senza_descrizione"
SENZA_PREZZO = "senza_prezzo"
SENZA_PEZZI_PER_COLLO = "senza_pezzi_per_collo"
NON_DISPONIBILE = "non_disponibile"
MOTIVO_NON_DICHIARATO = "motivo_non_dichiarato"
# Rows skipped before they ever become a record: they have no `usable` field
# to check, so they need their own counters to show up in the totals at all.
SENZA_EAN = "senza_ean"
NON_E_RIGA_PRODOTTO = "non_e_una_riga_prodotto"
FUORI_DAL_FILTRO = "riga_fuori_dal_filtro"


def conta_non_ordinabili(records: list[dict[str, Any]]) -> dict[str, int]:
    """Count read rows that can't be ordered, grouped by declared reason.

    The registry's `row_type` wins over the technical reason: a free-goods
    threshold row isn't "no price", it's a threshold row.
    """

    return dict(Counter(
        str(record.get("row_type") or record.get("unusable_reason") or MOTIVO_NON_DICHIARATO)
        for record in records
        if isinstance(record, dict) and record.get("usable") is False
    ))


def rapporto_di_lettura(
    *,
    righe_lette: int,
    records: list[dict[str, Any]],
    saltate: Counter[str] | None = None,
) -> dict[str, Any]:
    """Summarize a read of a known-schema price list.

    Same keys as `prepare_manifest_sources.reading_report`, since the page
    and the orchestrator only understand one report shape.
    """

    saltate = saltate or Counter()
    return {
        "sheet_rows": righe_lette,
        "rows_excluded": dict(saltate),
        "rows_kept": len(records),
        "rows_not_orderable": conta_non_ordinabili(records),
    }


def normalize_ean(value: Any) -> str:
    """Normalize only safe spreadsheet artifacts; never search orphan shared strings."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text


def decimal_value(value: Any, *, italian: bool = False) -> Decimal | None:
    """Return the number inside a cell, or None.

    NaN and infinity are values a spreadsheet can produce on its own (a
    division by zero carried through a formula chain), and no arithmetic
    should be done with them: they're treated as unread data. Letting them
    through would be worse than stopping — they don't raise, they multiply
    and sum like ordinary numbers, and a "NaN" total at the bottom of an
    order goes unnoticed.

    With `italian` on, the dot is the thousands separator, the convention
    most of the shipped supplier lists use. But text with exactly one dot, no
    comma, and one or two trailing digits isn't written that way: "1.25"
    reads as one euro twenty-five, and stripping the dot would silently turn
    it into 125 euros. That case is treated as unreadable and the row is
    excluded and counted as SENZA_PREZZO: a missing product is visible, a
    fabricated price isn't. "1.250" stays ambiguous and is read as one
    thousand two hundred fifty.

    Cells that are already numeric in the spreadsheet (not text) bypass this
    parsing entirely and are returned unchanged.
    """

    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        numero = value
    elif isinstance(value, (int, float)):
        numero = Decimal(str(value))
    else:
        text = str(value).strip()
        if italian:
            if re.fullmatch(r"[-+]?\d+\.\d{1,2}", text):
                return None
            text = text.replace(".", "").replace(",", ".")
        try:
            numero = Decimal(text)
        except InvalidOperation:
            return None
    return numero if numero.is_finite() else None


def fattore_d_ordine(value: Any, *, italian: bool = False) -> Decimal | None:
    """Return how many pieces make up an order unit, only if the number is real.

    A missing, zero, or unreadable factor is never coerced to one. The order
    total is the piece price times this number: defaulting to one to avoid
    stopping would price a whole carton as a single piece, making it look
    like the cheapest offer and win the comparison on a fabricated price. A
    row missing this number is still read, but marked unorderable.

    Anything that isn't a number at all is already rejected by
    `decimal_value`; this function only adds that zero and negative values
    don't count as an order factor either.
    """

    try:
        numero = decimal_value(value, italian=italian)
    except (InvalidOperation, TypeError, ValueError):
        # A bool falls into the numeric branch and fails conversion there:
        # the one case where `decimal_value` raises instead of returning None.
        return None
    if numero is None or numero <= 0:
        return None
    return numero


def json_decimal(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return format(value.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP), "f")


def active_rows(path: Path, sheet_name: str | None = None) -> tuple[Any, Iterable[tuple[Any, ...]]]:
    """Read every row of the sheet and close the workbook before returning.

    Returning openpyxl's lazy row generator instead would keep the file open
    for the process's whole lifetime; on Windows that file could then be
    neither deleted nor replaced, and `app/catalog_search.py` calls these
    readers inside the long-running server process, not a subprocess that
    would release the handle on exit.
    """

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        return sheet, list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


# Default column positions for the two dedicated readers below, used when
# nothing overrides them. A confirmed mapping can override them field by
# field, so a column correction made on the page reaches these readers
# without routing the document to the generic reader, which for some
# suppliers would mean losing display detection and offer handling.
COLONNE_PREDEFINITE_GESTIONALE = {
    "ean": 2,
    "description": 4,
    "unit": 5,
    "suggested_colli": 6,
    "source_quantity_ignored": 7,
    "last_unit_price": 8,
    "source_discount": 9,
    "vat": 10,
}
COLONNE_PREDEFINITE_BETULLA = {
    "ean": 1,
    "supplier_code": 2,
    "description": 4,
    "pieces_per_carton": 5,
    "unit_price_net": 6,
    "pallet": 7,
    "vat": 8,
}


def colonne_con_predefinite(
    predefinite: dict[str, int], scelte: dict[str, Any] | None
) -> dict[str, int]:
    """Return the columns to use: defaults, overridden by any manual choices.

    A choice that doesn't resolve to a valid column number is never ignored:
    silently keeping the default in that case would read the wrong column
    with nothing to show for it. It raises instead.
    """

    risultato = dict(predefinite)
    for campo, dichiarata in (scelte or {}).items():
        if campo not in predefinite:
            continue
        indice = indice_di_colonna(dichiarata)
        if indice is None:
            raise ValueError(
                f"La colonna indicata per «{campo}» ({dichiarata!r}) non è una colonna valida."
            )
        risultato[campo] = indice
    return risultato


def intestazioni_obbligatorie(adapter_id: str) -> list[str]:
    """Return the headers the registry declares required for an adapter.

    Already normalized: the registry writes them that way, and that's how
    `riconosci` compares them. A missing registry or entry returns an empty
    list and the caller requires nothing — same choice as
    `registro.adattatore`, since an unreadable registry must not block
    reading a price list.
    """

    firma = (registro.adattatore(adapter_id) or {}).get("header_signature")
    richieste = (firma or {}).get("required") if isinstance(firma, dict) else None
    if not isinstance(richieste, list):
        return []
    return [registro.normalizza(nome) for nome in richieste if registro.normalizza(nome)]


def pretendi_le_intestazioni(
    adapter_id: str, valori: Iterable[Any], path: Path, etichetta: str
) -> dict[str, Any]:
    """Check that the header row carries what the registry requires.

    Returns a map of normalized header -> value as written, used by readers
    that address columns by name (the CSV reader) and ignored by readers
    that address columns by position.

    The required list comes from the registry and the comparison goes
    through `registro.normalizza`, the same function used to recognize the
    document a moment earlier. Both sides must use the one function: two
    separate definitions of "the required headers" (one here, one in the
    registry) can drift apart the first time a supplier's casing changes,
    rejecting a perfectly good file with a message that doesn't even say
    what was wrong.

    This check itself is not redundant with recognition: a dedicated reader
    addresses columns by position, and not every caller goes through
    recognition first — a mapping written by hand can enable it without
    anyone having looked at the headers. What changed is that it now checks
    against what the registry declares, not a separate hard-coded list.

    And it reports what exactly is missing, rather than a bare "schema not
    recognized" that gives no hint of the actual mismatch.
    """

    presenti: dict[str, Any] = {}
    for valore in valori:
        if valore is None:
            continue
        chiave = registro.normalizza(valore)
        if chiave and chiave not in presenti:
            presenti[chiave] = valore
    mancanti = [nome for nome in intestazioni_obbligatorie(adapter_id) if nome not in presenti]
    if mancanti:
        lette = ", ".join(str(valore).strip() for valore in presenti.values()) or "nessuna"
        raise ValueError(
            f"Schema {etichetta} non riconosciuto in {path}: nella riga d'intestazione "
            f"mancano {', '.join(mancanti)}. Le intestazioni lette sono: {lette}."
        )
    return presenti


def read_gestionale(
    path: Path,
    report: dict[str, Any] | None = None,
    colonne: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    posizioni = colonne_con_predefinite(COLONNE_PREDEFINITE_GESTIONALE, colonne)
    workbook = load_workbook(path, read_only=True, data_only=True)
    candidate_names = ["Foglio1", *[name for name in workbook.sheetnames if name != "Foglio1"]]
    sheet = None
    for name in candidate_names:
        if name not in workbook.sheetnames:
            continue
        candidate = workbook[name]
        if any(str(candidate.cell(row, 1).value or "").strip().upper() == "C" for row in range(1, min(candidate.max_row, 200) + 1)):
            sheet = candidate
            break
    if sheet is None:
        workbook.close()
        raise ValueError(f"Schema gestionale non riconosciuto in {path}")
    # Same as `active_rows`: read everything, then close, or the management
    # export stays locked by the process that read it.
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    records = []
    righe_lette = 0
    saltate: Counter[str] = Counter()
    for row_number, row in enumerate(rows, start=1):
        righe_lette += 1
        if str(row[0] if len(row) > 0 else "").strip().upper() != "C":
            saltate[NON_E_RIGA_PRODOTTO] += 1
            continue
        records.append(
            {
                "source": "gestionale",
                "source_row": row_number,
                "ean": normalize_ean(cella(row, posizioni["ean"])),
                "description": str(cella(row, posizioni["description"]) or "").strip(),
                "unit": str(cella(row, posizioni["unit"]) or "").strip(),
                "suggested_colli": cella(row, posizioni["suggested_colli"]),
                "source_quantity_ignored": cella(row, posizioni["source_quantity_ignored"]),
                "manual_quantity": None,
                "last_unit_price": json_decimal(decimal_value(cella(row, posizioni["last_unit_price"]), italian=True)),
                "source_discount": cella(row, posizioni["source_discount"]),
                "vat": cella(row, posizioni["vat"]),
            }
        )
    if not records:
        raise ValueError(f"Nessuna riga prodotto gestionale trovata in {path}")
    if report is not None:
        report.update(rapporto_di_lettura(righe_lette=righe_lette, records=records, saltate=saltate))
    return records


def read_betulla(
    path: Path,
    report: dict[str, Any] | None = None,
    colonne: dict[str, Any] | None = None,
    order_column: str | None = None,
) -> list[dict[str, Any]]:
    posizioni = colonne_con_predefinite(COLONNE_PREDEFINITE_BETULLA, colonne)
    sheet, rows = active_rows(path)
    iterator = iter(rows)
    header = next(iterator, ())
    pretendi_le_intestazioni("betulla_v1", header, path, "BETULLA")
    records = []
    righe_lette = 0
    saltate: Counter[str] = Counter()
    for row_number, row in enumerate(iterator, start=2):
        righe_lette += 1
        current_ean = normalize_ean(cella(row, posizioni["ean"]))
        if not current_ean:
            saltate[SENZA_EAN] += 1
            continue
        prezzo = decimal_value(cella(row, posizioni["unit_price_net"]))
        # This supplier's price is per piece; the carton price depends on
        # pieces-per-carton. A row without that number is unorderable
        # (`SENZA_PEZZI_PER_COLLO`): left orderable, its carton would be
        # priced at zero and would also stay below any minimum-order threshold.
        pezzi = fattore_d_ordine(cella(row, posizioni["pieces_per_carton"]))
        motivo = SENZA_PREZZO if prezzo is None else (SENZA_PEZZI_PER_COLLO if pezzi is None else None)
        records.append(
            {
                "source": "betulla",
                "source_row": row_number,
                "ean": current_ean,
                "supplier_code": cella(row, posizioni["supplier_code"]),
                "description": str(cella(row, posizioni["description"]) or "").strip(),
                "pieces_per_carton": cella(row, posizioni["pieces_per_carton"]),
                "unit_price_net": json_decimal(prezzo),
                "pallet": cella(row, posizioni["pallet"]),
                "vat": cella(row, posizioni["vat"]),
                "order_column": str(order_column or "C"),
                "usable": motivo is None,
                # The reason field is only set when there is one: always
                # including it would read as "something's wrong" even on
                # healthy rows.
                **({} if motivo is None else {"unusable_reason": motivo}),
            }
        )
    if report is not None:
        report.update(rapporto_di_lettura(righe_lette=righe_lette, records=records, saltate=saltate))
    return records


def larice_discount(value: Any, ammessi: set[str] | None = None) -> tuple[Decimal, str, str | None]:
    """Parse a discount cell that's sometimes a number and sometimes a text code.

    Which codes are legitimate is declared by the adapter registry, not this
    function: a list hard-coded here would only be valid for today's price
    list and no one would update it when a new code appears.
    """

    if isinstance(value, (int, float, Decimal)):
        discount = Decimal(str(value))
        if discount > 1 and discount <= 100:
            discount /= 100
        if discount < 0 or discount > 1:
            raise ValueError(f"Percentuale sconto Larice non valida: {value}")
        return discount, "percentuale", None
    text = str(value or "").strip()
    noti = ammessi if ammessi is not None else set()
    warning = None if not text or text.upper() in noti else f"Codice sconto testuale inatteso: {text}"
    return Decimal("0"), "testo_nessuno_sconto", warning


# The columns this dedicated reader actually uses, with the name a user
# would recognize: the registry declares all of them, and when one is
# missing, reading stops and says which, instead of falling back to
# yesterday's column.
COLONNE_DEL_LETTORE_LARICE = {
    "supplier_code": "il codice dell'articolo",
    "pieces_per_carton": "i pezzi per collo",
    "pallet": "la pedana",
    "description": "le descrizioni",
    "unit_price_pre_discount": "i prezzi",
    "discount": "lo sconto",
    "vat": "l'IVA",
    "ean": "il codice a barre",
}


def indice_di_colonna(dichiarata: Any) -> int | None:
    """Resolve a registry-declared column to a 1-based number (A = 1).

    This supplier's price list has no header row: a column can only be
    addressed by letter ("R") or number (18), there's nothing to look up by
    name. A digit string isn't treated as a number either: elsewhere that
    would be a column header, and interpreting it as one here would silently
    read the wrong column.
    """

    if isinstance(dichiarata, bool):
        return None
    if isinstance(dichiarata, int):
        return dichiarata if dichiarata >= 1 else None
    lettere = str(dichiarata or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", lettere):
        return None
    indice = 0
    for lettera in lettere:
        indice = indice * 26 + (ord(lettera) - ord("A") + 1)
    return indice


def elenco_in_italiano(voci: list[str]) -> str:
    """Join items as a natural-language Italian list, not a bracketed one."""

    if len(voci) < 2:
        return voci[0] if voci else ""
    return ", ".join(voci[:-1]) + " e " + voci[-1]


def colonne_di_larice(adattatore: dict[str, Any], path: Path) -> dict[str, int]:
    """Return where the registry says this supplier's price list columns are.

    If the positions were hard-coded here, the day the supplier moved one —
    or the user confirmed a schema variation and the registry learned a new
    mapping — this reader would keep reading the old column. This reader
    handles prices: the failure mode isn't a few missing rows, it's an
    entire price list of wrong prices that look plausible.

    A column the registry doesn't declare is never guessed: it stops and
    says which one is missing.
    """

    colonne: dict[str, int] = {}
    mancanti: list[str] = []
    for campo, nome_per_l_utente in COLONNE_DEL_LETTORE_LARICE.items():
        indice = indice_di_colonna(registro.posizione_del_campo(adattatore, {}, campo))
        if indice is None:
            mancanti.append(nome_per_l_utente)
        else:
            colonne[campo] = indice
    if mancanti:
        raise ValueError(
            f"Il listino LARICE «{path.name}» non è stato letto: non so più dove il listino "
            f"tiene {elenco_in_italiano(mancanti)}. Va detto di nuovo dove stanno quelle colonne "
            "prima di rifare il confronto."
        )
    return colonne


def cella(row: tuple[Any, ...], colonna: int) -> Any:
    """Return the cell value at a 1-based column index, as the registry addresses it (A = 1)."""

    return row[colonna - 1] if len(row) >= colonna else None


def read_larice(
    path: Path,
    report: dict[str, Any] | None = None,
    colonne: dict[str, Any] | None = None,
    order_column: str | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    adattatore = registro.adattatore("larice_v1")
    codici = registro.codici_di_riga(adattatore)
    ammessi = registro.codici_ammessi(codici)
    # The registry columns are the rule; a mapping confirmed on the page
    # overrides them without switching readers — the generic reader would
    # lose display detection and free-goods thresholds for this supplier.
    colonne = colonne_con_predefinite(colonne_di_larice(adattatore, path), colonne)
    sheet, rows = active_rows(path)
    records = []
    warnings = []
    righe_lette = 0
    saltate: Counter[str] = Counter()
    for row_number, row in enumerate(rows, start=1):
        righe_lette += 1
        current_ean = normalize_ean(cella(row, colonne["ean"]))
        if not current_ean or current_ean.upper() in {"EAN", "#N/A"}:
            saltate[SENZA_EAN] += 1
            continue
        pre_price = decimal_value(cella(row, colonne["unit_price_pre_discount"]))
        sconto_grezzo = cella(row, colonne["discount"])
        discount, discount_type, warning = larice_discount(sconto_grezzo, ammessi)
        post_price = pre_price * (Decimal("1") - discount) if pre_price is not None else None
        if warning:
            warnings.append({"source": "larice", "source_row": row_number, "ean": current_ean, "warning": warning})
        descrizione = str(cella(row, colonne["description"]) or "").strip()
        motivo = SENZA_DESCRIZIONE if not descrizione else (SENZA_PREZZO if post_price is None else None)
        records.append(
            {
                "source": "larice",
                "source_row": row_number,
                "ean": current_ean,
                "supplier_code": cella(row, colonne["supplier_code"]),
                "description": descrizione,
                "pieces_per_carton": cella(row, colonne["pieces_per_carton"]),
                "pallet": cella(row, colonne["pallet"]),
                "unit_price_pre_discount": json_decimal(pre_price),
                "discount_raw": sconto_grezzo,
                "discount_type": discount_type,
                "discount_rate": json_decimal(discount),
                "unit_price_net": json_decimal(post_price),
                "vat": cella(row, colonne["vat"]),
                "order_column": str(order_column or "D"),
                "usable": motivo is None,
                **({} if motivo is None else {"unusable_reason": motivo}),
            }
        )
    # Rows the registry declares non-purchasable (free-goods threshold
    # rewards) are marked unusable here for what they are, not because they
    # happen to lack a price.
    registro.applica_codici_di_riga(records, codici)
    if report is not None:
        report.update(rapporto_di_lettura(righe_lette=righe_lette, records=records, saltate=saltate))
    return records, warnings


def integrate_larice_displays(
    records: list[dict[str, Any]],
    analysis: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Mark component rows non-orderable and normalize detected parent offers."""
    classifications = {item["source_row"]: item for item in analysis.get("row_classifications", [])}
    component_lookup: dict[int, dict[str, Any]] = {}
    for bundle in [*analysis.get("display_offers", []), *analysis.get("rejected_bundle_candidates", [])]:
        for component in bundle.get("components", []):
            component_lookup[int(component["source_row"])] = component

    for record in records:
        classification = classifications.get(record.get("source_row"))
        if not classification or classification.get("row_type") != "COMPONENT":
            continue
        record["row_type"] = "DISPLAY_COMPONENT" if classification.get("display_detected") else "BUNDLE_COMPONENT"
        record["usable"] = False
        record["display_offer_id"] = classification.get("display_offer_id")
        record["display_parent_row"] = classification.get("parent_row")
        component = component_lookup.get(int(record["source_row"]))
        if component:
            record["component_description"] = component.get("description")
            record["component_quantity"] = component.get("quantity")

    offers = []
    for raw in analysis.get("display_offers", []):
        offer = dict(raw)
        # A display's price is the piece price times the parent carton's
        # pieces-per-carton. If that number can't be read, the display has
        # no real price: defaulting it to one would price it like a single
        # piece, so it would always win the comparison.
        parent_pack = fattore_d_ordine(raw.get("pieces_per_carton"))
        parent_pre = decimal_value(raw.get("parent_price_pre_discount")) if parent_pack is not None else None
        parent_post = decimal_value(raw.get("parent_price_post_discount")) if parent_pack is not None else None
        if parent_pack is None:
            # Also clear the piece price: downstream, a missing display
            # price falls back to it, which would make the offer look cheap
            # through the back door.
            offer["parent_price_pre_discount"] = None
            offer["parent_price_post_discount"] = None
        offer.update(
            {
                "supplier": "larice",
                "supplier_id": "larice",
                "source_row": raw.get("source_rows", {}).get("parent_row"),
                "declared_units": raw.get("declared_quantity"),
                "list_price_per_display": json_decimal(parent_pre * parent_pack if parent_pre is not None else None),
                "net_price_per_display": json_decimal(parent_post * parent_pack if parent_post is not None else None),
                "quantity_reconciled": (
                    raw.get("sum_pieces") == raw.get("declared_quantity")
                    if raw.get("sum_pieces_complete") and raw.get("declared_quantity") is not None
                    else None
                ),
                "price_reconciled": bool(raw.get("component_parent_price_match_basis")) if raw.get("component_price_coverage") else None,
                "order_column": "D",
                "order_multiplier": "1.0000",
                "usable": parent_pack is not None,
                **({} if parent_pack is not None else {"unusable_reason": SENZA_PEZZI_PER_COLLO}),
            }
        )
        offers.append(offer)

    summary = {
        "rows_scanned": analysis.get("rows_scanned"),
        "active_rows": analysis.get("active_rows"),
        "classification_counts": analysis.get("classification_counts"),
        "display_offer_count": len(offers),
        # A display left out of the order is counted like any other dropped
        # row, with the same vocabulary of reasons.
        "display_offers_not_orderable": conta_non_ordinabili(offers),
        "confidence_counts": dict(Counter(offer.get("confidence") for offer in offers)),
        "component_rows_excluded_from_orderable": len(analysis.get("component_rows_excluded_from_orderable", [])),
        "rejected_bundle_candidates": len(analysis.get("rejected_bundle_candidates", [])),
    }
    return records, offers, summary


def read_noce(path: Path, report: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    records = []
    righe_lette = 0
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        # Columns are resolved by normalized name, then read through this
        # map. Making only the header check tolerant would be worse than the
        # failure it avoids: the check would pass, every `row.get("ean")`
        # would return None, and the price list would read as zero usable
        # rows with nothing to say why. A supplier vanishing silently is the
        # worst failure this program can have.
        colonne_csv = pretendi_le_intestazioni(
            "noce_csv_v1", reader.fieldnames or (), path, "CSV Noce"
        )

        def campo(riga: dict[str, Any], nome: str) -> Any:
            """Return a declared column's value, however the supplier wrote its header.

            Returns None when the column is absent: "variation" only appears
            when this supplier flags a variation, and the registry doesn't
            declare it required.
            """

            chiave = colonne_csv.get(registro.normalizza(nome))
            return riga.get(chiave) if chiave is not None else None

        # The row number recorded is the file's real line number, not the
        # product's position in the list: a description containing a
        # newline (legal inside a quoted CSV field) is enough to make the
        # two diverge, shifting every following row. Whoever checks the
        # order against the supplier's file would land on the wrong line.
        # `line_num` counts file lines and, after reading a record, points
        # at its last line; the first line is the one right after the
        # previous record ended, tracked separately here.
        prima_riga_del_prodotto = reader.line_num + 1
        for row in reader:
            row_number = prima_riga_del_prodotto
            prima_riga_del_prodotto = reader.line_num + 1
            righe_lette += 1
            unit_text = str(campo(row, "unit") or "").strip()
            unit_match = re.fullmatch(r"x\s*([0-9]+(?:[.,][0-9]+)?)", unit_text, flags=re.IGNORECASE)
            # "x 0" parses as a multiplier but isn't a valid one.
            order_multiplier = fattore_d_ordine(unit_match.group(1), italian=True) if unit_match else None
            availability = str(campo(row, "availability") or "").strip()
            is_available = availability.casefold().startswith("disponibile")
            price = decimal_value(campo(row, "price"), italian=True)
            records.append(
                {
                    "source": "noce",
                    "source_row": row_number,
                    "catalog_page": campo(row, "catalog_page"),
                    "ean": normalize_ean(campo(row, "ean")),
                    "description": str(campo(row, "product") or "").strip(),
                    "packaging": campo(row, "packaging"),
                    "availability": availability,
                    "variation": campo(row, "variation"),
                    "unit_price_net": json_decimal(price),
                    "unit": unit_text,
                    "order_multiplier": json_decimal(order_multiplier),
                    # Noce appends promotion terms to the availability text.
                    # "Disponibile ACQUISTA ..." is still an orderable item and
                    # must not be discarded merely because an offer follows it.
                    "usable": price is not None and order_multiplier is not None and is_available,
                    **(
                        {}
                        if price is not None and order_multiplier is not None and is_available
                        else {"unusable_reason": (
                            SENZA_PREZZO if price is None
                            else SENZA_PEZZI_PER_COLLO if order_multiplier is None
                            else NON_DISPONIBILE
                        )}
                    ),
                }
            )
    if report is not None:
        report.update(rapporto_di_lettura(righe_lette=righe_lette, records=records))
    return records


def source_stats(records: list[dict[str, Any]]) -> dict[str, Any]:
    eans = [record["ean"] for record in records if record.get("ean")]
    counts = Counter(eans)
    return {
        "rows": len(records),
        "rows_with_ean": len(eans),
        "distinct_eans": len(counts),
        "duplicate_ean_values": sum(1 for count in counts.values() if count > 1),
    }


def price_stats(records: list[dict[str, Any]]) -> dict[str, Any]:
    """Return a price list's typical per-piece price, to compare against the previous run.

    Feeds an orchestrator check that warns without stopping the run: a price
    list where every price shifted at once — a supplier switching net/gross
    pricing, a misread column — still has the usual row count, so nothing
    else would catch it.

    Uses the median, not the mean: a handful of thousand-euro items would be
    enough to skew the mean of a list with thousands of rows and fire the
    warning every week.
    """

    prezzi = sorted(
        valore
        for valore in (decimal_value(record.get("unit_price_net"), italian=False) for record in records)
        if valore is not None and valore > 0
    )
    if not prezzi:
        return {"usable": 0, "median": None}
    meta = len(prezzi) // 2
    mediana = prezzi[meta] if len(prezzi) % 2 else (prezzi[meta - 1] + prezzi[meta]) / 2
    return {"usable": len(prezzi), "median": json_decimal(Decimal(mediana))}


def solo_cifre(valore: Any) -> str:
    """Reduce a barcode to the form used for comparison.

    Must produce the same result as `conferme.codice_confrontabile`, where
    equivalences are written: the same rule duplicated in two places is
    technical debt, paid here because the pipeline scripts don't import
    `app/`. Tests (`test_prepare_sources`) check the two stay in sync.
    """

    testo = str(valore or "").strip()
    if testo.endswith(".0") and testo[:-2].isdigit():
        testo = testo[:-2]
    return "".join(carattere for carattere in testo if carattere.isdigit())


def mappa_delle_uguaglianze(classi: Any) -> dict[str, list[str]]:
    """Turn "groups of equal codes" into "for each code, the others in its group".

    Each group arrives as a list of codes declared equal. This reshapes that
    into the question `build_matching` actually needs: given a product's
    EAN, which other codes count as the same product. The code itself is
    excluded from its own answer, since exact matching already finds it, and
    including it would double-count the same rows.
    """

    mappa: dict[str, list[str]] = {}
    for gruppo in classi or []:
        codici = [solo_cifre(codice) for codice in gruppo or [] if solo_cifre(codice)]
        if len(set(codici)) < 2:
            continue
        for codice in codici:
            altri = mappa.setdefault(codice, [])
            for altro in codici:
                # Duplicates are removed here, and only here (`build_matching`
                # relies on that and adds no guard of its own). A repeated
                # code — within one group, or across two groups declaring the
                # same pair — would collect the same price-list row twice,
                # turning a certain match into `EAN_AMBIGUO`, an AI question
                # instead of a resolved match.
                if altro != codice and altro not in altri:
                    altri.append(altro)
    return mappa


def build_matching(
    master: list[dict[str, Any]],
    sources: dict[str, list[dict[str, Any]]],
    *,
    uguaglianze: Any = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    """Match every management-list product to supplier rows, by EAN.

    `uguaglianze` are groups of barcodes the user has declared to be the same
    item (`conferme.MagazzinoConferme.classi`). Without it, this function
    does exactly what it always did: this is why exact matching still looks
    up the raw code and equivalences use a second index, instead of
    normalizing every code and changing the basis of matches that already
    work.

    Motivating case: a product sits in the management list under one EAN
    that only one supplier uses for it, while three other suppliers carry
    the same item under a different EAN (a barcode variant per package
    batch). No scoring heuristic can infer that from the product name alone,
    but once a human has declared the two codes equal, all four suppliers
    match as native `EAN_ESATTO` and no downstream code needs a special case
    for it.
    """

    per_codice = mappa_delle_uguaglianze(uguaglianze)
    indexes: dict[str, dict[str, list[dict[str, Any]]]] = {}
    # This second index exists only for equivalences: keys are codes reduced
    # to digits, since a hand-declared code and one read from a spreadsheet
    # are rarely written the same way.
    indici_per_cifre: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for source_name, records in sources.items():
        index: dict[str, list[dict[str, Any]]] = defaultdict(list)
        per_cifre: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            if record.get("ean"):
                index[record["ean"]].append(record)
                chiave = solo_cifre(record["ean"])
                if chiave:
                    per_cifre[chiave].append(record)
        indexes[source_name] = index
        indici_per_cifre[source_name] = per_cifre

    matching = []
    semantic_queue = []
    exact_presence_counts = Counter()
    exact_unique_usable_counts = Counter()
    equivalence_counts = Counter()
    master_matched_by_presence = set()
    for product in master:
        product_result = {"gestionale": product, "suppliers": {}}
        altri_codici = per_codice.get(solo_cifre(product.get("ean")), []) if product.get("ean") else []
        for source_name, index in indexes.items():
            candidates = index.get(product["ean"], []) if product.get("ean") else []
            # Rows found via a declared equivalence, appended after the
            # native ones. There is no duplicate guard here, and that's
            # deliberate: `mappa_delle_uguaglianze` already removes repeated
            # codes within a group and excludes the product's own code, and a
            # price-list row has exactly one EAN, so no row can arrive by both
            # paths. Do not add one here without removing the redundancy at
            # its source: a check that looks like a safeguard but never
            # triggers is worse than none, since the next reader relies on it.
            da_uguaglianza = [
                candidate
                for altro in altri_codici
                for candidate in indici_per_cifre[source_name].get(altro, [])
            ]
            if da_uguaglianza:
                candidates = [*candidates, *da_uguaglianza]
                equivalence_counts[source_name] += 1
                product_result.setdefault("uguaglianze", {})[source_name] = sorted(
                    {solo_cifre(candidate.get("ean")) for candidate in da_uguaglianza}
                )
            usable_candidates = [candidate for candidate in candidates if candidate.get("usable", True)]
            if candidates:
                exact_presence_counts[source_name] += 1
                master_matched_by_presence.add(product["ean"])
            if len(usable_candidates) == 1:
                status = "EAN_ESATTO"
                exact_unique_usable_counts[source_name] += 1
            elif len(usable_candidates) > 1:
                status = "EAN_AMBIGUO"
            elif candidates:
                status = "EAN_PRESENTE_NON_UTILIZZABILE"
            else:
                status = "EAN_ASSENTE"
            voce = {"status": status, "candidates": candidates, "usable_candidates": usable_candidates}
            # Records where the match came from: without this, a row matched
            # via a human declaration is indistinguishable from one matched
            # by barcode, and whoever investigates a wrong order has no way
            # to trace it back.
            if (product_result.get("uguaglianze") or {}).get(source_name):
                voce["via_uguaglianza"] = product_result["uguaglianze"][source_name]
            product_result["suppliers"][source_name] = voce
            if status != "EAN_ESATTO":
                semantic_queue.append(
                    {
                        "gestionale_source_row": product["source_row"],
                        "ean": product["ean"],
                        "description": product["description"],
                        "supplier": source_name,
                        "reason": status,
                    }
                )
        matching.append(product_result)

    audit = {
        "master": source_stats(master),
        "sources": {name: source_stats(records) for name, records in sources.items()},
        "price_summary": {name: price_stats(records) for name, records in sources.items()},
        "exact_ean_presence": dict(exact_presence_counts),
        "exact_unique_usable": dict(exact_unique_usable_counts),
        "matched_by_at_least_one_exact_ean": len(master_matched_by_presence),
        "missing_from_all_exact_ean": len(master) - len(master_matched_by_presence),
        # How many matches came from a declared equivalence, and how many
        # declarations were active. A gain that isn't counted is as much an
        # unaudited change as a drop that isn't; this is the only place that
        # shows whether those declarations are still earning their keep.
        "matched_by_declared_equivalence": dict(equivalence_counts),
        "declared_equivalences": len([gruppo for gruppo in uguaglianze or [] if len(gruppo or []) > 1]),
    }
    return matching, semantic_queue, audit


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--gestionale", type=Path, required=True)
    parser.add_argument("--betulla", type=Path, required=True)
    parser.add_argument("--larice", type=Path, required=True)
    parser.add_argument("--noce", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    master = read_gestionale(args.gestionale)
    larice, warnings = read_larice(args.larice)
    larice, display_offers, display_summary = integrate_larice_displays(larice, analyse_workbook(args.larice))
    sources = {
        "betulla": read_betulla(args.betulla),
        "larice": larice,
        "noce": read_noce(args.noce),
    }
    matching, semantic_queue, audit = build_matching(master, sources)
    audit["warnings"] = warnings
    audit["displays"] = {"larice": display_summary}
    write_json(args.output / "normalized_sources.json", {"gestionale": master, **sources, "display_offers": {"larice": display_offers}})
    write_json(args.output / "display_offers.json", {"larice": display_offers})
    write_json(args.output / "matching_result.json", matching)
    write_json(args.output / "semantic_queue.json", semantic_queue)
    write_json(args.output / "audit.json", audit)
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
