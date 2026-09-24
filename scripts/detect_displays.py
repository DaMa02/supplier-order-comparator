#!/usr/bin/env python3
"""Detect orderable display offers and their non-orderable components in Larice XLSX files.

The detector is deliberately deterministic.  It scans the complete active worksheet,
recognises parent/component blocks from the documented Larice columns, and then uses
independent semantic, structural, quantity, and price evidence to decide whether a
bundle is an actual display.  Component rows are always marked non-orderable.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

import registro


# Zero-based indexes for the documented Larice A:R layout.
GROUP = 0
INDICATOR = 1
SUPPLIER_CODE = 2
HEADER_DESCRIPTION = 3
PIECES_PER_CARTON = 4
PALLET = 5
PARENT_DESCRIPTION = 6
PRICE_PRE_DISCOUNT = 14
DISCOUNT = 15
VAT = 16
EAN = 17
MIN_COLUMNS = 18

DISPLAY_TERM_RE = re.compile(r"(?<!\w)(ESPO|ESPOSITORE|DISPLAY|EXPO)(?!\w)", re.IGNORECASE)
NEGATED_DISPLAY_RE = re.compile(
    r"(?<!\w)(?:NO|NON|SENZA)\s+(?:ESPO|ESPOSITORE|DISPLAY|EXPO)(?!\w)",
    re.IGNORECASE,
)
DECLARED_QUANTITY_RE = re.compile(r"(?<!\w)X\s*(\d{1,5})(?!\d)", re.IGNORECASE)
GTIN_LENGTHS = {8, 12, 13, 14}
MONEY_QUANTUM = Decimal("0.0001")


def _present(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return text


def _normalised_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", _text(value))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def _decimal(value: Any, *, italian: bool = False) -> Decimal | None:
    if not _present(value) or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if italian:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _integer(value: Any) -> int | None:
    number = _decimal(value, italian=True)
    if number is None or number <= 0 or number != number.to_integral_value():
        return None
    integer = int(number)
    return integer if integer <= 100_000 else None


def _json_decimal(value: Decimal | None) -> str | None:
    if value is None:
        return None
    return format(value.quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP), "f")


def _codici_ammessi() -> set[str]:
    """Return the discount column's accepted codes, as declared in the adapter registry.

    Reads the same declaration the price-list reader uses; a second
    hand-maintained copy here would risk diverging, flagging the same code as
    valid in one place and suspicious in the other.
    """

    return registro.codici_ammessi(registro.codici_di_riga(registro.adattatore("larice_v1")))


def _discount(value: Any, ammessi: set[str] | None = None) -> tuple[Decimal, str, str | None]:
    number = _decimal(value)
    if number is not None:
        if number > 1 and number <= 100:
            number /= Decimal("100")
        if number < 0 or number > 1:
            return Decimal("0"), "invalid", f"Percentuale sconto non valida: {value}"
        return number, "percentage", None
    text = _text(value)
    noti = _codici_ammessi() if ammessi is None else ammessi
    warning = None if not text or text.upper() in noti else f"Codice sconto testuale inatteso: {text}"
    return Decimal("0"), "text_no_discount", warning


def _gtin_valid(value: str) -> bool:
    if not value.isdigit() or len(value) not in GTIN_LENGTHS:
        return False
    digits = [int(character) for character in value]
    payload = digits[:-1]
    weighted = sum(digit * (3 if (len(payload) - index) % 2 else 1) for index, digit in enumerate(payload))
    return (10 - weighted % 10) % 10 == digits[-1]


def _active(row: Sequence[Any]) -> bool:
    return any(_present(value) for value in row)


def _padded(row: Sequence[Any]) -> tuple[Any, ...]:
    values = tuple(row[:MIN_COLUMNS])
    return values + (None,) * (MIN_COLUMNS - len(values))


def _is_orderable(row: Sequence[Any]) -> bool:
    return _present(row[SUPPLIER_CODE]) or (
        _present(row[INDICATOR])
        and _present(row[PARENT_DESCRIPTION])
        and _decimal(row[PRICE_PRE_DISCOUNT]) is not None
    )


def _component_description(row: Sequence[Any]) -> str:
    # Larice uses H/I for some manufacturers and I/J for others.  Search only
    # the description area, before monetary columns, and require alphabetic text.
    for value in row[6:14]:
        text = _text(value)
        if text and re.search(r"[A-Za-zÀ-ÖØ-öø-ÿ]", text):
            return text
    return ""


def _component_quantity(row: Sequence[Any]) -> int | None:
    # Known files place the component quantity in H or I.  Looking through N
    # tolerates a one-column structural shift without ever reading the price in O.
    for value in row[7:14]:
        quantity = _integer(value)
        if quantity is not None:
            return quantity
    return None


def _component_candidate(row: Sequence[Any]) -> bool:
    if _is_orderable(row):
        return False
    description = _component_description(row)
    ean = _identifier(row[EAN])
    quantity = _component_quantity(row)
    return bool(description and (ean or quantity is not None))


def _declared_quantity(text: str) -> int | None:
    matches = {int(match) for match in DECLARED_QUANTITY_RE.findall(text)}
    return next(iter(matches)) if len(matches) == 1 else None


def _money_matches(left: Decimal | None, right: Decimal | None) -> bool:
    if left is None or right is None:
        return False
    tolerance = max(Decimal("0.02"), abs(left) * Decimal("0.001"))
    return abs(left - right) <= tolerance


def _composition_fingerprint(components: Sequence[dict[str, Any]]) -> tuple[str, str, list[dict[str, Any]]]:
    aggregated: dict[str, dict[str, Any]] = {}
    all_ean = True
    all_quantities = True
    for component in components:
        ean = component["ean"]
        quantity = component["quantity"]
        description_key = _normalised_text(component["description"])
        identifier = f"EAN:{ean}" if ean else f"DESC:{description_key}"
        all_ean = all_ean and bool(ean)
        all_quantities = all_quantities and quantity is not None
        entry = aggregated.setdefault(
            identifier,
            {"identifier": identifier, "ean": ean or None, "description_key": description_key, "quantity": 0},
        )
        if quantity is None:
            entry["quantity"] = None
        elif entry["quantity"] is not None:
            entry["quantity"] += quantity

    signature = sorted(aggregated.values(), key=lambda item: item["identifier"])
    basis = (
        "ean_quantity"
        if all_ean and all_quantities
        else "ean_without_complete_quantity"
        if all_ean
        else "mixed_identifier_quantity"
        if all_quantities
        else "mixed_identifier_without_complete_quantity"
    )
    # Descriptions are audit metadata, not part of the identity when an EAN is
    # available.  Hash only the stable identifier and quantity so equivalent
    # manufacturer compositions match across suppliers despite wording/order.
    canonical_signature = [
        {"identifier": entry["identifier"], "quantity": entry["quantity"]} for entry in signature
    ]
    canonical = json.dumps(
        {"version": 1, "components": canonical_signature},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest(), basis, signature


def _component_record(source_row: int, row: Sequence[Any]) -> dict[str, Any]:
    price = _decimal(row[PRICE_PRE_DISCOUNT])
    discount_rate, discount_type, warning = _discount(row[DISCOUNT])
    quantity = _component_quantity(row)
    extended_pre = price * quantity if price is not None and quantity is not None else None
    extended_post = extended_pre * (Decimal("1") - discount_rate) if extended_pre is not None else None
    ean = _identifier(row[EAN])
    return {
        "source_row": source_row,
        "description": _component_description(row),
        "ean": ean,
        "ean_valid_gtin": _gtin_valid(ean),
        "quantity": quantity,
        "unit_price_pre_discount": _json_decimal(price),
        "discount_raw": row[DISCOUNT],
        "discount_type": discount_type,
        "discount_rate": _json_decimal(discount_rate),
        "unit_price_post_discount": _json_decimal(price * (Decimal("1") - discount_rate) if price is not None else None),
        "extended_pre_discount": _json_decimal(extended_pre),
        "extended_post_discount": _json_decimal(extended_post),
        "warning": warning,
        "orderable": False,
    }


def _evaluate_bundle(
    rows: Sequence[tuple[Any, ...]],
    parent_index: int,
    component_indexes: Sequence[int],
    header_index: int | None,
    source_file: str,
    sheet_name: str,
) -> dict[str, Any]:
    parent = rows[parent_index]
    parent_row = parent_index + 1
    components = [_component_record(index + 1, rows[index]) for index in component_indexes]
    group_name = _text(parent[GROUP])
    parent_description = _text(parent[PARENT_DESCRIPTION]) or group_name
    semantic_text = f"{group_name} {parent_description}"
    negative = bool(NEGATED_DISPLAY_RE.search(semantic_text))
    positive_matches = sorted({match.upper() for match in DISPLAY_TERM_RE.findall(semantic_text)})
    explicit_display = bool(positive_matches) and not negative

    quantities = [component["quantity"] for component in components]
    quantity_complete = all(quantity is not None for quantity in quantities)
    sum_pieces = sum(quantity for quantity in quantities if quantity is not None) if any(quantity is not None for quantity in quantities) else None
    name_quantity = _declared_quantity(semantic_text)
    parent_pack = _integer(parent[PIECES_PER_CARTON])
    if name_quantity is not None:
        declared_quantity = name_quantity
        declared_quantity_source = "name_x_quantity"
    elif parent_pack is not None and parent_pack > 1:
        declared_quantity = parent_pack
        declared_quantity_source = "parent_pieces_per_carton"
    elif quantity_complete and sum_pieces is not None:
        declared_quantity = sum_pieces
        declared_quantity_source = "derived_from_components"
    else:
        declared_quantity = None
        declared_quantity_source = None

    parent_price_pre = _decimal(parent[PRICE_PRE_DISCOUNT])
    parent_discount_rate, parent_discount_type, parent_warning = _discount(parent[DISCOUNT])
    parent_price_post = (
        parent_price_pre * (Decimal("1") - parent_discount_rate) if parent_price_pre is not None else None
    )
    parent_order_total_pre = parent_price_pre * parent_pack if parent_price_pre is not None and parent_pack is not None else None
    parent_order_total_post = parent_price_post * parent_pack if parent_price_post is not None and parent_pack is not None else None

    priced_extended_pre = [
        Decimal(component["extended_pre_discount"])
        for component in components
        if component["extended_pre_discount"] is not None
    ]
    priced_extended_post = [
        Decimal(component["extended_post_discount"])
        for component in components
        if component["extended_post_discount"] is not None
    ]
    component_price_complete = len(priced_extended_pre) == len(components)
    component_extended_partial_pre = sum(priced_extended_pre, Decimal("0")) if priced_extended_pre else None
    component_extended_partial_post = sum(priced_extended_post, Decimal("0")) if priced_extended_post else None
    component_extended_pre = component_extended_partial_pre if component_price_complete else None
    component_extended_post = component_extended_partial_post if component_price_complete else None

    price_match_basis = None
    if _money_matches(component_extended_pre, parent_price_pre):
        price_match_basis = "parent_price_cell"
    elif _money_matches(component_extended_pre, parent_order_total_pre):
        price_match_basis = "parent_order_total"

    fingerprint, fingerprint_basis, signature = _composition_fingerprint(components)
    evidence: list[dict[str, Any]] = []
    score = 0

    def add(code: str, points: int, detail: str) -> None:
        nonlocal score
        score += points
        evidence.append({"code": code, "points": points, "detail": detail})

    if explicit_display:
        add("EXPLICIT_DISPLAY_TERM", 4, ", ".join(positive_matches))
    if negative:
        add("NEGATED_DISPLAY_TERM", -10, "Il testo dichiara esplicitamente che non e un espositore")
    add("PARENT_WITH_COMPONENT_BLOCK", 2, f"{len(components)} componenti consecutivi senza codice articolo")
    if header_index is not None:
        add("DEDICATED_HEADER", 1, f"Riga intestazione {header_index + 1}")
    child_eans = [component["ean"] for component in components if component["ean"]]
    if not _identifier(parent[EAN]) and child_eans:
        add("PARENT_EAN_EMPTY_CHILD_EANS_PRESENT", 1, f"{len(child_eans)} EAN sulle righe figlie")
    if len(set(child_eans)) >= 2:
        add("MULTIPLE_COMPONENT_GTINS", 1, f"{len(set(child_eans))} EAN distinti")
    if declared_quantity is not None and quantity_complete and sum_pieces == declared_quantity:
        add("DECLARED_QUANTITY_MATCHES_COMPONENT_SUM", 2, f"{declared_quantity} pezzi")
    if parent_pack is not None and parent_pack > 1 and quantity_complete and sum_pieces == parent_pack:
        add("PARENT_PACK_MATCHES_COMPONENT_SUM", 1, f"PzCt {parent_pack}")
    if price_match_basis:
        add("COMPONENT_SUM_MATCHES_PARENT_PRICE", 2, price_match_basis)

    detected = explicit_display and not negative and score >= 7 and len(components) >= 2
    confidence = "HIGH" if detected and score >= 10 else "MEDIUM" if detected else "REJECTED" if negative else "LOW"
    offer_id = f"LARICE:{_identifier(parent[SUPPLIER_CODE]) or parent_row}"
    source_rows = {
        "header_row": header_index + 1 if header_index is not None else None,
        "parent_row": parent_row,
        "component_rows": [index + 1 for index in component_indexes],
        "first_row": header_index + 1 if header_index is not None else parent_row,
        "last_row": component_indexes[-1] + 1,
    }
    price_delta = component_extended_pre - parent_price_pre if component_extended_pre is not None and parent_price_pre is not None else None

    return {
        "record_type": "DISPLAY_OFFER" if detected else "BUNDLE_CANDIDATE",
        "display_detected": detected,
        "offer_id": offer_id,
        "source": "larice",
        "source_file": source_file,
        "sheet_name": sheet_name,
        "group_name": group_name,
        "description": parent_description,
        "supplier_code": _identifier(parent[SUPPLIER_CODE]),
        "indicator": _text(parent[INDICATOR]),
        "ean": _identifier(parent[EAN]),
        "pieces_per_carton": parent_pack,
        "pallet": _integer(parent[PALLET]),
        "vat": parent[VAT],
        "declared_quantity": declared_quantity,
        "declared_quantity_source": declared_quantity_source,
        "sum_pieces": sum_pieces,
        "sum_pieces_complete": quantity_complete,
        "component_quantity_coverage": f"{sum(quantity is not None for quantity in quantities)}/{len(components)}",
        "parent_price_pre_discount": _json_decimal(parent_price_pre),
        "parent_discount_raw": parent[DISCOUNT],
        "parent_discount_type": parent_discount_type,
        "parent_discount_rate": _json_decimal(parent_discount_rate),
        "parent_price_post_discount": _json_decimal(parent_price_post),
        "parent_order_total_pre_discount": _json_decimal(parent_order_total_pre),
        "parent_order_total_post_discount": _json_decimal(parent_order_total_post),
        "component_extended_sum_pre_discount": _json_decimal(component_extended_pre),
        "component_extended_sum_post_discount": _json_decimal(component_extended_post),
        "component_extended_partial_pre_discount": _json_decimal(component_extended_partial_pre),
        "component_extended_partial_post_discount": _json_decimal(component_extended_partial_post),
        "component_price_coverage": f"{len(priced_extended_pre)}/{len(components)}",
        "component_parent_price_match_basis": price_match_basis,
        "component_parent_price_delta": _json_decimal(price_delta),
        "composition_fingerprint": f"sha256:{fingerprint}",
        "composition_fingerprint_basis": fingerprint_basis,
        "composition_signature": signature,
        "components": components,
        "score": score,
        "confidence": confidence,
        "evidence": evidence,
        "warning": parent_warning,
        "source_rows": source_rows,
    }


def _base_classification(source_row: int, row: Sequence[Any]) -> dict[str, Any]:
    orderable = _is_orderable(row)
    ean = _identifier(row[EAN])
    if orderable:
        row_type = "STANDARD"
    elif not ean:
        row_type = "HEADER"
    else:
        row_type = "STANDARD"
    return {
        "source_row": source_row,
        "row_type": row_type,
        "orderable": orderable,
        "group_name": _text(row[GROUP]),
        "supplier_code": _identifier(row[SUPPLIER_CODE]),
        "ean": ean,
        "parent_row": None,
        "display_detected": False,
        "display_offer_id": None,
    }


def analyse_rows(rows: Iterable[Sequence[Any]], *, source_file: str = "", sheet_name: str = "") -> dict[str, Any]:
    """Analyse all worksheet rows already loaded from columns A:R."""
    materialised = [_padded(row) for row in rows]
    active_indexes = [index for index, row in enumerate(materialised) if _active(row)]
    classifications = {
        index: _base_classification(index + 1, materialised[index]) for index in active_indexes
    }
    bundle_candidates: list[dict[str, Any]] = []
    consumed_components: set[int] = set()

    for parent_index, parent in enumerate(materialised):
        if parent_index in consumed_components or not _active(parent) or not _is_orderable(parent):
            continue
        group_key = _normalised_text(parent[GROUP])
        if not group_key:
            continue
        component_indexes: list[int] = []
        following = parent_index + 1
        while following < len(materialised):
            candidate = materialised[following]
            if not _active(candidate) or _normalised_text(candidate[GROUP]) != group_key or _is_orderable(candidate):
                break
            if not _component_candidate(candidate):
                break
            component_indexes.append(following)
            following += 1
        if len(component_indexes) < 2:
            continue

        header_index = None
        preceding = parent_index - 1
        if preceding >= 0:
            header = materialised[preceding]
            if (
                _active(header)
                and _normalised_text(header[GROUP]) == group_key
                and not _is_orderable(header)
                and not _identifier(header[EAN])
            ):
                header_index = preceding

        candidate_record = _evaluate_bundle(
            materialised,
            parent_index,
            component_indexes,
            header_index,
            source_file,
            sheet_name,
        )
        bundle_candidates.append(candidate_record)
        offer_id = candidate_record["offer_id"]
        display_detected = candidate_record["display_detected"]
        classifications[parent_index].update(
            {
                "row_type": "PARENT",
                "orderable": True,
                "display_detected": display_detected,
                "display_offer_id": offer_id if display_detected else None,
            }
        )
        if header_index is not None:
            classifications[header_index].update(
                {
                    "row_type": "HEADER",
                    "orderable": False,
                    "parent_row": parent_index + 1,
                    "display_detected": display_detected,
                    "display_offer_id": offer_id if display_detected else None,
                }
            )
        for component_index in component_indexes:
            consumed_components.add(component_index)
            classifications[component_index].update(
                {
                    "row_type": "COMPONENT",
                    "orderable": False,
                    "parent_row": parent_index + 1,
                    "display_detected": display_detected,
                    "display_offer_id": offer_id if display_detected else None,
                }
            )

    ordered_classifications = [classifications[index] for index in sorted(classifications)]
    display_offers = [candidate for candidate in bundle_candidates if candidate["display_detected"]]
    rejected_bundles = [candidate for candidate in bundle_candidates if not candidate["display_detected"]]
    counts = Counter(classification["row_type"] for classification in ordered_classifications)
    orderable_rows = [
        classification["source_row"] for classification in ordered_classifications if classification["orderable"]
    ]
    component_rows = [
        classification["source_row"]
        for classification in ordered_classifications
        if classification["row_type"] == "COMPONENT"
    ]
    return {
        "source": "larice",
        "source_file": source_file,
        "sheet_name": sheet_name,
        "rows_scanned": len(materialised),
        "active_rows": len(active_indexes),
        "classification_counts": dict(sorted(counts.items())),
        "display_offer_count": len(display_offers),
        "display_offers": display_offers,
        "rejected_bundle_candidates": rejected_bundles,
        "row_classifications": ordered_classifications,
        "orderable_rows": orderable_rows,
        "component_rows_excluded_from_orderable": component_rows,
    }


def analyse_workbook(path: str | Path, sheet_name: str | None = None) -> dict[str, Any]:
    """Read the complete Larice worksheet and return display and row-level audit data."""
    workbook_path = Path(path).resolve()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Foglio non trovato: {sheet_name}")
            sheet = workbook[sheet_name]
        else:
            candidates = [sheet for sheet in workbook.worksheets if sheet.max_column >= MIN_COLUMNS]
            if not candidates:
                raise ValueError(f"Nessun foglio Larice con almeno {MIN_COLUMNS} colonne in {workbook_path}")
            sheet = candidates[0]
        rows = sheet.iter_rows(
            min_row=1,
            max_row=sheet.max_row,
            min_col=1,
            max_col=MIN_COLUMNS,
            values_only=True,
        )
        return analyse_rows(rows, source_file=str(workbook_path), sheet_name=sheet.title)
    finally:
        workbook.close()


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Listino Larice XLSX da analizzare in sola lettura")
    parser.add_argument("--sheet", help="Nome foglio; se omesso usa il primo foglio commerciale A:R")
    parser.add_argument("--output", type=Path, help="File JSON di output; se omesso scrive il JSON su stdout")
    parser.add_argument("--summary", action="store_true", help="Stampa solo un riepilogo compatto")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    result = analyse_workbook(args.workbook, args.sheet)
    if args.summary:
        value: Any = {
            "source_file": result["source_file"],
            "sheet_name": result["sheet_name"],
            "rows_scanned": result["rows_scanned"],
            "active_rows": result["active_rows"],
            "classification_counts": result["classification_counts"],
            "display_offer_count": result["display_offer_count"],
            "display_offers": [
                {
                    "supplier_code": offer["supplier_code"],
                    "description": offer["description"],
                    "declared_quantity": offer["declared_quantity"],
                    "sum_pieces": offer["sum_pieces"],
                    "parent_price_pre_discount": offer["parent_price_pre_discount"],
                    "parent_price_post_discount": offer["parent_price_post_discount"],
                    "component_extended_sum_pre_discount": offer["component_extended_sum_pre_discount"],
                    "score": offer["score"],
                    "confidence": offer["confidence"],
                    "source_rows": offer["source_rows"],
                }
                for offer in result["display_offers"]
            ],
            "rejected_bundle_candidates": [
                {
                    "description": candidate["description"],
                    "score": candidate["score"],
                    "confidence": candidate["confidence"],
                    "source_rows": candidate["source_rows"],
                }
                for candidate in result["rejected_bundle_candidates"]
            ],
        }
    else:
        value = result
    encoded = json.dumps(value, ensure_ascii=False, indent=2, default=str)
    if args.output:
        args.output.write_text(encoded + "\n", encoding="utf-8")
    else:
        print(encoded)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
