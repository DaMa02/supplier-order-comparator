#!/usr/bin/env python3
"""Parse an input manifest through known or declarative adapters.

Run ``validate_input_manifest.py`` first.  Known unchanged schemas reuse the
tested readers in ``prepare_sources.py``; varied/new schemas require an explicit
field mapping in the manifest and are parsed generically.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

# The .xls reader lives in the app folder and uses only the standard library.
APP_DIR = Path(__file__).resolve().parents[1] / "app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

from xls_reader import read_workbook  # noqa: E402

import registro
from inspect_sources import container_format
from prepare_sources import (
    FUORI_DAL_FILTRO,
    MOTIVO_NON_DICHIARATO,
    NON_DISPONIBILE,
    NON_E_RIGA_PRODOTTO,
    SENZA_DESCRIZIONE,
    SENZA_PEZZI_PER_COLLO,
    SENZA_PREZZO,
    build_matching,
    conta_non_ordinabili,
    decimal_value,
    fattore_d_ordine,
    larice_discount,
    indice_di_colonna,
    json_decimal,
    normalize_ean,
    read_betulla,
    read_gestionale,
    read_larice,
    read_noce,
    integrate_larice_displays,
    write_json,
)
from detect_displays import analyse_workbook
# The marker shape is validated by the manifest validator, not a second copy
# here: the reader and the validator must agree.
from validate_input_manifest import errori_del_marcatore


def normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def column_number(value: Any, header: list[Any] | None = None) -> int:
    """Resolve a mapping's declared column to a 1-based column number.

    When the sheet has a header row, the declared name is looked up there
    only, with no fallback to spreadsheet letters: names like "cat" or "Iva"
    are real column headers in one supplier's list and also valid Excel
    column references (CAT = 2074, IVA = 6657). With a letter fallback, a
    renamed column wouldn't raise an error, it would silently read an empty
    column and let thousands of unrelated rows into the comparison unnoticed.
    Stopping with a clear message costs a click; reading the wrong column
    costs a wrong order.

    Letters are accepted only where there is no header row to look names up in.
    """
    if isinstance(value, int) and value >= 1:
        return value
    text = str(value or "").strip()
    if header is not None:
        wanted = normalized(text)
        matches = [index for index, item in enumerate(header, start=1) if normalized(item) == wanted]
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise ValueError(f"Intestazione duplicata nella mappatura: {text}")
        letti = [str(item).strip() for item in header if str(item or "").strip()]
        elenco = ", ".join(f"«{nome}»" for nome in letti[:20]) or "nessuna"
        if len(letti) > 20:
            elenco += f" e altre {len(letti) - 20}"
        raise ValueError(
            f"Colonna «{text}» non trovata fra le intestazioni dichiarate: il documento non ha "
            f"la forma attesa e non viene letto a metà. Intestazioni lette: {elenco}. "
            "Per indicare una colonna che non ha un nome si dichiara il suo numero (per esempio 5), "
            "non la lettera."
        )
    if re.fullmatch(r"[A-Za-z]+", text):
        return column_index_from_string(text.upper())
    raise ValueError(f"Colonna non risolta: {value!r}")


def mapped_value(row: tuple[Any, ...] | list[Any], columns: dict[str, int], field: str) -> Any:
    index = columns.get(field)
    return row[index - 1] if index and index <= len(row) else None


def _testo_confrontabile(valore: Any) -> str:
    """Normalize a cell's text for comparison (collapse whitespace, casefold).

    Extra spaces and case don't make a different marker: a trailing space in
    the file shouldn't defeat the match.
    """

    return " ".join(str(valore if valore is not None else "").split()).casefold()


def colonna_del_marcatore(spec: Any) -> int:
    """Resolve a data-start marker's column: a number, or a letter.

    Letters are allowed here, unlike in `column_number`, where they're
    rejected because names like "cat" or "Iva" are both real column headers
    and valid Excel column references. The marker doesn't name a data field:
    it sits outside the table, on a row with no header, so position is the
    only way to address it.
    """

    if isinstance(spec, bool):
        raise ValueError("data_start_marker.column deve essere una lettera di colonna o il suo numero")
    if isinstance(spec, int) and spec >= 1:
        return spec
    testo = str(spec or "").strip()
    if re.fullmatch(r"[1-9][0-9]*", testo):
        return int(testo)
    if re.fullmatch(r"[A-Za-z]{1,3}", testo):
        return column_index_from_string(testo.upper())
    raise ValueError(
        f"data_start_marker.column: «{spec}» non è una colonna. Indicare la lettera "
        "(per esempio «A») oppure il numero della colonna (per esempio 1)."
    )


def riga_dopo_il_marcatore(rows: list[tuple[Any, ...]], marker: dict[str, Any], nome_file: str) -> int:
    """Return where products start per the marker, recomputed on this file.

    There is no fallback: if the marker isn't found, or is found more than
    once, reading stops and reports what it searched for and where. Falling
    back to last week's row number is exactly the defect this feature closes:
    a promotional block can grow or shrink week to week, and a frozen row
    number would silently pull in phantom rows or drop real ones.
    """

    problemi = errori_del_marcatore(marker)
    if problemi:
        raise ValueError(f"In «{nome_file}» la regola dell'inizio dei dati non è scritta bene: "
                         + "; ".join(problemi))
    indice = colonna_del_marcatore(marker.get("column"))
    lettera = get_column_letter(indice)
    esatto = str(marker.get("equals") or "").strip()
    contenuto = str(marker.get("contains") or "").strip()
    cercato = esatto or contenuto
    atteso = _testo_confrontabile(cercato)
    scarto = int(marker.get("offset", 1))
    come = "è esattamente" if esatto else "contiene"

    trovate = [
        numero
        for numero, riga in enumerate(rows, start=1)
        if indice <= len(riga) and (
            _testo_confrontabile(riga[indice - 1]) == atteso if esatto
            else atteso in _testo_confrontabile(riga[indice - 1])
        )
    ]
    dove = f"in «{nome_file}», colonna {lettera}"
    if not trovate:
        raise ValueError(
            f"Non trovo dove comincia il listino: {dove} nessuna riga {come} «{cercato}». "
            f"La configurazione di questo fornitore dice che i prodotti cominciano {scarto} "
            f"riga/e dopo quella riga. Il documento non viene letto a metà: un numero di riga "
            f"deciso la volta scorsa taglierebbe l'elenco nel punto sbagliato. Apri «{nome_file}» e "
            f"controlla se quella scritta c'è ancora; se il fornitore l'ha tolta o cambiata, togli il "
            f"documento dal passo 1 «Importa i dati» e ricaricalo corretto."
        )
    if len(trovate) > 1:
        elenco = ", ".join(str(numero) for numero in trovate[:10])
        coda = f" e altre {len(trovate) - 10}" if len(trovate) > 10 else ""
        raise ValueError(
            f"Non so dove comincia il listino: {dove} ci sono {len(trovate)} righe in cui il "
            f"testo {come} «{cercato}» (righe {elenco}{coda}). Il punto in cui cominciano i "
            f"prodotti dev'essere uno solo: serve una scritta che compaia una volta sola, "
            f"oppure la riga esatta. Apri «{nome_file}»: se sono due pagine incollate una sotto "
            f"l'altra, togli la seconda intestazione e ricaricalo dal passo 1 «Importa i dati»."
        )
    riga = trovate[0] + scarto
    if riga > len(rows):
        raise ValueError(
            f"{dove} la riga «{cercato}» è la {trovate[0]}, e {scarto} riga/e più in basso "
            f"il documento è già finito: ha {len(rows)} righe. Dopo quella scritta non c'è "
            "nessun prodotto da leggere."
        )
    return riga


def prima_riga_dei_dati(rows: list[tuple[Any, ...]], mapping: dict[str, Any],
                        header_row: int, nome_file: str) -> int:
    """Return the first product row: the marker rule wins over a fixed number.

    When the mapping declares a `data_start_marker`, it is resolved fresh
    against this document every time. `data_start_row` stays in the mapping
    because the order writer needs it, but this function never reads it as a
    fallback when a marker is present: a silent fallback to yesterday's row
    number is the defect this function exists to avoid.
    """

    marker = mapping.get("data_start_marker")
    if marker not in (None, "", {}):
        return riga_dopo_il_marcatore(rows, marker, nome_file)
    return int(mapping.get("data_start_row") or (header_row + 1 if header_row else 1))


def selected_sheet(workbook: Any, mapping: dict[str, Any]) -> Any:
    selector = mapping.get("sheet")
    if selector in (None, "", "FIRST"):
        return workbook.worksheets[0]
    if isinstance(selector, int):
        if selector < 0 or selector >= len(workbook.worksheets):
            raise ValueError(f"Indice foglio non valido: {selector}")
        return workbook.worksheets[selector]
    if str(selector) not in workbook.sheetnames:
        raise ValueError(f"Foglio non trovato: {selector}")
    return workbook[str(selector)]


def selected_xls_sheet(path: Path, mapping: dict[str, Any]) -> list[list[tuple[Any, bool]]]:
    """Select a sheet inside an .xls file, with the same rules as selected_sheet."""
    sheets = read_workbook(path)
    if not sheets:
        raise ValueError(f"Il file {path.name} non contiene fogli di dati")
    selector = mapping.get("sheet")
    if selector in (None, "", "FIRST"):
        return sheets[0].rows
    if isinstance(selector, int):
        if selector < 0 or selector >= len(sheets):
            raise ValueError(f"Indice foglio non valido: {selector}")
        return sheets[selector].rows
    for sheet in sheets:
        if sheet.name == str(selector):
            return sheet.rows
    raise ValueError(f"Foglio non trovato: {selector}")


def mapped_rows(path: Path, mapping: dict[str, Any]) -> tuple[int, dict[str, int], list[tuple[Any, ...]], list[tuple[bool, ...]]]:
    """Read the source grid, picking the reader from the file's bytes, not its extension.

    The extension decides nothing: an .xls renamed to .xlsx is still an .xls.
    Bold is also read per cell, because one supplier flags its offer prices
    that way; both formats support it, otherwise re-saving the price list as
    .xlsx would silently drop that signal. Bold comes for free from .xls
    alongside the values; from .xlsx it costs 5-45% more read time, so it's
    only read there when the mapping declares it needs that signal
    (`offer_from_bold`).
    """
    if container_format(path) == "xls":
        griglia = selected_xls_sheet(path, mapping)
        rows: list[tuple[Any, ...]] = [tuple(valore for valore, _grassetto in riga) for riga in griglia]
        bold: list[tuple[bool, ...]] = [tuple(grassetto for _valore, grassetto in riga) for riga in griglia]
    else:
        # Pass the content, not the path: openpyxl rejects a path ending in
        # .xls even when the bytes are genuinely .xlsx.
        with path.open("rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                sheet = selected_sheet(workbook, mapping)
                if mapping.get("offer_from_bold"):
                    rows = []
                    bold = []
                    for riga in sheet.iter_rows():
                        rows.append(tuple(cella.value for cella in riga))
                        # A cell that was never written has no font at all;
                        # getattr defaults it to "not bold" instead of
                        # crashing the whole read.
                        bold.append(tuple(
                            bool(getattr(getattr(cella, "font", None), "bold", False))
                            for cella in riga
                        ))
                else:
                    rows = list(sheet.iter_rows(values_only=True))
                    bold = []
            finally:
                workbook.close()
    header_row = int(mapping.get("header_row") or 0)
    if header_row < 0:
        raise ValueError(
            f"La mappatura dichiara la riga di intestazione {header_row}: deve essere un numero "
            "di riga a partire da 1, oppure 0 se il documento non ha intestazioni."
        )
    if header_row > len(rows):
        # Without this check, the failure would be a raw "list index out of
        # range" that tells no one what to do.
        raise ValueError(
            f"Il foglio «{mapping.get('sheet') or 'primo foglio'}» di {path.name} non arriva alla riga "
            f"{header_row}, dove dovrebbero esserci le intestazioni: il file non ha la forma attesa."
        )
    header = list(rows[header_row - 1]) if header_row else None
    raw_columns = mapping.get("columns") or mapping.get("field_mapping") or {}
    columns = {field: column_number(spec, header) for field, spec in raw_columns.items() if spec not in (None, "")}
    data_start = prima_riga_dei_dati(rows, mapping, header_row, path.name)
    return data_start, columns, rows, bold


def last_data_row(rows: list[tuple[Any, ...]], columns: dict[str, int], mapping: dict[str, Any], data_start: int) -> int:
    """Return where the data actually ends, per the mapping's rule, not a guess.

    One supplier's list continues for hundreds of rows past the last
    product: rows that only carry a leftover Excel formula in the total
    column. Without a boundary, those would become phantom products. The
    rule names which fields identify a real product; the last row with at
    least one of them is the end of the data.
    """
    fields = (mapping.get("data_end_rule") or {}).get("last_row_with_any")
    if not fields:
        return len(rows)
    for row_number in range(len(rows), data_start - 1, -1):
        row = rows[row_number - 1]
        if any(str(mapped_value(row, columns, field) or "").strip() for field in fields):
            return row_number
    return data_start - 1


def excluded_row_label(row: tuple[Any, ...], columns: dict[str, int], mapping: dict[str, Any]) -> str | None:
    """Return the declared reason a row is excluded from the comparison, or None.

    For example, one supplier's list mixes in a food category the store
    doesn't stock; keeping those rows would mean proposing thousands of
    irrelevant items. How many rows are dropped is recorded in the audit
    report: a silent drop isn't a choice, it's data loss.
    """
    for rule in mapping.get("exclude_rows") or []:
        actual = str(mapped_value(row, columns, rule.get("field")) or "").strip()
        if "equals" in rule:
            if actual.casefold() == str(rule["equals"]).strip().casefold():
                return str(rule.get("label") or rule["equals"])
        elif "regex" in rule:
            if re.search(str(rule["regex"]), actual, flags=re.IGNORECASE):
                return str(rule.get("label") or rule["regex"])
        else:
            raise ValueError("exclude_rows deve usare equals o regex")
    return None


def non_e_una_riga_prodotto(row: tuple[Any, ...], columns: dict[str, int], mapping: dict[str, Any]) -> bool:
    """Return True when a row carries no product data at all.

    No description, no EAN, and no price: this isn't a badly read product,
    it's something else written into the table — a section title, a divider
    between a promotional block and the price list, a supplier's note. The
    master (management-software) reader has always dropped these rows
    (`read_mapped_master`, same `NON_E_RIGA_PRODOTTO` label); the supplier
    reader didn't, and let them through as unorderable products.

    Measured on a real price list: divider rows had empty description, EAN
    and price, but their text sat in the `supplier_code` column, so the older
    check (which required all three fields including the code to be empty)
    missed them and let them into the read products. A divider above the
    actual list does little harm; one that ends up mixed in among products
    becomes a catalog entry with the divider text in place of a product code.

    The supplier code is deliberately excluded from counting as product data:
    on its own it can't be ordered (a row with no description is already
    dropped), can't be matched, and has no price. It is exactly the field
    dividers write their text into.

    A price counts as present if the cell holds any value, even one that
    later fails to parse: a badly written price is a row to flag
    (`senza_prezzo`), not a row that doesn't exist.
    """

    if str(mapped_value(row, columns, "description") or "").strip():
        return False
    if normalize_ean(mapped_value(row, columns, "ean")):
        return False
    for campo in ("unit_price_net", "unit_price_pre_discount"):
        valore = mapped_value(row, columns, campo)
        if valore is not None and str(valore).strip():
            return False
    return True


def expiry_from_description(description: str, mapping: dict[str, Any]) -> tuple[str, dict[str, Any] | None]:
    """Split an expiry date written as an HTML fragment out of the description.

    One supplier writes something like "PRODUCT ML.500<br> Scadenza
    30/08/2026": that fragment isn't part of the item name and would pollute
    description matching, so it's parsed out separately and the description
    stays clean. The date itself isn't trusted blindly (real lists have
    contained implausible years); an absurd or calendar-impossible date
    becomes a warning, not an error that stops reading the whole file.
    """
    rule = mapping.get("expiry_from_description")
    if not rule:
        return description, None
    pattern = re.compile(str(rule.get("regex") or ""), re.IGNORECASE)
    if not {"giorno", "mese", "anno"} <= set(pattern.groupindex):
        raise ValueError("expiry_from_description: la regola deve nominare i gruppi giorno, mese e anno")
    match = pattern.search(description or "")
    if not match:
        return description, {"expiry_raw": None, "expiry_date": None, "expiry_plausible": None, "expiry_warning": None}

    cleaned = (description[: match.start()] + description[match.end():]).strip()
    raw = match.group(0).strip()
    try:
        expiry = date(int(match.group("anno")), int(match.group("mese")), int(match.group("giorno")))
    except ValueError:
        return cleaned, {
            "expiry_raw": raw,
            "expiry_date": None,
            "expiry_plausible": False,
            "expiry_warning": f"Scadenza scritta male nella descrizione: «{raw}». Il campo va verificato a mano.",
        }

    # Plausible means close to today: comparing by year keeps the rule from
    # depending on the exact day the program happens to run.
    window = rule.get("plausible_window_years") or {}
    back = int(window.get("back", 2))
    ahead = int(window.get("ahead", 10))
    today = date.today()
    plausible = today.year - back <= expiry.year <= today.year + ahead
    warning = None
    if not plausible:
        warning = (
            f"Scadenza fuori dal credibile: «{raw}». Il dato è stato letto lo stesso, "
            "ma va verificato sul listino del fornitore."
        )
    return cleaned, {
        "expiry_raw": raw,
        "expiry_date": expiry.isoformat(),
        "expiry_plausible": plausible,
        "expiry_warning": warning,
    }


def row_allowed(row: tuple[Any, ...], columns: dict[str, int], mapping: dict[str, Any]) -> bool:
    rule = mapping.get("row_filter") or {}
    if not rule:
        return True
    field = rule.get("field")
    actual = str(mapped_value(row, columns, field) or "").strip()
    if "equals" in rule:
        return actual.casefold() == str(rule["equals"]).strip().casefold()
    if "regex" in rule:
        return re.search(str(rule["regex"]), actual, flags=re.IGNORECASE) is not None
    raise ValueError("row_filter deve usare equals o regex")


def reading_report(
    rows: list[tuple[Any, ...]],
    data_start: int,
    data_end: int,
    excluded: Counter[str],
    records: list[dict[str, Any]],
) -> dict[str, Any]:
    """Summarize a read: how much was read and how much was dropped.

    For whoever reviews the result: if excluded rows aren't reported
    anywhere, no one notices a price list came in only half-read.
    """
    return {
        "sheet_rows": len(rows),
        "data_start_row": data_start,
        "data_end_row": data_end,
        "rows_ignored_after_data_end": max(0, len(rows) - data_end),
        "rows_excluded": dict(excluded),
        "rows_kept": len(records),
        "expiry_dates_read": sum(1 for record in records if record.get("expiry_date")),
        "expiry_dates_to_check": sum(1 for record in records if record.get("expiry_plausible") is False),
    }


def read_mapped_master(
    path: Path,
    mapping: dict[str, Any],
    report: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    data_start, columns, rows, _bold = mapped_rows(path, mapping)
    required = {"ean", "description", "last_unit_price"}
    missing = required - set(columns)
    if missing:
        raise ValueError(f"Mappatura master incompleta: {sorted(missing)}")
    records = []
    excluded: Counter[str] = Counter()
    for source_row, row in enumerate(rows[data_start - 1:], start=data_start):
        if not row_allowed(row, columns, mapping):
            excluded[FUORI_DAL_FILTRO] += 1
            continue
        description = str(mapped_value(row, columns, "description") or "").strip()
        ean = normalize_ean(mapped_value(row, columns, "ean"))
        if not description and not ean:
            excluded[NON_E_RIGA_PRODOTTO] += 1
            continue
        records.append({
            "source": "gestionale",
            "source_row": source_row,
            "ean": ean,
            "description": description,
            "unit": str(mapped_value(row, columns, "unit") or "").strip(),
            "suggested_colli": mapped_value(row, columns, "suggested_colli"),
            "source_quantity_ignored": mapped_value(row, columns, "source_quantity_ignored"),
            "manual_quantity": None,
            "last_unit_price": json_decimal(decimal_value(mapped_value(row, columns, "last_unit_price"), italian=bool(mapping.get("italian_numbers")))),
            "source_discount": mapped_value(row, columns, "source_discount"),
            "vat": mapped_value(row, columns, "vat"),
        })
    if not records:
        raise ValueError(f"Nessuna riga master letta da {path}")
    if report is not None:
        report.update(reading_report(rows, data_start, len(rows), excluded, records))
    return records


def supplier_record(
    supplier_id: str,
    source_row: int,
    row: tuple[Any, ...],
    columns: dict[str, int],
    mapping: dict[str, Any],
    bold: tuple[bool, ...] | None = None,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    description = str(mapped_value(row, columns, "description") or "").strip()
    description, expiry = expiry_from_description(description, mapping)
    ean = normalize_ean(mapped_value(row, columns, "ean"))
    supplier_code = mapped_value(row, columns, "supplier_code")
    if "supplier_code" in (mapping.get("text_columns") or ()):
        # One supplier's item code has leading zeros: if it were coerced to a
        # number, "0000000449070" would shrink to 449070 and never recover
        # them. normalize_ean only strips spreadsheet artifacts and leaves
        # the rest of the text as written.
        supplier_code = normalize_ean(supplier_code)
    if not description and not ean and supplier_code in (None, ""):
        return None, None

    italian = bool(mapping.get("italian_numbers"))
    direct_price = decimal_value(mapped_value(row, columns, "unit_price_net"), italian=italian)
    pre_price = decimal_value(mapped_value(row, columns, "unit_price_pre_discount"), italian=italian)
    discount_raw = mapped_value(row, columns, "discount")
    discount_rate = Decimal("0")
    discount_type = "none"
    warning = None
    if pre_price is not None:
        if mapping.get("discount_mode") == "numeric_percent_or_text_zero":
            discount_rate, discount_type, warning = larice_discount(
                discount_raw, registro.codici_ammessi(registro.codici_di_riga(mapping))
            )
        elif discount_raw not in (None, ""):
            numeric_discount = decimal_value(discount_raw, italian=italian)
            if numeric_discount is None:
                raise ValueError(f"Sconto non interpretabile alla riga {source_row}: {discount_raw!r}")
            discount_rate = numeric_discount / 100 if numeric_discount > 1 else numeric_discount
            discount_type = "percentuale"
        direct_price = pre_price * (Decimal("1") - discount_rate)

    pieces = decimal_value(mapped_value(row, columns, "pieces_per_carton"), italian=italian)
    multiplier = decimal_value(mapped_value(row, columns, "order_multiplier"), italian=italian)
    # A fixed default from the supplier profile only counts if it's a number
    # greater than zero: a zero, or text where a number was expected, isn't
    # an order factor and must not stand in for one.
    if pieces is None:
        pieces = fattore_d_ordine(mapping.get("pieces_per_carton_default"))
    if multiplier is None:
        multiplier = fattore_d_ordine(mapping.get("order_multiplier_default"))
    # The multiplier still wins when present: a badly written multiplier is
    # not silently replaced by pieces-per-carton, since that would mean
    # ordering against a number nobody declared.
    factor = fattore_d_ordine(multiplier if multiplier is not None else pieces)
    available_field = mapped_value(row, columns, "availability")
    availability = str(available_field or "").strip()
    available = True
    if mapping.get("available_values"):
        available = availability.casefold() in {str(value).casefold() for value in mapping["available_values"]}

    # Same condition as the earlier drop check, but able to say what was
    # missing: this is what makes a dropped row auditable.
    if not description:
        motivo: str | None = SENZA_DESCRIZIONE
    elif direct_price is None:
        motivo = SENZA_PREZZO
    elif factor is None:
        motivo = SENZA_PEZZI_PER_COLLO
    elif not available:
        motivo = NON_DISPONIBILE
    else:
        motivo = None

    record = {
        "source": supplier_id,
        "source_row": source_row,
        "ean": ean,
        "supplier_code": supplier_code,
        "description": description,
        "pieces_per_carton": json_decimal(pieces),
        "order_multiplier": json_decimal(multiplier),
        "unit_price_pre_discount": json_decimal(pre_price),
        "discount_raw": discount_raw,
        "discount_type": discount_type,
        "discount_rate": json_decimal(discount_rate),
        "unit_price_net": json_decimal(direct_price),
        "pallet": mapped_value(row, columns, "pallet"),
        "packaging": mapped_value(row, columns, "packaging"),
        "availability": availability,
        "unit": str(mapped_value(row, columns, "unit") or "").strip(),
        "vat": mapped_value(row, columns, "vat"),
        "order_column": mapping.get("order_column"),
        "usable": motivo is None,
    }
    if motivo is not None:
        record["unusable_reason"] = motivo
    if "category" in columns:
        # The field the filter checks is also shown on rows that pass it:
        # that's how the user verifies the filter did the right thing.
        record["category"] = str(mapped_value(row, columns, "category") or "").strip()
    if expiry is not None:
        record.update({field: value for field, value in expiry.items() if field != "expiry_warning"})

    bold_field = mapping.get("offer_from_bold")
    if bold_field or "offer_flag" in columns:
        # This list flags an offer two ways: an explicit column and bold
        # price text. Either one is enough, and both are kept in the record.
        bold_index = columns.get(str(bold_field)) if bold_field else None
        bold_price = bool(bold_index and bold and bold_index <= len(bold) and bold[bold_index - 1])
        declared = str(mapped_value(row, columns, "offer_flag") or "").strip()
        record["offer_flag"] = declared
        record["offer_price_bold"] = bold_price
        record["offer"] = bold_price or declared.casefold() in {"si", "sì", "s", "true", "1"}

    messages = [text for text in (warning, (expiry or {}).get("expiry_warning")) if text]
    warning_record = (
        {"source": supplier_id, "source_row": source_row, "ean": ean, "warning": " ".join(messages)}
        if messages
        else None
    )
    return record, warning_record


def read_mapped_xlsx_supplier(
    path: Path,
    supplier_id: str,
    mapping: dict[str, Any],
    report: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Read a supplier price list from a spreadsheet: .xlsx or .xls.

    Still returns two values because ``app/catalog_search.py`` also calls
    this function; read stats (data boundary, dropped rows) go into
    ``report`` when the caller passes one.
    """
    data_start, columns, rows, bold = mapped_rows(path, mapping)
    required = {"description", "unit_price_net"} if "unit_price_net" in columns else {"description", "unit_price_pre_discount"}
    missing = required - set(columns)
    if missing:
        raise ValueError(f"Mappatura fornitore incompleta: {sorted(missing)}")
    if not ({"pieces_per_carton", "order_multiplier"} & set(columns)):
        dichiarati = {"pieces_per_carton_default", "order_multiplier_default"} & set(mapping)
        if not dichiarati:
            raise ValueError("Manca il fattore d'ordine o un default esplicito")
        if not any(fattore_d_ordine(mapping.get(chiave)) is not None for chiave in dichiarati):
            # A fixed default that isn't a number greater than zero would
            # leave every row of the list unorderable, one by one; better to
            # say so once, up front.
            raise ValueError(
                f"Per {path.name} non è indicata nessuna colonna con i pezzi per collo, e il valore "
                "fisso messo al suo posto non è un numero maggiore di zero: così nessun prodotto di "
                "questo listino potrebbe essere ordinato. Indicare la colonna dei pezzi per collo, "
                "oppure un valore fisso valido (per esempio 6)."
            )
    data_end = last_data_row(rows, columns, mapping, data_start)
    records = []
    warnings = []
    excluded: Counter[str] = Counter()
    for source_row, row in enumerate(rows[data_start - 1:data_end], start=data_start):
        if not row_allowed(row, columns, mapping):
            excluded[FUORI_DAL_FILTRO] += 1
            continue
        label = excluded_row_label(row, columns, mapping)
        if label:
            excluded[label] += 1
            continue
        if non_e_una_riga_prodotto(row, columns, mapping):
            excluded[NON_E_RIGA_PRODOTTO] += 1
            continue
        record, warning = supplier_record(
            supplier_id,
            source_row,
            row,
            columns,
            mapping,
            bold=bold[source_row - 1] if source_row - 1 < len(bold) else None,
        )
        if record:
            records.append(record)
        if warning:
            warnings.append(warning)
    registro.applica_codici_di_riga(records, registro.codici_di_riga(mapping))
    if report is not None:
        report.update(reading_report(rows, data_start, data_end, excluded, records))
        report["rows_not_orderable"] = conta_non_ordinabili(records)
    return records, warnings


def mapped_standalone_displays(
    records: list[dict[str, Any]],
    supplier_id: str,
    mapping: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any] | None]:
    """Separate standalone display rows declared by a reviewed supplier mapping.

    A new supplier can list an already assembled display on a single orderable
    row, without the child rows available in the Larice canvass.  The mapping
    must opt in explicitly: a product is never reclassified merely because its
    name happens to contain a similar word.
    """

    rules = mapping.get("display_detection")
    if not isinstance(rules, dict) or not rules.get("description_regex"):
        return records, [], None

    try:
        description_pattern = re.compile(str(rules["description_regex"]), re.IGNORECASE)
        excluded_pattern = (
            re.compile(str(rules["exclude_description_regex"]), re.IGNORECASE)
            if rules.get("exclude_description_regex")
            else None
        )
        code_pattern = (
            re.compile(str(rules["supplier_code_regex"]), re.IGNORECASE)
            if rules.get("supplier_code_regex")
            else None
        )
        declared_pattern = (
            re.compile(str(rules["declared_units_regex"]), re.IGNORECASE)
            if rules.get("declared_units_regex")
            else None
        )
    except re.error as exc:
        raise ValueError(f"Regola espositore non valida per {supplier_id}: {exc}") from exc

    expected_factor = decimal_value(rules.get("order_factor_equals"), italian=bool(mapping.get("italian_numbers")))
    standard: list[dict[str, Any]] = []
    displays: list[dict[str, Any]] = []
    rejected = 0
    for record in records:
        description = str(record.get("description") or "").strip()
        supplier_code = str(record.get("supplier_code") or "").strip()
        if not description_pattern.search(description):
            standard.append(record)
            continue
        if excluded_pattern and excluded_pattern.search(description):
            rejected += 1
            standard.append(record)
            continue
        if code_pattern and not code_pattern.search(supplier_code):
            rejected += 1
            standard.append(record)
            continue
        # The display's order factor is the source row's: the multiplier
        # when present, otherwise pieces-per-carton — same precedence as
        # `supplier_record`, since it's the same row. It goes through the
        # same guard as everything else: zero isn't a factor.
        grezzo = record.get("order_multiplier")
        if grezzo is None:
            grezzo = record.get("pieces_per_carton")
        factor = fattore_d_ordine(grezzo)
        if expected_factor is not None and factor != expected_factor:
            rejected += 1
            standard.append(record)
            continue

        declared_units: Decimal | None = None
        if declared_pattern:
            match = declared_pattern.search(description)
            if match:
                numbers = re.findall(r"[0-9]+(?:[.,][0-9]+)?", match.group(0))
                values = [decimal_value(value, italian=True) for value in numbers]
                if values and all(value is not None for value in values):
                    declared_units = sum((value for value in values if value is not None), Decimal("0"))

        evidence = [
            "Riga singola con nome espositore e codice fornitore dedicato.",
            "Non sono presenti righe componente contigue da riconciliare.",
        ]
        if declared_units is not None:
            evidence.append(f"Quantità dichiarata nel nome: {json_decimal(declared_units)} pezzi.")
        # A display is never more usable than its source row: if that row
        # can't be ordered (unreadable price, missing factor, a free-goods
        # line), neither can the display, for the same reason. Without this
        # check the display reached the comparison with no statement on
        # whether it could actually be bought.
        if record.get("usable") is False:
            motivo = str(record.get("row_type") or record.get("unusable_reason") or MOTIVO_NON_DICHIARATO)
        elif factor is None:
            motivo = SENZA_PEZZI_PER_COLLO
        else:
            motivo = None
        displays.append({
            "supplier": supplier_id,
            "source_row": record.get("source_row"),
            "supplier_code": record.get("supplier_code"),
            "description": description,
            "ean": record.get("ean") or "",
            "unit_price_net": record.get("unit_price_net"),
            "net_price_per_display": record.get("unit_price_net"),
            "declared_units": json_decimal(declared_units),
            "components": [],
            "confidence": str(rules.get("confidence") or "MEDIA").upper(),
            "quantity_reconciled": False,
            "price_reconciled": False,
            "evidence": evidence,
            "usable": motivo is None,
            **({} if motivo is None else {"unusable_reason": motivo}),
        })

    audit = {
        "rows_scanned": len(records),
        "standalone_display_count": len(displays),
        "standard_rows": len(standard),
        "rejected_candidates": rejected,
        "component_rows_available": False,
        # Same as the other supplier's displays: whatever stays out of the
        # order is counted, with its reason recorded.
        "display_offers_not_orderable": conta_non_ordinabili(displays),
    }
    return standard, displays, audit


def read_mapped_csv_supplier(
    path: Path,
    supplier_id: str,
    mapping: dict[str, Any],
    report: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Read a supplier price list from CSV, using the same declared rules as the spreadsheet reader.

    The data boundary, excluded rows and expiry-in-description parsing all
    apply here too: a rule that only works for one format is a trap for
    whoever maintains this later.

    `header_row = 0` declares the file has no header row: columns are then
    addressed by number or letter, as in the spreadsheet branch. Without this
    escape hatch, a headerless CSV couldn't be mapped at all, since its first
    data row would be mistaken for the header.

    If reading produces no orderable row at all, it stops and says so: a
    price list opened with the wrong delimiter looks, from the result alone,
    exactly like a supplier that carries none of the products being compared.
    """
    encoding = mapping.get("encoding", "utf-8-sig")
    delimiter = mapping.get("delimiter")
    with path.open("r", encoding=encoding, newline="") as stream:
        sample = stream.read(65536)
        stream.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if not delimiter else None
        reader = csv.reader(stream, dialect) if dialect else csv.reader(stream, delimiter=delimiter)
        rows = []
        # The row number recorded is the file's real line number, not the
        # record's position in the list: a quoted field can contain an
        # embedded newline, so the two counts diverge, and whoever checks
        # the order against the supplier's file needs the real line. `line_num`
        # after a record gives its last line; the first line is right after
        # the previous record ended.
        righe_fisiche: list[int] = []
        prima_riga_del_prodotto = 1
        for grezza in reader:
            rows.append(grezza)
            righe_fisiche.append(prima_riga_del_prodotto)
            prima_riga_del_prodotto = reader.line_num + 1
    grezzo = mapping.get("header_row")
    header_row = 1 if grezzo in (None, "") else int(grezzo)
    if header_row < 0 or header_row > len(rows):
        raise ValueError(
            f"La mappatura dichiara la riga di intestazione {header_row}, ma {path.name} ha "
            f"{len(rows)} righe: usare 0 se il documento non ha intestazioni."
        )
    header = rows[header_row - 1] if header_row else None
    raw_columns = mapping.get("columns") or mapping.get("field_mapping") or {}
    if header is not None and len(header) <= 1 and len(raw_columns) > 1:
        # When not declared, the delimiter is guessed by sniffing the text;
        # a wrong guess leaves the whole file in one column, and the
        # supplier would appear to carry none of the products compared.
        raise ValueError(
            f"Il documento {path.name} non si è aperto in colonne: nella riga delle intestazioni "
            f"c'è un campo solo, mentre per questo fornitore ne sono indicate {len(raw_columns)}. "
            "Di solito vuol dire che il separatore delle colonne non è quello previsto: indicare "
            "quello giusto (per esempio «;» oppure «,») nella configurazione del fornitore."
        )
    columns = {field: column_number(spec, header) for field, spec in raw_columns.items() if spec not in (None, "")}
    tuples = [tuple(raw) for raw in rows]
    data_start = prima_riga_dei_dati(tuples, mapping, header_row, path.name)
    data_end = last_data_row(tuples, columns, mapping, data_start)
    records = []
    warnings = []
    excluded: Counter[str] = Counter()
    for posizione, row in enumerate(tuples[data_start - 1:data_end], start=data_start):
        if not row_allowed(row, columns, mapping):
            excluded[FUORI_DAL_FILTRO] += 1
            continue
        label = excluded_row_label(row, columns, mapping)
        if label:
            excluded[label] += 1
            continue
        if non_e_una_riga_prodotto(row, columns, mapping):
            excluded[NON_E_RIGA_PRODOTTO] += 1
            continue
        record, warning = supplier_record(
            supplier_id, righe_fisiche[posizione - 1], row, columns, mapping
        )
        if record:
            records.append(record)
        if warning:
            warnings.append(warning)
    # This check runs before the registry rules: those drop rows that were
    # read perfectly well (a bonus item, a free-goods line) and say nothing
    # about whether the read itself succeeded.
    if not any(record.get("usable") for record in records):
        letto = (
            "non è stata letta nessuna riga di prodotto" if not records
            else f"sono state lette {len(records)} righe, ma nessuna è risultata ordinabile"
        )
        raise ValueError(
            f"Da {path.name} {letto}: il documento non ha la forma attesa, oppure il separatore "
            "delle colonne o la riga delle intestazioni non sono quelli indicati. Conviene "
            "controllare il file e la configurazione di questo fornitore: un listino letto a vuoto "
            "lo farebbe risultare senza nessuno dei prodotti cercati, che è una cosa diversa da "
            "un listino non letto."
        )
    registro.applica_codici_di_riga(records, registro.codici_di_riga(mapping))
    if report is not None:
        report.update(reading_report(tuples, data_start, data_end, excluded, records))
        report["rows_not_orderable"] = conta_non_ordinabili(records)
    return records, warnings


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--adapters", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    # Groups of barcodes the user has declared to be the same item. The
    # server writes this file from the confirmations store: the pipeline is
    # a chain of steps reading declared artifacts, not processes opening a
    # database directly, which is why this phase is independently testable.
    parser.add_argument("--equivalenze", type=Path)
    return parser.parse_args()


def leggi_uguaglianze(percorso: Path | None) -> list[list[str]]:
    """Return the user-declared groups of equal codes, or none.

    Not asking for the argument is legitimate: it means this run has none.
    Asking and not finding the file is not: these are human decisions, and
    silently losing them would revert manually matched products to their
    unmatched state with nothing to say so — the worst way to lose a
    decision. Same rule, and same reason, as `--displays` in
    `build_review_data.py`.

    The document is `{"classi": [[code, code, …], …]}`, or directly the list
    of groups, so a hand-written test fixture doesn't fail over one key.
    """

    if percorso is None:
        return []
    if not percorso.exists():
        raise ValueError(
            f"Il file delle uguaglianze fra codici non esiste: {percorso}. È stato chiesto, "
            "quindi qualcuno doveva scriverlo: senza, i prodotti abbinati a mano tornerebbero "
            "a non avere quel fornitore, e nessuno lo direbbe."
        )
    contenuto = json.loads(percorso.read_text(encoding="utf-8"))
    gruppi = contenuto.get("classi") if isinstance(contenuto, dict) else contenuto
    if not isinstance(gruppi, list):
        raise ValueError(f"Il file delle uguaglianze {percorso} non contiene una lista di gruppi")
    return [
        [str(codice) for codice in gruppo]
        for gruppo in gruppi
        if isinstance(gruppo, list) and len(gruppo) > 1
    ]


def colonne_corrette(
    decision: dict[str, Any], adapter: dict[str, Any], adapter_id: str
) -> dict[str, Any]:
    """Return the columns the user manually corrected for a recognized document.

    Only the decision's own corrections: a registry adapter's `field_mapping`
    describes a different way to read the same supplier, not a correction,
    and feeding it to the dedicated readers would silently conflate the two.

    `read_noce` reads a CSV and addresses columns by name, so it has nothing
    to correct by position and is excluded.
    """

    # By `adattatore_base`, not the raw id: `betulla_v1__locale` (a confirmed
    # mapping layered on the shipped one) uses the same reader and the same
    # corrections. Matching on the full id would mean that learning a
    # variant silently drops the corrections that variant exists to carry.
    base = registro.adattatore_base(adapter_id)
    if base == "noce_csv_v1":
        return {}
    mappatura = decision.get("field_mapping")
    if not isinstance(mappatura, dict) or not mappatura:
        return {}
    colonne = mappatura.get("columns")
    extra: dict[str, Any] = {}
    if isinstance(colonne, dict) and colonne:
        extra["colonne"] = colonne_per_posizione(colonne, adapter, adapter_id)
    if base in {"betulla_v1", "larice_v1"}:
        ordine = mappatura.get("order_column")
        if ordine:
            extra["order_column"] = str(ordine)
    return extra


def colonne_per_posizione(
    colonne: dict[str, Any], adapter: dict[str, Any], adapter_id: str
) -> dict[str, Any]:
    """Return the decision's columns with names resolved to positions.

    The mapping confirmed on the page addresses a column by name whenever the
    header is unambiguous (`schema_mapping.specifica_colonna`), but the
    dedicated readers read by position. Passing names straight through to
    them would fail the following week with "column X is not valid" — the
    supplier dropped from the comparison by exactly the mapping meant to keep
    it in.

    Positions aren't guessed from the document: they're what the adapter
    measured the day the variant was confirmed (`header_signature.columns`,
    by normalized name). Name is checked before letter, because a header like
    "EAN" is also a valid (and very different) Excel column letter.

    Numbers and letters pass through unchanged. A name the adapter's
    fingerprint doesn't recognize stops here: letting it through would mean
    silently reading the default column instead — prices that merely look
    plausible.
    """

    posizioni = ((adapter or {}).get("header_signature") or {}).get("columns")
    posizioni = posizioni if isinstance(posizioni, dict) else {}
    risolte: dict[str, Any] = {}
    for campo, dichiarata in colonne.items():
        if not isinstance(dichiarata, str):
            risolte[campo] = dichiarata
            continue
        indice = posizioni.get(registro.normalizza(dichiarata))
        if indice is None and indice_di_colonna(dichiarata) is None:
            raise ValueError(
                f"La colonna «{dichiarata}» indicata per «{campo}» non è fra le intestazioni con "
                f"cui lo schema «{adapter_id}» è stato imparato: il listino è cambiato ancora, e "
                "la mappatura di questo fornitore va rifatta dalla pagina Importa."
            )
        risolte[campo] = dichiarata if indice is None else indice
    return risolte


LETTORI_A_SCHEMA_NOTO: dict[str, Callable[..., Any]] = {
    "gestionale_v1": read_gestionale,
    "betulla_v1": read_betulla,
    "larice_v1": read_larice,
    "noce_csv_v1": read_noce,
}


def lettore_dedicato(state: Any, adapter_id: Any) -> Callable[..., Any] | None:
    """Return the dedicated reader for this decision, or `None` for the generic reader.

    This is the one place that picks which reader opens a document. The
    product-search catalog calls this rather than keeping its own
    `if supplier == "..."` logic: if the two answered differently, a document
    read generically here but sent to a dedicated reader by the catalog would
    silently vanish from the catalog viewer.

    The criterion is the decision, not the supplier: the same supplier, in
    two different runs, can arrive with a recognized schema or a varied one
    confirmed by hand, and those are two different reads of the same name.
    """

    if str(state or "") != "SCHEMA_NOTO":
        return None
    # By `adattatore_base`: `larice_v1__locale` is still that supplier and
    # must go through its dedicated reader. Without this, learning a variant
    # on one of the suppliers with a dedicated reader would route it to the
    # generic reader instead, which reads the same columns but skips
    # everything else — no display detection, no free-goods thresholds, no
    # master row marker.
    return LETTORI_A_SCHEMA_NOTO.get(registro.adattatore_base(adapter_id))


def main() -> int:
    args = parse_args()
    manifest = load_json(args.manifest)
    # The merged registry, not just the shipped file the pipeline passes
    # here: a locally learned adapter (`betulla_v1__locale`) isn't in the
    # shipped file, and looking it up there returns nothing, i.e. no
    # fallback `supplier_id` and no column positions to reread a variant
    # confirmed the week before. Same invariant as in
    # `validate_input_manifest.py`: reading the registry always goes through
    # `registro`.
    adapters = {str(item.get("id") or ""): item for item in registro.adattatori(args.adapters)}
    files = manifest.get("files") or manifest.get("profiles") or []
    master = None
    sources: dict[str, list[dict[str, Any]]] = {}
    warnings: list[dict[str, Any]] = []
    display_offers: dict[str, list[dict[str, Any]]] = {}
    display_audit: dict[str, dict[str, Any]] = {}
    input_audit = []

    for item in files:
        decision = item.get("ai_preflight") or {}
        state = decision.get("state")
        if state in {"FILE_NON_PERTINENTE", "AMBIGUO"}:
            continue
        path = Path(item["path"])
        role = decision.get("role")
        adapter_id = decision.get("adapter_id")
        supplier_id = decision.get("supplier_id")
        adapter = adapters.get(adapter_id, {})
        mapping = decision.get("field_mapping") or adapter.get("field_mapping") or {}
        if role == "supplier" and not supplier_id:
            supplier_id = adapter.get("supplier_id")
        records: Any
        local_warnings: list[dict[str, Any]] = []
        reading: dict[str, Any] = {}

        lettore = lettore_dedicato(state, adapter_id)
        if lettore is not None:
            # Dedicated readers report read stats too: without `report`, the
            # excluded-rows panel stayed empty for exactly the largest price
            # lists (the ones with a dedicated reader).
            #
            # Manually corrected columns are fed into the dedicated reader
            # itself, never by routing the document to the generic reader
            # instead. The generic reader reads the same columns but skips
            # everything else — no display detection, no free-goods
            # thresholds, no master row marker. A column correction must not
            # change what the program knows how to do with that document.
            extra = colonne_corrette(decision, adapter, adapter_id)
            result = lettore(path, report=reading, **extra)
            if registro.adattatore_base(adapter_id) == "larice_v1":
                records, local_warnings = result
                records, supplier_displays, supplier_display_audit = integrate_larice_displays(records, analyse_workbook(path))
                display_offers[supplier_id or "larice"] = supplier_displays
                display_audit[supplier_id or "larice"] = supplier_display_audit
            else:
                records = result
        elif role == "master":
            records = read_mapped_master(path, mapping, report=reading)
        # File format is decided by the file's bytes: a spreadsheet renamed
        # to .csv is still a spreadsheet, and reading it line by line as text
        # would silently produce an empty price list.
        elif container_format(path) == "csv":
            records, local_warnings = read_mapped_csv_supplier(path, supplier_id, mapping, report=reading)
        else:
            records, local_warnings = read_mapped_xlsx_supplier(path, supplier_id, mapping, report=reading)

        if role == "supplier" and mapping.get("display_detection"):
            records, supplier_displays, supplier_display_audit = mapped_standalone_displays(records, supplier_id, mapping)
            if supplier_displays:
                display_offers[supplier_id] = supplier_displays
                display_audit[supplier_id] = supplier_display_audit or {}

        if role == "master":
            if master is not None:
                raise ValueError("Manifest con più di un master")
            master = records
        elif role == "supplier":
            if not supplier_id:
                adapter = adapters.get(adapter_id, {})
                supplier_id = adapter.get("supplier_id")
            if not supplier_id:
                raise ValueError(f"supplier_id mancante per {path}")
            if supplier_id in sources:
                raise ValueError(f"Fornitore duplicato: {supplier_id}")
            for record in records:
                record["source"] = supplier_id
            sources[supplier_id] = records
        warnings.extend(local_warnings)
        # Rows can also become unorderable after the read itself (a
        # display's component rows, a supplier's standalone displays), so
        # this count is recomputed here, once all mutations are done, and is
        # the same number reported in both places in the audit.
        non_ordinabili = conta_non_ordinabili(records if isinstance(records, list) else [])
        if reading:
            reading["rows_not_orderable"] = non_ordinabili
        input_audit.append({
            "path": str(path.resolve()),
            "state": state,
            "role": role,
            "supplier_id": supplier_id,
            "adapter_id": adapter_id,
            "records": len(records),
            # How many rows the program decided not to make orderable, and
            # with what declared reason: free-goods threshold rows, a
            # display's components, but also an unreadable price or a
            # missing order factor. A drop that isn't counted isn't a choice.
            "rows_not_orderable": non_ordinabili,
            "content_format": container_format(path),
            "reading": reading,
        })

    if master is None:
        raise ValueError("Master gestionale non trovato nel manifest")
    if not sources:
        raise ValueError("Nessun fornitore incluso nel manifest")

    matching, semantic_queue, audit = build_matching(master, sources, uguaglianze=leggi_uguaglianze(args.equivalenze))
    audit["warnings"] = warnings
    audit["inputs"] = input_audit
    audit["displays"] = display_audit
    args.output.mkdir(parents=True, exist_ok=True)
    write_json(args.output / "normalized_sources.json", {"gestionale": master, **sources, "display_offers": display_offers})
    write_json(args.output / "display_offers.json", display_offers)
    write_json(args.output / "matching_result.json", matching)
    write_json(args.output / "semantic_queue.json", semantic_queue)
    write_json(args.output / "audit.json", audit)
    print(json.dumps(audit, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
