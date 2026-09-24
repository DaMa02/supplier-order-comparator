#!/usr/bin/env python3
"""Verify that the copy delivered to a supplier is faithful to their price list.

This module exists because a writer reporting success only proves it wrote
the cells it meant to write — it says nothing about what openpyxl may have
lost or altered elsewhere while re-serializing the whole workbook. Copy
fidelity can't depend on the good behavior of the library that writes the
file, so every fill gets an independent proof: reopen both files and
compare them cell by cell.

READING BOTH FILES
Both are loaded in full, never with `read_only`. openpyxl's fast mode
trusts the `<dimension>` element and row order declared in the file; a
legitimate price list with a conservative `<dimension>` or out-of-order
rows would make a perfect copy look wrong. A full load ignores the
declaration and places every cell at its real row number. Cost on real
price lists: under a second per workbook.

THE ONLY DIFFERENCES ALLOWED
Only inside the order column, from the first data row down, and only of
two kinds:

1. the plan's quantity, on the rows the plan touches;
2. zeroing a quantity that was already there — without this, stale
   quantities from a previous run would ship as phantom order lines.

Everything else: not one cell. A deleted section header, a changed EAN, a
shifted price, a missing sheet — any of those and the copy for that
supplier is not delivered, with the reason stated.

WHAT THE PLAN ASKS FOR MUST BE THERE
Looking only for unwanted differences answers "is there anything extra?"
and misses the mirror question, "is everything that's needed there?". An
order line the writer fails to write isn't a difference — the copy's cell
stays identical to the original — so a diff-only check would call the copy
faithful while it ships with an order missing. The plan is the contract:
every row `quantita` requests must be in the copy, showing the number the
plan says. Not "must have changed": if the source price list already had
that number — the same order as the previous run — the copy is correct
without there being any difference to count. What matters is the value the
supplier reads.

These are two distinct failures and the message names them separately (a
cell changed outside the order column vs. a missing order row), because
they point to different places to investigate.

A quantity is a number. Same rule `app/xls_writer.py` applies to its `.xls`
target, and it's what distinguishes a quantity to zero from a section
header to leave alone.

NUMBER FORMAT IS PART OF THE COMPARISON, WHERE THE COPY SHOWS A VALUE
The supplier doesn't read the file's internal value, they read the
displayed number, and the number format decides how it's shown: a price of
`1.75` under format `0` displays as `2`; a quantity under format `;;;`
displays as blank. So where the copy shows a visible value, its number
format must match the source price list. Where the copy is blank, no
format displays anything, and a format mismatch there doesn't block a
correct order.

WHAT THIS COMPARISON DELIBERATELY DOES NOT CHECK
Anything that doesn't change a displayed value: colors and fonts, column
widths, merged cells, images, hyperlinks, autofilters, print settings,
document properties, hidden sheets. Also cached formula results: formulas
are compared as text (`data_only=False`), and the library recalculates
results anyway on export. The library reshuffles these parts of the file
regardless of content, and telling a real loss from a harmless
normalization there would mean parsing the OpenXML by hand. What's
defended here is what the supplier actually reads: numbers, text, and
their numeric formatting.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Mapping

# How many differences get named in the message before falling back to a
# count: a page-long warning goes unread.
ESEMPI_NEL_MESSAGGIO = 3

# Number format of cells that don't declare one.
FORMATO_PREDEFINITO = "General"


class ConfrontoImpossibile(Exception):
    """The two documents can't even be opened for comparison."""


@dataclass
class EsitoFedelta:
    """What was compared and what doesn't match.

    Two separate measures for two separate failures: `quante_rifiutate`
    counts unwanted changes (cells changed where they shouldn't be),
    `righe_mancanti` lists what's missing (rows the plan requested that the
    copy doesn't carry). The copy is delivered only if both are empty.
    """

    celle_confrontate: int = 0
    differenze_ammesse: int = 0
    quante_rifiutate: int = 0
    esempi_rifiutati: list[str] = field(default_factory=list)
    # Rows the plan requested on the order sheet, and which of them the
    # copy doesn't carry. Row numbers as the user sees them.
    righe_richieste: int = 0
    righe_mancanti: list[int] = field(default_factory=list)

    @property
    def fedele(self) -> bool:
        return self.quante_rifiutate == 0 and not self.righe_mancanti


def _vuoto(valore: Any) -> bool:
    """None, an empty string and whitespace-only all count as an empty cell.

    Not just a convenience: the source price list can carry an empty
    shared string, which Excel displays identically to a truly empty cell,
    while the copy rewrites it as a truly empty cell. Treating that as a
    difference would block a correct fill over something invisible to the
    user.
    """

    if valore is None:
        return True
    return isinstance(valore, str) and not valore.strip()


def _numero(valore: Any) -> float | None:
    """The value as a number, if it is one. A boolean is not a quantity."""

    if isinstance(valore, bool) or valore is None:
        return None
    if isinstance(valore, (int, float)):
        return float(valore)
    if isinstance(valore, str):
        testo = valore.strip().replace(",", ".")
        if not testo:
            return None
        try:
            return float(testo)
        except ValueError:
            return None
    return None


def lettera_di_colonna(numero: int) -> str:
    lettere = ""
    while numero > 0:
        numero, resto = divmod(numero - 1, 26)
        lettere = chr(ord("A") + resto) + lettere
    return lettere


def numero_di_colonna(lettera: str) -> int:
    testo = str(lettera or "").strip().upper()
    if not testo or not testo.isalpha():
        raise ConfrontoImpossibile(f"Colonna d'ordine non valida: {lettera!r}.")
    numero = 0
    for carattere in testo:
        numero = numero * 26 + (ord(carattere) - ord("A") + 1)
    return numero


def _mostra(valore: Any) -> str:
    if _vuoto(valore):
        return "vuota"
    testo = str(valore)
    return f"«{testo}»" if len(testo) <= 40 else f"«{testo[:40]}…»"


def confronta_copia(
    originale: Path,
    copia: Path,
    *,
    colonna_ordine: str,
    prima_riga: int,
    quantita: Mapping[int, int],
    foglio_ordine: str | None = None,
) -> EsitoFedelta:
    """Compare the copy with the source price list, cell by cell.

    `quantita` are the rows the plan asked to write (row number as the user
    sees it -> carton count). It's both the tolerance (a difference is
    allowed there) and the contract (the quantity must be there); see
    `_righe_del_piano_mancanti`. `foglio_ordine` names the sheet holding
    the order column: on every other sheet, no difference is allowed, not
    even in the column with the same letter.
    """

    from openpyxl import load_workbook

    esito = EsitoFedelta()
    colonna = numero_di_colonna(colonna_ordine)

    try:
        libro_originale = load_workbook(originale, read_only=False, data_only=False)
    except Exception as errore:  # noqa: BLE001 - whatever the cause, this is the message
        raise ConfrontoImpossibile(
            f"non riesco a riaprire il listino di partenza «{Path(originale).name}» "
            f"per confrontarlo con la copia: {errore}"
        ) from errore
    try:
        try:
            libro_copia = load_workbook(copia, read_only=False, data_only=False)
        except Exception as errore:  # noqa: BLE001
            raise ConfrontoImpossibile(
                f"non riesco a riaprire la copia «{Path(copia).name}» per confrontarla "
                f"con il listino di partenza: {errore}"
            ) from errore
        try:
            nomi_originale = list(libro_originale.sheetnames)
            nomi_copia = list(libro_copia.sheetnames)
            if nomi_originale != nomi_copia:
                esito.quante_rifiutate = 1
                esito.esempi_rifiutati.append(
                    "i fogli della copia non sono quelli del listino di partenza "
                    f"({', '.join(nomi_copia) or 'nessuno'} invece di {', '.join(nomi_originale)})"
                )
                return esito

            # The order sheet is the one the caller declares; `FIRST` or a
            # missing name mean "the first one", same as for the writer.
            atteso = str(foglio_ordine or "").strip()
            if not atteso or atteso.upper() == "FIRST" or atteso not in nomi_originale:
                atteso = nomi_originale[0] if nomi_originale else ""

            for nome in nomi_originale:
                _confronta_foglio(
                    libro_originale[nome],
                    libro_copia[nome],
                    nome=nome if len(nomi_originale) > 1 else "",
                    colonna_ordine=colonna if nome == atteso else None,
                    prima_riga=prima_riga,
                    quantita=quantita,
                    esito=esito,
                )
        finally:
            libro_copia.close()
    finally:
        libro_originale.close()
    return esito


def _celle_significative(foglio: Any) -> dict[tuple[int, int], tuple[Any, str]]:
    """`{(row, column): (value, format)}` for every cell that carries something.

    A full load already puts a file with out-of-order rows back in order
    (the parser stores each cell at its real coordinate, while `iter_rows`
    always starts from row 1). The key still uses `cell.row`, the row
    number the cell itself declares, so the comparison doesn't depend on
    read order even if the loading strategy changes later. A cell with no
    value and the default format carries nothing and is left out of the map.
    """

    mappa: dict[tuple[int, int], tuple[Any, str]] = {}
    for riga in foglio.iter_rows():
        for cella in riga:
            valore = cella.value
            formato = cella.number_format
            if valore is None and formato == FORMATO_PREDEFINITO:
                continue
            mappa[(cella.row, cella.column)] = (valore, formato)
    return mappa


def _confronta_foglio(
    foglio_originale: Any,
    foglio_copia: Any,
    *,
    nome: str,
    colonna_ordine: int | None,
    prima_riga: int,
    quantita: Mapping[int, int],
    esito: EsitoFedelta,
) -> None:
    originale = _celle_significative(foglio_originale)
    copia = _celle_significative(foglio_copia)
    vuota = (None, FORMATO_PREDEFINITO)
    for chiave in sorted(originale.keys() | copia.keys()):
        riga, colonna = chiave
        valore_originale, formato_originale = originale.get(chiave, vuota)
        valore_copia, formato_copia = copia.get(chiave, vuota)
        esito.celle_confrontate += 1
        if valore_originale != valore_copia and not (
            _vuoto(valore_originale) and _vuoto(valore_copia)
        ):
            if (
                colonna == colonna_ordine
                and riga >= prima_riga
                and _differenza_ammessa(valore_originale, valore_copia, quantita.get(riga))
            ):
                esito.differenze_ammesse += 1
            else:
                _rifiuta(esito, nome, colonna, riga, valore_originale, valore_copia)
                continue
        # Number format only matters where the copy shows a value: on an
        # empty cell no format displays anything, so flagging it there
        # would reject a correct order over formatting nobody sees.
        if formato_originale != formato_copia and not _vuoto(valore_copia):
            _rifiuta_formato(esito, nome, colonna, riga, formato_originale, formato_copia)

    # So far this has only checked for unwanted extras. Now for what must
    # be present: the diff loop above can't detect a quantity that was
    # never written, because an untouched cell isn't a difference.
    if colonna_ordine is not None:
        esito.righe_richieste += len(quantita)
        esito.righe_mancanti.extend(
            _righe_del_piano_mancanti(copia, colonna_ordine=colonna_ordine, quantita=quantita)
        )


def _quantita_del_piano_nella_cella(valore: Any, attesa: int) -> bool:
    """Whether the cell shows the cartons the plan requests for that row.

    The single definition of "the plan is satisfied here", shared by both
    checks: the one that allows a difference while scanning cells, and the
    one that verifies at the end that every plan row is actually present.
    """

    return _numero(valore) == float(attesa)


def _righe_del_piano_mancanti(
    copia: Mapping[tuple[int, int], tuple[Any, str]],
    *,
    colonna_ordine: int,
    quantita: Mapping[int, int],
) -> list[int]:
    """Rows the plan requested that aren't in the copy.

    Checks the value the copy actually carries, not whether it differs
    from the price list: if the source price list already had that
    number — the same order as the previous run, unchanged — the copy is
    correct even without a difference, and requiring one would reject a
    correct order.

    Zero quantities: `quantita` comes from the plan, and the code that
    builds it already discards rows with zero cartons, and the writer
    rejects any quantity below 1 too. If a zero-quantity row did arrive
    here anyway, the general rule still applies — the copy must show what
    the plan asks, i.e. the number 0 — so the copy is rejected rather than
    delivered with a cell that doesn't match the plan. This is the same
    reading `_differenza_ammessa` already gives `attesa == 0`.

    Zeroing doesn't enter this count: emptying a cell that held a number is
    allowed precisely when the plan requests nothing for that row, so those
    rows are absent from `quantita`.
    """

    mancanti: list[int] = []
    for riga in sorted(quantita):
        valore, _formato = copia.get((riga, colonna_ordine), (None, FORMATO_PREDEFINITO))
        if not _quantita_del_piano_nella_cella(valore, quantita[riga]):
            mancanti.append(riga)
    return mancanti


def _differenza_ammessa(valore_originale: Any, valore_copia: Any, attesa: int | None) -> bool:
    """The plan's quantity, or the zeroing of a quantity that was already there."""

    if attesa is not None:
        return _quantita_del_piano_nella_cella(valore_copia, attesa)
    # No quantity requested for this row: the only allowed change is
    # emptying a cell that held a number. A section header may not.
    return _vuoto(valore_copia) and _numero(valore_originale) is not None


def _dove(nome_foglio: str, colonna: int, riga: int) -> str:
    dove = f"{lettera_di_colonna(colonna)}{riga}"
    if nome_foglio:
        dove += f" del foglio «{nome_foglio}»"
    return dove


def _rifiuta(
    esito: EsitoFedelta,
    nome_foglio: str,
    colonna: int,
    riga: int,
    valore_originale: Any,
    valore_copia: Any,
) -> None:
    esito.quante_rifiutate += 1
    if len(esito.esempi_rifiutati) >= ESEMPI_NEL_MESSAGGIO:
        # Past the first examples, keep counting without writing more: a
        # reader needs the first cases and a total, not the full list.
        return
    esito.esempi_rifiutati.append(
        f"{_dove(nome_foglio, colonna, riga)} nel listino di partenza è "
        f"{_mostra(valore_originale)} e nella copia è {_mostra(valore_copia)}"
    )


def _rifiuta_formato(
    esito: EsitoFedelta,
    nome_foglio: str,
    colonna: int,
    riga: int,
    formato_originale: str,
    formato_copia: str,
) -> None:
    esito.quante_rifiutate += 1
    if len(esito.esempi_rifiutati) >= ESEMPI_NEL_MESSAGGIO:
        return
    esito.esempi_rifiutati.append(
        f"{_dove(nome_foglio, colonna, riga)} ha cambiato formato numerico "
        f"(«{formato_originale}» nel listino di partenza, «{formato_copia}» nella copia): "
        "il numero mostrato non sarebbe più quello"
    )


def _elenco_di_righe(righe: list[int]) -> str:
    """Format a row list: "row 7", "rows 7 and 9", "rows 7, 9 and 12, and 4 more"."""

    mostrate = [str(riga) for riga in righe[:ESEMPI_NEL_MESSAGGIO]]
    if len(mostrate) == 1:
        elenco = f"riga {mostrate[0]}"
    else:
        elenco = "righe " + ", ".join(mostrate[:-1]) + f" e {mostrate[-1]}"
    restano = len(righe) - len(mostrate)
    if not restano:
        return elenco
    return elenco + (", e un'altra" if restano == 1 else f", e altre {restano}")


def _frase_delle_righe_mancanti(fornitore: str, esito: EsitoFedelta, *, insieme: bool) -> str:
    """The opposite failure: not extra cells, but missing order rows.

    `insieme` is true when the copy also has unwanted cell changes; in that
    case this sentence attaches to the first one instead of restarting from
    the supplier's name.
    """

    quante = len(esito.righe_mancanti)
    if esito.righe_richieste == 1:
        conto = "manca l'unica riga d'ordine richiesta"
    else:
        verbo = "manca" if quante == 1 else "mancano"
        conto = f"{verbo} {quante} delle {esito.righe_richieste} righe d'ordine richieste"
    apertura = "Inoltre alla copia" if insieme else f"Alla copia per {fornitore}"
    return (
        f"{apertura} {conto}: la quantità del piano non è arrivata nella sua cella "
        f"({_elenco_di_righe(esito.righe_mancanti)})"
    )


def frase_di_rifiuto(fornitore: str, esito: EsitoFedelta) -> str:
    """What the user reads when the copy isn't faithful.

    The two failures get two different sentences, because they point to
    different places to investigate: a cell changed outside the order
    column means the library altered the price list, a missing order row
    means an incomplete order. A copy with both failures states both.
    """

    parti: list[str] = []
    # Handles the case with no named failures too (mismatched sheets count
    # as a single difference).
    if esito.quante_rifiutate or not esito.righe_mancanti:
        quante = esito.quante_rifiutate
        celle = "cella è diversa" if quante == 1 else "celle sono diverse"
        testa = (
            f"La copia per {fornitore} non è fedele al listino di partenza: {quante} {celle} "
            "fuori dalle celle dell'ordine"
        )
        if esito.esempi_rifiutati:
            testa += " (" + "; ".join(esito.esempi_rifiutati) + ")"
        parti.append(testa)
    if esito.righe_mancanti:
        parti.append(_frase_delle_righe_mancanti(fornitore, esito, insieme=bool(parti)))
    return ". ".join(parti) + ". La copia non viene consegnata; il listino di partenza è rimasto invariato."
