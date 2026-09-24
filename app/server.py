#!/usr/bin/env python3
"""Local single-user server for the interactive supplier comparison UI.

The server binds to 127.0.0.1, stores state locally, never overwrites uploaded
files, and creates an audited JSON order plan.  Supplier XLSX copies remain the
responsibility of the deterministic writer in the skill pipeline.
"""

from __future__ import annotations

import argparse
import base64
import binascii
import hashlib
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import threading
import traceback
from copy import deepcopy
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse
from uuid import uuid4


APP_DIR = Path(__file__).resolve().parent
STATIC_DIR = APP_DIR / "static"
SKILL_ROOT = APP_DIR.parent
SCRIPTS_DIR = SKILL_ROOT / "scripts"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import ai_client  # noqa: E402
# The key-redaction regex lives in one place, its own module. Copying it here
# would mean two patterns to keep in sync, and the one that gets forgotten is
# always the one that leaks the key.
from ai_client import oscura as senza_la_chiave  # noqa: E402
from catalog_search import SupplierCatalog  # noqa: E402
# The order-quantity column can be moved from the page; the rules for what
# that declaration becomes live in this module, and `launcher.source_rule`
# is the single source of truth for whether the column is writable.
import colonna_ordine  # noqa: E402
# User confirmations live in a SQLite store that survives the weekly
# recompute: identity is the item (barcode + normalized name), never the row
# and never the price. None of that module's rules are duplicated here —
# what counts as "the same item" is decided in one place, that module.
from conferme import (  # noqa: E402
    MagazzinoConferme,
    MagazzinoNonUtilizzabile,
    codice_confrontabile,
    impronta_prodotto,
)
# The dated output folder, the audit trail, the list of compiled orders and
# the path safeguards all live in `consegna`; none of those rules are
# reimplemented here, since two copies of the same safeguard tend to diverge.
import consegna  # noqa: E402
# The list of products no supplier carries: sheet, name and reason live in
# their own module. Only the entry filter stays here, since `find_offer`/
# `offer_is_available` are the ones that know how to answer it.
import da_reperire as da_reperire_modulo  # noqa: E402
from build_review_data import impronta_articolo  # noqa: E402
# What counts as an offer — whose it is, whether it can be ordered, its unit
# cost — is decided by its own module, for every caller: `offer_is_available`
# is the single authority for orderability.
from offerta import (  # noqa: E402
    STATO_RIFIUTATO_UTENTE,
    find_offer,
    numero as number,
    offer_is_available,
    offer_pricing,
    offer_supplier_id,
)
from inspect_sources import profile_file  # noqa: E402
import order_history  # noqa: E402
import registro  # noqa: E402
# Writes a file so it's never found half-written, even after a power loss —
# the shared implementation for atomic-and-durable writes across the app.
import scrittura_sicura  # noqa: E402
import versione_del_codice  # noqa: E402
from pipeline_jobs import (  # noqa: E402
    STATI_IN_CORSO,
    ConfigurazionePipeline,
    LavoroGiaInCorso,
    PipelineJobManager,
)
from promotion_bridge import PromotionService  # noqa: E402


MAX_JSON_BYTES = 150 * 1024 * 1024
MAX_UPLOAD_BYTES = 60 * 1024 * 1024
# `.xls` is accepted because supplier Noce ships Excel 97-2003 files with no
# `.xlsx` equivalent. The extension is only a first filter on the name; what
# the file actually is comes from its first bytes, in `profile_file`.
ALLOWED_UPLOAD_SUFFIXES = {".xlsx", ".xls", ".csv"}
FORMATO_DELL_ESTENSIONE = {".xlsx": "xlsx", ".xls": "xls", ".csv": "csv"}
NOME_DEL_FORMATO = {"xlsx": "Excel (.xlsx)", "xls": "Excel 97-2003 (.xls)", "csv": "CSV"}
SAFE_FILE_RE = re.compile(r"[^A-Za-z0-9._() -]+")


def load_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return deepcopy(default)
    return json.loads(path.read_text(encoding="utf-8"))


def atomic_json(path: Path, value: Any) -> None:
    """Write JSON atomically and durably, via `scrittura_sicura`.

    `os.replace` is atomic for metadata, not for data: without an fsync
    before the replace, a power loss at the wrong instant can leave
    `state.json` present but empty.
    """

    scrittura_sicura.scrivi_json(path, value)


def _centesimi_come_in_pagina(value: float) -> float:
    """Round the way the page displays an amount, replicated here.

    Python's `round()` rounds the binary double and ties to even;
    `Intl.NumberFormat` in the browser rounds the shortest decimal
    representation and ties away from zero — on 14.665 the former gives
    14.66 while the screen shows 14.67. Every amount the server states that
    the page also draws (the summary totals, the row sums) must go through
    this, or the two disagree. `Decimal(str(x))` uses the same shortest
    decimal representation as the browser.
    """

    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def safe_upload_name(value: str) -> str:
    name = Path(str(value or "")).name.strip()
    name = SAFE_FILE_RE.sub("_", name).strip(" .")
    # The 180-char truncation happens before the extension check, not after:
    # truncating an already-approved name can cut the dot and extension off a
    # very long name, and the file then lands on disk with no extension even
    # though it passed the check. The failure is silent — the upload still
    # succeeds, since the reader picks the format from the first bytes, not
    # the name — and the price list quietly drops out of the comparison,
    # because `candidate_files` in `scripts/inspect_sources.py` filters by
    # suffix and an extensionless file doesn't match.
    suffix = Path(name).suffix.casefold()
    if len(name) > 180:
        name = name[:180 - len(suffix)] + suffix
    # The check stays after the truncation: a suffix longer than 180 chars
    # could otherwise produce an empty name, which must fall into this same
    # validation path as any other invalid name.
    if not name or name in {".", ".."}:
        raise ValueError("Nome del documento non valido")
    suffix = Path(name).suffix.casefold()
    if suffix not in ALLOWED_UPLOAD_SUFFIXES:
        raise ValueError(
            f"Formato non accettato: {suffix or 'senza estensione'}. "
            "Si caricano fogli di calcolo Excel (.xlsx o .xls) e file CSV."
        )
    return name


def upload_role(value: Any) -> str:
    """Role chosen in the import box, in the pipeline's own vocabulary."""

    role = str(value or "").strip().casefold()
    if role in {"management", "master"}:
        return "master"
    if role in {"suppliers", "supplier"}:
        return "supplier"
    raise ValueError("Scegli se il documento è l'elenco del gestionale o un listino fornitore")


def profiled_upload_role(profile: Any) -> str:
    if not isinstance(profile, dict):
        return ""
    role = str(profile.get("upload_role") or (profile.get("ai_preflight") or {}).get("role") or "").strip().casefold()
    if role in {"master", "supplier"}:
        return role
    adattatore_id = str((profile.get("deterministic_hint") or {}).get("adapter_id") or "")
    if adattatore_id:
        try:
            adattatore = registro.adattatore(adattatore_id)
        except (KeyError, ValueError):
            return ""
        if not adattatore:
            return ""
        return "master" if adattatore.get("kind") == "master" else "supplier"
    return ""


def unique_destination(directory: Path, name: str) -> Path:
    candidate = directory / name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    for index in range(2, 10000):
        candidate = directory / f"{stem} ({index}){suffix}"
        if not candidate.exists():
            return candidate
    raise ValueError(f"Troppe versioni del documento {name}")


def supplier_label(supplier_id: str) -> str:
    """Return this supplier's display name.

    The name comes from the adapter registry — the same source the supplier
    is created from — not from a list hardcoded here: a hardcoded list would
    show a learned supplier as its raw id, underscores included, in messages
    and history, while the registry already carries its `display_name`.
    """

    return registro.nome_del_fornitore(supplier_id)


def frase(value: Any) -> str:
    """Close a sentence with a period.

    Technical reasons come back from libraries without punctuation, and
    concatenating them with the next sentence produces run-on text.
    """

    testo = str(value or "").strip()
    if not testo:
        return ""
    return testo if testo[-1] in ".!?" else f"{testo}."


def normalize_header(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split()).upper()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ---------------------------------------------------------------------------
# Shared rules on offers, prices and thresholds.
#
# They live here, outside individual endpoints, because state validation,
# order compilation and the move-quantity preview must all read prices the
# same way. Two copies of the same rule tend to diverge, leaving the user
# with a preview that doesn't match the compiled order.
# ---------------------------------------------------------------------------


def nessuna_offerta_utilizzabile(product: Any) -> bool:
    """True when no supplier can fulfil this product.

    Not "the user hasn't chosen yet" but "there's nothing to choose from". A
    match still pending confirmation does NOT count here — an offer awaiting
    confirmation is still a usable offer, and the product stays a case to
    verify, not an unavailable one. This is the condition that separates a
    quantity with no supplier — a valid state that lands in the to-be-sourced
    list — from a quantity on an offer that can't be ordered, which stays an
    error.
    """

    if not isinstance(product, dict):
        return False
    return not any(offer_is_available(offer) for offer in product.get("offers") or [])


def offer_needs_confirmation(product: Any, offer: Any) -> bool:
    """True when the user must confirm by hand before this can be ordered.

    Combines both sources: the product (an uncertain match flagged by the
    pipeline) and the individual offer (an inexact match, a low-confidence
    display). This is the same condition that triggers CONFERMA_MANCANTE
    when saving state, so the move-quantity preview can warn ahead of time
    instead of letting the following save fail.
    """

    requires = bool(isinstance(product, dict) and (product.get("requiresConfirmation") or product.get("requires_confirmation")))
    if isinstance(offer, dict):
        requires = requires or bool(offer.get("requiresConfirmation") or offer.get("requires_confirmation"))
    return requires


def supplier_threshold(definition: Any) -> float:
    if not isinstance(definition, dict):
        return 0.0
    return number(definition.get("minimumOrder") or definition.get("minimum_order") or definition.get("thresholdNet")) or 0.0


def meets_threshold(total: float, threshold: float) -> bool:
    """Compilation rule: a supplier with nothing ordered has no threshold to meet."""

    return not (0 < total < threshold)


# The two ways to compile an order. The `.xlsx` copy is written by the Node
# writer; the document that must be patched in place — a handful of bytes
# per cell, everything else identical — is written by `app/xls_writer.py`,
# in Python.
PATCH_IN_POSIZIONE = "patch_xls_in_posizione"


def procedura_di_scrittura(regola: Any) -> str:
    """How this supplier's price list gets compiled: the registry rule decides.

    The choice between the two write paths comes only from the adapter
    registry (`order_write.mode`, carried into the write rule by
    `launcher.source_rule` as `compilazione`), never from a hardcoded
    supplier name. A hardcoded check would break in both directions: a new
    supplier shipping `.xls` — the exact case the adapter-learning flow is
    meant to support — would fall into the Node-writer branch and be
    rejected as "not available in XLSX format", blaming the supplier's file
    instead of the configuration, which is the only thing that can actually
    be fixed. And a known in-place-patch supplier switching to `.xlsx` would
    still get routed to the in-place patch by the stale hardcoded name.

    A supplier that declares nothing goes through the Node writer: that's
    the default path, and the one every `.xlsx` price list takes.
    """

    if not isinstance(regola, dict):
        return ""
    return str(regola.get("compilazione") or "").strip()


# Keys that `state.json` receives from a partial-update route — a candidate
# response, a manually matched row, a header discount, an added product —
# and that `validate_snapshot` must therefore copy forward from disk instead
# of rebuilding from the page's snapshot: the page never sends them, and
# rebuilding state without them would erase them.
#
# The value given here is what the key defaults to when nothing is on disk
# yet. Adding a route that writes its own key into `state` means adding it
# here too: `tests/test_abbinamento_a_mano.py` reads the routes and compares
# the two lists, and a key missing from either one fails that test.
CHIAVI_DI_STATO_RICOPIATE: dict[str, Any] = {
    "manualProducts": [],
    "matchOverrides": [],
    "supplierDiscounts": {},
    "manualMatches": [],
}


class ReviewStore:
    def __init__(self, review_path: Path, state_path: Path, upload_dir: Path, output_dir: Path, writer_config: Path | None = None, history_path: Path | None = None, orders_dir: Path | None = None, conferme_path: Path | None = None) -> None:
        self.review_path = review_path.resolve()
        self.state_path = state_path.resolve()
        self.upload_dir = upload_dir.resolve()
        self.output_dir = output_dir.resolve()
        # Compiled orders live next to `outputs`, not inside it. Nothing
        # writes into `outputs` at runtime; the folder and the `/outputs/`
        # route only serve artifacts from past runs, and removing them is a
        # separate decision.
        self.orders_dir = orders_dir.resolve() if orders_dir else self.output_dir.parent / "ordini"
        self.writer_config = writer_config.resolve() if writer_config else None
        # History lives outside the current run's folder: it must survive
        # the weekly recompute of the comparison.
        self.history_path = history_path.resolve() if history_path else self.state_path.parent.parent / "history" / "orders.json"
        # Confirmations sit next to history, for the same reason: a given
        # answer applies to the item, not to the week it was given. Inside
        # the run folder it would vanish on the next recompute, which is
        # exactly the failure this store exists to prevent.
        self.conferme_path = conferme_path.resolve() if conferme_path else self.history_path.parent / "conferme.db"
        # Opened lazily, on first use, not here: a file that fails to open
        # must not stop the server from starting. The failure is reported
        # once, on the page, and everything else keeps working without
        # confirmation memory.
        self._conferme: MagazzinoConferme | None = None
        self._conferme_guasto = ""
        self.lock = threading.RLock()
        self.catalog = SupplierCatalog()
        self.promotion_service = PromotionService()
        self.upload_profiles_path = self.upload_dir / "upload_profiles.json"
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.orders_dir.mkdir(parents=True, exist_ok=True)
        # One job at a time, whatever it is: there is a single lock, held by
        # whoever is about to replace `review_data.json`.
        self.lucchetto_lavori = threading.Lock()
        self.pipeline_jobs = PipelineJobManager(
            ConfigurazionePipeline(
                data_dir=self.state_path.parent,
                uploads_dir=self.upload_dir,
                review_path=self.review_path,
                state_path=self.state_path,
            ),
            lucchetto_lavori=self.lucchetto_lavori,
            lucchetto_dati=self.lock,
            su_confronto_attivato=self.riconfigura_compilazione,
            uguaglianze_dichiarate=self.uguaglianze_in_vigore,
            rifiuti_dichiarati=self.rifiuti_in_vigore,
        )

    # ------------------------------------------------------- given confirmations
    #
    # Why a separate store, and why `state.json` alone isn't enough. A
    # product on the page is keyed `product:N`, its row number in the
    # management-software export: across two real exports, of 457
    # identifiers present in both, 449 pointed to a different item. A
    # confirmation tied to that identifier would either get lost on
    # recompute or — worse — get silently reapplied to merchandise the user
    # never looked at. Here the confirmation is tied to the item itself, and
    # survives into next week's price list.

    def magazzino_conferme(self) -> MagazzinoConferme | None:
        """Open the confirmation store lazily; `None` if it fails to open.

        The module's exception is caught here and only here, as its
        docstring requires: reporting "no confirmations" for a file that
        fails to open would silently re-ask every question, and the user
        would reconfirm by hand believing the program never knew the answer.
        The reason is kept in `_conferme_guasto` and surfaced among the
        comparison's warnings.
        """

        if self._conferme is not None or self._conferme_guasto:
            return self._conferme
        try:
            self._conferme = MagazzinoConferme(self.conferme_path)
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Il file delle conferme non si apre."
            self._conferme = None
        return self._conferme

    def _magazzino_solo_se_c_e(self) -> MagazzinoConferme | None:
        """Open the confirmation store, but never create it if it doesn't exist.

        Keeps the invariant that a program with no confirmations creates no
        `conferme.db`. SQLite keeps the file open for the life of the
        connection, and on Windows an open file locks the folder that
        contains it — a real risk for tests that clean up a temp folder
        while a connection is still open. Equalities are read on every
        comparison read: always opening the store would reintroduce that
        same failure mode. If nothing has ever been declared, there is
        nothing to read.
        """

        if self._conferme is None and not self.conferme_path.exists():
            return None
        return self.magazzino_conferme()

    def esporta_le_conferme(self) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Everything the user has confirmed, history included.

        Two lists because the store has two tables and both are permanent
        memory: confirmations ("yes, same item") and declared barcode
        equalities. Exporting only one would make "copy" mean half the file.

        `conferme.db` is the program's permanent memory — a confirmed match
        stays valid for next week's price list too — and since `app/data/`
        is entirely git-ignored it has no backup copy of its own.
        `MagazzinoConferme.esporta` is the documented, supported way to get a
        readable copy of an otherwise-opaque `.db` file without copying an
        open SQLite file by hand.

        The store is never created if it doesn't already exist: a program
        with no confirmations must not end up with an empty `conferme.db`
        just from opening Settings. If it fails to open, this returns the
        empty lists, same as equalities do; the reason is already among the
        comparison's warnings.
        """

        magazzino = self._magazzino_solo_se_c_e()
        if magazzino is None:
            return [], []
        try:
            return magazzino.esporta(), magazzino.esporta_uguaglianze()
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Le conferme date non si leggono."
            return [], []

    def uguaglianze_dichiarate(self) -> list[dict[str, Any]]:
        """Declared equalities in force, one per pair, for the page.

        Pairs, not groups: this removes exactly what was declared, no more.
        If the store fails to open, this returns the empty list — the reason
        is already among the comparison's warnings, and failing to render
        the page over one missing memory would be disproportionate.
        """

        magazzino = self._magazzino_solo_se_c_e()
        if magazzino is None:
            return []
        try:
            return magazzino.uguaglianze()
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Le uguaglianze fra codici non si leggono."
            return []

    def elenco_delle_uguaglianze(self, query: str = "") -> dict[str, Any]:
        """Declared "these two codes are the same item" statements, made readable.

        A bare pair of thirteen-digit codes gives the reader nothing to
        judge the declaration against — it needs both product names to be
        checkable at a glance.

        Both names are already available with no new column needed: the
        store saves a fingerprint for each item (`impronta_prodotto` for the
        management-software side, `impronta_articolo` for the price-list
        row), and the normalized name is embedded in that fingerprint. This
        unpacks the two fingerprints into their component fields. The names
        are normalized — uppercase, punctuation stripped — the same form the
        program matches on, not what the supplier prints.

        Search looks at everything shown: codes, names, supplier and reason.
        A list that grows by one row a week accumulates dozens of entries
        within a year, and free-text search is what keeps them usable.
        """

        cercato = " ".join(str(query or "").split()).casefold()
        dichiarate = self.uguaglianze_dichiarate()
        voci = []
        for riga in dichiarate:
            voce = self._uguaglianza_leggibile(riga)
            # `cercabile` is the search text, not a field of the
            # declaration: it's popped here and never reaches the page.
            testo = voce.pop("cercabile")
            if cercato and cercato not in testo:
                continue
            voci.append(voce)
        return {
            "ok": True,
            "query": " ".join(str(query or "").split()),
            "totale": len(dichiarate),
            "uguaglianze": voci,
            "motivo": self._conferme_guasto or "",
        }

    @staticmethod
    def _uguaglianza_leggibile(riga: Any) -> dict[str, Any]:
        """Unpack a declaration into names, not just the two codes.

        The two fingerprints have different shapes because they answer
        different questions: the management-software one is
        `code|name`, the price-list row one is
        `supplier|code|supplier_code|name`. Fields are unpacked by position,
        and a missing field stays empty — a fingerprint from an older format
        must not make the row disappear from the list.
        """

        riga = riga if isinstance(riga, dict) else {}
        codici = [str(voce) for voce in (riga.get("codici") or [])]
        pezzi_articolo = str(riga.get("articolo") or "").split("|")
        pezzi_offerta = str(riga.get("offerta") or "").split("|")
        gestionale = {
            "codice": pezzi_articolo[0] if len(pezzi_articolo) > 0 else "",
            "nome": pezzi_articolo[1] if len(pezzi_articolo) > 1 else "",
        }
        fornitore_id = pezzi_offerta[0] if len(pezzi_offerta) > 0 else ""
        listino = {
            "fornitoreId": fornitore_id,
            "fornitore": supplier_label(fornitore_id) if fornitore_id else "",
            "codice": pezzi_offerta[1] if len(pezzi_offerta) > 1 else "",
            "codiceFornitore": pezzi_offerta[2] if len(pezzi_offerta) > 2 else "",
            "nome": pezzi_offerta[3] if len(pezzi_offerta) > 3 else "",
        }
        motivo = str(riga.get("motivo") or "")
        cercabile = " ".join([
            *codici, gestionale["codice"], gestionale["nome"],
            listino["fornitore"], listino["fornitoreId"], listino["codice"],
            listino["codiceFornitore"], listino["nome"], motivo,
        ]).casefold()
        return {
            "codici": codici,
            "gestionale": gestionale,
            "listino": listino,
            "motivo": motivo,
            "dal": str(riga.get("valida_dal") or ""),
            "cercabile": cercabile,
        }

    def uguaglianze_in_vigore(self) -> list[list[str]]:
        """Groups of barcodes declared equal, for the matching chain.

        Called by `pipeline_jobs` at the start of reading price lists, and
        its result is written into the run folder. If the store fails to
        open, this returns no groups: the reason is already in
        `_conferme_guasto` and surfaces among the comparison's warnings, and
        blocking a recompute over this supplementary memory would be
        disproportionate.
        """

        magazzino = self._magazzino_solo_se_c_e()
        if magazzino is None:
            return []
        try:
            return magazzino.classi()
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Le uguaglianze fra codici non si leggono."
            return []

    def rifiuti_in_vigore(self) -> dict[tuple[str, str], dict[str, Any]]:
        """Declared rejections, indexed by (supplier, management-software item).

        A single query, not one per product: this runs on every comparison
        read, and the page autosaves 450ms after each edit.

        Goes through `_magazzino_solo_se_c_e`, not `magazzino_conferme`, for
        the same reason as equalities: a program with no confirmations must
        not end up with an empty `conferme.db` just from opening a page, and
        on Windows an open SQLite file locks the folder that contains it.

        If the store fails to open, this returns no rejections; the reason
        is already in `_conferme_guasto` and surfaces among the comparison's
        warnings. The failure direction is deliberately the safe one: with
        no rejection memory, the offer becomes visible again and the
        question is asked again, rather than silently staying hidden.
        """

        magazzino = self._magazzino_solo_se_c_e()
        if magazzino is None:
            return {}
        try:
            voci = magazzino.rifiuti()
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Gli abbinamenti rifiutati non si leggono."
            return {}
        return {
            (str(voce.get("fornitore") or ""), str(voce.get("articolo") or "")): voce
            for voce in voci
        }

    def spegni_le_offerte_rifiutate(self, review: dict[str, Any]) -> None:
        """Apply "not the same item": the offer drops out of the comparison.

        The negative counterpart of `dichiara_la_conferma`, producing the
        one effect that matters: the offer stops being usable. Nothing else
        is touched, because `offerta.offer_is_available` is the single
        source of truth for "can this be ordered", already used by
        `nessuna_offerta_utilizzabile`, `validate_snapshot`, the
        `da_reperire` loop in `compile`, and `pipeline_jobs._ripulisci_stato`.

        Only applied when today's row fingerprint matches the one the
        rejection was recorded against — same discipline as
        `conferma_in_vigore`: next week's price list may match that product
        to a different row that nobody has judged yet, and applying the old
        rejection to it would be a decision on merchandise never seen.

        Called here, not in `review()`. `dichiara_la_conferma` lives in
        `review()` on purpose, because a confirmation re-applied from the
        store during a save would silently undo a revocation. A rejection
        doesn't have that problem: there's no snapshot field for it, it is
        only ever set or cleared through its own route, so the store is the
        sole authority for both reading and writing it. It must run on the
        shared read path, or `validate_snapshot` would still see the offer
        as live and reject the user's only remaining answer with
        `OFFERTA_NON_VALIDA`.

        Not destructive: `base_review()` re-reads `review_data.json` from
        disk on every call, so clearing the rejection restores the offer
        exactly as it was — status, prices, confirmation requirement.
        """

        rifiuti = self.rifiuti_in_vigore()
        if not rifiuti:
            return
        for product in review.get("products") or []:
            if not isinstance(product, dict):
                continue
            articolo = impronta_prodotto(product)
            # An item that can't be fingerprinted has no rejection to look up.
            if not articolo:
                continue
            spente: list[dict[str, Any]] = []
            for offer in product.get("offers") or []:
                if not isinstance(offer, dict) or not offer_is_available(offer):
                    continue
                voce = rifiuti.get((offer_supplier_id(offer).strip().casefold(), articolo))
                if voce is None or str(voce.get("offerta") or "") != impronta_articolo(offer):
                    continue
                spente.append(offer)
                offer["available"] = False
                offer["status"] = STATO_RIFIUTATO_UTENTE
                offer["matchStatus"] = STATO_RIFIUTATO_UTENTE
                # The question is no longer open: still asking it on an
                # offer that can't be chosen would keep the blocking
                # "confirmation required" state on a product with nothing
                # left to confirm.
                offer["requiresConfirmation"] = False
                offer["confirmed"] = False
                # What the page needs to be able to say: what was rejected
                # and when. Without it, the offer grid would claim this
                # supplier doesn't carry the item — false, the row is there,
                # it's just rejected.
                offer["rifiutata"] = {
                    "since": str(voce.get("valida_dal") or ""),
                    "description": str(offer.get("description") or ""),
                }
            self._riscegli_dopo_il_no(product, spente)

    def _riscegli_dopo_il_no(self, product: dict[str, Any], spente: list[dict[str, Any]]) -> None:
        """After a rejection: same-code rows need confirmation, and if the
        rejection turned off the selected offer, reselect a new one.

        Suspicion lives on individual offers, not on the product. If the
        rejected row isn't the item, another row with the same barcode at a
        different supplier is suspect too — and must ask for confirmation no
        matter which path picks it: the fallback selection done here, a
        header discount applied later by `review()`, a manual move, or the
        move-quantity preview. All of them go through
        `offer_needs_confirmation`, which combines product- and offer-level
        flags; marking the flag only on the product would leave the
        fallback selection computed here unflagged, letting a same-barcode
        row slip into the order with no confirmation.

        The selection is redone here because the comparison result comes
        from the matching chain, which has no notion of rejections: without
        this, a product would keep pointing at the disabled row after a
        rejection, `dichiara_la_conferma` would look up the confirmation
        under the wrong supplier, and a confirmation given on the new
        selection would never get reapplied. The product-level question,
        which belonged to the disabled row, is cleared here — from this
        point on, offers decide.
        """

        codici = {
            codice: str(offer.get("supplierName") or offer_supplier_id(offer)).upper()
            for offer in spente
            if (codice := codice_confrontabile(offer.get("ean")))
        }
        for offer in product.get("offers") or []:
            if not isinstance(offer, dict) or not offer_is_available(offer):
                continue
            codice = codice_confrontabile(offer.get("ean"))
            if codice not in codici:
                continue
            offer["requiresConfirmation"] = True
            offer["confirmed"] = False
            offer["confirmationMessage"] = (
                f"Hai detto che la riga di {codici[codice]} non è questo articolo, e questa ha lo "
                f"stesso codice a barre ({codice}): controlla che sia davvero lo stesso articolo."
            )
        scelta = str(product.get("selectedSupplierId") or "")
        if not any(offer_supplier_id(offer) == scelta for offer in spente):
            return
        migliore = self._offerta_piu_conveniente(product)
        product["selectedSupplierId"] = offer_supplier_id(migliore) if migliore else ""
        product["confirmed"] = False
        product["requiresConfirmation"] = False
        product["confirmationMessage"] = ""

    def chiudi(self) -> None:
        """Release the confirmations file.

        Needed because SQLite keeps the file open for the life of the
        connection, and on Windows an open file can't be deleted or
        renamed — nor can the folder that contains it. A store that's no
        longer used and never released leaves its `conferme.db` locked,
        which matters most for the full test suite, where many stores are
        created in temp folders that must be cleaned up afterward.

        Safe to call more than once; the store simply reopens on the next
        confirmation lookup.
        """

        if self._conferme is not None:
            self._conferme.chiudi()
            self._conferme = None

    def conferma_in_vigore(self, product: Any, offer: Any) -> dict[str, Any] | None:
        """Return the confirmation covering this offer of this item, if any.

        The store returns a match even when today's supplier row is no
        longer the one that was confirmed — that's a deliberate choice, so
        callers can tell "the price-list format changed" apart from "the
        item was swapped under the same row". The freshness check therefore
        happens here: only a confirmation whose fingerprint matches the
        current price-list row is applied.

        A rejected confirmation (`accettata` false) is not a confirmation:
        this returns `None`, since rejections are handled by a separate
        lookup, not this one.
        """

        if not isinstance(offer, dict):
            return None
        fornitore = offer_supplier_id(offer)
        articolo = impronta_prodotto(product)
        impronta = impronta_articolo(offer)
        # Fingerprints are checked before opening the store: an item that
        # can't be fingerprinted has nothing to look up, and opening the
        # store just to find that out would create a `conferme.db` for a
        # program that has never recorded any confirmation.
        if not fornitore or not articolo or not impronta:
            return None
        magazzino = self.magazzino_conferme()
        if magazzino is None:
            return None
        try:
            voce = magazzino.cerca(fornitore, articolo)
        except MagazzinoNonUtilizzabile as exc:
            self._conferme_guasto = frase(exc) or "Il file delle conferme non si legge più."
            return None
        if voce is None or not voce.get("accettata") or str(voce.get("offerta") or "") != impronta:
            return None
        return voce

    def dichiara_la_conferma(self, product: Any) -> None:
        """Set, on the product, whether a confirmation already exists and since when.

        These are the two facts the page needs to show — that an answer was
        already given, and when — and they're computed here because a
        confirmation lives in the store, not in this week's saved state.

        `confirmation` is informational and can be absent; `confirmed`
        becomes true even with no matching entry in this week's state,
        which is the whole point of the store: next week's
        management-software export is a different file, the product has a
        different row number and its saved decision is gone — but the item
        is the same, and the earlier answer still applies.
        """

        if not isinstance(product, dict):
            return
        offerta = find_offer(product, str(product.get("selectedSupplierId") or ""))
        # Only where a confirmation was actually needed: a certain match was
        # never a question, so reporting it as confirmed would be a fact
        # this store never recorded.
        if offerta is None or not offer_needs_confirmation(product, offerta):
            return
        voce = self.conferma_in_vigore(product, offerta)
        if voce is None:
            return
        product["confirmed"] = True
        product["confirmation"] = {
            "supplierId": offer_supplier_id(offerta),
            "since": voce.get("valida_dal") or "",
            # The identity the confirmation applies to, spelled out: this is
            # what explains why it survives a new price list, and without it
            # the page couldn't show that to the user.
            "article": voce.get("articolo") or "",
        }

    def ricorda_le_conferme(
        self, clean: dict[str, Any], review: dict[str, Any], precedente: dict[str, Any]
    ) -> None:
        """Persist to the store what the user just answered.

        A revocation is recognized by comparing against the previous state,
        never from a bare `confirmed: false`. The page sends `false` in two
        unrelated situations: the user unchecked the confirmation, or the
        recompute expired it because the price-list row now points to a
        different item. Treating both the same way would clear the store
        exactly in the week it's supposed to help. So a confirmation is only
        forgotten when the previous state held a confirmation on the same
        item from the same supplier — that's the only case that reads
        unambiguously as the user changing their mind.

        A store failure does not stop the save: it's recorded and surfaced
        on the page instead. A failing `PUT /api/state` must not lose every
        answer that comes after it, only the one it couldn't persist.
        """

        prodotti = {str(item.get("id")): item for item in review.get("products") or []}
        prima = {
            str(item.get("id")): item
            for item in (precedente.get("products") or [])
            if isinstance(item, dict)
        }
        quando = str(clean.get("updatedAt") or "")
        if not quando:
            return

        # First check whether there's anything to remember, and only then
        # open the store: opening it unconditionally would create a
        # `conferme.db` for a program that has never recorded a
        # confirmation, and on Windows an open file locks its folder.
        da_scrivere: list[tuple[bool, str, str, str, str]] = []
        for decisione in clean.get("products") or []:
            identificativo = str(decisione.get("id") or "")
            prodotto = prodotti.get(identificativo)
            fornitore = str(decisione.get("selectedSupplierId") or "")
            if prodotto is None or not fornitore:
                continue
            offerta = find_offer(prodotto, fornitore)
            # Only remembered where a confirmation was actually required.
            # The comparison also sets `confirmed` on certain matches
            # (`confirmed = not requiresConfirmation`), so without this
            # filter the store would fill up with answers the user never
            # actually gave, burying the few real ones. A memory the user
            # can't tell apart from noise isn't a memory.
            if offerta is None or not offer_needs_confirmation(prodotto, offerta):
                continue
            articolo = impronta_prodotto(prodotto)
            impronta = impronta_articolo(offerta)
            if not articolo or not impronta:
                continue
            if decisione.get("confirmed"):
                da_scrivere.append((True, fornitore, articolo, impronta, str(offerta.get("method") or "")))
                continue
            vecchia = prima.get(identificativo) or {}
            if vecchia.get("confirmed") and str(vecchia.get("confirmedArticle") or "") == impronta:
                da_scrivere.append((False, fornitore, articolo, impronta, ""))
        if not da_scrivere:
            return

        magazzino = self.magazzino_conferme()
        if magazzino is None:
            return
        try:
            for accettata, fornitore, articolo, impronta, motivo in da_scrivere:
                if accettata:
                    magazzino.ricorda(
                        fornitore=fornitore,
                        articolo=articolo,
                        offerta=impronta,
                        accettata=True,
                        motivo=motivo,
                        quando=quando,
                    )
                else:
                    # Only clear the entry if it's currently a yes.
                    # `dimentica` closes whatever row is currently in force,
                    # and that row could be a rejection instead of a
                    # confirmation: an autosave that reached this branch
                    # unconditionally would erase the rejection, bringing
                    # the question back on a match the user had already
                    # dismissed. This is a belt-and-suspenders check — the
                    # rejection route also clears the selected supplier, so
                    # this loop would skip the product anyway — but it's the
                    # cheapest safeguard available.
                    voce_in_vigore = magazzino.cerca(fornitore, articolo)
                    if voce_in_vigore is None or voce_in_vigore.get("accettata"):
                        magazzino.dimentica(fornitore, articolo, quando=quando)
        except (MagazzinoNonUtilizzabile, ValueError) as exc:
            self._conferme_guasto = frase(exc) or "Le conferme non sono state salvate."

    def riconfigura_compilazione(self, review: dict[str, Any]) -> None:
        """After a recompute, price lists are different files: rebuild the write config.

        Without this, order compilation would use last week's writer
        configuration — pointing at last week's price list — with this
        week's row numbers. The writer would only catch the mismatch when
        the file path is unchanged, since that's the only case where it
        checks the fingerprint; a price list uploaded under a new name
        would not be caught.
        """

        if self.writer_config is None:
            return
        from launcher import prepare_writer_config  # lazy import: the server must start even without it

        setup = prepare_writer_config(review, write=True, destinazione=self.writer_config)
        if setup.config_path is None:
            # `prepare_writer_config` can return without writing and without
            # raising when a requirement is missing (Node absent, script
            # missing). Without this explicit raise, that path fails
            # silently and the COMPILAZIONE_DA_RICONFIGURARE warning never
            # surfaces. The run-id guard would still block compilation, but
            # the goal is to report the failure immediately.
            raise RuntimeError(setup.message)

    def avvia_pipeline(self) -> dict[str, Any]:
        return self.pipeline_jobs.avvia()

    def stato_pipeline(self) -> dict[str, Any]:
        return self.pipeline_jobs.stato()

    def schemi_pendenti(self) -> dict[str, Any]:
        return self.pipeline_jobs.schemi_pendenti()

    def colonne_dei_documenti(self) -> dict[str, Any]:
        return self.pipeline_jobs.colonne_dei_documenti()

    def valida_schemi(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.pipeline_jobs.valida_schemi(payload)

    def conferma_schemi(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.pipeline_jobs.conferma_schemi(payload)

    # The column picker opened by hand on an already-uploaded price list.
    def colonne_del_documento(self, nome: Any) -> dict[str, Any]:
        return self.pipeline_jobs.colonne_del_documento(nome)

    def prova_colonne(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.pipeline_jobs.prova_colonne_del_documento(payload)

    def salva_colonne(self, payload: dict[str, Any]) -> dict[str, Any]:
        return self.pipeline_jobs.salva_colonne_del_documento(payload)

    def colonne_d_ordine(self, supplier_id: str) -> dict[str, Any]:
        """List the columns available for the order-quantity column, per supplier.

        Every column in the sheet is shown, including the ones the pipeline
        already reads: hiding them would make the document look like it has
        fewer columns than it does, and a user looking for "the one after
        the price column" counts what's visible. Columns already in use are
        flagged, not removed from the list.
        """

        documento = self.pipeline_jobs.documento_del_fornitore(supplier_id)
        effettiva = documento["effettiva"]
        chiave = str(supplier_id or "").strip().casefold()
        return {
            "ok": True,
            "supplierId": chiave,
            "supplierName": supplier_label(chiave),
            "fileName": effettiva.get("fileName"),
            "sheet": effettiva.get("sheet"),
            "headerRow": effettiva.get("headerRow"),
            "dataStartRow": effettiva.get("dataStartRow"),
            "attuale": effettiva.get("orderColumn"),
            "colonne": colonna_ordine.colonne_per_la_scelta(effettiva, documento["colonne"]),
        }

    def cambia_colonna_d_ordine(self, payload: Any) -> dict[str, Any]:
        """Move the column the order quantities get written into.

        The check that decides whether the new column is valid is not
        reimplemented here: it's `launcher.source_rule`, the same function
        that drives order compilation. The declaration to write is built,
        tried against the real document, and the registry is only touched
        if it passes. A second, independent rule here could disagree with
        the one that actually compiles the order, silently breaking
        compilation for that supplier.

        Written in four places, because four different readers need to
        agree on the same fact:

        1. the adapter registry, which governs every future week;
        2. the registry's confirmed field mapping, for suppliers that
           declare `from_field_mapping` — `source_rule` refuses to run if
           the two disagree, so updating only one would leave the supplier
           unreadable;
        3. the comparison currently loaded, otherwise the new column
           would only take effect on the next recompute, not now;
        4. the manual override for that document, when one exists: it
           takes precedence over the registry, and leaving it stale would
           silently revert to the old column on the next recompute.
        """

        from launcher import resolve_source_path, source_rule  # noqa: PLC0415

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        supplier_id = str(payload.get("supplierId") or "").strip().casefold()
        if not supplier_id:
            raise ValueError("Indicare il fornitore di cui cambiare la colonna d'ordine")
        indice = colonna_ordine.indice_scelto(payload.get("colonna"))
        if indice is None:
            raise ValueError("Indicare la colonna in cui scrivere l'ordine")
        if self.pipeline_jobs.in_corso():
            raise LavoroGiaInCorso(
                "Un confronto è in corso: la colonna d'ordine si cambia quando ha finito."
            )
        lettera = colonna_ordine.lettera_di_indice(indice)
        nome = supplier_label(supplier_id)

        with self.lock:
            documento = self.pipeline_jobs.documento_del_fornitore(supplier_id)
            effettiva = documento["effettiva"]
            colonne = {
                int(voce["colonna"]): voce
                for voce in documento["colonne"]
                if isinstance(voce, dict) and voce.get("colonna")
            }
            voce = colonne.get(indice)
            if voce is None:
                raise ValueError(
                    f"Il foglio «{effettiva.get('sheet') or ''}» di {nome} non arriva alla colonna {lettera}."
                )
            occupata = colonna_ordine.campo_che_occupa(effettiva, indice)
            if occupata:
                raise ValueError(
                    f"La colonna {lettera} è quella che leggo come {occupata}: "
                    "l'ordine la azzererebbe e la riscriverebbe, e quel dato andrebbe perso. "
                    "Scegline una libera."
                )
            formule = int(voce.get("formule") or 0)
            if formule:
                raise ValueError(
                    f"La colonna {lettera} di {nome} contiene {formule} "
                    f"{'formula' if formule == 1 else 'formule'}: scrivere l'ordine le cancella. "
                    "Scegline una senza calcoli dentro."
                )

            adattatore = documento["adattatore"]
            if not isinstance(adattatore, dict) or not adattatore:
                raise ValueError(
                    f"Il confronto in uso non dice con quale adattatore ha letto il listino {nome}."
                )
            scrittura = adattatore.get("order_write")
            if not isinstance(scrittura, dict) or not scrittura:
                raise ValueError(
                    f"Il registro non dichiara come si scrive l'ordine dentro il listino {nome}: "
                    "la colonna si sceglie quando ti chiedo le colonne del documento."
                )
            intestazione = voce.get("intestazione")
            # Some price lists have no header row at all, so there is no
            # cell above the column to confirm as empty in that case.
            # Requiring one would block this with an error that blames the
            # registry for something it never claimed.
            c_e_intestazione = bool(effettiva.get("headerRow"))
            candidato = dict(adattatore)
            # Keeping the adapter's original id here would write a full
            # copy of the row that was just sent to the learned registry,
            # and from that point no further update to that adapter would
            # ever arrive — the same failure mode as a stale fingerprint
            # left shadowed under a learned copy. The new entry instead gets
            # its own id, following the same suffix convention the guided
            # mapping flow uses, kept in one place.
            candidato["id"] = registro.identificativo_da_scrivere(
                adattatore.get("id"), self.pipeline_jobs.configurazione.adapters_path
            )
            if candidato["id"] != adattatore.get("id"):
                # Recorded in the entry itself, not only in the id's suffix,
                # the same way `impara_adattatore` does it: a reader coming
                # back later can tell where it came from without knowing the
                # suffix convention.
                candidato["derivato_da"] = adattatore.get("id")
            candidato["order_write"] = colonna_ordine.dichiarazione_aggiornata(
                scrittura, lettera, intestazione, c_e_intestazione=c_e_intestazione
            )
            da_mappatura = bool(candidato["order_write"].get("from_field_mapping"))
            if da_mappatura:
                candidato["field_mapping"] = colonna_ordine.mappatura_aggiornata(
                    adattatore.get("field_mapping"), lettera, intestazione,
                    c_e_intestazione=c_e_intestazione,
                )

            review = load_json(self.review_path, {})
            scheda = self._scheda_del_fornitore(review, supplier_id)
            if scheda is None:
                raise ValueError(f"Nel confronto in uso non c'è nessun listino di {nome}.")
            sorgente = resolve_source_path(
                scheda.get("sourcePath")
                or scheda.get("source_path")
                or scheda.get("originalPath")
                or scheda.get("path")
            )
            if sorgente is None or not sorgente.is_file():
                raise ValueError(f"Il listino di {nome} non è più al suo posto sul disco.")
            prova = dict(scheda)
            if da_mappatura:
                prova["fieldMapping"] = colonna_ordine.mappatura_aggiornata(
                    scheda.get("fieldMapping") or scheda.get("field_mapping"),
                    lettera,
                    intestazione,
                    c_e_intestazione=c_e_intestazione,
                )
            regola, motivo = source_rule(supplier_id, sorgente, prova, candidato)
            if regola is None:
                # The message is the real check's own message, not a copy
                # written here: it's the same one that would show up on the
                # next startup, and it must read identically in both places.
                raise ValueError(motivo or f"{nome}: la colonna {lettera} non è scrivibile.")

            registro.scrivi_adattatore(candidato, self.pipeline_jobs.configurazione.adapters_path)
            if da_mappatura:
                scheda["fieldMapping"] = prova["fieldMapping"]
                atomic_json(self.review_path, review)
            self._aggiorna_decisione_a_mano(
                scheda,
                prova.get("fieldMapping") if da_mappatura else None,
                lettera,
                intestazione,
                c_e_intestazione=c_e_intestazione,
            )
            # From this point the column has already been changed: the
            # declaration is written and passed verification against the
            # document. If rebuilding the write configuration fails — Node
            # missing, folder not writable — letting the exception propagate
            # would tell the user nothing happened, when almost everything
            # did. The response instead states what changed, what didn't,
            # and how to recover.
            avviso = ""
            try:
                self.riconfigura_compilazione(review)
            except Exception as exc:  # noqa: BLE001 - report whatever the cause is
                avviso = (
                    f"La colonna è cambiata, ma la configurazione di scrittura non si è "
                    f"rifatta: {frase(exc)} Si rifà da sola al prossimo confronto, oppure "
                    "chiudendo e riaprendo il programma."
                )

        titolo = " ".join(str(intestazione or "").split())
        return {
            "ok": True,
            "supplierId": supplier_id,
            "supplierName": nome,
            "colonna": lettera,
            "intestazione": titolo,
            "regola": regola,
            "avviso": avviso,
            "message": (
                f"L'ordine di {nome} si scrive nella colonna {lettera}"
                + (f" («{titolo}»)" if titolo else ", che non ha intestazione")
                + (". Vale da subito, anche per la compilazione di adesso." if not avviso
                   else ". " + avviso)
            ),
        }

    @staticmethod
    def _scheda_del_fornitore(review: Any, supplier_id: str) -> dict[str, Any] | None:
        """The document entry for this supplier's price list in the comparison."""

        cercato = str(supplier_id or "").strip().casefold()
        for voce in (review or {}).get("files") or []:
            if not isinstance(voce, dict):
                continue
            chiave, _nome = ReviewStore._fornitore_della_voce(voce)
            if chiave == cercato:
                return voce
        return None

    def _aggiorna_decisione_a_mano(
        self,
        scheda: dict[str, Any],
        mappatura: dict[str, Any] | None,
        lettera: str,
        intestazione: Any,
        *,
        c_e_intestazione: bool = True,
    ) -> None:
        """Keep any manual override for this document in sync with the new column.

        A manual override takes precedence over the adapter registry (see
        `pipeline_jobs`). Leaving it stale would mean seeing the new column
        today and silently reverting to the old one on the next recompute.
        """

        percorso = self.pipeline_jobs.configurazione.decisioni_manuali_path
        documento = load_json(percorso, None)
        decisioni = documento.get("decisions") if isinstance(documento, dict) else documento
        if not isinstance(decisioni, list) or not decisioni:
            return
        nome = str(scheda.get("name") or "").strip().casefold()
        impronta = str(scheda.get("sourceSha256") or scheda.get("source_sha256") or "").strip().casefold()
        cambiate = False
        for decisione in decisioni:
            if not isinstance(decisione, dict):
                continue
            suo_nome = str(decisione.get("file_name") or "").strip().casefold()
            sua_impronta = str(decisione.get("file_sha256") or "").strip().casefold()
            if suo_nome != nome and not (impronta and sua_impronta == impronta):
                continue
            decisione["field_mapping"] = colonna_ordine.mappatura_aggiornata(
                mappatura if mappatura is not None else decisione.get("field_mapping"),
                lettera,
                intestazione,
                c_e_intestazione=c_e_intestazione,
            )
            cambiate = True
        if cambiate:
            atomic_json(percorso, {"decisions": decisioni})

    def review_with_manual_products(self, state: dict[str, Any] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
        """Return the comparison decorated with manual state. `state` lets a
        caller decorate against state it already holds in memory.

        Reading state from disk is right for every caller except one: a
        caller that needs to see the comparison as it will be AFTER a
        decision it hasn't written yet. `set_supplier_discount` is that
        case — it reassigns supplier selection based on a discount it just
        set — and it must pass its in-memory state through explicitly
        rather than rebuilding the comparison on its own, which would skip
        manual matches and answered proposals.
        """

        review = self.base_review()
        if state is None:
            state = load_json(self.state_path, {})
        known_ids = {str(item.get("id") or "") for item in review.get("products") or []}
        for product in state.get("manualProducts") or []:
            product_id = str(product.get("id") or "")
            if not product_id or product_id in known_ids:
                continue
            review.setdefault("products", []).append(deepcopy(product))
            known_ids.add(product_id)
        for product in review.get("products") or []:
            is_display = str(product.get("itemType") or product.get("kind") or "").casefold() == "display"
            product["quantityLabel"] = "espositori" if is_display else "colli"
            product["orderUnitLabel"] = "espositori" if is_display else "colli"
        self.apply_match_overrides(review, state)
        # Manually chosen rows from the price-list viewer: right after the
        # answers to the automatic matching, since they're the same kind of
        # decision — "this row is my product" — and before discounts, which
        # must also apply to an offer that was just matched in.
        self._applica_abbinamenti_manuali(review, state)
        # Declared rejections: after both positive answers — the accepted
        # proposal and the manually chosen row — since a rejection is the
        # last word on that offer. There's no conflict with a manual match:
        # that rewrites the row's barcode and description, so its
        # fingerprint changes and an old rejection no longer covers it. If
        # the user manually picks the very row they had previously
        # rejected, the rejection is cleared by `abbina_riga_di_listino`
        # instead: two human answers on the same product, the more recent
        # one wins.
        self.spegni_le_offerte_rifiutate(review)
        # Here, and only here. This is the single point every price-touching
        # read goes through: the page, the summary totals, the
        # move-quantity preview, the order plan. Applying the discount later
        # — inside `offer_pricing`, say — would mean discounted prices in
        # the service's own totals and full prices on screen, two different
        # numbers for the same row. Applying it earlier doesn't work either:
        # inside `base_review` the catalog rewrites offers with fresh prices
        # (`catalog_search.enrich_review`), which would silently drop the
        # discount on just the products the catalog happens to know.
        self.applica_sconti_fornitore(review, self.sconti_del_confronto(state, review))
        return review, state

    # The header discount a supplier applies across their whole price list
    # (e.g. "this supplier discounts everything by 6%"). It doesn't live in
    # the adapter registry — that describes how a document is READ — but in
    # this week's state, alongside the other decisions made this week.
    @staticmethod
    def sconti_del_confronto(state: Any, review: Any) -> dict[str, float]:
        """Discounts valid for THIS comparison, already validated.

        They expire on their own at the next recompute, with no explicit
        notice: a discount is set right after uploading price lists and is
        meant to apply silently from then on. The expiry check happens here
        rather than in a periodic cleanup: a discount tied to a different
        run simply never applies, so there's no window where it could apply
        by mistake.
        """

        run = str(((review or {}).get("run") or {}).get("id") or "")
        voci = (state or {}).get("supplierDiscounts")
        if not run or not isinstance(voci, dict):
            return {}
        validi: dict[str, float] = {}
        for chiave, voce in voci.items():
            if not isinstance(voce, dict) or str(voce.get("runId") or "") != run:
                continue
            rate = number(voce.get("rate"))
            if rate is None or not 0 < rate < 1:
                continue
            validi[str(chiave).strip().casefold()] = float(rate)
        return validi

    @staticmethod
    def applica_sconti_fornitore(review: dict[str, Any], sconti: dict[str, float]) -> None:
        """Apply the discount rate to every price of that supplier's offers."""

        if not sconti:
            return
        for product in review.get("products") or []:
            if not isinstance(product, dict):
                continue
            for offer in product.get("offers") or []:
                if not isinstance(offer, dict):
                    continue
                rate = sconti.get(str(offer.get("supplierId") or "").strip().casefold())
                if not rate:
                    continue
                for campo in ("unitPriceNet", "orderUnitPriceNet", "price", "pricePerPiece", "unitPricePreDiscount"):
                    valore = number(offer.get(campo))
                    if valore is not None:
                        offer[campo] = round(valore * (1 - rate), 4)
                # The "vs. last paid" price-difference fields are computed
                # upstream from the full, undiscounted price: keeping them
                # here would point at a number the screen no longer shows.
                # Dropped here, the page recomputes them from the price it
                # actually has.
                offer.pop("lastPriceDifference", None)
                offer.pop("lastPriceDifferencePct", None)
                offer["supplierDiscountRate"] = rate

    @staticmethod
    def _candidate_matches_warning(warning: dict[str, Any], supplier_id: str, candidate_key: str) -> bool:
        return (
            str(warning.get("code") or "") == "RIFIUTO_CON_CANDIDATO_FORTE"
            and str(warning.get("supplierId") or "") == supplier_id
            and str(warning.get("candidateKey") or "") == candidate_key
        )

    def apply_match_overrides(self, review: dict[str, Any], state: dict[str, Any]) -> None:
        """Apply human corrections to the current run's flagged rejections.

        A decision only applies when the run, product, supplier and the
        proposed row's fingerprint all match. Next week's price list can
        reuse the same row number for a different item; in that case the
        fingerprint changes and the old decision stays inert.
        """

        review_run = str((review.get("run") or {}).get("id") or "")
        if str(state.get("runId") or "") != review_run:
            return
        overrides = {
            (str(item.get("productId") or ""), str(item.get("supplierId") or "")): item
            for item in state.get("matchOverrides") or []
            if isinstance(item, dict) and str(item.get("runId") or "") == review_run
        }
        for product in review.get("products") or []:
            product_id = str(product.get("id") or "")
            for offer in product.get("offers") or []:
                supplier_id = str(offer.get("supplierId") or "")
                candidate = offer.get("rejectedCandidate")
                if not isinstance(candidate, dict):
                    continue
                candidate_key = str(candidate.get("candidateKey") or "")
                override = overrides.get((product_id, supplier_id))
                if not candidate_key or not override or str(override.get("candidateKey") or "") != candidate_key:
                    continue
                accepted = override.get("accepted") is True
                offer["candidateDecision"] = "accepted" if accepted else "rejected"
                if accepted and candidate.get("available") is True:
                    preserved = deepcopy(candidate)
                    offer.update(deepcopy(candidate))
                    offer["rejectedCandidate"] = preserved
                    offer["available"] = True
                    offer["status"] = "SEMANTICO_CONFERMATO_UTENTE"
                    offer["matchStatus"] = "SEMANTICO_CONFERMATO_UTENTE"
                    offer["method"] = "CORREZIONE_UTENTE"
                    offer["confidence"] = "UTENTE"
                    offer["requiresConfirmation"] = False
                    offer["confirmed"] = True
                product["warnings"] = [
                    warning for warning in product.get("warnings") or []
                    if not (
                        isinstance(warning, dict)
                        and self._candidate_matches_warning(warning, supplier_id, candidate_key)
                    )
                ]

        # The overall count written by the pipeline builder can't know
        # about answers that arrived afterward, so it's rebuilt here from
        # the cases that are still open.
        review["warnings"] = [
            warning for warning in review.get("warnings") or []
            if not (isinstance(warning, dict) and warning.get("code") == "RIFIUTI_CON_CANDIDATO_FORTE")
        ]
        aperti = [
            product for product in review.get("products") or []
            if any(
                isinstance(warning, dict) and warning.get("code") == "RIFIUTO_CON_CANDIDATO_FORTE"
                for warning in product.get("warnings") or []
            )
        ]
        if aperti:
            review.setdefault("warnings", []).append({
                "code": "RIFIUTI_CON_CANDIDATO_FORTE",
                "severity": "warning",
                "blocking": False,
                "title": "Proposte da controllare",
                "message": (
                    (
                        "1 prodotto ha una proposta di un fornitore da confermare o rifiutare. "
                        if len(aperti) == 1
                        else f"{len(aperti)} prodotti hanno una proposta di un fornitore da confermare o rifiutare. "
                    )
                    + "Li trovi con il filtro «Da confermare»."
                ),
                "count": len(aperti),
            })

    def _profilo_ha_ancora_la_sua_copia(self, profile: Any) -> bool:
        """Check whether the profile describes a document still actually present in uploads.

        A profile is not the document, it's a record about it. If the file
        is deleted from the folder directly — outside the app — the record
        stays behind and keeps claiming the file exists. This is the single
        check that tells a real record apart from a stale one, kept in one
        place because every caller needs the same answer: without it,
        `upload` could reject a re-upload as "already present" for a file
        that's no longer on disk, with the page's own ghost-filtering
        leaving nothing visible to delete to break out of that state.
        """

        if not isinstance(profile, dict):
            return False
        name = str(profile.get("file_name") or "")
        path = consegna.file_sicuro(self.upload_dir, name) if name else None
        if path is None:
            return False
        try:
            return Path(str(profile.get("path") or "")).resolve() == path.resolve()
        except OSError:
            return False

    def _profili_veri_e_fantasmi(
        self, profiles_doc: dict[str, Any]
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Split profiles into those with a matching file on disk and those without."""

        veri: list[dict[str, Any]] = []
        fantasmi: list[dict[str, Any]] = []
        for profile in profiles_doc.get("profiles") or []:
            if not isinstance(profile, dict):
                continue
            (veri if self._profilo_ha_ancora_la_sua_copia(profile) else fantasmi).append(profile)
        return veri, fantasmi

    def _safe_upload_profiles(self, profiles_doc: dict[str, Any]) -> dict[str, dict[str, Any]]:
        """Profiles whose path really matches a copy in the uploads folder."""

        veri, _ = self._profili_veri_e_fantasmi(profiles_doc)
        return {str(profile.get("file_name") or ""): profile for profile in veri}

    @staticmethod
    def _fornitore_della_voce(voce: dict[str, Any]) -> tuple[str, str]:
        """Return the supplier key and name for a document entry, or empty strings.

        Document entries don't all have the same shape: ones produced by a
        recompute carry `supplierId`, while entries from an older
        comparison still pointing at a source file outside uploads carry
        only `supplier`, the full name. Reading just one of the two misses
        half the cases, so this fallback lives here, written once.
        """

        if str(voce.get("role") or "").strip().casefold() == "master":
            return "", ""
        etichetta = str(voce.get("supplier") or "").strip()
        chiave = str(voce.get("supplierId") or "").strip().casefold()
        if not chiave and etichetta.casefold() not in {
            "", "gestionale", "da confermare", "da riconoscere",
        }:
            chiave = etichetta.casefold()
        if not chiave:
            return "", ""
        return chiave, (etichetta or chiave.upper())

    @staticmethod
    def _spegni_i_fornitori_senza_documento(
        review: dict[str, Any], rimossi: dict[str, str]
    ) -> None:
        """A supplier whose price list was removed stops being orderable.

        `delete_upload` removes the file's document entry and `base_review`
        filters `review["files"]`, but neither one touches `suppliers` or
        `products`: without this step, a supplier would stay among every
        product's offers, stay selected, and stay counted in the summary
        total after their price list was deleted.

        Nothing is deleted here: `available` is turned off, which is
        already the flag the rest of the program uses to mean "cannot be
        ordered from here". The page reassigns the best offer among what's
        still available and reports the change, validation rejects an
        attempt to order from the removed supplier anyway, and if the price
        list comes back the comparison is restored intact — this is a
        derived, reversible state, not a deletion.

        `rimossi` contains only suppliers that had a document and lost it.
        A supplier never seen in this comparison stays orderable: nothing
        is known about it, and disabling it would be a guess.
        """

        if not rimossi:
            return

        toccati: set[str] = set()
        for prodotto in review.get("products") or []:
            if not isinstance(prodotto, dict):
                continue
            for offerta in prodotto.get("offers") or []:
                if not isinstance(offerta, dict):
                    continue
                nome = str(offerta.get("supplierId") or offerta.get("supplier_id") or "").strip().casefold()
                if nome not in rimossi:
                    continue
                offerta["available"] = False
                offerta["warning"] = (
                    f"Il listino {rimossi[nome]} è stato eliminato: questa offerta non è "
                    "più ordinabile finché non ricarichi il listino e rifai il confronto."
                )
                toccati.add(nome)

        for voce in review.get("suppliers") or []:
            if isinstance(voce, dict) and str(voce.get("id") or "").strip().casefold() in rimossi:
                voce["documentMissing"] = True

        if not toccati:
            return
        nomi = sorted(rimossi[nome] for nome in toccati)
        elenco = ", ".join(nomi)
        review.setdefault("warnings", []).append({
            "code": "LISTINO_ELIMINATO",
            "severity": "warning",
            "blocking": False,
            "suppliers": sorted(toccati),
            "title": (
                f"{elenco}: listino eliminato" if len(nomi) == 1
                else f"{len(nomi)} listini eliminati"
            ),
            "message": (
                f"Hai eliminato il listino di {elenco} dopo l'ultimo confronto. Le sue offerte "
                "restano visibili ma non si possono più ordinare, e i prodotti che gli erano "
                "assegnati sono passati a un altro fornitore. Ricarica il listino e rifai il "
                "confronto per rimetterlo in gioco."
            ),
        })

    def base_review(self) -> dict[str, Any]:
        review = load_json(self.review_path, {
            "run": {"id": "", "status": "awaiting_files", "createdAt": datetime.now(tz=timezone.utc).isoformat(), "label": "Nuova run"},
            "files": [],
            "suppliers": [],
            "products": [],
            "warnings": [],
        })
        if not isinstance(review, dict):
            raise ValueError("review_data.json non contiene un oggetto")
        profiles = load_json(self.upload_profiles_path, {"profiles": []})
        safe_profiles = self._safe_upload_profiles(profiles)
        visible_files = []
        fornitori_senza_documento: dict[str, str] = {}
        for file_entry in review.get("files") or []:
            if not isinstance(file_entry, dict):
                continue
            name = str(file_entry.get("name") or "")
            profile = safe_profiles.get(name)
            ruolo_scelto = profiled_upload_role(profile)
            if ruolo_scelto:
                file_entry["role"] = ruolo_scelto
                file_entry["kind"] = "Gestionale" if ruolo_scelto == "master" else "Listino"
                file_entry["supplier"] = "Gestionale" if ruolo_scelto == "master" else str(file_entry.get("supplier") or "Da confermare")
            source = str(file_entry.get("sourcePath") or "")
            inside_uploads = False
            if source:
                try:
                    inside_uploads = Path(source).resolve().parent == self.upload_dir
                except OSError:
                    inside_uploads = False
            if inside_uploads:
                # The live comparison stays valid, but the Import page
                # describes the copies still available for the next
                # recompute. A deleted copy must not reappear just because
                # the previous run's JSON still mentions it.
                path = consegna.file_sicuro(self.upload_dir, name)
                if path is None:
                    # This branch knows something worth acting on: this
                    # supplier had a document and it is gone. That's the
                    # difference between "unknown" and "removed", and only
                    # the latter justifies disabling its offers. A supplier
                    # that was never in this list at all doesn't reach this
                    # branch, and stays orderable — nothing is known about it.
                    chiave, etichetta = self._fornitore_della_voce(file_entry)
                    if chiave:
                        fornitori_senza_documento[chiave] = etichetta
                    continue
                if profile is not None:
                    file_entry["deletable"] = True
                    file_entry["uploadName"] = name
            # A comparison created before the guided import flow existed
            # may still point at the original file the user picked, outside
            # the uploads folder. The page must be able to remove it from
            # its own list without deleting that file from disk.
            if str(file_entry.get("role") or "").casefold() in {"supplier", "master"}:
                file_entry["deletable"] = True
                file_entry["uploadName"] = str(file_entry.get("name") or "")
            visible_files.append(file_entry)
        review["files"] = visible_files
        self._spegni_i_fornitori_senza_documento(review, fornitori_senza_documento)
        existing_names = {str(item.get("name") or "").casefold() for item in review.get("files") or []}
        for name, profile in safe_profiles.items():
            if name.casefold() in existing_names:
                continue
            hint = profile.get("deterministic_hint") or {}
            details = profile.get("details") or {}
            ruolo_scelto = profiled_upload_role(profile)
            gestionale = ruolo_scelto == "master"
            rows = details.get("active_range", {}).get("nonempty_rows")
            if details.get("format") == "xlsx":
                rows = sum(sheet.get("active_range", {}).get("nonempty_rows", 0) for sheet in details.get("sheets") or [])
            review.setdefault("files", []).append({
                "id": profile.get("profile_id"),
                "name": name,
                "kind": "Gestionale" if gestionale else "Documento acquisito",
                "supplier": "Gestionale" if gestionale else "Da confermare",
                "role": ruolo_scelto or None,
                "status": "review",
                "schemaState": hint.get("state") or "AMBIGUO",
                "rows": rows or 0,
                # This message must not name an internal tool the user has
                # no reason to know about, and must state what actually
                # happens: columns are recognized automatically by the
                # recompute, and the page asks for them only when that fails.
                "message": "Letto. Le colonne verranno riconosciute al prossimo confronto; se non le riconosce, te le chiede.",
                "sourcePath": str((self.upload_dir / name).resolve()),
                "deletable": True,
                "uploadName": name,
            })
        try:
            self.catalog.enrich_review(review)
        except Exception as exc:
            review.setdefault("warnings", []).append({
                "code": "CATALOGO_NON_AGGIORNATO",
                "severity": "warning",
                "blocking": False,
                "title": "La ricerca nei listini può non trovare tutto",
                "message": (
                    "Il confronto dei prezzi resta buono. Quello che può mancare è "
                    "«+ Aggiungi un prodotto», che cerca nei listini la merce che non è "
                    f"nell'elenco del gestionale. Il motivo, per chi ripara: {exc}"
                ),
            })
        # A supplier that fails to read doesn't block the rest of the
        # comparison, which is exactly why it must be reported here —
        # otherwise it silently drops out of product search with no
        # indication anywhere.
        for failure in self.catalog.load_errors:
            review.setdefault("warnings", []).append({
                "code": "CATALOGO_FORNITORE_NON_LETTO",
                "supplier": failure.get("supplier"),
                "severity": "warning",
                "blocking": False,
                "title": f"Listino {failure.get('supplierName') or 'fornitore'} non letto",
                "message": (
                    f"{frase(failure.get('message')) or 'Lettura non riuscita.'} "
                    "Gli altri fornitori restano nel confronto; questo non compare nella ricerca "
                    "prodotti. Per rimetterlo dentro, ricarica il listino dal passo 1 «Importa i dati»."
                ),
            })
        return review

    def review(self) -> dict[str, Any]:
        with self.lock:
            review, state = self.review_with_manual_products()
            decisions = {str(item.get("id")): item for item in state.get("products") or [] if item.get("id") is not None}
            riallineati: list[str] = []
            for product in review.get("products") or []:
                decision = decisions.get(str(product.get("id")))
                if not decision:
                    product["excluded"] = False
                else:
                    product["quantity"] = decision.get("quantity", product.get("quantity", 0))
                    # The saved supplier is restored only when the user
                    # explicitly chose it. If it was a default — the
                    # cheapest offer at the time — the comparison's current
                    # cheapest offer wins instead, i.e. the cheapest offer
                    # NOW.
                    #
                    # Without this distinction, a stale default would look
                    # like a user choice and survive a recompute that made
                    # it wrong: a new, cheaper price list gets added, the
                    # page shows the cheaper supplier, but the order still
                    # goes to the old one with nothing pointing that out.
                    #
                    # Decisions saved before this distinction existed carry
                    # no explicit marker and are treated as automatic; any
                    # that change supplier as a result are collected below
                    # so the page can report it.
                    salvato = str(decision.get("selectedSupplierId") or "")
                    # A single if/else with no early exit: an early
                    # `continue` here would skip the `quantitySource`
                    # restore below, and a product with both a manually
                    # chosen supplier and a manually entered quantity would
                    # revert to reporting "value from the management
                    # software" — and then get zeroed out by the command
                    # that resets only default quantities.
                    if str(decision.get("selectedSupplierSource") or "") == "utente":
                        product["selectedSupplierId"] = decision.get("selectedSupplierId", product.get("selectedSupplierId"))
                        product["confirmed"] = bool(decision.get("confirmed"))
                    else:
                        # Recomputed here rather than taken from
                        # `selectedSupplierId` in the comparison: that value
                        # comes from the matching chain, which knows nothing
                        # about header discounts. `review_with_manual_products`
                        # has just applied them to prices, so the cheapest
                        # offer here is the one the page actually shows —
                        # the only number worth reselecting on. Using the
                        # chain's own value would discard the discount-aware
                        # reselection.
                        migliore = self._offerta_piu_conveniente(product)
                        product["selectedSupplierId"] = str(migliore.get("supplierId") or "") if migliore else ""
                        if str(product.get("selectedSupplierId") or "") == salvato:
                            product["confirmed"] = bool(decision.get("confirmed"))
                        else:
                            # A supplier change means a different item: the
                            # earlier confirmation doesn't cover it, same as
                            # after a manual move.
                            product["confirmed"] = False
                            riallineati.append(str(product.get("name") or product.get("id") or ""))
                    product["excluded"] = bool(decision.get("excluded"))
                    # Without this restore, a product the user already
                    # edited would revert to "value from the management
                    # software" after a reload, and get zeroed out by
                    # mistake by the command that resets only default
                    # quantities.
                    if decision.get("quantitySource") in {"gestionale", "utente"}:
                        product["quantitySource"] = decision["quantitySource"]
                self.dichiara_la_conferma(product)
            # When a product's supplier changes without the user asking for
            # it, the explanation must land at the same time as the change.
            # Some products move to another supplier on their own here,
            # correctly — it's cheaper — but it still needs to be reported:
            # if the user was keeping one of those suppliers for their own
            # reason, this is the only point where they can notice and
            # restore it.
            # Only reported when the saved decisions belong to THIS
            # comparison. On a new run they're from a different week, the
            # supplier hasn't "moved" from anyone's point of view — it's a
            # different comparison — and the warning would just be noise
            # over something the program did correctly.
            stessa_run = str(state.get("runId") or "") == str((review.get("run") or {}).get("id") or "")
            if riallineati and stessa_run:
                quanti = len(riallineati)
                esempi = ", ".join(f"«{nome}»" for nome in riallineati[:3] if nome)
                coda = f" ({esempi}{', e altri' if quanti > 3 else ''})" if esempi else ""
                review.setdefault("warnings", []).append({
                    "code": "FORNITORE_PIU_CONVENIENTE_RIPRESO",
                    "severity": "warning",
                    "blocking": False,
                    "title": (
                        f"{quanti} prodotti sono passati al fornitore più conveniente"
                        if quanti > 1 else
                        "Un prodotto è passato al fornitore più conveniente"
                    ),
                    "message": (
                        f"Il confronto è cambiato da quando li avevi salvati, e su {'questi' if quanti > 1 else 'questo'} "
                        f"adesso costa meno un altro fornitore{coda}. Il fornitore che avevi scelto tu a mano "
                        "non è stato toccato: si sono spostati solo quelli che il programma aveva scelto da sé. "
                        "Se uno di questi lo tenevi apposta, riscegli il fornitore dalla riga del prodotto."
                    ),
                    "count": quanti,
                })
            try:
                review = self.promotion_service.decorate(review, decisions)
            except Exception as exc:
                # Promotions are supplementary: if the promotion engine
                # fails, the user must still see prices, quantities and
                # suppliers. A failure here must never block the page.
                print(f"[AVVISO] promozioni non calcolate — {type(exc).__name__}: {exc}")
                review.setdefault("warnings", []).append({
                    "code": "PROMOZIONI_NON_CALCOLATE",
                    "severity": "warning",
                    "blocking": False,
                    "title": "Non ho letto le offerte dei fornitori",
                    "message": (
                        "Prezzi, quantità e confronto restano buoni. Quello che manca è "
                        "quanto devi comprare per avere gli omaggi: quel conto non c'è. "
                        f"{frase(exc) or 'Il guasto non si è descritto.'}"
                    ),
                })
            # A supplier already flagged by the catalog doesn't get a second
            # warning — same file, same failure. But the first warning must
            # also mention this loss, or free-goods thresholds silently
            # disappear with nothing pointing it out.
            per_fornitore = {
                str(item.get("supplier") or ""): item
                for item in review.get("warnings") or []
                if item.get("supplier")
            }
            for failure in self.promotion_service.load_errors:
                esistente = per_fornitore.get(str(failure.get("supplier") or ""))
                if esistente is not None:
                    esistente["message"] = frase(esistente.get("message")) + " Mancano anche le sue soglie con omaggio."
                    continue
                review.setdefault("warnings", []).append({
                    "code": "CONDIZIONI_COMMERCIALI_NON_LETTE",
                    "supplier": failure.get("supplier"),
                    "severity": "warning",
                    "blocking": False,
                    "title": f"Offerte {failure.get('supplierName') or 'fornitore'} non lette",
                    "message": (
                        f"{frase(failure.get('message')) or 'Lettura non riuscita.'} "
                        "Prezzi e confronto restano validi; mancano le soglie con omaggio di questo fornitore."
                    ),
                })
            # A confirmation store that fails to open is never silent:
            # without this warning every already-answered question would
            # come back, and the user would reanswer by hand believing the
            # program never knew anything. The rest of the comparison stays
            # valid.
            if self._conferme_guasto:
                review.setdefault("warnings", []).append({
                    "code": "CONFERME_NON_DISPONIBILI",
                    "severity": "warning",
                    "blocking": False,
                    "title": "Le conferme già date non sono disponibili",
                    "message": (
                        f"{self._conferme_guasto} Prezzi, quantità e fornitori restano validi: "
                        "quello che manca è la memoria delle corrispondenze già confermate, "
                        "quindi le domande possono tornare anche se avevi già risposto."
                    ),
                })
            self.decorate_pending_orders(review)
            # The summary shows its totals right on first load, without
            # waiting for a save: the browser doesn't recompute them on its
            # own, and any rounding discrepancy is stated up front.
            review["orderSummary"] = self.order_summary_from_state(review, state)
            review["state"] = {
                "currentStep": max(1, min(3, int(number(state.get("currentStep")) or 1))),
                "acceptBelowThreshold": bool(state.get("acceptBelowThreshold")),
                "summaryGrouping": state.get("summaryGrouping") if state.get("summaryGrouping") in {"supplier", "product"} else "supplier",
                # The version this snapshot was built from; echoed back on
                # every save, which is how the service detects that another
                # snapshot was saved in the meantime.
                "stateVersion": int(number(state.get("stateVersion")) or 0),
                # Header discounts, as percentages, only for display in the
                # form field: the prices the page already has are already
                # discounted, so it must not recompute them.
                "supplierDiscounts": {
                    chiave: round(valore * 100, 4)
                    for chiave, valore in self.sconti_del_confronto(state, review).items()
                },
            }
            # Declared code equalities are not part of this payload: they
            # live in Settings, with names and search, under their own
            # route (`GET /api/matches/uguaglianze`), which is reachable
            # even with no comparison loaded.
            return review

    def read_order_history(self) -> dict[str, Any]:
        """Read order history, applying expiry and resaving it if it changed."""

        with self.lock:
            history, changed = order_history.read_history(self.history_path)
            if changed:
                try:
                    order_history.save_history(self.history_path, history)
                except OSError as exc:
                    # Expiry is just cleanup: if the save fails, the
                    # in-memory list is still correct for this read.
                    print(f"[AVVISO] Storico ordini non aggiornato su disco: {exc}")
            return history

    @staticmethod
    def avviso_ordini_scaduti(scaduti: list[dict[str, Any]]) -> dict[str, Any]:
        """Build the message reporting expired orders, without misstating who answered.

        The question stops being asked a fixed number of days after the
        last interaction; staying silent about that would let missing
        merchandise quietly drop out of view. "No answer" is stated only
        for orders that were genuinely never answered: an order answered
        "not arrived yet" several times must not be reported the same way
        as one nobody ever looked at.
        """

        frasi = []
        for voce in scaduti:
            momento = order_history.parse_moment(voce.get("createdAt"))
            # With no creation date on record, this states that fact rather
            # than fabricating a count.
            quando = (f" del {consegna.data_leggibile(momento.astimezone())}" if momento
                      else " (la voce non porta la data di creazione)")
            risposta = order_history.parse_moment(voce.get("answeredAt"))
            esito = (f", l'ultima risposta è del {consegna.data_leggibile(risposta.astimezone())}"
                     if risposta else ", mai risposto")
            frasi.append(f"l'ordine {voce.get('supplierName') or voce.get('supplier')}{quando}{esito}")
        elenco = "; ".join(frasi)
        return {
            "code": "ORDINI_SCADUTI_SENZA_RISPOSTA",
            "severity": "warning",
            "blocking": False,
            "title": "Ordini di cui non chiedo più notizie",
            "message": (
                f"Di {elenco} non ti chiedo più se la merce è arrivata: sono passati "
                f"{order_history.EXPIRY_DAYS} giorni da quando ne abbiamo parlato "
                "l'ultima volta. Se non è mai arrivata, va riordinata."
            ),
            "count": len(scaduti),
            "orders": scaduti,
        }

    def decorate_pending_orders(self, review: dict[str, Any]) -> None:
        """Flag on each product any orders not yet received.

        Matching is by item identity: the barcode when there is one,
        otherwise the product's stable identifier (displays have no
        barcode of their own). `product:<row>` identifiers are excluded:
        they depend on position in the weekly export, and matching on them
        would produce wrong matches with no way to tell.
        """

        products = review.get("products") or []
        try:
            history = self.read_order_history()
        except Exception as exc:
            for product in products:
                product["pendingOrders"] = []
            review.setdefault("warnings", []).append({
                "code": "STORICO_ORDINI_NON_LEGGIBILE",
                "severity": "warning",
                "blocking": False,
                "title": "Storico degli ordini non disponibile",
                "message": f"Il confronto resta utilizzabile, ma non posso ricordare la merce ordinata e non ancora ricevuta. Dettaglio: {exc}",
            })
            return
        order_history.attach_pending_orders(
            products,
            history,
            exclude_run_id=str((review.get("run") or {}).get("id") or ""),
        )
        scaduti = order_history.expired_orders(history)
        if scaduti:
            review.setdefault("warnings", []).append(self.avviso_ordini_scaduti(scaduti))

    def current_run_id(self) -> str:
        """Id of the currently open run: its own orders are excluded from pending-order flags."""

        try:
            review = load_json(self.review_path, {})
        except Exception:
            return ""
        return str(((review or {}).get("run") or {}).get("id") or "")

    def history_pending(self) -> dict[str, Any]:
        with self.lock:
            history = self.read_order_history()
            return {
                "ok": True,
                "pending": order_history.pending_summary(history, exclude_run_id=self.current_run_id()),
            }

    def answer_history_order(self, payload: Any) -> dict[str, Any]:
        if not isinstance(payload, dict):
            raise ValueError("Risposta non valida")
        order_id = str(payload.get("orderId") or "").strip()
        if not order_id:
            raise ValueError("Ordine non indicato")
        # "Will never arrive" closes the question permanently: a third
        # possible answer, needed when the goods will never show up and
        # marking the order received would be false.
        closed = payload.get("closed") is True
        received = payload.get("received")
        if not closed and not isinstance(received, bool):
            raise ValueError("Indicare se la merce è stata ricevuta: sì oppure no")
        with self.lock:
            history = self.read_order_history()
            registrata = (
                order_history.close_order(history, order_id)
                if closed
                else order_history.answer_order(history, order_id, bool(received))
            )
            if registrata is None:
                raise ValueError("Ordine non presente nello storico")
            order_history.save_history(self.history_path, history)
            return {
                "ok": True,
                "pending": order_history.pending_summary(history, exclude_run_id=self.current_run_id()),
            }

    def answer_rejected_candidate(self, payload: Any) -> dict[str, Any]:
        """Accept or reject a proposed row, only within the run that produced it."""

        if not isinstance(payload, dict):
            raise ValueError("Risposta non valida")
        run_id = str(payload.get("runId") or "").strip()
        product_id = str(payload.get("productId") or "").strip()
        supplier_id = str(payload.get("supplierId") or "").strip()
        candidate_key = str(payload.get("candidateKey") or "").strip()
        accepted = payload.get("accepted")
        if not all((run_id, product_id, supplier_id, candidate_key)) or not isinstance(accepted, bool):
            raise ValueError("Indicare la proposta e scegliere sì oppure no")

        with self.lock:
            review, state = self.review_with_manual_products()
            expected_run = str((review.get("run") or {}).get("id") or "")
            if run_id != expected_run:
                raise ValueError("Il confronto è cambiato: ricarica la pagina prima di rispondere")
            saved_run = str(state.get("runId") or "")
            if saved_run and saved_run != run_id:
                raise ValueError("Le scelte salvate appartengono a un altro confronto: ricarica la pagina")

            product = next(
                (item for item in review.get("products") or [] if str(item.get("id") or "") == product_id),
                None,
            )
            offer = next(
                (
                    item for item in (product or {}).get("offers") or []
                    if str(item.get("supplierId") or "") == supplier_id
                ),
                None,
            )
            candidate = (offer or {}).get("rejectedCandidate")
            if not isinstance(candidate, dict) or str(candidate.get("candidateKey") or "") != candidate_key:
                raise ValueError("Questa proposta non è più disponibile: ricarica il confronto")
            if accepted and candidate.get("available") is not True:
                raise ValueError("La riga proposta non ha prezzo e confezione utilizzabili")

            overrides = [
                item for item in state.get("matchOverrides") or []
                if isinstance(item, dict)
                and not (
                    str(item.get("productId") or "") == product_id
                    and str(item.get("supplierId") or "") == supplier_id
                )
            ]
            overrides.append({
                "runId": run_id,
                "productId": product_id,
                "supplierId": supplier_id,
                "candidateKey": candidate_key,
                "accepted": accepted,
                "answeredAt": datetime.now(tz=timezone.utc).isoformat(),
            })
            state["schemaVersion"] = int(number(state.get("schemaVersion")) or 1)
            state["runId"] = run_id
            state["matchOverrides"] = overrides
            state["updatedAt"] = datetime.now(tz=timezone.utc).isoformat()
            atomic_json(self.state_path, state)
            return {
                "ok": True,
                "message": "Prodotto collegato al fornitore." if accepted else "Proposta esclusa.",
            }

    def _avanza_la_versione(self, state: dict[str, Any], origine: str) -> int:
        """Save state with the version bumped by one, and return the new version.

        The returned number must reach the caller, not just get written to
        disk: every save states the version it started from, and a save
        against a stale version is rejected. A caller that bumps the
        version here without returning it to the client leaves the
        client's own copy one behind, and its next save — from that same
        client — gets rejected as conflicting with a save that never
        actually happened.
        """

        versione = int(number(state.get("stateVersion")) or 0) + 1
        state["stateVersion"] = versione
        state["stateVersionOrigin"] = origine
        atomic_json(self.state_path, state)
        return versione

    def abbina_riga_di_listino(self, payload: Any) -> dict[str, Any]:
        """Manually match a price-list row to a product: one click, two effects.

        Immediately: the offer enters the current comparison, the same as
        accepting an automatic-matching proposal.

        Permanently: the two barcodes are recorded as the same item.
        Remembering just the row position would only last a week — that
        position shifts with every new price list — while a barcode
        equality holds indefinitely and across every supplier, turning the
        match into a native `EAN_ESATTO` match on the next recompute.

        The permanent effect isn't always possible: a handful of products
        have no barcode in the management-software export, and some
        suppliers' display rows carry no barcode either. There, the match
        only holds for this comparison, and the response says so — treating
        it as permanent would be a promise broken by the next recompute.
        """

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        product_id = str(payload.get("productId") or "").strip()
        supplier_id = str(payload.get("supplierId") or "").strip().casefold()
        source_row = payload.get("sourceRow")
        if not product_id or not supplier_id or source_row in (None, ""):
            raise ValueError("Indicare il prodotto, il fornitore e la riga del listino")

        with self.lock:
            review, state = self.review_with_manual_products()
            run_id = str((review.get("run") or {}).get("id") or "")
            # `productId` and `sourceRow` are positional: a client stuck on
            # last week's comparison would point, on the new comparison, at
            # two arbitrary items — and declare them equal permanently and
            # across every supplier. Same check and same message as the
            # rejection route, since both answer the same question. A
            # request without `runId` is still accepted, so a stale cached
            # page doesn't break outright.
            dichiarata = str(payload.get("runId") or "").strip()
            if dichiarata and run_id and dichiarata != run_id:
                raise ValueError("Il confronto è cambiato: ricarica la pagina prima di rispondere")
            product = next(
                (item for item in review.get("products") or [] if str(item.get("id") or "") == product_id),
                None,
            )
            if product is None:
                raise ValueError("Questo prodotto non è nel confronto: ricarica la pagina")
            if not any(str(voce.get("id") or "").casefold() == supplier_id for voce in review.get("suppliers") or []):
                raise ValueError("Questo fornitore non è nel confronto")
            record, offerta = self.catalog.offerta_dalla_riga(review, supplier_id, source_row)

            abbinamenti = [
                voce for voce in state.get("manualMatches") or []
                if isinstance(voce, dict)
                and not (
                    str(voce.get("productId") or "") == product_id
                    and str(voce.get("supplierId") or "") == supplier_id
                )
            ]
            abbinamenti.append({
                "runId": run_id,
                "productId": product_id,
                "supplierId": supplier_id,
                "sourceRow": record.get("source_row"),
                "chosenAt": datetime.now(tz=timezone.utc).isoformat(),
            })
            state["manualMatches"] = abbinamenti
            # An earlier rejection of the automatic proposal for this
            # supplier must not turn off the offer just chosen manually:
            # both are human answers on the same product, and the more
            # recent one wins. It's removed outright, rather than left to
            # coexist with the risk of ambiguity about which one applies.
            rifiuti_tolti = 0
            rimasti = []
            for voce in state.get("matchOverrides") or []:
                if (
                    isinstance(voce, dict)
                    and str(voce.get("productId") or "") == product_id
                    and str(voce.get("supplierId") or "") == supplier_id
                ):
                    rifiuti_tolti += 1
                    continue
                rimasti.append(voce)
            state["matchOverrides"] = rimasti

            state["runId"] = run_id
            state["schemaVersion"] = int(number(state.get("schemaVersion")) or 1)
            state["updatedAt"] = datetime.now(tz=timezone.utc).isoformat()
            versione = self._avanza_la_versione(state, "abbinamento")

            # An earlier "not the same item" rejection on this supplier is
            # cleared too: both are human answers on the same product, and
            # the more recent one wins. This matters specifically when the
            # user browses the price list and manually picks the exact row
            # they had previously rejected: the fingerprint is the same, and
            # without this the rejection would silently turn the offer back
            # off right away.
            #
            # Only cleared when the current entry is a rejection: `dimentica`
            # doesn't distinguish, and would just as happily clear a
            # confirmation nobody revoked.
            articolo_del_prodotto = impronta_prodotto(product)
            magazzino_dei_no = self.magazzino_conferme()
            if magazzino_dei_no is not None and articolo_del_prodotto:
                try:
                    in_vigore = magazzino_dei_no.cerca(supplier_id, articolo_del_prodotto)
                    if in_vigore is not None and not in_vigore.get("accettata"):
                        magazzino_dei_no.dimentica(
                            supplier_id, articolo_del_prodotto,
                            quando=datetime.now(tz=timezone.utc).isoformat(),
                        )
                except MagazzinoNonUtilizzabile as exc:
                    # Today's match stays valid regardless: the failure is
                    # recorded and surfaced by the comparison, same as the
                    # equality-declaration failure below.
                    self._conferme_guasto = frase(exc) or "Il rifiuto precedente non è stato tolto."

            ricordata, motivo = self._ricorda_l_uguaglianza(product, record, supplier_id)
            return {
                "ok": True,
                "stateVersion": versione,
                "productId": product_id,
                "supplierId": supplier_id,
                "sourceRow": record.get("source_row"),
                "offerta": offerta,
                "uguaglianzaRicordata": ricordata,
                "message": motivo,
                "rifiutiTolti": rifiuti_tolti,
            }

    def _ricorda_l_uguaglianza(
        self, product: Any, record: Any, supplier_id: str
    ) -> tuple[bool, str]:
        """Record that the two barcodes are the same item.

        Returns `(recorded, message to show)`. Never raises: if the store
        fails to open, today's match stays valid regardless, but the part
        that's lost — it holding for next week too — must be reported, not
        hidden behind an apparent success.
        """

        codice_prodotto = codice_confrontabile((product or {}).get("ean"))
        codice_riga = codice_confrontabile((record or {}).get("ean"))
        nome = str((record or {}).get("description") or "").strip()
        if not codice_prodotto or not codice_riga:
            manca = "il prodotto" if not codice_prodotto else "la riga del listino"
            return False, (
                f"Abbinato per questo confronto. Non vale per i prossimi: {manca} non ha un "
                "codice a barre, e senza non c'è niente da dichiarare uguale."
            )
        if codice_prodotto == codice_riga:
            return False, "Abbinato: i due codici a barre erano già lo stesso."
        magazzino = self.magazzino_conferme()
        if magazzino is None:
            return False, (
                "Abbinato per questo confronto. Non vale per i prossimi: "
                + (self._conferme_guasto or "il file delle conferme non si apre.")
            )
        try:
            magazzino.unisci(
                codice_prodotto,
                codice_riga,
                articolo=impronta_prodotto(product),
                offerta=impronta_articolo({**(record or {}), "supplierId": supplier_id}),
                motivo=f"scelta a mano dal listino {supplier_label(supplier_id)}, riga {record.get('source_row')}",
                quando=datetime.now(tz=timezone.utc).isoformat(),
            )
        except (MagazzinoNonUtilizzabile, ValueError) as exc:
            return False, f"Abbinato per questo confronto. Non vale per i prossimi: {frase(exc)}"
        return True, (
            f"Abbinato a «{nome}». D'ora in poi {codice_prodotto} e {codice_riga} valgono come lo "
            "stesso articolo, per tutti i fornitori."
        )

    def rifiuta_l_abbinamento(self, payload: Any) -> dict[str, Any]:
        """Reject a proposed match ("not the same item") and undo it.

        Without this route, a user unable to confirm a proposed match had no
        way out: confirming orders the wrong item, leaving it unconfirmed
        blocks compilation on "confirmation required", and excluding the
        product from the order zeroes its quantity — which also drops it
        from the to-be-sourced list, since that list skips zero-quantity
        products.

        Recorded in the confirmation store as a row with `accettata` false;
        `MagazzinoConferme.ricorda` already supports writing that shape, so
        no new storage is introduced. It is deliberately not written into
        `state.json`, which lives under the current run and is discarded on
        recompute, while this answer must survive it: the semantic proposal
        would otherwise come back every week, and a rejection that expired
        on each recompute would put the user back in the same dead end.

        The rejection applies to the pair (supplier, management-software
        item), anchored to the fingerprint of the price-list row it was
        said about. It is not a negated equality: an equality would apply
        to every supplier forever, and one supplier's row not being the
        right item says nothing about what another supplier carries.

        Only this product's own decision is touched in state: the selected
        supplier is cleared and its confirmation drops, but the quantity is
        left untouched — it's still the number needed to source the item
        elsewhere.
        """

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        product_id = str(payload.get("productId") or "").strip()
        supplier_id = str(payload.get("supplierId") or "").strip().casefold()
        rifiutata = payload.get("rifiutata")
        if not product_id or not supplier_id or not isinstance(rifiutata, bool):
            raise ValueError("Indicare il prodotto, il fornitore e la risposta")

        with self.lock:
            review, state = self.review_with_manual_products()
            run_id = str((review.get("run") or {}).get("id") or "")
            dichiarata = str(payload.get("runId") or "").strip()
            if dichiarata and run_id and dichiarata != run_id:
                raise ValueError("Il confronto è cambiato: ricarica la pagina prima di rispondere")
            product = next(
                (item for item in review.get("products") or [] if str(item.get("id") or "") == product_id),
                None,
            )
            if product is None:
                raise ValueError("Questo prodotto non è nel confronto: ricarica la pagina")
            # `find_offer` runs on the already-decorated comparison: if the
            # rejection is already in force, the offer is still here with
            # `available` false, and is found all the same. This is what
            # makes the undo path work.
            offerta = find_offer(product, supplier_id)
            if offerta is None:
                raise ValueError("Questo fornitore non ha nessuna riga per questo prodotto")
            articolo = impronta_prodotto(product)
            impronta = impronta_articolo(offerta)
            # Fingerprints checked before opening the store, as in
            # `conferma_in_vigore`: a row that can't be fingerprinted can't
            # be recorded, and opening the store just to find that out would
            # create a `conferme.db` for a program that has none yet.
            if not articolo or not impronta:
                raise ValueError(
                    "Questa riga non si può rifiutare: prodotto o riga del listino non hanno "
                    "né codice a barre né nome, e una risposta agganciata al vuoto tornerebbe "
                    "su merce a caso"
                )
            magazzino = self.magazzino_conferme()
            if magazzino is None:
                raise ValueError(self._conferme_guasto or "Il file delle conferme non si apre.")
            quando = datetime.now(tz=timezone.utc).isoformat()
            nome_fornitore = supplier_label(supplier_id)
            try:
                if rifiutata:
                    magazzino.ricorda(
                        fornitore=supplier_id,
                        articolo=articolo,
                        offerta=impronta,
                        accettata=False,
                        motivo="non è lo stesso articolo",
                        quando=quando,
                    )
                    # This message is a toast that disappears after a few
                    # seconds, so it confirms the action taken rather than
                    # restating the full rule for how long a rejection lasts
                    # — that explanation is already shown, persistently, on
                    # the rejected-supplier panel.
                    messaggio = f"{nome_fornitore} esce da questo prodotto. La quantità resta."
                else:
                    # Only cleared when the current entry is a rejection.
                    # `dimentica` doesn't distinguish: calling it
                    # unconditionally would just as happily clear a
                    # confirmation nobody revoked.
                    voce = magazzino.cerca(supplier_id, articolo)
                    if voce is not None and not voce.get("accettata"):
                        magazzino.dimentica(supplier_id, articolo, quando=quando)
                    messaggio = (
                        f"{nome_fornitore} torna fra i fornitori di questo prodotto: "
                        "la domanda è di nuovo aperta."
                    )
            except (MagazzinoNonUtilizzabile, ValueError) as exc:
                # Raised here, unlike `ricorda_le_conferme`, which records
                # the failure and continues. There, saving quantities must
                # not fail over a memory write; here the memory write IS the
                # answer, and a rejection the user believes was recorded but
                # wasn't would silently put them back in the same dead end.
                raise ValueError(f"La risposta non è stata registrata: {frase(exc)}") from exc

            # State: only this product's own decision, and only on keys
            # that `validate_snapshot` reconstructs from the snapshot
            # anyway. No new keys, so `CHIAVI_DI_STATO_RICOPIATE` needs no
            # changes.
            voci: list[Any] = []
            cambiato = False
            for voce in state.get("products") or []:
                if (
                    rifiutata
                    and isinstance(voce, dict)
                    and str(voce.get("id") or "") == product_id
                    and str(voce.get("selectedSupplierId") or "").strip().casefold() == supplier_id
                ):
                    ripulita = {**voce, "selectedSupplierId": "", "confirmed": False}
                    ripulita.pop("confirmedArticle", None)
                    cambiato = True
                    voci.append(ripulita)
                    continue
                voci.append(voce)
            # The version is returned even when nothing was written: the
            # client still needs to stay in sync, and a silent response
            # would leave it guessing.
            versione = int(number(state.get("stateVersion")) or 0)
            if cambiato:
                state["products"] = voci
                state["runId"] = run_id
                state["schemaVersion"] = int(number(state.get("schemaVersion")) or 1)
                state["updatedAt"] = quando
                versione = self._avanza_la_versione(state, "rifiuto")
            return {
                "ok": True,
                "stateVersion": versione,
                "productId": product_id,
                "supplierId": supplier_id,
                "rifiutata": bool(rifiutata),
                "message": messaggio,
            }

    def togli_uguaglianza(self, payload: Any) -> dict[str, Any]:
        """Remove a "these two codes are the same item" declaration.

        Must be a single click: until it's removed, the declaration applies
        to every future comparison and every supplier. The row isn't
        deleted — if it already caused a wrong order, it's the only record
        explaining why — it's just closed.
        """

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        codici = payload.get("codici")
        if not isinstance(codici, (list, tuple)) or len(codici) != 2:
            raise ValueError("Indicare i due codici da separare")
        magazzino = self.magazzino_conferme()
        if magazzino is None:
            raise ValueError(self._conferme_guasto or "Il file delle conferme non si apre.")
        tolta = magazzino.separa(
            codici[0], codici[1], quando=datetime.now(tz=timezone.utc).isoformat(),
        )
        return {
            "ok": True,
            "tolta": tolta,
            # The updated list comes back already formatted: the page that
            # called this is the same one that displays it, and forcing a
            # second round trip for the same rows would be wasted work.
            "uguaglianze": self.elenco_delle_uguaglianze()["uguaglianze"],
            "message": (
                "Dichiarazione tolta: dal prossimo confronto quei due codici tornano a essere "
                "due articoli diversi."
                if tolta
                else "Quella dichiarazione non era in vigore."
            ),
        }

    def _applica_abbinamenti_manuali(self, review: dict[str, Any], state: dict[str, Any]) -> None:
        """Turn manually matched rows into offers within this comparison.

        Only applies to the run the match was made in: on recompute the row
        number no longer means anything, and it's the barcode equality —
        a separate mechanism, kept in the confirmation store — that keeps
        the match alive across weeks.
        """

        review_run = str((review.get("run") or {}).get("id") or "")
        scelte = {
            (str(voce.get("productId") or ""), str(voce.get("supplierId") or "")): voce
            for voce in state.get("manualMatches") or []
            if isinstance(voce, dict) and str(voce.get("runId") or "") == review_run
        }
        if not scelte:
            return
        for product in review.get("products") or []:
            product_id = str(product.get("id") or "")
            for offer in product.get("offers") or []:
                voce = scelte.get((product_id, str(offer.get("supplierId") or "").casefold()))
                if voce is None:
                    continue
                try:
                    _record, offerta = self.catalog.offerta_dalla_riga(
                        review, str(offer.get("supplierId") or ""), voce.get("sourceRow"),
                    )
                except ValueError:
                    # The price list changed underneath this match: the
                    # offer isn't reinstated, and the product stays as it
                    # was. Not a failure worth stopping over — the match
                    # stays recorded and becomes useful again on recompute,
                    # which resolves it by code instead of row number.
                    continue
                offer.update(deepcopy(offerta))
                offer["available"] = True
                offer["status"] = "SCELTO_A_MANO"
                offer["matchStatus"] = "SCELTO_A_MANO"
                offer["method"] = "SCELTA_UTENTE"
                offer["confidence"] = "UTENTE"
                offer["requiresConfirmation"] = False
                offer["confirmed"] = True
                offer["sceltaManuale"] = True

    def set_supplier_discount(self, payload: Any) -> dict[str, Any]:
        """Discount a supplier's offers and reassign products to the cheapest one.

        This applies a header discount across every offer of a supplier,
        then silently reselects the cheapest offer per product. Reselection
        happens right here, once, when the discount is set — not on every
        comparison read, which would redo the selection later too, undoing
        decisions made in the meantime.
        """

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        supplier_id = str(payload.get("supplierId") or "").strip().casefold()
        if not supplier_id:
            raise ValueError("Indicare il fornitore")
        percento = number(payload.get("percent"))
        if percento is None or percento < 0 or percento >= 100:
            raise ValueError("Lo sconto è una percentuale fra 0 e 99")

        with self.lock:
            review, state = self.review_with_manual_products()
            run_id = str((review.get("run") or {}).get("id") or "")
            if not run_id:
                raise ValueError("Non c'è nessun confronto su cui applicare uno sconto")
            fornitori = {str(voce.get("id") or "").strip().casefold() for voce in review.get("suppliers") or []}
            if supplier_id not in fornitori:
                raise ValueError("Questo fornitore non è nel confronto")

            sconti = dict(state.get("supplierDiscounts") or {})
            if percento <= 0:
                sconti.pop(supplier_id, None)
            else:
                sconti[supplier_id] = {
                    "rate": round(float(percento) / 100, 6),
                    "runId": run_id,
                    "setAt": datetime.now(tz=timezone.utc).isoformat(),
                }
            state["supplierDiscounts"] = sconti
            state["schemaVersion"] = int(number(state.get("schemaVersion")) or 1)
            state["runId"] = run_id

            # The comparison is re-read with the discount just decided,
            # because reselection must run against the prices the user will
            # actually see, not the ones from before this change.
            #
            # It's re-read through the same shared decoration path used by
            # every other caller, passed the in-memory state — the discount
            # isn't on disk yet, `_avanza_la_versione` writes it below.
            # Rebuilding the comparison from `base_review()` alone here would
            # miss answered proposals and manually matched rows, and
            # reselection would run against offers the page doesn't
            # actually show.
            scontato, _ = self.review_with_manual_products(state)

            decisioni = {str(item.get("id")): item for item in state.get("products") or [] if isinstance(item, dict)}
            spostati = 0
            for product in scontato.get("products") or []:
                migliore = self._offerta_piu_conveniente(product)
                if migliore is None:
                    continue
                product_id = str(product.get("id") or "")
                decisione = decisioni.get(product_id)
                if decisione is None:
                    # A product with no saved decision is not silently
                    # skipped: nothing guarantees every product already has
                    # one, and skipping it here would make the discount
                    # reassign only part of the comparison with no
                    # indication why. A decision is created instead.
                    #
                    # `quantitySource` is copied from the comparison, not
                    # invented: it's what tells the pipeline whether the
                    # quantity comes from the management software, and
                    # getting it wrong freezes the quantity out of future
                    # recomputes (`pipeline_jobs._ripulisci_stato`).
                    origine = str(product.get("quantitySource") or "")
                    decisione = {
                        "id": product_id,
                        "quantity": int(number(product.get("quantity")) or 0),
                        "quantitySource": origine if origine in {"gestionale", "utente"} else "gestionale",
                        "excluded": bool(product.get("excluded")),
                        "confirmed": False,
                        "selectedSupplierId": str(product.get("selectedSupplierId") or ""),
                    }
                    decisioni[product_id] = decisione
                    state.setdefault("products", []).append(decisione)
                # A manually chosen supplier is left untouched: the
                # discount changes prices, not decisions the user made
                # explicitly. This checks the same marker `review()`
                # respects, so a discount never overrides a supplier the
                # user picked on purpose.
                if str(decisione.get("selectedSupplierSource") or "") == "utente":
                    continue
                if str(decisione.get("selectedSupplierId") or "") == migliore["supplierId"]:
                    continue
                decisione["selectedSupplierId"] = migliore["supplierId"]
                # The confirmation applies to the item it was checked
                # against: a different supplier means a different item, and
                # the earlier confirmation doesn't cover it.
                decisione["confirmed"] = not bool(migliore.get("requiresConfirmation"))
                decisione.pop("confirmedArticle", None)
                spostati += 1

            state["updatedAt"] = datetime.now(tz=timezone.utc).isoformat()
            versione = self._avanza_la_versione(state, "sconto")
            return {
                "ok": True,
                "stateVersion": versione,
                "supplierId": supplier_id,
                "percent": round(float(percento), 4),
                "reassigned": spostati,
            }

    @staticmethod
    def _offerta_piu_conveniente(product: Any) -> dict[str, Any] | None:
        """Return the available offer with the lowest unit price.

        Ranked on unit price, consistently with the rest of the program:
        that's the commercial decision the comparison is built around.
        """

        migliori = [
            offer for offer in (product or {}).get("offers") or []
            if offer_is_available(offer)
            and number(offer.get("unitPriceNet")) is not None
        ]
        if not migliori:
            return None
        return min(migliori, key=lambda offer: number(offer.get("unitPriceNet")) or 0.0)

    def delete_compilation(self, payload: Any) -> dict[str, Any]:
        """Delete a compiled order and all its entries from order history."""

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        name = str(payload.get("cartella") or "").strip()
        if not name:
            raise ValueError("Compilazione non indicata")

        with self.lock:
            visible = consegna.elenco(self.orders_dir, etichetta_fornitore=supplier_label)
            if name not in {str(item.get("cartella") or "") for item in visible}:
                raise ValueError("Compilazione non trovata")
            folder = consegna.cartella_sicura(self.orders_dir, name)
            if folder is None:
                raise ValueError("Compilazione non trovata")

            history_existed = self.history_path.exists()
            original_history = self.history_path.read_bytes() if history_existed else None
            history, _changed = order_history.read_history(self.history_path)
            removed = order_history.remove_compilation(history, name)
            quarantine = self.orders_dir / f".eliminazione-{uuid4().hex}"
            folder.rename(quarantine)
            history_saved = False
            try:
                order_history.save_history(self.history_path, history)
                history_saved = True
                shutil.rmtree(quarantine)
            except Exception:
                # Both parts must stay in sync: if one fails, folder and
                # history both revert to how they were.
                try:
                    if history_saved:
                        if original_history is None:
                            self.history_path.unlink(missing_ok=True)
                        else:
                            # This branch runs when something has already
                            # gone wrong, so it uses the same safe-write
                            # helper as everywhere else: a temp file of its
                            # own, flushed to disk before replacing.
                            scrittura_sicura.scrivi_bytes(self.history_path, original_history)
                    if quarantine.exists() and not folder.exists():
                        quarantine.rename(folder)
                except Exception as rollback_error:
                    raise RuntimeError(
                        "Eliminazione interrotta e ripristino non riuscito; non usare lo storico finché non viene controllato"
                    ) from rollback_error
                raise

            return {
                "ok": True,
                "message": "Compilazione eliminata.",
                "promemoriaRimossi": len(removed),
                "compilazioni": consegna.elenco(self.orders_dir, etichetta_fornitore=supplier_label),
            }

    def delete_upload(self, payload: Any) -> dict[str, Any]:
        """Remove a price list from the app; only deletes copies the app manages."""

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        name = str(payload.get("name") or "").strip()
        try:
            nome_sicuro = safe_upload_name(name)
        except ValueError:
            raise ValueError("Listino non trovato") from None
        if nome_sicuro != name or name == self.upload_profiles_path.name:
            raise ValueError("Listino non trovato")

        with self.lock:
            # A copy inside the uploads folder is actually deleted; a
            # reference to an external original is only detached from the page.
            path = consegna.file_sicuro(self.upload_dir, name)
            profiles_doc = load_json(
                self.upload_profiles_path,
                {"schema_version": 1, "profiles": [], "errors": []},
            )
            profiles = profiles_doc.get("profiles") or []
            profile = next(
                (item for item in profiles if isinstance(item, dict) and str(item.get("file_name") or "") == name),
                None,
            )
            active_review = load_json(self.review_path, {})
            file_attivi = active_review.get("files") if isinstance(active_review, dict) else []
            if not isinstance(file_attivi, list):
                file_attivi = []
            riferimenti_documento = [
                voce for voce in file_attivi
                if isinstance(voce, dict)
                and str(voce.get("name") or "") == name
                and str(voce.get("role") or "").casefold() in {"supplier", "master"}
            ]
            ruolo_documento = profiled_upload_role(profile) or next((
                str(voce.get("role") or "").casefold() for voce in riferimenti_documento
                if str(voce.get("role") or "").casefold() in {"master", "supplier"}
            ), "supplier")
            if profile is not None:
                if path is None:
                    raise ValueError("Il file caricato non è più disponibile")
                try:
                    declared_path = Path(str(profile.get("path") or "")).resolve()
                except OSError as exc:
                    raise ValueError("Il profilo del listino non è valido") from exc
                if declared_path != path.resolve():
                    raise ValueError("Il profilo del listino non corrisponde al file caricato")
            elif riferimenti_documento:
                path = None
            else:
                raise ValueError("Listino non trovato")

            original_profiles = deepcopy(profiles_doc)
            original_review = deepcopy(active_review)
            review_changed = bool(riferimenti_documento)
            if review_changed:
                active_review["files"] = [voce for voce in file_attivi if voce not in riferimenti_documento]
                # Removing the entry also removes the only evidence that
                # this supplier had a document: `base_review` would have
                # nothing left to infer that removal from. So the "supplier
                # lost its price list" rule is applied here, in this same
                # function, while the evidence still exists.
                # (`base_review` still handles files removed outside the app.)
                fornitori_tolti: dict[str, str] = {}
                gestionale_tolto = False
                for voce in riferimenti_documento:
                    if str(voce.get("role") or "").strip().casefold() == "master":
                        gestionale_tolto = True
                        continue
                    chiave, etichetta = self._fornitore_della_voce(voce)
                    if chiave:
                        fornitori_tolti[chiave] = etichetta
                self._spegni_i_fornitori_senza_documento(active_review, fornitori_tolti)
                if gestionale_tolto and active_review.get("products"):
                    # Deleting the management-software export is a
                    # deliberate action — useful after uploading the wrong
                    # one — but the comparison's products come from it.
                    # They stay visible, and their current source must be
                    # explained.
                    active_review.setdefault("warnings", []).append({
                        "code": "GESTIONALE_ELIMINATO",
                        "severity": "warning",
                        "blocking": False,
                        "title": "Elenco dei prodotti eliminato",
                        "message": (
                            f"Hai tolto l'elenco del gestionale. I {len(active_review['products'])} "
                            "prodotti che vedi vengono ancora dall'ultimo confronto: carica il nuovo "
                            "elenco e rifai il confronto prima di preparare gli ordini."
                        ),
                    })
            if profile is not None:
                profiles_doc["profiles"] = [item for item in profiles if item is not profile]
                profiles_doc["errors"] = [
                    item for item in profiles_doc.get("errors") or []
                    if not isinstance(item, dict)
                    or str(item.get("file_name") or item.get("name") or "") != name
                ]
                profiles_doc["generated_at"] = datetime.now(tz=timezone.utc).isoformat()
            quarantine = self.upload_dir / f".eliminazione-{uuid4().hex}{path.suffix}" if path else None
            if path is not None and quarantine is not None:
                path.rename(quarantine)
            profiles_saved = False
            review_saved = False
            try:
                if profile is not None:
                    atomic_json(self.upload_profiles_path, profiles_doc)
                    profiles_saved = True
                if review_changed:
                    atomic_json(self.review_path, active_review)
                    review_saved = True
                if quarantine is not None:
                    quarantine.unlink()
            except Exception:
                try:
                    if review_saved:
                        atomic_json(self.review_path, original_review)
                    if profiles_saved:
                        atomic_json(self.upload_profiles_path, original_profiles)
                    if quarantine is not None and path is not None and quarantine.exists() and not path.exists():
                        quarantine.rename(path)
                except Exception as rollback_error:
                    raise RuntimeError(
                        "Eliminazione interrotta e ripristino non riuscito; non ricalcolare finché il listino non viene controllato"
                    ) from rollback_error
                raise

            # What changed, not just that something changed: this is what
            # lets the page name the removed price list when explaining
            # that the prices shown are from an earlier comparison.
            stato_pipeline = self.pipeline_jobs.input_modificato(
                "eliminato",
                documenti=[name],
                fornitori=sorted(fornitori_tolti.values()) if review_changed else [],
            )
            return {
                "ok": True,
                "message": (
                    f"{'Elenco' if ruolo_documento == 'master' else 'Listino'} rimosso dall'app. Il file originale non è stato cancellato."
                    if profile is None
                    else f"{'Elenco' if ruolo_documento == 'master' else 'Listino'} eliminato. Il confronto attuale resta visibile fino al prossimo confronto."
                ),
                "pipeline": stato_pipeline,
            }

    def nuova_comparazione(self, payload: Any = None) -> dict[str, Any]:
        """Clear documents and the current comparison to start over.

        Removes uploaded price lists and the management-software export so
        the next weekly run starts clean, without requiring a manual,
        file-by-file cleanup first. Without it, two price lists from the
        same supplier left in the comparison would silently collapse into
        just one, picked by file modification date (`renderDocumentoConteso`
        explains this on the page, but it has to be read to be noticed).

        What this deliberately does NOT touch, which is the whole point of
        the function: confirmations and equalities (`conferme.db`), learned
        supplier adapters (`adattatori_imparati.json`), orders still
        awaiting delivery (`orders.json`), and past compilations. The first
        two are memory no recompute can rebuild — confirmations have no
        backup copy outside this store — and clearing the third would cause
        merchandise already on its way to be reordered, which is the one
        way a cleanup command could itself produce a wrong order. Adding a
        line here that touches `history/` changes that guarantee, it isn't
        "more cleanup".

        Only copies disappear: input documents are read-only and the
        program keeps its own copy of them, so the files the user actually
        uploaded from stay untouched wherever they are.
        """

        del payload  # the request carries no fields: this clears everything or nothing
        with self.lock:
            if self.pipeline_jobs.in_corso():
                raise LavoroGiaInCorso(
                    "Un confronto è in corso: la comparazione nuova si comincia quando ha finito."
                )

            profiles_doc = load_json(
                self.upload_profiles_path,
                {"schema_version": 1, "profiles": [], "errors": []},
            )
            profili = [voce for voce in (profiles_doc.get("profiles") or []) if isinstance(voce, dict)]
            review = load_json(self.review_path, {})
            voci_file = review.get("files") if isinstance(review, dict) else []
            if not isinstance(voci_file, list):
                voci_file = []

            # Names to remove come from both sources: a document can be in
            # the upload profiles but not the comparison (uploaded after the
            # last recompute), or the other way around (from an earlier
            # comparison).
            nomi: list[str] = []
            for voce in profili:
                nome = str(voce.get("file_name") or "").strip()
                if nome and nome not in nomi:
                    nomi.append(nome)
            fornitori: dict[str, str] = {}
            elenchi = 0
            for voce in voci_file:
                if not isinstance(voce, dict):
                    continue
                ruolo = str(voce.get("role") or "").strip().casefold()
                if ruolo not in {"master", "supplier"}:
                    continue
                nome = str(voce.get("name") or "").strip()
                if nome and nome not in nomi:
                    nomi.append(nome)
                if ruolo == "master":
                    elenchi += 1
                    continue
                chiave, etichetta = self._fornitore_della_voce(voce)
                if chiave:
                    fornitori[chiave] = etichetta

            # Only copies inside `uploads/`: a document the user linked
            # directly from disk is only detached from the page, exactly
            # like `delete_upload` does, and its original is never touched.
            copie: list[Path] = []
            for nome in nomi:
                if nome == self.upload_profiles_path.name:
                    continue
                percorso = consegna.file_sicuro(self.upload_dir, nome)
                if percorso is not None and percorso.is_file():
                    copie.append(percorso)

            profili_originali = deepcopy(profiles_doc)
            vuoto = {
                "schema_version": profiles_doc.get("schema_version", 1),
                "profiles": [],
                "errors": [],
                "generated_at": datetime.now(tz=timezone.utc).isoformat(),
            }

            # Everything moves to quarantine first, then gets written, and
            # only at the end gets actually deleted — the same discipline as
            # `delete_upload`, and more important here since many files are
            # involved: a partial failure could otherwise leave a
            # comparison describing documents that no longer exist.
            quarantena: list[tuple[Path, Path]] = []
            profili_scritti = False
            try:
                for percorso in copie:
                    riparo = self.upload_dir / f".nuova-comparazione-{uuid4().hex}{percorso.suffix}"
                    percorso.rename(riparo)
                    quarantena.append((percorso, riparo))
                for percorso in (self.review_path, self.state_path):
                    if percorso.is_file():
                        riparo = percorso.with_name(f".nuova-comparazione-{uuid4().hex}.json")
                        percorso.rename(riparo)
                        quarantena.append((percorso, riparo))
                atomic_json(self.upload_profiles_path, vuoto)
                profili_scritti = True
            except Exception:
                try:
                    if profili_scritti:
                        atomic_json(self.upload_profiles_path, profili_originali)
                    for percorso, riparo in reversed(quarantena):
                        if riparo.exists() and not percorso.exists():
                            riparo.rename(percorso)
                except Exception as rimessa_fallita:
                    raise RuntimeError(
                        "Pulizia interrotta e ripristino non riuscito; controlla i documenti caricati prima di ricalcolare"
                    ) from rimessa_fallita
                raise

            for _, riparo in quarantena:
                try:
                    riparo.unlink()
                except OSError:
                    # The work is already done: `uploads/` no longer lists
                    # it and the comparison is gone. A leftover quarantine
                    # file is clutter, not a failure, and reporting it here
                    # wouldn't help anyone.
                    pass

            # No "stale prices" banner here: that banner exists to say the
            # prices on screen are from an earlier comparison, and there are
            # no prices or comparison pages left to look at. State just
            # returns to awaiting new files.
            stato_pipeline = self.pipeline_jobs.input_modificato()

            # "Has this arrived?" questions are reopened now, rather than
            # waiting for their normal delay. Someone who answers "not yet"
            # gets that question deferred by a fixed interval, which is fine
            # when a timer is the only thing tracking time — but starting a
            # fresh comparison is a stronger signal than the timer: it's the
            # moment someone actually checks whether last week's order
            # arrived, so pending questions should surface now rather than
            # stay deferred past it.
            domande = 0
            storico_illeggibile = ""
            try:
                storico, _scaduti = order_history.read_history(self.history_path)
                domande = order_history.riapri_le_domande(storico)
                if domande:
                    order_history.save_history(self.history_path, storico)
            except Exception as exc:  # noqa: BLE001 - cleanup already completed
                # Not swallowed silently: the new comparison has genuinely
                # started and isn't rolled back for this, but "couldn't
                # reopen the pending questions" is information worth
                # surfacing — an unreadable history is functionally the
                # same as not knowing what merchandise is still expected.
                storico_illeggibile = frase(exc) or "Lo storico degli ordini non si è lasciato leggere."

            listini = max(0, len(nomi) - elenchi)
            return {
                "ok": True,
                "message": "Comparazione nuova: carica l'elenco del gestionale e i listini di questa settimana.",
                "tolti": {"elenco": elenchi, "listini": listini, "documenti": len(nomi)},
                "fornitori": sorted(fornitori.values()),
                "domandeRiaperte": domande,
                "storicoIlleggibile": storico_illeggibile,
                "pipeline": stato_pipeline,
            }

    def record_order_history(
        self,
        plan: dict[str, Any],
        *,
        order_key: str = "",
        delivered: set[str] | None = None,
    ) -> list[str]:
        """Record the just-compiled order plan as pending delivery.

        Called only after compilation genuinely succeeded — the price-list
        copies are written to disk — and only for suppliers whose copy
        exists: an order that couldn't actually be sent must not later
        become a "has this arrived?" question.

        The compiled copies already exist by this point; a failure here
        must not undo a successful compilation, so it's only reported.
        """

        try:
            history, _changed = order_history.read_history(self.history_path)
            order_history.record_plan(
                history,
                plan,
                supplier_name=supplier_label,
                order_key=order_key,
                delivered=delivered,
            )
            order_history.save_history(self.history_path, history)
        except Exception as exc:
            message = f"L'ordine non è stato registrato nello storico delle consegne: {exc}"
            print(f"[AVVISO] {message}")
            return [message]
        return []

    def search_products(self, query: str, limit: int = 20) -> dict[str, Any]:
        with self.lock:
            review, _state = self.review_with_manual_products()
            results = self.catalog.search(review, query, limit)
            return {"ok": True, "query": query, "results": results, "count": len(results)}

    def sfoglia_listino(
        self,
        fornitore: str,
        *,
        query: str = "",
        da: int = 0,
        quante: int = 50,
        riga: Any = None,
    ) -> dict[str, Any]:
        """Return one page of a supplier's price list, as the pipeline read it.

        Answers a question no matching score can resolve on its own: does
        this supplier carry the same product under a different name? Two
        price lists can describe the same item with different wording and
        a different barcode, in a way that isn't obvious from either
        description alone — seeing the raw price-list row makes it clear.

        A price list that fails to read does NOT fail this request. The
        list of browsable suppliers (computed on the line below) is always
        returned regardless, and the requested supplier's own failure is
        reported through `problema`, with the real reason, rather than
        failing the whole call and leaving the supplier picker empty even
        when every other price list is fine.
        """

        with self.lock:
            review, _state = self.review_with_manual_products()
            # The supplier list is computed before the requested page, not
            # inside the same expression as the result: computing it lazily
            # there would mean never computing it at all when `sfoglia`
            # raises.
            fornitori = self.catalog.fornitori_sfogliabili(review)
            chiesto = str(fornitore or "").strip().casefold()
            # No supplier requested means "pick one for me": the "browse
            # price lists" button on a product with no offers at all
            # carries no supplier id, and without this fallback that case
            # would resolve to an empty supplier name.
            # An explicit choice is never overridden: matched rows carry
            # "this is it", tied to the supplier of the window that's open.
            # Silently switching to a different supplier than the one
            # requested would let a manual match land on the wrong price
            # list, producing a wrong order.
            if not chiesto and fornitori:
                chiesto = str(fornitori[0].get("id") or "")
            try:
                pagina = self.catalog.sfoglia(
                    review, chiesto, query=query, da=da, quante=quante, riga=riga,
                )
            except ValueError:
                pagina = self._listino_non_sfogliabile(chiesto, fornitori, quante=quante)
            return {"ok": True, "fornitori": fornitori, **pagina}

    def _listino_non_sfogliabile(
        self, chiesto: str, fornitori: list[dict[str, Any]], *, quante: int
    ) -> dict[str, Any]:
        """Build the empty-page response for a price list that isn't there, with the real reason.

        `SupplierCatalog.load_errors` knows why that supplier dropped out
        — file moved, schema not recognized, columns not declared, no
        orderable row — and reports it in plain language, rather than a
        generic "no price list loaded" message that says neither what
        happened nor what to do about it.
        """

        motivo = next(
            (
                frase(voce.get("message"))
                for voce in self.catalog.load_errors
                if str(voce.get("supplier") or "").casefold() == chiesto
            ),
            "",
        )
        if not motivo:
            motivo = (
                f"Il listino {supplier_label(chiesto)} non fa parte di questo confronto."
                if chiesto
                else "Questo confronto non ha nessun listino da sfogliare."
            )
        seguito = (
            "Gli altri listini si sfogliano dalla tendina qui sopra."
            if fornitori
            else "Ricarica i listini dal passo 1 «Importa i dati» e rifai il confronto."
        )
        return {
            "supplier": chiesto,
            "supplierName": supplier_label(chiesto) if chiesto else "",
            "righe": [],
            "da": 0,
            # The same cap `SupplierCatalog.sfoglia` applies: the page uses
            # `quante` to count pages, and two different limits would
            # eventually disagree.
            "quante": max(1, min(int(quante or 50), 200)),
            "trovate": 0,
            "totale": 0,
            "scartate": 0,
            "rigaCercata": None,
            "problema": {"fornitore": chiesto, "messaggio": f"{motivo} {seguito}".strip()},
        }

    def add_manual_product(self, payload: Any) -> dict[str, Any]:
        catalog_id = str(payload.get("catalogId") or "") if isinstance(payload, dict) else ""
        if not catalog_id:
            raise ValueError("Prodotto non indicato")
        with self.lock:
            review, state = self.review_with_manual_products()
            product = self.catalog.get(review, catalog_id)
            product_id = str(product.get("id") or "")
            current_ids = {str(item.get("id") or "") for item in review.get("products") or []}
            if product_id in current_ids:
                raise ValueError("Il prodotto è già presente nell'elenco")
            ean = str(product.get("ean") or "")
            if ean and any(str(item.get("ean") or "") == ean for item in review.get("products") or []):
                raise ValueError("Un prodotto con lo stesso codice è già presente nell'elenco")
            product.pop("catalogId", None)
            product["addedManually"] = True
            product["quantity"] = 0
            product["excluded"] = False
            state.setdefault("manualProducts", []).append(product)
            state["updatedAt"] = datetime.now(tz=timezone.utc).isoformat()
            atomic_json(self.state_path, state)
            return {"ok": True, "message": "Prodotto aggiunto all'elenco.", "product": product}

    def validate_snapshot(
        self, snapshot: Any, *, for_compile: bool
    ) -> tuple[dict[str, Any], list[dict[str, Any]], dict[str, Any]]:
        """Return the cleaned state, the errors, and the comparison they were decided against.

        Returning the comparison is not just a convenience: without it, callers
        would rebuild it right afterward, and that means re-reading and re-parsing
        `review_data.json` (the largest file in the app, ~2.4 MB on a real
        comparison) a second time. The page autosaves 450 ms after every edit, so
        that cost would double per quantity change and the first build would be
        discarded.

        This is safe because nothing the comparison depends on can change between
        the two calls: `review_data.json` is frozen while it's in use (the pipeline
        now takes the same lock the routes do to replace it), and the four state
        keys that feed into the comparison (`manualProducts`, `matchOverrides`,
        `manualMatches`, `supplierDiscounts`, listed in `CHIAVI_DI_STATO_RICOPIATE`)
        are copied verbatim from the state read here. It isn't a cache with a
        staleness window — the comparison is handed off within the same call, never
        held between requests.
        """

        if not isinstance(snapshot, dict):
            raise ValueError("Snapshot non valido")
        review, existing_state = self.review_with_manual_products()
        run_id = str(snapshot.get("runId") or "")
        expected_run = str((review.get("run") or {}).get("id") or "")
        if expected_run and run_id != expected_run:
            raise ValueError("La run dello snapshot non coincide con la run caricata")

        # Two tabs open on the same comparison can overwrite each other's work:
        # `save_state` rewrites the whole state, and the run check alone can't tell
        # them apart (same run). Each tab declares the version it started from; if
        # the one on disk is newer, someone else already saved, and that has to be
        # reported instead of silently discarding quantities. A snapshot without a
        # declared version isn't rejected — it's an old tab left open, and blocking
        # its save would be worse than the small risk it carries.
        versione_sul_disco = int(number(existing_state.get("stateVersion")) or 0)
        versione_dichiarata = number(snapshot.get("stateVersion"))
        if versione_dichiarata is not None and int(versione_dichiarata) != versione_sul_disco:
            # Who wrote that version matters: a compilation rewrites the state
            # before producing the copies, so a tab that falls behind by one number
            # may not be behind any other tab at all — it's behind its own
            # compilation. Telling it "another tab saved after you" would send the
            # user looking for a colleague who isn't there, and "retry" can't work
            # because it saves again first.
            # There's a third origin besides another tab and a compilation:
            # `state.json` may not exist on disk at all while the tab still
            # declares a version, meaning "Start new comparison" was pressed in the
            # meantime. That also isn't "another tab saved after you".
            svuotato = versione_sul_disco == 0 and not self.state_path.is_file()
            return {}, [{
                "code": "STATO_SOVRASCRITTO",
                "stateVersion": versione_sul_disco,
                "origin": "svuotato" if svuotato else str(existing_state.get("stateVersionOrigin") or "scheda"),
            }], review

        products_by_id = {str(product.get("id")): product for product in review.get("products") or []}
        decisions = snapshot.get("products")
        if not isinstance(decisions, list):
            raise ValueError("products deve essere una lista")
        seen = set()
        normalized = []
        errors = []
        for decision in decisions:
            product_id = str(decision.get("id") or "")
            if not product_id or product_id in seen:
                errors.append({"code": "ID_PRODOTTO_NON_VALIDO", "productId": product_id})
                continue
            seen.add(product_id)
            product = products_by_id.get(product_id)
            if product is None:
                errors.append({"code": "PRODOTTO_SCONOSCIUTO", "productId": product_id})
                continue
            quantity_value = number(decision.get("quantity"))
            if quantity_value is None or quantity_value < 0 or int(quantity_value) != quantity_value or quantity_value > 100000:
                errors.append({
                    "code": "QUANTITA_NON_VALIDA",
                    "productId": product_id,
                    # The name travels with the error: the page lists stuck
                    # products one by one, and an id isn't something to hunt for
                    # by hand among hundreds of rows.
                    "productName": str(product.get("name") or product.get("description") or product_id),
                })
                continue
            quantity = int(quantity_value)
            supplier_id = str(decision.get("selectedSupplierId") or "")
            confirmed = bool(decision.get("confirmed"))
            excluded = bool(decision.get("excluded"))
            if excluded:
                quantity = 0
            offer = find_offer(product, supplier_id)
            product_name = str(product.get("name") or product.get("description") or product_id)
            # A positive quantity with no supplier available at all is a valid
            # state: the product doesn't enter the plan or any supplier's price
            # list, but goes into the "items to source" list produced at compile
            # time. The rule that an order line without a usable offer can't be
            # sent still applies to the other two cases: a chosen supplier whose
            # offer can't be ordered, and no supplier chosen when one is actually
            # available. Those are pending decisions, and letting them through
            # silently would send a wrong order.
            da_reperire = not supplier_id and nessuna_offerta_utilizzabile(product)
            offerta_da_sistemare = (
                (not offer_is_available(offer)) if supplier_id else not da_reperire
            )
            if quantity > 0 and offerta_da_sistemare:
                # The rejection names which product it's about, so the user isn't
                # left hunting through hundreds of rows.
                #
                # Two distinct codes: "the chosen supplier has an offer that can't
                # be ordered" is an error to fix now, while "no supplier chosen yet
                # among those who have it" is a pending decision (typically after
                # rejecting a match when another supplier does carry the item). A
                # single shared code would block autosave for every product in the
                # snapshot, including the "go back" action that saves before
                # responding. This blocks compile, not save.
                errors.append({
                    "code": "OFFERTA_NON_VALIDA" if supplier_id else "FORNITORE_DA_SCEGLIERE",
                    "productId": product_id,
                    "productName": product_name,
                    "supplierId": supplier_id,
                    "supplierName": supplier_label(supplier_id) if supplier_id else "",
                })
            # With no supplier chosen there's nothing to confirm: requiring a
            # confirmation on a product no supplier carries would be a checkbox the
            # user can never tick, and would block compilation forever.
            requires = bool(supplier_id) and offer_needs_confirmation(product, offer)
            # The confirmation store is deliberately not consulted here.
            # `review()` seeds `product.confirmed` from it, the page echoes that
            # value back unchanged in the snapshot, and that's what arrives here.
            # If the store could also answer "yes" at this point, a confirmation
            # the user just unchecked would be reinstated by yesterday's memory in
            # the same save that removed it — revocation would be impossible. One
            # authority at a time: the store on read, the snapshot on save.
            if quantity > 0 and requires and not confirmed:
                errors.append({
                    "code": "CONFERMA_MANCANTE",
                    "productId": product_id,
                    "productName": product_name,
                    "supplierId": supplier_id,
                    "supplierName": supplier_label(supplier_id) if supplier_id else "",
                })
            quantity_source = str(decision.get("quantitySource") or "")
            if quantity_source not in {"gestionale", "utente"}:
                quantity_source = "gestionale"
            # Who chose this supplier: the app's default, or the user.
            #
            # This is inferred rather than asked of the page, and the inference is
            # exact: the app has already picked the best offer on its own, so a
            # choice that matches the best offer can only be the default — picking
            # it by hand wouldn't change anything. A choice that differs from the
            # best offer is a deliberate one, and stays deliberate.
            #
            # `review()` relies on this: a recalculated comparison picks the best
            # offer again, and the saved decision then overrides it. Without this
            # inference, a default computed last week would survive a recalculation
            # that made it wrong, disguised as a user choice — a saved default could
            # silently outlive the price change that invalidated it.
            migliore = self._offerta_piu_conveniente(product)
            scelta_dell_utente = bool(supplier_id) and (
                migliore is None or supplier_id != str(migliore.get("supplierId") or "")
            )
            voce = {
                "id": product_id,
                "quantity": quantity,
                "selectedSupplierId": supplier_id,
                "selectedSupplierSource": "utente" if scelta_dell_utente else "automatico",
                "confirmed": confirmed,
                "excluded": excluded,
                "quantitySource": quantity_source,
            }
            # What was confirmed, not just that something was. Without this, after
            # a later recalculation the checkbox would stay ticked even if the
            # product had since matched a different price-list row (different EAN,
            # different description): the user's confirmation would then cover an
            # item they never actually saw.
            if confirmed and offer:
                voce["confirmedArticle"] = impronta_articolo(offer)
            normalized.append(voce)

        # A compilation made entirely of items to source is not an empty order to
        # reject: it's a week where the management software asks for goods no
        # supplier carries, and the correct output is the list of those items.
        # `ORDINE_VUOTO` still applies when there is neither an order nor a list.
        if for_compile and not any(item["quantity"] > 0 for item in normalized):
            errors.append({"code": "ORDINE_VUOTO"})
        clean = {
            "schemaVersion": 1,
            # Bumped on every full state rewrite (`save_state`, `compile`): this
            # is how the next tab notices it isn't the only one. Writes that touch
            # only part of the state (a candidate answer, a manually added product)
            # don't bump it, because `validate_snapshot` copies those fields from
            # disk and can't lose them.
            "stateVersion": versione_sul_disco + 1,
            # Who wrote this version: `compile` resets it to "compilazione" after
            # calling this function, so a tab that finds itself behind can tell
            # whether it lost to another tab or to its own compilation.
            "stateVersionOrigin": "scheda",
            "runId": run_id,
            "currentStep": max(1, min(3, int(number(snapshot.get("currentStep")) or 1))),
            "acceptBelowThreshold": bool(snapshot.get("acceptBelowThreshold")),
            "summaryGrouping": snapshot.get("summaryGrouping") if snapshot.get("summaryGrouping") in {"supplier", "product"} else "supplier",
            "updatedAt": datetime.now(tz=timezone.utc).isoformat(),
            "products": normalized,
        }
        # Candidate decisions are tied to the run and to their fingerprint;
        # the header discount and manually matched rows are each written by their
        # own route. None of the three travel in the tab's snapshot, so the
        # quantity autosave must not touch them — they're copied verbatim from
        # disk, key by key.
        for chiave, quando_manca in CHIAVI_DI_STATO_RICOPIATE.items():
            clean[chiave] = existing_state.get(chiave) or deepcopy(quando_manca)
        return clean, errors, review

    def order_summary(
        self,
        review: dict[str, Any],
        selections: dict[str, dict[str, Any]],
    ) -> dict[str, Any]:
        """Compute the summary totals: header total, sum of the displayed lines, and their gap.

        The header total and the sum of the displayed lines don't always match.
        The order total is computed on full-precision prices (promotions and price
        lists carry up to six decimals), while each line on the page is shown
        rounded to the cent — summing the displayed lines can differ from the
        header total by a few cents. Neither number is wrong; showing one without
        the other is, because a manual check then finds a gap it can't explain.
        """

        products_by_id = {str(item.get("id")): item for item in review.get("products") or []}
        supplier_defs = {str(item.get("id")): item for item in review.get("suppliers") or []}
        righe: dict[str, dict[str, Any]] = {}
        for decision in selections.values():
            quantity = int(number(decision.get("quantity")) or 0)
            if quantity <= 0 or decision.get("excluded"):
                continue
            product = products_by_id.get(str(decision.get("id") or ""))
            if product is None:
                continue
            supplier = str(decision.get("selectedSupplierId") or "")
            pricing = offer_pricing(find_offer(product, supplier))
            if pricing is None:
                continue
            line_net = round(quantity * pricing["orderUnitPriceNet"], 4)
            voce = righe.setdefault(supplier, {
                "supplierId": supplier,
                "supplierName": supplier_label(supplier),
                "lineCount": 0,
                "_intero": 0.0,
                "_arrotondato": 0.0,
            })
            voce["lineCount"] += 1
            voce["_intero"] = round(voce["_intero"] + line_net, 4)
            # The line sum must count the way the screen displays it. Python's
            # `round()` rounds the double and ties go to even; the browser's
            # `Intl.NumberFormat` rounds the decimal representation and ties go up
            # — on 14.665 the former gives 14.66, the page shows 14.67. The summary
            # text compares the header total to "what the lines show"; using a
            # different rounding here would state a number that never appeared on
            # screen.
            voce["_arrotondato"] = round(voce["_arrotondato"] + _centesimi_come_in_pagina(line_net), 2)

        suppliers = []
        for supplier in sorted(righe, key=lambda item: (supplier_label(item), item)):
            voce = righe[supplier]
            totale = _centesimi_come_in_pagina(voce.pop("_intero"))
            somma_righe = round(voce.pop("_arrotondato"), 2)
            voce["totalNet"] = totale
            voce["linesTotalNet"] = somma_righe
            voce["roundingDifference"] = round(totale - somma_righe, 2)
            voce["threshold"] = round(supplier_threshold(supplier_defs.get(supplier, {})), 2)
            suppliers.append(voce)

        totale_generale = _centesimi_come_in_pagina(sum(item["totalNet"] for item in suppliers))
        somma_generale = _centesimi_come_in_pagina(sum(item["linesTotalNet"] for item in suppliers))
        return {
            "suppliers": suppliers,
            "totalNet": totale_generale,
            "linesTotalNet": somma_generale,
            "roundingDifference": round(totale_generale - somma_generale, 2),
        }

    def order_summary_from_state(self, review: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
        """Same summary, built from the choices saved on disk."""

        selections = {
            str(item.get("id") or ""): item
            for item in (state or {}).get("products") or []
            if isinstance(item, dict)
        }
        return self.order_summary(review, selections)

    def save_state(self, snapshot: Any) -> dict[str, Any]:
        with self.lock:
            # The previous state is needed to tell a revocation from an
            # expiration apart, so it's read now, before it gets replaced.
            precedente = load_json(self.state_path, {}) if self.state_path.is_file() else {}
            clean, errors, review = self.validate_snapshot(snapshot, for_compile=False)
            # A missing confirmation blocks compiling, not saving. If it also
            # blocked saving, a handful of unconfirmed products would silently
            # stop autosave for every other product in the review until all
            # pending proposals were answered.
            #
            # The gate stays where it's useful: `compile` validates with
            # `for_compile=True`, and there `CONFERMA_MANCANTE` does stop
            # everything, with a message naming the product and supplier. The user
            # already knows what's pending from the page's own "to confirm"
            # filter, so repeating it here would just be a second notification for
            # the same thing.
            # `FORNITORE_DA_SCEGLIERE` is treated the same way: it's a pending
            # decision, not a broken order, so it must not block autosave for the
            # rest of the products either — though it still blocks compilation.
            non_bloccanti = {"CONFERMA_MANCANTE", "FORNITORE_DA_SCEGLIERE"}
            bloccanti = [errore for errore in errors if errore.get("code") not in non_bloccanti]
            if bloccanti:
                raise SnapshotError(bloccanti)
            atomic_json(self.state_path, clean)
            if str((precedente or {}).get("runId") or "") != str(clean.get("runId") or ""):
                # The one case where the comparison `validate_snapshot` built
                # isn't the one we'd get now: `apply_match_overrides` (the answers
                # given to candidates) only applies when the state declares the
                # same run as the comparison. If the previous state declared a
                # different run, it skipped those overrides while the one just
                # written doesn't — rare, and silent otherwise, so it's rebuilt
                # here.
                review, _stato = self.review_with_manual_products()
            self.ricorda_le_conferme(clean, review, precedente if isinstance(precedente, dict) else {})
            selections = {str(item["id"]): item for item in clean["products"]}
            decorated = self.promotion_service.decorate(review, selections)
            promotion_states = {
                str(item.get("id")): item.get("state") or {}
                for item in decorated.get("promotions") or []
                if item.get("id")
            }
            return {
                "ok": True,
                "savedAt": clean["updatedAt"],
                "stateVersion": clean["stateVersion"],
                "promotionSummary": decorated.get("promotionSummary") or {"counts": {}},
                "promotionStates": promotion_states,
                # The summary totals and the header/lines rounding gap.
                "orderSummary": self.order_summary(review, selections),
            }

    def omaggi_per_fornitore(
        self,
        review: dict[str, Any],
        selections: dict[str, dict[str, Any]],
    ) -> dict[str, int]:
        """Count the threshold-gift rewards earned by these choices.

        The count follows the same promotion rules (single-order and repeatable
        thresholds); the gift's monetary value doesn't factor in here — it's
        informational only.
        """

        from promotions import (  # deferred import: the module lives under `scripts`
            KIND_THRESHOLD_GIFT,
            STATUS_EARNED,
            calculate_promotion_state,
        )

        conteggi: dict[str, int] = {}
        with self.promotion_service.lock:
            promozioni = self.promotion_service.detect(review)
        for promozione in promozioni:
            if str(promozione.get("kind") or "") != KIND_THRESHOLD_GIFT:
                continue
            stato = calculate_promotion_state(promozione, review, selections=selections)
            if stato.get("status") != STATUS_EARNED:
                continue
            fornitore = str(promozione.get("supplier") or "")
            conteggi[fornitore] = conteggi.get(fornitore, 0) + max(1, int(stato.get("reward_count") or 1))
        return conteggi

    def move_preview(self, payload: Any) -> dict[str, Any]:
        """Quote what moving all of one supplier's products to another supplier would cost.

        Writes nothing — no state, no history, no final documents. It only shows
        what changing supplier would cost, so the user can decide from the
        numbers; the actual move happens through the normal state save. The
        computation runs on the server, never the browser, because prices and
        totals are always the server's decision.

        The carton (or display) count never changes; the supplier changes, and
        with it the order-unit price and the delivered piece count.
        """

        if not isinstance(payload, dict):
            raise ValueError("Richiesta non valida")
        from_supplier = str(payload.get("from") or "").strip()
        if not from_supplier:
            raise ValueError("Fornitore di partenza non indicato")

        with self.lock:
            # Same checks as the other endpoints: wrong run, unknown product or
            # impossible quantity are rejected here. Not CONFERMA_MANCANTE though
            # — a preview writes nothing, and a missing confirmation on a product
            # from a different supplier shouldn't block seeing this one's numbers.
            # Confirmations that will be needed after the move are still declared
            # per assignment, in `needsConfirmation`.
            clean, errors, review = self.validate_snapshot(payload, for_compile=False)
            blocking = [error for error in errors if error.get("code") != "CONFERMA_MANCANTE"]
            if blocking:
                raise SnapshotError(blocking)
            products_by_id = {str(item.get("id")): item for item in review.get("products") or []}
            supplier_defs = {str(item.get("id")): item for item in review.get("suppliers") or []}
            # Current choices, in the shape promotions understands, for counting
            # gifts before and after the move.
            selezioni_prima = {str(item["id"]): dict(item) for item in clean["products"]}
            omaggi_prima = self.omaggi_per_fornitore(review, selezioni_prima)

            # First pass: line-by-line net total of the whole order, same
            # arithmetic as compilation. Zero-quantity and excluded products don't
            # count (validation zeroes out excluded items).
            totals_before: dict[str, float] = {}
            current_total = 0.0
            movable: list[dict[str, Any]] = []
            for decision in clean["products"]:
                if decision["quantity"] <= 0:
                    continue
                product = products_by_id.get(decision["id"])
                if product is None:
                    continue
                supplier = decision["selectedSupplierId"]
                pricing = offer_pricing(find_offer(product, supplier))
                line_net = round(decision["quantity"] * pricing["orderUnitPriceNet"], 4) if pricing else 0.0
                totals_before[supplier] = round(totals_before.get(supplier, 0.0) + line_net, 4)
                current_total = round(current_total + line_net, 4)
                if supplier != from_supplier:
                    continue
                alternatives: dict[str, dict[str, Any]] = {}
                for offer in product.get("offers") or []:
                    destination = offer_supplier_id(offer)
                    if not destination or destination == from_supplier or destination in alternatives:
                        continue
                    if not offer_is_available(offer):
                        continue
                    destination_pricing = offer_pricing(offer)
                    if destination_pricing is None:
                        continue
                    alternatives[destination] = {
                        "supplierId": destination,
                        "pricing": destination_pricing,
                        "needsConfirmation": offer_needs_confirmation(product, offer),
                    }
                movable.append({
                    "id": decision["id"],
                    "name": str(product.get("name") or product.get("description") or decision["id"]),
                    "product": product,
                    "quantity": decision["quantity"],
                    "lineNet": line_net,
                    "factor": pricing["factor"] if pricing else 1.0,
                    "alternatives": alternatives,
                })

            def left_behind_reason(item: dict[str, Any], supplier_id: str | None) -> str:
                """Why the product stays where it is. It's neither zeroed nor removed."""

                if supplier_id:
                    examined = [find_offer(item["product"], supplier_id)]
                else:
                    examined = [
                        offer for offer in item["product"].get("offers") or []
                        if offer_supplier_id(offer) not in {"", from_supplier}
                    ]
                if any(offer_is_available(offer) for offer in examined):
                    return "PREZZO_NON_DISPONIBILE"
                return "NESSUNA_OFFERTA"

            def build_option(option_id: str, kind: str, label: str, chosen: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
                assignments: list[dict[str, Any]] = []
                left_behind: list[dict[str, Any]] = []
                totals_after = dict(totals_before)
                delta = 0.0
                # Same carton count, different delivered pieces, because cartons
                # from different suppliers don't hold the same amount of goods.
                # Without these counts the cost difference alone is misleading: an
                # option that "saves money" could deliver half the goods.
                pieces_before = 0.0
                pieces_after = 0.0
                moved_net_before = 0.0
                moved_net_after = 0.0
                for item in movable:
                    choice = chosen.get(item["id"])
                    if choice is None:
                        left_behind.append({
                            "productId": item["id"],
                            "productName": item["name"],
                            "reason": left_behind_reason(item, option_id if kind == "supplier" else None),
                        })
                        continue
                    destination = choice["supplierId"]
                    new_line = round(item["quantity"] * choice["pricing"]["orderUnitPriceNet"], 4)
                    totals_after[from_supplier] = round(totals_after.get(from_supplier, 0.0) - item["lineNet"], 4)
                    totals_after[destination] = round(totals_after.get(destination, 0.0) + new_line, 4)
                    delta = round(delta + new_line - item["lineNet"], 4)
                    new_factor = choice["pricing"]["factor"]
                    previous_pieces = item["quantity"] * item["factor"]
                    new_pieces = item["quantity"] * new_factor
                    pieces_before += previous_pieces
                    pieces_after += new_pieces
                    moved_net_before = round(moved_net_before + item["lineNet"], 4)
                    moved_net_after = round(moved_net_after + new_line, 4)
                    assignments.append({
                        "productId": item["id"],
                        "productName": item["name"],
                        "toSupplierId": destination,
                        "previousFactor": item["factor"],
                        "newFactor": new_factor,
                        "factorChanged": new_factor != item["factor"],
                        "previousPieces": int(round(previous_pieces)),
                        "newPieces": int(round(new_pieces)),
                        "needsConfirmation": bool(choice["needsConfirmation"]),
                        "previousLineNet": round(item["lineNet"], 2),
                        "newLineNet": round(new_line, 2),
                    })
                if not assignments:
                    # An option that moves nothing isn't offered: it would be a
                    # choice indistinguishable from doing nothing.
                    return None
                # Threshold gifts are earned on a single supplier's order, so
                # moving goods away can lose them; recomputed here with the
                # post-move choices, same promotion rules.
                selezioni_dopo = {key: dict(value) for key, value in selezioni_prima.items()}
                for assignment in assignments:
                    voce = selezioni_dopo.get(assignment["productId"])
                    if voce is not None:
                        voce["selectedSupplierId"] = assignment["toSupplierId"]
                omaggi_dopo = self.omaggi_per_fornitore(review, selezioni_dopo)
                totale_omaggi_prima = sum(omaggi_prima.values())
                totale_omaggi_dopo = sum(omaggi_dopo.values())
                # Lost and gained are counted per supplier and then summed:
                # losing 2 gifts from one supplier and gaining 2 from another isn't
                # "zero" — they're different goods from different suppliers, and a
                # plain difference of the grand totals would silently net them out.
                fornitori_omaggi = set(omaggi_prima) | set(omaggi_dopo)
                omaggi_persi = sum(
                    max(0, omaggi_prima.get(nome, 0) - omaggi_dopo.get(nome, 0))
                    for nome in fornitori_omaggi
                )
                omaggi_guadagnati = sum(
                    max(0, omaggi_dopo.get(nome, 0) - omaggi_prima.get(nome, 0))
                    for nome in fornitori_omaggi
                )
                totals_rows = []
                for supplier in [from_supplier] + sorted({item["toSupplierId"] for item in assignments}):
                    threshold = supplier_threshold(supplier_defs.get(supplier, {}))
                    before = totals_before.get(supplier, 0.0)
                    after = totals_after.get(supplier, 0.0)
                    totals_rows.append({
                        "supplierId": supplier,
                        "supplierName": supplier_label(supplier),
                        "netTotalBefore": round(before, 2),
                        "netTotalAfter": round(after, 2),
                        "giftsBefore": int(omaggi_prima.get(supplier, 0)),
                        "giftsAfter": int(omaggi_dopo.get(supplier, 0)),
                        "threshold": round(threshold, 2),
                        # meets_threshold follows compile's rule that a supplier
                        # with no order at all isn't "below threshold". Without
                        # hadOrderBefore, a reader could mistake that `true` for
                        # "the threshold was met" and report a supplier that
                        # started at zero as having dropped.
                        "hadOrderBefore": before > 0,
                        "meetsThresholdBefore": meets_threshold(before, threshold),
                        "meetsThresholdAfter": meets_threshold(after, threshold),
                    })
                cost_before = round(moved_net_before / pieces_before, 4) if pieces_before else 0.0
                cost_after = round(moved_net_after / pieces_after, 4) if pieces_after else 0.0
                return {
                    "id": option_id,
                    "kind": kind,
                    "label": label,
                    "movedCount": len(assignments),
                    "movableCount": len(movable),
                    "deltaNet": round(delta, 2),
                    "deliveredPiecesBefore": int(round(pieces_before)),
                    "deliveredPiecesAfter": int(round(pieces_after)),
                    "deltaPieces": int(round(pieces_after - pieces_before)),
                    "costPerPieceBefore": cost_before,
                    "costPerPieceAfter": cost_after,
                    "deltaCostPerPiece": round(cost_after - cost_before, 4),
                    # The gift count, never its monetary value: this app doesn't
                    # price gifts.
                    "giftsBefore": totale_omaggi_prima,
                    "giftsAfter": totale_omaggi_dopo,
                    "giftsLost": omaggi_persi,
                    "giftsGained": omaggi_guadagnati,
                    "assignments": assignments,
                    "leftBehind": left_behind,
                    "supplierTotalsAfter": totals_rows,
                }

            options = []
            # "Best alternative" is chosen on price per piece, never per carton:
            # cartons from different suppliers hold different quantities, so the
            # cheapest carton can still be the most expensive per piece. Ties go
            # to the supplier that sorts first alphabetically, so the same input
            # always produces the same answer.
            best_choice = {
                item["id"]: min(
                    item["alternatives"].values(),
                    key=lambda alternative: (alternative["pricing"]["unitPriceNet"], alternative["supplierId"]),
                )
                for item in movable if item["alternatives"]
            }
            best_option = build_option("best", "best", "Migliore alternativa per ciascun prodotto", best_choice)
            if best_option is not None:
                options.append(best_option)

            supplier_options = []
            for destination in sorted({key for item in movable for key in item["alternatives"]}):
                chosen = {
                    item["id"]: item["alternatives"][destination]
                    for item in movable if destination in item["alternatives"]
                }
                option = build_option(destination, "supplier", supplier_label(destination), chosen)
                if option is not None:
                    supplier_options.append(option)
            # The display order can't be the total spent: orders that deliver
            # different amounts of goods aren't comparable on total alone, same
            # pitfall as comparing offers directly. Options that move the most
            # products off the source supplier come first (that's the point of
            # this command), then options with the lowest cost per piece.
            supplier_options.sort(key=lambda option: (
                -option["movedCount"],
                option["deltaCostPerPiece"],
                option["id"],
            ))
            options.extend(supplier_options)

            return {
                "ok": True,
                "from": from_supplier,
                "fromName": supplier_label(from_supplier),
                "movableCount": len(movable),
                "currentNetTotal": round(current_total, 2),
                "options": options,
            }

    def upload(self, payload: Any) -> dict[str, Any]:
        files = payload.get("files") if isinstance(payload, dict) else None
        if not isinstance(files, list) or not files:
            raise ValueError("Nessun documento ricevuto")
        saved = []
        letti_per_il_contenuto: list[str] = []
        rinominati: list[str] = []
        # Names and extensions are all validated before writing a single byte: if
        # the second document in the batch is rejected, the first must not be left
        # on disk as an orphaned copy nobody knows about.
        da_scrivere: list[tuple[str, bytes, str]] = []
        for item in files:
            if not isinstance(item, dict):
                raise ValueError("Documento non valido")
            name = safe_upload_name(str(item.get("name") or ""))
            ruolo_scelto = upload_role(item.get("role") or "suppliers")
            encoded = item.get("data")
            if not isinstance(encoded, str):
                raise ValueError(f"Contenuto mancante per {name}")
            try:
                content = base64.b64decode(encoded, validate=True)
            except (binascii.Error, ValueError) as exc:
                raise ValueError(f"Base64 non valido per {name}") from exc
            if not content:
                raise ValueError(f"«{name}» è vuoto: non contiene nessun dato.")
            if len(content) > MAX_UPLOAD_BYTES:
                raise ValueError(f"«{name}» supera i {MAX_UPLOAD_BYTES // (1024 * 1024)} MB consentiti.")
            da_scrivere.append((name, content, ruolo_scelto))

        with self.lock:
            profiles_doc = load_json(self.upload_profiles_path, {"schema_version": 1, "profiles": [], "errors": []})
            # "Already present" means the copy is on disk, not just remembered.
            # Profiles whose file is gone from disk are dropped from the registry
            # here: otherwise a document deleted by hand from the folder could
            # neither be seen (the page filters it out) nor re-uploaded (the
            # fingerprint check would call it a duplicate and write nothing),
            # leaving the comparison permanently stuck.
            profili_veri, fantasmi = self._profili_veri_e_fantasmi(profiles_doc)
            if fantasmi:
                profiles_doc["profiles"] = profili_veri
                nomi_fantasma = ", ".join(sorted(str(item.get("file_name") or "?") for item in fantasmi))
                print(
                    f"[AVVISO] {len(fantasmi)} documenti non sono più nella cartella dei "
                    f"caricamenti e sono stati tolti dal registro: {nomi_fantasma}"
                )
            profili_per_hash = {
                item.get("sha256"): item for item in profili_veri
                if item.get("sha256")
            }
            known_hashes = set(profili_per_hash)
            scritti: list[Path] = []
            try:
                for name, content, ruolo_scelto in da_scrivere:
                    destination = unique_destination(self.upload_dir, name)
                    destination.write_bytes(content)
                    scritti.append(destination)
                    try:
                        profile = profile_file(destination)
                    except Exception as exc:
                        # The document never makes it to its destination, so the
                        # reason has to be stated here: otherwise the user would
                        # see the raw reader error, in English and unactionable.
                        # The exception type stays in the technical detail and on
                        # the console — it's the only thing that distinguishes a
                        # malformed file from a bug in the app itself.
                        dettaglio = f"{type(exc).__name__}: {exc}" if str(exc).strip() else type(exc).__name__
                        print(f"[AVVISO] «{name}» non profilato — {dettaglio}")
                        raise ValueError(
                            f"«{name}» non è stato riconosciuto: non è un foglio di calcolo Excel "
                            "né un CSV leggibile. Noce manda un .xls, gli altri fornitori un "
                            f".xlsx o un .csv. Dettaglio tecnico: {dettaglio}"
                        ) from exc
                    profile["upload_role"] = ruolo_scelto
                    profile.setdefault("ai_preflight", {})["role"] = ruolo_scelto
                    if profile.get("sha256") in known_hashes:
                        destination.unlink(missing_ok=True)
                        scritti.pop()
                        esistente = profili_per_hash.get(profile.get("sha256"))
                        if isinstance(esistente, dict):
                            esistente["upload_role"] = ruolo_scelto
                            esistente.setdefault("ai_preflight", {})["role"] = ruolo_scelto
                        saved.append({
                            "name": str((esistente or {}).get("file_name") or name),
                            "status": "duplicate",
                            "role": ruolo_scelto,
                            "message": "Documento già acquisito; aggiornata la sua destinazione nell'app.",
                        })
                        continue
                    profiles_doc.setdefault("profiles", []).append(profile)
                    known_hashes.add(profile.get("sha256"))
                    profili_per_hash[profile.get("sha256")] = profile
                    if destination.name != name:
                        # A corrected price list, re-uploaded, doesn't replace the
                        # old one — they coexist, and without this line the user
                        # wouldn't know which of the two they're confirming.
                        rinominati.append(f"«{name}» era già presente: salvato come «{destination.name}»")
                    formato = str(profile.get("content_format") or "")
                    atteso = FORMATO_DELL_ESTENSIONE.get(Path(destination.name).suffix.casefold())
                    if formato and atteso and formato != atteso:
                        # A renamed price list still reads fine, but stating this
                        # saves the user from puzzling over columns that aren't
                        # what the extension led them to expect.
                        letti_per_il_contenuto.append(
                            f"«{destination.name}» dentro è un {NOME_DEL_FORMATO.get(formato, formato)}"
                        )
                    saved.append({
                        "name": destination.name,
                        "status": "profiled",
                        "contentFormat": formato,
                        "schemaState": (profile.get("deterministic_hint") or {}).get("state"),
                        "adapterId": (profile.get("deterministic_hint") or {}).get("adapter_id"),
                        "role": ruolo_scelto,
                    })
            except Exception:
                for copia in scritti:
                    copia.unlink(missing_ok=True)
                raise
            profiles_doc["generated_at"] = datetime.now(tz=timezone.utc).isoformat()
            atomic_json(self.upload_profiles_path, profiles_doc)
        acquisiti = sum(1 for voce in saved if voce.get("status") == "profiled")
        duplicati = len(saved) - acquisiti
        if acquisiti:
            message = "1 documento caricato e letto." if acquisiti == 1 else f"{acquisiti} documenti caricati e letti."
            message += " Premi «Confronta i listini» quando hai finito di caricare."
        else:
            message = (
                "1 documento era già presente: nessuna nuova copia creata."
                if duplicati == 1
                else f"{duplicati} documenti erano già presenti: nessuna nuova copia creata."
            )
        if duplicati and acquisiti:
            message += f" Altri {duplicati} erano già presenti e non sono stati ricopiati." if duplicati > 1 else " Un altro era già presente e non è stato ricopiato."
        if rinominati:
            message += " " + "; ".join(rinominati) + "."
        if letti_per_il_contenuto:
            coda = "letto per quello che è" if len(letti_per_il_contenuto) == 1 else "letti per quello che sono"
            message += " " + "; ".join(letti_per_il_contenuto) + f": {coda}, non per il nome."
        # The names of the newly uploaded documents and, when the registry
        # recognizes them, the supplier names: this is what the banner in the
        # page names, since a supplier name is more useful than a filename.
        nomi_caricati = [
            str(voce.get("name") or "") for voce in saved
            if str(voce.get("status") or "") == "profiled"
        ]
        fornitori_caricati: list[str] = []
        for voce in saved:
            adattatore_id = str(voce.get("adapterId") or "")
            if str(voce.get("status") or "") != "profiled" or not adattatore_id:
                continue
            try:
                adattatore = registro.adattatore(adattatore_id)
            except (KeyError, ValueError):
                continue
            # Real suppliers only. The management-software export has an adapter
            # of its own (`gestionale_v1`) but no `supplier_id`, and
            # `supplier_label("")` falls back to a generic placeholder string. A
            # plain `if etichetta` guard isn't enough to exclude it, since that
            # placeholder is a non-empty string; the registry's `kind` field is
            # what actually distinguishes it.
            supplier_id = str((adattatore or {}).get("supplier_id") or "")
            if not supplier_id or (adattatore or {}).get("kind") == "master":
                continue
            etichetta = supplier_label(supplier_id)
            if etichetta and etichetta not in fornitori_caricati:
                fornitori_caricati.append(etichetta)
        stato_pipeline = self.pipeline_jobs.input_modificato(
            "caricato", documenti=nomi_caricati, fornitori=sorted(fornitori_caricati),
        )
        return {
            "ok": True,
            "status": "PROFILED_AI_PENDING",
            "files": saved,
            "message": message,
            "pipeline": stato_pipeline,
        }

    # The registry keys that turn a registry declaration into a usable write
    # rule on its own, without the user-confirmed column mapping.
    CHIAVI_REGOLA_DI_SCRITTURA = ("sheet", "header_row", "data_start_row", "order_column", "expected_header")

    @staticmethod
    def default_write_rule(supplier: str, percorso: Path | None = None) -> dict[str, Any] | None:
        """Return the write rule the registry declares for this supplier, if any.

        Read from the registry rather than duplicated here, so a supplier known
        only through the learned-adapter path still inherits a usable rule instead
        of being permanently left without one.

        Returns `None` when the registry declares `from_field_mapping`: there,
        sheet and row come from the user-confirmed column mapping instead, so no
        *default* rule exists and the real one must come from configuration.
        """

        chiave = str(supplier or "").strip().casefold()
        if not chiave:
            return None
        for voce in registro.adattatori(percorso):
            if str(voce.get("supplier_id") or "").strip().casefold() != chiave:
                continue
            regola = registro.scrittura_ordine(voce)
            if not regola or regola.get("from_field_mapping"):
                continue
            essenziale = {
                nome: regola[nome]
                for nome in ReviewStore.CHIAVI_REGOLA_DI_SCRITTURA
                if regola.get(nome) not in (None, "")
            }
            if essenziale.get("order_column") and essenziale.get("sheet"):
                essenziale.setdefault("data_start_row", 2)
                return essenziale
        return None

    def configurazione_di_un_altra_run(self, config: dict[str, Any], plan: dict[str, Any]) -> str:
        """Check that the write configuration belongs to the plan being compiled; refuse to write otherwise.

        After a recalculation the price lists are different files, with
        different names and row numbers. If `writer_config.json` isn't rewritten
        (the process dies mid-run, the restart never happened, reconfiguration
        fails), compilation would take last week's price list and write this
        week's row numbers into it — a quantity would end up on the wrong row's
        cell, undetected. The file fingerprint alone isn't enough: it only
        confirms the file hasn't changed since it was last verified, which stays
        true even when it's the wrong-but-internally-consistent file, and it's
        only checked for some suppliers.

        The comparison is made against the plan itself, not a live comparison
        re-read from disk: re-reading `review_data.json` here would open a race
        where a recalculation finishes between building the plan and this check,
        so the configuration and the freshly re-read comparison would agree on
        the new run while the plan — the rows about to be written — still belongs
        to the old one. The plan carries its own run id, fixed at creation time,
        which is what this check actually protects; `validate_snapshot`
        guarantees plan/comparison consistency when the plan is built.

        The check itself is a string comparison, applied to every compilable
        supplier. It fails closed: a configuration that declares no run (written
        before the field existed) and a plan with no run (the active comparison
        declares no recalculation) are both treated as "can't answer", and an
        unanswerable question doesn't get written into a supplier's price list.
        """

        attesa = str(plan.get("run_id") or "")
        dichiarata = str(config.get("run_id") or "")
        if attesa and dichiarata and attesa == dichiarata:
            return ""
        rimedio = (
            "Premi «Ricalcola il confronto» nella pagina di importazione: la configurazione "
            "viene riscritta alla fine del confronto. Anche riavviare il programma la rifà."
        )
        if not attesa:
            return (
                "Il confronto attivo non dichiara da dove viene: non creo copie dei listini, "
                "perché non posso sapere a quale listino appartengono le righe del piano. "
                + rimedio
            )
        if not dichiarata:
            return (
                "La configurazione per creare le copie dei listini non dice a quale confronto "
                "appartiene: non creo copie, perché potrebbe puntare ai listini di una volta "
                "precedente e le righe finirebbero sbagliate. " + rimedio
            )
        # Without tracking which came first, it's not knowable which of the two
        # is stale, so the message doesn't claim to know.
        return (
            "La configurazione per creare le copie dei listini non appartiene allo stesso "
            "confronto di questo piano ordini: non creo copie, perché scriverebbe le quantità "
            "nelle righe dei listini di un'altra run. " + rimedio
        )

    def writer_configuration_issues(self, plan: dict[str, Any]) -> list[str]:
        """Validate every selected supplier before the writer can create a copy.

        The JSON plan is intentionally still useful when a new supplier has not
        completed the separate XLSX-writing verification.  In particular this
        prevents a BETULLA/Larice partial result when CIPRESSO is selected too.
        """

        if self.writer_config is None:
            return []
        if not self.writer_config.is_file():
            return ["La configurazione per creare le copie dei listini non è disponibile."]
        try:
            config = load_json(self.writer_config, {})
        except (OSError, json.JSONDecodeError) as exc:
            return [f"La configurazione per creare le copie non è leggibile: {exc}."]
        if not isinstance(config, dict):
            return ["La configurazione per creare le copie non è valida."]

        selected = {
            str(item.get("supplier") or "").casefold()
            for item in plan.get("orders") or []
            if isinstance(item, dict) and item.get("supplier")
        }
        if not selected:
            return ["Il piano non contiene fornitori da compilare."]

        problema_di_run = self.configurazione_di_un_altra_run(config, plan)
        if problema_di_run:
            # Return immediately, skipping the rest of the checks: if the
            # configuration belongs to another run, everything else in it refers
            # to different files, and a longer issue list would bury the only
            # one that matters.
            return [problema_di_run]

        files_value = config.get("supplier_files")
        rules_value = config.get("supplier_write_rules")
        if not isinstance(files_value, dict):
            return ["Manca l'elenco dei listini da compilare nella configurazione."]
        if rules_value is None:
            rules_value = {}
        if not isinstance(rules_value, dict):
            return ["Le regole di scrittura dei listini non sono valide."]
        supplier_files = {str(key).casefold(): value for key, value in files_value.items()}
        supplier_rules = {str(key).casefold(): value for key, value in rules_value.items()}
        issues: list[str] = []

        # A supplier compiled in place doesn't go through Node: its copy is a
        # Python patch on the supplier's own document. An order made entirely of
        # such suppliers compiles without the writer, so requiring Node there
        # would reject a compilation that's actually possible. Whether a supplier
        # is patched in place is declared by its write rule, not by its name.
        in_posizione = {
            supplier for supplier in selected
            if procedura_di_scrittura(supplier_rules.get(supplier)) == PATCH_IN_POSIZIONE
        }
        requires_writer = bool(set(selected) - in_posizione)
        if requires_writer:
            node_value = config.get("node_executable")
            script_value = config.get("writer_script")
            if not node_value or not Path(str(node_value)).is_file():
                issues.append("Sul computer manca quello che serve per scrivere i fogli di calcolo: non posso creare le copie dei listini.")
            if not script_value or not Path(str(script_value)).is_file():
                issues.append("Non è disponibile lo strumento di compilazione dei listini.")

        for supplier in sorted(selected):
            if supplier in in_posizione:
                issues.extend(self.problemi_in_posizione(supplier, supplier_files, supplier_rules))
                continue
            source_value = supplier_files.get(supplier)
            if not isinstance(source_value, str) or not source_value.strip():
                issues.append(f"{supplier_label(supplier)} è selezionato, ma manca il suo listino configurato per la compilazione.")
                continue
            source = Path(source_value.strip()).expanduser()
            if not source.is_absolute():
                source = self.writer_config.parent / source
            source = source.resolve()
            if not source.is_file():
                issues.append(f"Il listino configurato per {supplier_label(supplier)} non è disponibile.")
                continue
            if source.suffix.casefold() != ".xlsx":
                # The document itself isn't wrong: the configuration just doesn't
                # say how to compile it. Blaming the price list would send the
                # user chasing a different file from the supplier, which fixes
                # nothing — the write procedure is declared per supplier instead,
                # as it is for the in-place `.xls` case.
                if not isinstance(supplier_rules.get(supplier), dict):
                    # No rule at all: same message a missing `.xlsx` rule gets
                    # further below.
                    issues.append(
                        f"{supplier_label(supplier)} è selezionato, ma manca la regola di "
                        "scrittura verificata del suo listino."
                    )
                else:
                    # Neither the document nor necessarily the configuration is
                    # wrong here, so both possibilities are stated. The actionable
                    # fix is a plain Excel re-save, so that's what's suggested.
                    issues.append(
                        f"Il listino di {supplier_label(supplier)} non è un .xlsx e la sua "
                        "configurazione non dichiara come compilarlo: se il fornitore lo manda "
                        "in Excel 97-2003, aprilo con Excel e salvalo come «Cartella di lavoro "
                        "di Excel (.xlsx)», poi ricaricalo."
                    )
                continue

            rule = supplier_rules.get(supplier)
            if rule is None:
                rule = self.default_write_rule(supplier)
            if not isinstance(rule, dict):
                issues.append(f"{supplier_label(supplier)} è selezionato, ma manca la regola di scrittura verificata.")
                continue
            order_column = str(rule.get("order_column") or rule.get("orderColumn") or "").strip().upper()
            sheet_name = str(rule.get("sheet_name") or rule.get("sheet") or "").strip()
            data_start_value = rule.get("data_start_row")
            if data_start_value is None:
                data_start_value = rule.get("dataStartRow")
            if data_start_value is None:
                data_start_value = 2
            data_start = number(data_start_value)
            if not re.fullmatch(r"[A-Z]{1,3}", order_column) or not sheet_name:
                issues.append(f"La regola di scrittura per {supplier_label(supplier)} è incompleta.")
                continue
            if data_start is None or data_start < 1 or int(data_start) != data_start:
                issues.append(f"La riga iniziale della regola di scrittura per {supplier_label(supplier)} non è valida.")
                continue

            # This check runs for whichever supplier declares it, not for one
            # hardcoded name: it defends against deleting the wrong price list,
            # reloading the right one under the same filename, and compiling
            # without recalculating — which would write the new list's row
            # numbers against the old comparison's quantities. The registry
            # decides who gets checked (whoever declares a fingerprint and
            # expected header), the code just executes it.
            etichetta_fornitore = supplier_label(supplier)
            header_row = number(rule.get("header_row") or rule.get("headerRow"))
            expected_header = normalize_header(
                rule.get("expected_header")
                or rule.get("expectedHeader")
                or rule.get("order_header")
                or rule.get("orderHeader")
            )
            expected_hash = str(rule.get("source_sha256") or rule.get("sourceSha256") or "").strip().casefold()

            # The fingerprint: declared by `source_rule` for any supplier.
            if expected_hash:
                if not re.fullmatch(r"[a-f0-9]{64}", expected_hash):
                    issues.append(f"L'impronta dichiarata del listino {etichetta_fornitore} non è valida.")
                    continue
                if sha256_file(source) != expected_hash:
                    issues.append(
                        f"Il listino {etichetta_fornitore} è cambiato dopo l'ultimo confronto: "
                        "non creo copie finché non rifai il confronto."
                    )
                    continue

            # The verified header: only checked for suppliers that declare one.
            if expected_header:
                if sheet_name.casefold() == "first":
                    issues.append(f"Per {etichetta_fornitore} manca il nome esatto del foglio verificato.")
                    continue
                if header_row is None or header_row < 1 or int(header_row) != header_row:
                    issues.append(f"Per {etichetta_fornitore} manca la riga dell'intestazione verificata.")
                    continue

            if not expected_header and sheet_name.casefold() == "first":
                # With no sheet name and no declared header there's nothing to
                # reopen the file for: stick with the lighter check above.
                continue

            try:
                from openpyxl import load_workbook
                from openpyxl.utils import column_index_from_string

                workbook = load_workbook(source, read_only=True, data_only=False)
                try:
                    if sheet_name not in workbook.sheetnames:
                        issues.append(f"Il foglio verificato del listino {etichetta_fornitore} non è presente.")
                        continue
                    sheet = workbook[sheet_name]
                    if expected_header:
                        indice = column_index_from_string(order_column)
                        letto = normalize_header(sheet.cell(int(header_row), indice).value)
                        if letto != expected_header:
                            issues.append(
                                f"L'intestazione {order_column} del listino {etichetta_fornitore} "
                                f"non coincide più con {expected_header}."
                            )
                            continue
                    for order in plan.get("orders") or []:
                        if str(order.get("supplier") or "").casefold() != supplier:
                            continue
                        source_row = number(order.get("supplier_source_row"))
                        if (
                            source_row is None
                            or source_row < data_start
                            or int(source_row) != source_row
                            or source_row > sheet.max_row
                        ):
                            issues.append(
                                f"Una riga selezionata per {etichetta_fornitore} non appartiene "
                                "al listino verificato."
                            )
                            break
                finally:
                    workbook.close()
            except Exception as exc:
                issues.append(
                    f"Non riesco a verificare il foglio {etichetta_fornitore} prima della compilazione: {exc}."
                )
        return issues

    def problemi_in_posizione(
        self, supplier: str, supplier_files: dict[str, Any], supplier_rules: dict[str, Any]
    ) -> list[str]:
        """Return what blocks compiling this supplier in place, checked before writing.

        The expensive check — that the order column is still all fixed-length
        numbers — runs here, ahead of the write itself, not only at write time:
        it's the condition the whole patch depends on, and has to be redone for
        every file, since one week the supplier could put a formula in that
        column and invalidate it.

        Applies to whichever supplier declares `patch_xls_in_posizione` in its
        write rule, not to a hardcoded name.
        """

        from xls_writer import XlsError as ErroreXls, controlla_colonna_ordine

        nome = supplier_label(supplier)
        problemi: list[str] = []
        raw = supplier_files.get(supplier)
        if not isinstance(raw, str) or not raw.strip():
            return [f"{nome} è selezionato, ma manca il suo listino configurato per la compilazione."]
        source = Path(raw.strip()).expanduser()
        if not source.is_absolute() and self.writer_config is not None:
            source = self.writer_config.parent / source
        source = source.resolve()
        if not source.is_file():
            return [f"Il listino configurato per {nome} non è disponibile."]
        regola = supplier_rules.get(supplier)
        if not isinstance(regola, dict):
            return [f"{nome} è selezionato, ma manca la regola di scrittura verificata del suo listino."]
        colonna = str(regola.get("order_column") or "").strip().upper()
        if not re.fullmatch(r"[A-Z]{1,3}", colonna):
            problemi.append(f"La colonna d'ordine verificata di {nome} non è dichiarata nel registro.")
        impronta = str(regola.get("source_sha256") or "").strip().casefold()
        if impronta and sha256_file(source) != impronta:
            problemi.append(f"Il listino {nome} è cambiato dopo la verifica: va ricontrollato.")
        if problemi:
            return problemi
        try:
            esito = controlla_colonna_ordine(
                source,
                colonna_ordine=colonna,
                foglio=str(regola.get("sheet") or "") or None,
                prima_riga=int(number(regola.get("data_start_row")) or 1),
            )
        except ErroreXls as exc:
            return [f"Il listino {nome} non si lascia leggere per la compilazione: {frase(exc)}"]
        if not esito.get("compilabile"):
            quante = len(esito.get("celle_di_altro_tipo") or {})
            problemi.append(
                f"La colonna d'ordine {colonna} del listino {nome} ha {quante} celle che non "
                "sono numeri a lunghezza fissa: il loro documento non si può compilare in "
                f"posizione e l'ordine {nome} va preparato a mano."
            )
        return problemi

    @staticmethod
    def righe_di_listino_contese(orders: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """Find two plan rows that would land in the same price-list cell.

        Both writers (the Node writer and `xls_writer` for the in-place case)
        sum the quantities of two plan rows that point to the same price-list
        row — correct when two management-software items map to the same
        supplier item.

        It's wrong when the two rows order different units, which the app can
        produce on its own: a display's offer carries the row number of its
        parent carton (`display_offer`), and that parent carton stays orderable
        on its own too. Ordering 4 cartons and 6 displays of the same product
        would then write `10` into a single cell, and the supplier would read ten
        of something nobody actually ordered.

        Summing isn't safe here and guessing isn't either, so this stops and
        reports which two products are contending for the row. The user decides
        the number to write, by moving one of the two rows or zeroing it out.
        """

        per_riga: dict[tuple[str, int], list[dict[str, Any]]] = {}
        for order in orders:
            riga = number(order.get("supplier_source_row"))
            fornitore = str(order.get("supplier") or "").strip().casefold()
            if riga is None or not fornitore:
                continue
            per_riga.setdefault((fornitore, int(riga)), []).append(order)
        errori: list[dict[str, Any]] = []
        for (fornitore, riga), voci in sorted(per_riga.items()):
            if len(voci) < 2:
                continue
            unita = {str(voce.get("desired_quantity_unit") or "") for voce in voci}
            fattori = {round(float(number(voce.get("quantity_factor")) or 0), 6) for voce in voci}
            if len(unita) == 1 and len(fattori) == 1:
                # The same supplier item bought for two management-software
                # items: summing is exactly what's wanted here.
                continue
            altri = ", ".join(
                f"«{str(voce.get('description') or voce.get('product_id'))}»" for voce in voci[1:]
            )
            errori.append({
                "code": "RIGA_LISTINO_CONTESA",
                "productId": str(voci[0].get("product_id") or ""),
                "productName": str(voci[0].get("description") or voci[0].get("product_id") or ""),
                "supplierId": fornitore,
                "supplierName": supplier_label(fornitore),
                "sourceRow": riga,
                "otherProducts": altri,
            })
        return errori

    def compile(self, snapshot: Any) -> dict[str, Any]:
        with self.lock:
            # Not while the pipeline is running. `compile` holds the data lock,
            # not the pipeline lock: until the pipeline's last stage replaces the
            # live comparison, compiling now would be legal but would produce the
            # price lists of the *previous* comparison — last week's prices —
            # while the page shows a progress bar claiming an update is underway.
            # The user believes they're compiling what's being built. Waiting
            # costs nothing: the pipeline finishes on its own and the button comes
            # back.
            if self.pipeline_jobs.in_corso():
                raise LavoroGiaInCorso(
                    "Il confronto si sta aggiornando: i listini si compilano quando ha finito, "
                    "altrimenti sarebbero quelli di prima."
                )
            clean, errors, review = self.validate_snapshot(snapshot, for_compile=True)
            if errors:
                raise SnapshotError(errors)
            selections = {str(item["id"]): item for item in clean["products"]}
            review = self.promotion_service.decorate(review, selections)
            products_by_id = {str(product.get("id")): product for product in review.get("products") or []}
            supplier_defs = {str(item.get("id")): item for item in review.get("suppliers") or []}
            totals: dict[str, float] = {supplier: 0.0 for supplier in supplier_defs}
            orders = []
            # Products the management software asks for that no supplier
            # carries. They don't enter the plan or any supplier's price list —
            # there's no row to write them on — but aren't lost: they go into a
            # separate list, the only thing that can be done with them.
            da_reperire: list[dict[str, Any]] = []
            for decision in clean["products"]:
                if decision["quantity"] <= 0:
                    continue
                product = products_by_id[decision["id"]]
                supplier = decision["selectedSupplierId"]
                if not supplier and nessuna_offerta_utilizzabile(product):
                    da_reperire.append({
                        "ean": str(product.get("ean") or ""),
                        "descrizione": str(product.get("name") or product.get("description") or ""),
                        "quantita": decision["quantity"],
                        "unita": "espositori" if str(
                            product.get("itemType") or product.get("kind") or ""
                        ).casefold() == "display" else "colli",
                        "ultimo_prezzo": product.get("lastUnitPrice"),
                        "motivo": da_reperire_modulo.motivo(product.get("offers") or []),
                    })
                    continue
                offer = find_offer(product, supplier)
                pricing = offer_pricing(offer)
                if pricing is None:
                    errors.append({
                        "code": "PREZZO_NON_VALIDO",
                        "productId": decision["id"],
                        "productName": str(product.get("name") or product.get("description") or decision["id"]),
                        "supplierId": supplier,
                        "supplierName": supplier_label(supplier) if supplier else "",
                    })
                    continue
                offer = offer or {}
                factor = pricing["factor"]
                unit_price = pricing["unitPriceNet"]
                order_price = pricing["orderUnitPriceNet"]
                is_display = str(product.get("itemType") or product.get("kind") or "").casefold() == "display"
                desired_quantity = decision["quantity"]
                # The user enters cartons (or displays) directly: no
                # pieces-to-cartons rounding, for any item type or supplier.
                order_quantity = desired_quantity
                delivered_pieces = int(round(order_quantity * factor))
                excess_pieces = 0
                line_total = round(order_quantity * order_price, 4)
                totals[supplier] = round(totals.get(supplier, 0.0) + line_total, 4)
                orders.append({
                    "product_id": decision["id"],
                    "item_type": str(product.get("itemType") or product.get("kind") or "product").lower(),
                    "gestionale_source_row": product.get("sourceRow") or product.get("source_row"),
                    "ean": product.get("ean") or "",
                    "description": product.get("name") or product.get("description") or "",
                    "supplier": supplier,
                    "supplier_source_row": offer.get("sourceRow") or offer.get("source_row"),
                    "supplier_ean": offer.get("ean") or "",
                    "supplier_description": offer.get("description") or product.get("name") or "",
                    "quantity": order_quantity,
                    "desired_quantity": desired_quantity,
                    "desired_quantity_unit": "espositori" if is_display else "colli",
                    "delivered_pieces": delivered_pieces,
                    "excess_pieces": excess_pieces,
                    "unit_price_net": round(unit_price, 6),
                    "quantity_factor": factor,
                    "order_unit_price_net": round(order_price, 6),
                    "line_total_net": line_total,
                    "match_method": offer.get("method") or offer.get("matchStatus") or offer.get("match_status"),
                    "match_confidence": offer.get("confidence"),
                    "confirmed": decision["confirmed"],
                    "promotions": offer.get("promotions") or [],
                })
            errors.extend(self.righe_di_listino_contese(orders))
            if errors:
                raise SnapshotError(errors)

            below = []
            for supplier, total in totals.items():
                threshold = supplier_threshold(supplier_defs.get(supplier, {}))
                if not meets_threshold(total, threshold):
                    below.append({"supplier": supplier, "total_net": round(total, 2), "threshold_net": round(threshold, 2)})
            if below and not clean["acceptBelowThreshold"]:
                raise SnapshotError([{"code": "SOGLIA_NON_CONFERMATA", **item} for item in below])


            plan = {
                "schema_version": 1,
                "created_at": datetime.now(tz=timezone.utc).isoformat(),
                "source_review": str(self.review_path),
                "source_state": str(self.state_path),
                # The plan's run id is the one from the comparison that resolved
                # these offers, not the one the browser's snapshot declared: when
                # the comparison declares no run, the two can diverge
                # (`validate_snapshot` has nothing to check against), and the plan
                # must not "belong" to a run on the client's word alone.
                "run_id": str((review.get("run") or {}).get("id") or ""),
                "threshold_override_confirmed": clean["acceptBelowThreshold"],
                "below_threshold": below,
                "totals_net": {supplier: round(total, 2) for supplier, total in totals.items()},
                "promotion_summary": review.get("promotionSummary") or {},
                "promotions": review.get("promotions") or [],
                "orders": orders,
            }
            # The state is rewritten by compile, not by the tab: if something
            # goes wrong from here on, a tab that finds itself behind a version
            # needs to know it was its own compilation that overtook it.
            clean["stateVersionOrigin"] = "compilazione"
            atomic_json(self.state_path, clean)
            # Writing starts here. The folder is created only after the last
            # validation, threshold included, so a rejected compilation doesn't
            # leave behind an empty folder that the compilations list would show
            # as a completed delivery.
            momento = datetime.now().astimezone()
            cartella = consegna.crea_cartella(self.orders_dir, momento)
            # The folder counts as a real delivery only once this reaches the
            # end. The audit file is written last, and until it exists the work
            # is incomplete: the compilations list would otherwise show it as
            # delivered, with a downloadable zip whose copies never passed the
            # fidelity check.
            consegna_completata = False
            # History is touched before the audit, and the audit can still fail
            # (full disk, antivirus, a locked file). If the folder were simply
            # deleted on failure, history would be left rewritten with the new,
            # never-delivered order and without the real one — which
            # `record_plan` drops because it's superseded. The two must stay in
            # sync, the same way they do when a compilation is deleted.
            storico_di_prima = self.history_path.read_bytes() if self.history_path.exists() else None
            try:
                plan_path = cartella / consegna.NOME_PIANO
                atomic_json(plan_path, plan)
                # If there's nothing to source, the file isn't created, and the
                # page has nothing new to show: a week where everything can be
                # ordered is the normal case, not something worth flagging.
                if da_reperire:
                    da_reperire_modulo.scrivi(cartella, da_reperire, momento)
                history_issues: list[str] = []
                writer_issues: list[str] = []
                # Kept separate even after being merged into `writer_issues`:
                # they describe opposite things (a copy that wasn't delivered vs.
                # a copy that was delivered with a row to double-check), and the
                # final message treats them differently.
                avvisi_writer: list[str] = []
                infedeli: list[str] = []
                copie: list[tuple[str, Path, Path]] = []
                problemi_consegna: list[str] = []
                status = "PLAN_READY"
                # Why the copies don't exist, when they don't. Empty means
                # "nothing was even attempted" — the actual sentence is composed
                # by `messaggio_della_compilazione`, the single place that decides
                # what the user reads.
                copie_non_create: list[str] = []
                if self.writer_config:
                    writer_issues = self.writer_configuration_issues(plan)
                    if writer_issues:
                        copie_non_create = list(writer_issues)
                    else:
                        try:
                            _generated, prodotte, avvisi_writer = self.run_writer(plan_path, cartella)
                        except ValueError as exc:
                            # The writer verifies everything before exposing the
                            # copies, so the plan stays deliverable even when a
                            # final price-list check fails.
                            #
                            # The reason goes into `writer_issues`, not just the
                            # final message: it's the same field filled by the
                            # branch above when the configuration is missing, and
                            # it's the only one the page reads to know that
                            # compilation didn't fully succeed. If it stayed only
                            # in the free-text message, a compilation that
                            # produced no copy at all could look like a full
                            # delivery.
                            writer_issues = [frase(exc) or "Le copie dei listini non sono state create."]
                            copie_non_create = list(writer_issues)
                        else:
                            # Before delivering anything: is the copy actually the
                            # supplier's price list, or something else? The writer
                            # only checks the cells it meant to write; this guard
                            # checks all of them. A copy that fails this never
                            # leaves this function.
                            prodotte, infedeli = self.scarta_copie_infedeli(plan, prodotte)
                            # Writer warnings come first: they're about the
                            # copies actually being delivered, and the reader
                            # needs to see them even when no copy was discarded.
                            writer_issues = [*avvisi_writer, *infedeli]
                            if prodotte:
                                # A rename failure can't fail an otherwise
                                # successful compilation: the copies exist and are
                                # correct, and an ugly filename is reported, not
                                # turned into an error.
                                copie, problemi_consegna = self.rinomina_listini(prodotte, momento)
                                status = "FILES_READY"
                            else:
                                # Here the reason is the discarded copies
                                # (`infedeli`), not the writer's own warnings:
                                # those say "the quantity was written anyway",
                                # which would contradict a message stating nothing
                                # was created.
                                copie_non_create = list(infedeli)
                # History is recorded here, after the writes, not before: a
                # compilation stopped by `writer_issues` (or a writer that never
                # starts) must not register the order in history, since the
                # program would then ask "did it arrive?" about goods nobody
                # could actually order. The rule is simple: only a supplier whose
                # price-list copy exists on disk enters history.
                if status == "FILES_READY":
                    history_issues = self.record_order_history(
                        plan,
                        order_key=cartella.name,
                        # Casefolded because `record_plan` groups the plan by
                        # casefolded supplier name; a supplier name with different
                        # casing here would never match its own delivered copy.
                        delivered={str(fornitore).strip().casefold()
                                   for fornitore, _sorgente, _copia in copie},
                    )
                elif status == "PLAN_READY":
                    # No writer means no copy, which means no order — that's the
                    # rule. But silently skipping history too (no reminders, no
                    # follow-up) would turn an ordinary case (missing Node,
                    # missing a price list) into a silent loss of tracking, so
                    # it's stated explicitly instead.
                    history_issues = [
                        "Questa compilazione non entra fra gli ordini da controllare: senza le "
                        "copie dei listini non c'è nessun ordine da mandare, quindi nessuna "
                        "domanda «è arrivata?» nascerà per lei. Il piano ordini JSON resta "
                        "scaricabile."
                    ]
                # The file list is read from disk, not from the list already in
                # hand: if the writer left behind something unexpected, the audit
                # must reflect it instead of describing an imagined folder.
                file_audit = self.descrivi_cartella(cartella, copie)
                message = self.messaggio_della_compilazione(
                    status=status,
                    writer_configurato=bool(self.writer_config),
                    file_audit=file_audit,
                    copie_non_create=copie_non_create,
                    infedeli=infedeli,
                    avvisi_writer=avvisi_writer,
                    problemi_consegna=problemi_consegna,
                    history_issues=history_issues,
                )
                fornitori_ordinati = sorted({str(item["supplier"]) for item in orders})
                totali_netti = {supplier: round(totals[supplier], 2) for supplier in fornitori_ordinati}
                # The audit is written last: its presence means everything else
                # in the folder was already in place.
                consegna.scrivi_audit(cartella, {
                    "schema_audit": consegna.SCHEMA_AUDIT,
                    "cartella": cartella.name,
                    "run_id": clean["runId"],
                    "creato_il": momento.isoformat(),
                    "stato": status,
                    "messaggio": message,
                    "fornitori": fornitori_ordinati,
                    "totali_netti": totali_netti,
                    "totale_netto": round(sum(totali_netti.values()), 2),
                    "righe": len(orders),
                    "sotto_soglia": below,
                    "avvisi": [*writer_issues, *problemi_consegna, *history_issues],
                    "file": file_audit,
                })
                # From here on the folder is a real delivery and stays on disk.
                consegna_completata = True
                # The entry is re-read from the audit file just written, so the
                # compile response and `GET /api/ordini`'s listing can't disagree
                # about the same folder.
                voce = consegna.voce(cartella, etichetta_fornitore=supplier_label)
                return {
                    "ok": True,
                    "status": status,
                    "message": message,
                    # Compilation also rewrites the whole state: without
                    # returning the new version, the tab that just compiled would
                    # have its next save rejected as stale.
                    "stateVersion": clean["stateVersion"],
                    "cartella": cartella.name,
                    "outputs": [
                        {"name": item["nome"], "tipo": item["tipo"], "url": item["url"]}
                        for item in voce["file"]
                    ],
                    "zipUrl": voce["zipUrl"],
                    "zipNome": voce["zipNome"],
                    "totalsNet": plan["totals_net"],
                    "belowThreshold": below,
                    "writerIssues": writer_issues,
                    # Delivery warnings (e.g. a filename that stayed ugly
                    # because `os.rename` failed) are a dedicated field, not just
                    # embedded in `message`, so the page can flag them and the
                    # audit file preserves them across reloads.
                    "deliveryIssues": problemi_consegna,
                    "historyIssues": history_issues,
                }
            finally:
                if not consegna_completata:
                    shutil.rmtree(cartella, ignore_errors=True)
                    if storico_di_prima is None:
                        self.history_path.unlink(missing_ok=True)
                    else:
                        # Same helper used elsewhere: write to a temp file first,
                        # then replace, because this is the branch that has to
                        # work reliably when something has already gone wrong.
                        scrittura_sicura.scrivi_bytes(self.history_path, storico_di_prima)

    @staticmethod
    def spiegazione_del_writer(stderr: str | None, stdout: str | None) -> str:
        """Return what the user reads when the writer stops.

        The writer marks its own message with `ERRORE_COMPILAZIONE:` on a single
        line; that's what's extracted and shown. If it's missing — an unexpected
        failure the writer didn't anticipate — the user is told compilation
        stopped without dumping the raw Node stack trace into the page: technical
        details stay on the app's own console, where they're actually useful.
        """

        marca = "ERRORE_COMPILAZIONE:"
        righe = str(stderr or "").splitlines()
        for riga in reversed(righe):
            testo = riga.strip()
            if testo.startswith(marca):
                spiegazione = testo[len(marca):].strip()
                if spiegazione:
                    # Whatever surrounds the marked line (stack trace, the cause
                    # of an unexpected failure) is technical detail: it stays on
                    # the app's console, not in the page.
                    tecnico = "\n".join(
                        r for r in righe if not r.strip().startswith(marca)
                    ).strip()
                    if tecnico:
                        print(f"[WRITER] {tecnico[-4000:]}")
                    return spiegazione
        tecnico = (str(stderr or "") or str(stdout or "")).strip()
        if tecnico:
            print(f"[WRITER] {tecnico[-4000:]}")
        return (
            "lo strumento che crea le copie dei listini si è fermato senza spiegare perché. "
            "Il piano ordini è completo e si può scaricare."
        )

    def scarta_copie_infedeli(
        self, plan: dict[str, Any], prodotte: list[tuple[str, Path, Path]]
    ) -> tuple[list[tuple[str, Path, Path]], list[str]]:
        """Reopen every produced copy and compare it cell by cell against its source price list.

        This closes a gap the writer's own checks miss, since it only verifies
        the cells it meant to write — a corrupted copy that still passed those
        checks (garbled EAN cells, missing section titles) would otherwise be
        delivered as if nothing were wrong.

        The only differences allowed are in the order column: the plan's
        quantities, and zeroing out quantities that were already there. Any other
        difference fails compilation for that supplier — the copy is deleted and
        the reason is reported, rather than delivering a document that resembles
        the price list without actually being it.

        Returns the copies that can be delivered and the messages to show for the
        rest.
        """

        from copia_fedele import confronta_copia, frase_di_rifiuto

        config = load_json(self.writer_config, {}) if self.writer_config else {}
        regole_value = config.get("supplier_write_rules") if isinstance(config, dict) else None
        regole = regole_value if isinstance(regole_value, dict) else {}
        buone: list[tuple[str, Path, Path]] = []
        problemi: list[str] = []
        for fornitore, sorgente, copia in prodotte:
            # A supplier compiled in place gets back its own `.xls`: fidelity
            # there isn't a comparison, it's the untouched bytes, already
            # guaranteed by `app/xls_writer.py` on its own.
            if copia.suffix.casefold() != ".xlsx":
                buone.append((fornitore, sorgente, copia))
                continue
            regola = regole.get(fornitore)
            if not isinstance(regola, dict):
                regola = self.default_write_rule(fornitore) or {}
            colonna = str(regola.get("order_column") or regola.get("orderColumn") or "").strip()
            prima_riga = number(regola.get("data_start_row") or regola.get("dataStartRow")) or 2
            quantita: dict[int, int] = {}
            for order in plan.get("orders") or []:
                if str(order.get("supplier") or "").casefold() != fornitore:
                    continue
                riga = number(order.get("supplier_source_row"))
                colli = number(order.get("quantity"))
                if riga is None or colli is None:
                    continue
                quantita[int(riga)] = quantita.get(int(riga), 0) + int(colli)
            try:
                esito = confronta_copia(
                    sorgente,
                    copia,
                    colonna_ordine=colonna,
                    prima_riga=int(prima_riga),
                    quantita=quantita,
                    foglio_ordine=str(regola.get("sheet") or regola.get("sheet_name") or "") or None,
                )
            except Exception as errore:  # noqa: BLE001 - `ConfrontoImpossibile` compreso
                # A copy that can't be verified isn't delivered: same rule as
                # the rest of the app, where doubt is reported, not shipped.
                copia.unlink(missing_ok=True)
                problemi.append(
                    f"La copia per {supplier_label(fornitore)} non è stata verificata e quindi "
                    f"non viene consegnata: {frase(errore)}"
                )
                continue
            if esito.fedele:
                buone.append((fornitore, sorgente, copia))
                continue
            copia.unlink(missing_ok=True)
            problemi.append(frase_di_rifiuto(supplier_label(fornitore), esito))
        return buone, problemi

    @staticmethod
    def messaggio_dei_file(file_audit: list[dict[str, Any]]) -> str:
        """Summarize what was generated, counted straight from disk."""

        workbook_count = sum(1 for item in file_audit if item["tipo"] == "listino")
        parts = ["piano ordini JSON"]
        if workbook_count:
            parts.append(f"{workbook_count} copia del listino" if workbook_count == 1 else f"{workbook_count} copie dei listini")
        summary = " e ".join(parts)
        summary = summary[:1].upper() + summary[1:]
        verb = "generato" if len(parts) == 1 else "generati"
        return f"{summary} {verb}. Gli originali sono rimasti invariati; nessun ordine è stato inviato."

    @staticmethod
    def messaggio_della_compilazione(
        *,
        status: str,
        writer_configurato: bool,
        file_audit: list[dict[str, Any]],
        copie_non_create: list[str],
        infedeli: list[str],
        avvisi_writer: list[str],
        problemi_consegna: list[str],
        history_issues: list[str],
    ) -> str:
        """Build the message the user reads at the end of compilation, also stored in the audit file.

        This composes the message from a small table of cases instead of nested
        conditionals, since testing six branches of string concatenation by hand
        was error-prone and only checkable by actually running a compilation.

        The cases, in the order they're added:

        1. no writer configured: no copies, states what's missing to get them;
        2. no copies created: configuration exists but something stopped it,
           or every copy was discarded. `copie_non_create` carries the reason —
           in the discarded-copy case that's `infedeli`, not the writer's own
           warnings, which say "the quantity was written anyway" and would
           contradict a message stating nothing was created.

        What separates case 1 from case 2 is `writer_configurato`, not whether
        `copie_non_create` is empty — "nobody even tried" and "it tried and
        nothing came out of it" are different situations, and confusing them
        tells a user with a complete configuration to go complete it.

        3. copies created: summarizes what's in the folder;
        4. some copies discarded, others not: names which price list is
           missing, rather than just a file count;
        5. writer warnings: only the count. The text itself is already listed
           on the page via `writerIssues`; repeating it here would duplicate it
           under a heading that contradicts it, so only the count is kept, since
           the compilation history stores this message but not the full list;
        6. rename and history: an ugly filename or an order that doesn't
           enter the tracking history don't fail anything, but are worth stating.
        """

        if status == "FILES_READY":
            message = ReviewStore.messaggio_dei_file(file_audit)
            if infedeli:
                message = f"{message} Attenzione: " + " ".join(infedeli)
            if avvisi_writer:
                message = (
                    f"{message} Su un listino preparato c'è una segnalazione da leggere."
                    if len(avvisi_writer) == 1 else
                    f"{message} Sui listini preparati ci sono {len(avvisi_writer)} "
                    "segnalazioni da leggere."
                )
        elif writer_configurato:
            message = (
                "Piano ordini convalidato. Non sono state create copie dei listini: "
                + " ".join(copie_non_create)
            )
        else:
            message = (
                "Piano ordini convalidato. Per creare le copie dei listini occorre "
                "completare la configurazione di scrittura."
            )
        if problemi_consegna:
            message = f"{message} Attenzione: " + " ".join(problemi_consegna)
        if history_issues:
            # Without this warning, compilation would look fully successful and
            # the follow-up reminder about undelivered goods would never appear —
            # exactly what this function exists to prevent.
            message = (
                f"{message} Attenzione: questo ordine non è stato registrato fra quelli "
                "da controllare la prossima settimana. "
                + " ".join(history_issues)
            )
        return message

    @staticmethod
    def rinomina_listini(
        prodotte: list[tuple[str, Path, Path]], momento: datetime
    ) -> tuple[list[tuple[str, Path, Path]], list[str]]:
        """Rename each produced copy to its human-readable delivery name.

        Returns the copies with the name they actually ended up with, and the
        list of ones that couldn't get the readable name.

        An ugly filename isn't a reason to discard a correct order. On Windows,
        `os.rename` raises `PermissionError` if another process holds the file
        open (antivirus scanning it, a sync client, the user previewing it in
        Excel) — raising there would fail the whole compilation with a correct,
        deliverable price list sitting right there. So the copy is kept as-is and
        the failure is reported instead.

        An existing destination is never overwritten: inside a freshly created
        delivery folder, the user expects nothing in it to be touched again once
        written.
        """

        rinominate: list[tuple[str, Path, Path]] = []
        problemi: list[str] = []
        for fornitore, sorgente, prodotta in prodotte:
            # The copy carries its own extension; it isn't chosen here — a
            # supplier compiled in place keeps its original `.xls`.
            # The name used in the delivered filename is the supplier's display
            # name from the registry, not its internal identifier:
            # `consegna.nome_listino` uppercases what it's given, so passing the
            # raw identifier would produce a filename with underscores in it
            # instead of the readable supplier name.
            destinazione = prodotta.parent / consegna.nome_listino(
                supplier_label(fornitore), momento, prodotta.suffix
            )
            if destinazione == prodotta:
                rinominate.append((fornitore, sorgente, prodotta))
                continue
            if destinazione.exists():
                problemi.append(
                    f"«{prodotta.name}» è rimasto con questo nome: nella cartella "
                    f"c'era già un documento chiamato «{destinazione.name}»."
                )
                rinominate.append((fornitore, sorgente, prodotta))
                continue
            try:
                os.rename(prodotta, destinazione)
            except OSError as errore:
                problemi.append(
                    f"«{prodotta.name}» è rimasto con questo nome invece di diventare "
                    f"«{destinazione.name}»: {errore.strerror or errore}. "
                    "Il documento è completo e si può scaricare lo stesso."
                )
                rinominate.append((fornitore, sorgente, prodotta))
                continue
            rinominate.append((fornitore, sorgente, destinazione))
        return rinominate, problemi

    @staticmethod
    def descrivi_cartella(cartella: Path, copie: list[tuple[str, Path, Path]]) -> list[dict[str, Any]]:
        """Build the audit's `file` list by scanning the folder on disk.

        Reads the disk rather than the list of files just written: an audit
        built from its own intentions would describe what was supposed to
        happen, not what actually did. A `.xlsx` this code didn't rename itself
        stays classified as `altro` — its origin is unknown, so it isn't counted
        among the price lists to deliver.
        """

        listini = {prodotta.name: (fornitore, sorgente) for fornitore, sorgente, prodotta in copie}
        with os.scandir(cartella) as scansione:
            nomi = sorted(
                trovato.name for trovato in scansione
                if trovato.is_file() and consegna.e_documento(trovato.name)
            )
        righe: list[dict[str, Any]] = []
        for nome in nomi:
            try:
                byte: int | None = (cartella / nome).stat().st_size
            except OSError:
                byte = None
            if nome in listini:
                fornitore, sorgente = listini[nome]
                righe.append({
                    "nome": nome,
                    "tipo": "listino",
                    "fornitore": fornitore,
                    "origine": str(sorgente),
                    "byte": byte,
                })
                continue
            tipo = consegna.tipo_file(nome)
            righe.append({"nome": nome, "tipo": tipo if tipo != "listino" else "altro", "byte": byte})
        return righe

    @staticmethod
    def riepilogo_del_writer(stdout: str | None) -> dict[str, Any]:
        """Return the summary the Node writer declares at the end of its run.

        `stdout` as a whole isn't parsed as JSON, since the spreadsheet library
        writes its own lines into it. The writer marks its summary line with
        `RIEPILOGO_COMPILAZIONE:` specifically for this. An older writer that
        doesn't emit it just means nothing is known here, which isn't an error.
        """

        marca = "RIEPILOGO_COMPILAZIONE:"
        for riga in reversed(str(stdout or "").splitlines()):
            pulita = riga.strip()
            if not pulita.startswith(marca):
                continue
            try:
                letto = json.loads(pulita[len(marca):])
            except ValueError:
                return {}
            return letto if isinstance(letto, dict) else {}
        return {}

    @staticmethod
    def copie_dichiarate_dal_writer(
        riepilogo: dict[str, Any], destinazione: Path
    ) -> tuple[dict[str, Path], list[str]]:
        """Return what the writer claims to have produced, and what it couldn't verify.

        The copy's filename is declared by whoever wrote it, not recomputed from
        the supplier's internal identifier: the writer builds
        `ORDINE_<name>_<pricelist>.xlsx` from the registry's display name, so
        recomputing the filename from the identifier for a learned adapter would
        look for a file that doesn't exist under that name, even though the
        correct copy sits right there.

        Also returns the rows the writer wrote without being able to confirm
        they were the right ones — otherwise those warnings would only reach the
        app's own console, never the user.

        A copy declared outside the compilation's own folder is rejected: the
        service only delivers what's inside that folder, and a path pointing
        elsewhere would put an unverified file into the delivered zip.
        """

        copie: dict[str, Path] = {}
        avvisi: list[str] = []
        voci = riepilogo.get("supplier_copies")
        for voce in voci if isinstance(voci, list) else []:
            if not isinstance(voce, dict):
                continue
            fornitore = str(voce.get("supplier") or "").strip().casefold()
            if not fornitore:
                continue
            # The readable label is declared by the writer, taken from the
            # registry: recomputing it here would mean keeping two name tables.
            etichetta = str(voce.get("supplier_name") or "").strip() or supplier_label(fornitore)
            dichiarata = voce.get("destination")
            if voce.get("skipped") or not isinstance(dichiarata, str) or not dichiarata.strip():
                # A supplier configured but not ordered: the writer lists it for
                # completeness, but there's no copy to deliver for it.
                continue
            percorso = Path(dichiarata.strip()).expanduser()
            if not percorso.is_absolute():
                percorso = destinazione / percorso
            percorso = percorso.resolve()
            if percorso.parent != destinazione.resolve():
                raise ValueError(
                    f"Il writer dichiara la copia di {etichetta} fuori dalla cartella della "
                    f"compilazione ({percorso}): non viene consegnata."
                )
            copie[fornitore] = percorso
            dette = 0
            for testo in voce.get("warnings") if isinstance(voce.get("warnings"), list) else []:
                pulita = str(testo or "").strip()
                if pulita:
                    avvisi.append(pulita)
                    dette += 1
            # Unverifiable rows with no message of their own happen when the
            # write rule declares no column to check, meaning none of that
            # supplier's rows were verified at all. Without a message here, the
            # page would just say "ready", indistinguishable from a successful
            # check.
            non_verificabili = number(voce.get("unverifiable_rows")) or 0
            verificate = number(voce.get("verified_rows")) or 0
            silenziose = int(non_verificabili) - dette
            if silenziose > 0:
                totale = int(verificate) + int(non_verificabili)
                avvisi.append(
                    f"Di {etichetta} una riga su {totale} è stata scritta senza poter verificare "
                    "che fosse la riga giusta: controllala prima di mandare l'ordine."
                    if silenziose == 1 else
                    f"Di {etichetta} {silenziose} righe su {totale} sono state scritte senza poter "
                    "verificare che fossero le righe giuste: controllale prima di mandare l'ordine."
                )
        return copie, avvisi

    def run_writer(
        self, plan_path: Path, destinazione: Path
    ) -> tuple[list[Path], list[tuple[str, Path, Path]], list[str]]:
        """Return the expected documents, the produced copies, and the writer's warnings.

        Warnings are a separate third return value, not a field nested inside
        each copy: a copy delivered with unverified rows isn't a half-successful
        compilation — it's a successful one that has something the caller must
        surface.
        """

        if self.writer_config is None:
            raise ValueError("Configurazione di scrittura non disponibile")
        config = load_json(self.writer_config, {})
        if not isinstance(config, dict):
            raise ValueError("Configurazione di scrittura non valida")
        supplier_files_value = config.get("supplier_files")
        if not isinstance(supplier_files_value, dict):
            raise ValueError("Manca l'elenco dei listini da compilare")
        supplier_files = {str(key).casefold(): value for key, value in supplier_files_value.items()}
        rules_value = config.get("supplier_write_rules")
        supplier_rules = (
            {str(key).casefold(): value for key, value in rules_value.items()}
            if isinstance(rules_value, dict) else {}
        )
        plan = load_json(plan_path, {"orders": []})
        ordered_suppliers = {
            str(item.get("supplier") or "").casefold()
            for item in plan.get("orders") or []
            if isinstance(item, dict) and item.get("supplier")
        }
        # Whether a supplier compiles in place is declared by its write rule;
        # everyone else goes through the Node writer.
        in_posizione = {
            supplier for supplier in ordered_suppliers
            if procedura_di_scrittura(supplier_rules.get(supplier)) == PATCH_IN_POSIZIONE
        }
        ordered_workbooks = ordered_suppliers - in_posizione
        missing_sources = [supplier_label(supplier) for supplier in sorted(ordered_workbooks) if not supplier_files.get(supplier)]
        if missing_sources:
            raise ValueError("Manca il listino configurato per: " + ", ".join(missing_sources))
        source_paths: dict[str, Path] = {}
        for supplier in ordered_workbooks:
            raw_source = Path(str(supplier_files[supplier])).expanduser()
            if not raw_source.is_absolute():
                raw_source = self.writer_config.parent / raw_source
            source_paths[supplier] = raw_source.resolve()
        expected = [plan_path]
        # Supplier, source price list and produced copy travel together: all
        # three are needed later, for the readable rename and for the audit's
        # `origine` field, and reconstructing them from filenames alone would
        # mean guessing which price list a copy came from.
        prodotte: list[tuple[str, Path, Path]] = []
        avvisi: list[str] = []
        # Node is only needed for `.xlsx` copies; a supplier compiled in place
        # is patched here in Python instead. An order made entirely of in-place
        # suppliers shouldn't require a writer it doesn't need.
        if ordered_workbooks:
            node_value = config.get("node_executable")
            script_value = config.get("writer_script")
            if not node_value or not script_value:
                raise ValueError("Configurazione di scrittura incompleta")
            node = Path(str(node_value)).resolve()
            script = Path(str(script_value)).resolve()
            working = Path(config.get("working_directory") or script.parent).resolve()
            missing = [str(path) for path in (node, script, working, *source_paths.values()) if not path.exists()]
            if missing:
                raise ValueError("Configurazione di scrittura incompleta; percorsi assenti: " + ", ".join(missing))
            command = [
                str(node),
                str(script),
                "--plan", str(plan_path),
                "--config", str(self.writer_config),
                # The writer writes directly into the compilation's own folder,
                # not into a shared `outputs` folder: otherwise there would be a
                # window between writing and moving where the previous
                # compilation's files could already be overwritten.
                "--output-dir", str(destinazione),
            ]
            writer_environment = os.environ.copy()
            result = subprocess.run(
                command,
                cwd=working,
                env=writer_environment,
                capture_output=True,
                text=True,
                # Node writes UTF-8. Without declaring it, Python on Windows
                # would decode with the console's codepage instead, garbling the
                # writer's accented characters.
                encoding="utf-8",
                errors="replace",
                timeout=180,
                check=False,
            )
            if result.returncode != 0:
                raise ValueError(self.spiegazione_del_writer(result.stderr, result.stdout))
            riepilogo = self.riepilogo_del_writer(result.stdout)
            dichiarate, avvisi = self.copie_dichiarate_dal_writer(riepilogo, destinazione)
            for supplier in sorted(ordered_workbooks):
                source = source_paths[supplier]
                if riepilogo:
                    destination = dichiarate.get(supplier)
                    if destination is None:
                        # The writer reported a summary but didn't name this
                        # supplier: any file matching the expected name pattern
                        # would belong to someone else or a previous run.
                        raise ValueError(
                            f"Il writer non dichiara nessuna copia per {supplier_label(supplier)}, "
                            "che il piano ordina."
                        )
                else:
                    # A writer that doesn't emit a summary: falls back to the
                    # filename this service has always computed. Not knowing
                    # isn't a failure, but it's the only case where the filename
                    # is guessed rather than declared.
                    destination = destinazione / f"ORDINE_{supplier.upper()}_{source.stem}.xlsx"
                if not destination.is_file():
                    raise ValueError(f"Il writer non ha creato la copia prevista per {supplier_label(supplier)}")
                expected.append(destination)
                prodotte.append((supplier, source, destination))
        for supplier in sorted(in_posizione):
            fornitore, sorgente, copia = self.compila_in_posizione(supplier, plan, config, destinazione)
            expected.append(copia)
            prodotte.append((fornitore, sorgente, copia))
        return sorted(expected, key=lambda path: path.name.casefold()), prodotte, avvisi

    def compila_in_posizione(
        self, supplier: str, plan: dict[str, Any], config: dict[str, Any], destinazione: Path
    ) -> tuple[str, Path, Path]:
        """Send the supplier back its own document, with only the order column changed.

        Applies to any supplier that declares `patch_xls_in_posizione` in its
        write rule. This bypasses the Node writer, which imports and re-exports
        `.xlsx` and would rewrite the whole file from scratch — here only the
        order-column cells are patched inside a copy of the original `.xls`,
        everything else stays byte-identical.

        If the patch can't be applied, compilation for that supplier fails and
        says so, rather than falling back to a format it doesn't support.
        """

        from xls_writer import CompilazioneXlsError, compila_ordine  # import tardivo

        nome = supplier_label(supplier)
        regole = config.get("supplier_write_rules")
        regola = (regole or {}).get(supplier) if isinstance(regole, dict) else None
        if not isinstance(regola, dict):
            raise ValueError(
                f"{nome} è selezionato, ma manca la regola di scrittura verificata del suo listino."
            )
        sorgente = self.percorso_listino(config, supplier)
        impronta = str(regola.get("source_sha256") or "").strip().casefold()
        if impronta and sha256_file(sorgente) != impronta:
            raise ValueError(
                f"Il listino {nome} è cambiato dopo la verifica: non creo copie finché non "
                "viene ricontrollato."
            )
        righe: dict[int, int] = {}
        ean_attesi: dict[int, str] = {}
        for order in plan.get("orders") or []:
            if str(order.get("supplier") or "").casefold() != supplier:
                continue
            riga = number(order.get("supplier_source_row"))
            quantita = number(order.get("quantity"))
            if riga is None or int(riga) != riga or riga < 1 or quantita is None or int(quantita) != quantita or quantita < 1:
                raise ValueError(
                    f"Riga o quantità {nome} non valide: riga {order.get('supplier_source_row')!r}, "
                    f"quantità {order.get('quantity')!r}."
                )
            # Two plan rows for the same product are summed, same as the Node
            # writer does for other suppliers.
            righe[int(riga)] = righe.get(int(riga), 0) + int(quantita)
            ean_attesi[int(riga)] = str(order.get("supplier_ean") or "")
        if not righe:
            raise ValueError(f"Il piano dichiara {nome} ma non ha nessuna riga da compilare.")

        colonna_ean = self.colonna_ean_dichiarata(sorgente, regola)
        copia = destinazione / f"ORDINE_{supplier.upper()}_{sorgente.stem}{sorgente.suffix}"
        try:
            esito = compila_ordine(
                sorgente,
                copia,
                righe,
                colonna_ordine=str(regola.get("order_column") or "I"),
                foglio=str(regola.get("sheet") or "") or None,
                ean_attesi=ean_attesi,
                colonna_ean=colonna_ean,
                prima_riga=int(number(regola.get("data_start_row")) or 1),
            )
        except CompilazioneXlsError as exc:
            copia.unlink(missing_ok=True)
            raise ValueError(str(exc)) from exc
        print(f"[{supplier.upper()}] {json.dumps(esito, ensure_ascii=False)}")
        if not copia.is_file():
            raise ValueError(f"La copia del listino {nome} non è stata creata.")
        return supplier, sorgente, copia

    def percorso_listino(self, config: dict[str, Any], supplier: str) -> Path:
        files_value = config.get("supplier_files")
        raw = (files_value or {}).get(supplier) if isinstance(files_value, dict) else None
        if not isinstance(raw, str) or not raw.strip():
            raise ValueError(
                f"{supplier_label(supplier)} è selezionato, ma manca il suo listino configurato "
                "per la compilazione."
            )
        source = Path(raw.strip()).expanduser()
        if not source.is_absolute() and self.writer_config is not None:
            source = self.writer_config.parent / source
        source = source.resolve()
        if not source.is_file():
            raise ValueError(f"Il listino configurato per {supplier_label(supplier)} non è disponibile.")
        return source

    @staticmethod
    def colonna_ean_dichiarata(sorgente: Path, regola: dict[str, Any]) -> int | str | None:
        """Resolve which column carries the EAN, from the file's own header row.

        The registry declares the column by name (e.g. `codice_a_barre`), not by
        letter — some price lists use column names that are themselves valid
        Excel column references, so falling back to letter-based lookup could
        silently read the wrong column. Name-to-index resolution goes through the
        reader's own logic, in one place.
        """

        nome = regola.get("ean_column_name") or regola.get("ean_column")
        if nome in (None, ""):
            return None
        riga_intestazione = number(regola.get("header_row"))
        if riga_intestazione is None or riga_intestazione < 1:
            return nome if isinstance(nome, str) else None
        from prepare_manifest_sources import column_number
        from xls_reader import read_workbook

        fogli = read_workbook(sorgente)
        atteso = str(regola.get("sheet") or "")
        griglia = next((f.rows for f in fogli if f.name == atteso), fogli[0].rows if fogli else [])
        indice = int(riga_intestazione) - 1
        if not 0 <= indice < len(griglia):
            raise ValueError(
                f"Il listino NOCE non arriva alla riga {int(riga_intestazione)}, dove "
                "dovrebbero esserci le intestazioni."
            )
        intestazione = [valore for valore, _grassetto in griglia[indice]]
        return column_number(nome, intestazione)


def frase_di_un_errore(errore: dict[str, Any]) -> str:
    """Build the user-facing sentence for a failed check, naming the product.

    The bare error code isn't enough for anyone: the page shows `message`, and a
    generic "checks failed" in front of hundreds of rows doesn't say which one
    to look at.
    """

    codice = str(errore.get("code") or "")
    prodotto = str(errore.get("productName") or errore.get("productId") or "").strip()
    fornitore = str(errore.get("supplierName") or errore.get("supplierId") or "").strip()
    if codice == "FORNITORE_DA_SCEGLIERE":
        return (
            f"«{prodotto}»: non è stato scelto nessun fornitore, e qualcuno ce l'ha. "
            "Scegline uno, oppure metti la quantità a zero se non ti serve."
        )
    if codice == "OFFERTA_NON_VALIDA":
        if not fornitore:
            return (
                f"«{prodotto}»: non è stato scelto nessun fornitore. Scegline uno "
                "oppure metti la quantità a zero."
            )
        return (
            f"«{prodotto}»: l'offerta di {fornitore} non è utilizzabile (il listino non "
            "la riporta più disponibile). Scegli un altro fornitore per questo prodotto "
            "oppure metti la quantità a zero."
        )
    if codice == "CONFERMA_MANCANTE":
        coda = f" di {fornitore}" if fornitore else ""
        return f"«{prodotto}»: va confermata la corrispondenza con l'offerta{coda} prima di procedere."
    if codice == "SOGLIA_NON_CONFERMATA":
        return (
            f"{supplier_label(str(errore.get('supplier') or ''))}: l'ordine resta sotto la "
            f"soglia netta di {errore.get('threshold_net')} €. Serve la conferma esplicita."
        )
    if codice == "PREZZO_NON_VALIDO":
        coda = f" di {fornitore}" if fornitore else ""
        return f"«{prodotto}»: il prezzo dell'offerta{coda} non è utilizzabile per calcolare l'ordine."
    if codice == "RIGA_LISTINO_CONTESA":
        altri = str(errore.get("otherProducts") or "").strip()
        coda = f" e {altri}" if altri else ""
        return (
            f"«{prodotto}»{coda}: sono due ordini diversi sulla stessa riga "
            f"{errore.get('sourceRow')} del listino {fornitore}, e in quella cella ci va un "
            "numero solo. Succede fra un espositore e il collo da cui è ricavato. Tieni "
            "l'ordine che ti serve e metti l'altro a zero."
        )
    if codice == "PRODOTTO_SCONOSCIUTO":
        return f"«{prodotto}» non fa parte del confronto caricato: ricarica la pagina."
    if codice == "QUANTITA_NON_VALIDA":
        return f"«{prodotto}»: la quantità non è un numero intero di colli."
    if codice == "ID_PRODOTTO_NON_VALIDO":
        return "Una riga arriva senza identificativo o ripetuta: ricarica la pagina."
    if codice == "STATO_SOVRASCRITTO":
        if str(errore.get("origin") or "") == "svuotato":
            return (
                "Le tue scelte non sono state salvate, ed è giusto così: è cominciata una "
                "comparazione nuova, e parlavano del confronto di prima. Non c'è niente da "
                "recuperare — carica i documenti di questa settimana e ricomincia da lì."
            )
        if str(errore.get("origin") or "") == "compilazione":
            return (
                "Le tue scelte non sono state salvate: le ha già riscritte l'ultima "
                "compilazione, che salva lo stato prima di preparare i listini. Ricarica "
                "questa pagina — le scelte con cui hai compilato sono quelle salvate — e "
                "riprova da lì."
            )
        return (
            "Le tue scelte non sono state salvate: un'altra scheda del comparatore ha "
            "salvato dopo di te, e salvare adesso cancellerebbe quello che ha scritto. "
            "Ricarica questa pagina per vedere le scelte aggiornate, poi rifai le tue modifiche."
        )
    if codice == "ORDINE_VUOTO":
        return "Nessun prodotto ha una quantità maggiore di zero: non c'è niente da compilare."
    return f"Controllo non superato: {codice or 'senza codice'}."


def frase_degli_errori(errori: list[dict[str, Any]], *, massimo: int = 3) -> str:
    """Return the first few sentences plus a count of the rest: an unbounded list isn't readable."""

    frasi: list[str] = []
    for errore in errori or []:
        if not isinstance(errore, dict):
            continue
        frase_singola = frase_di_un_errore(errore)
        if frase_singola not in frasi:
            frasi.append(frase_singola)
    if not frasi:
        return "Controlli non superati."
    mostrate = frasi[:massimo]
    restanti = len(frasi) - len(mostrate)
    testo = " ".join(mostrate)
    if restanti > 0:
        testo = f"{testo} E altri {restanti} controlli non superati."
    return testo


class SnapshotError(ValueError):
    def __init__(self, errors: list[dict[str, Any]]) -> None:
        super().__init__(frase_degli_errori(errors))
        self.errors = errors


# ----------------------------------------------------------------------------
# The Settings page
# ----------------------------------------------------------------------------

# The only entries changeable from this page. `versione_prompt` and
# `versione_avversario` are deliberately excluded — they're app components
# chosen by measuring how many wrong `ALTA` outcomes they produce, not user
# preferences — and so are `max_tokens`, `temperature` and `base_url`, which are
# part of that same measurement or of the service address. Exposing them here
# would let the user invalidate that measurement without anything flagging it.
VOCI_IMPOSTAZIONI = ("model", "tetto_spesa_usd", "tetto_chiamate", "parallelismo", "timeout_secondi")

# OpenRouter's `GET /api/v1/models` responds even without a key and even with
# an expired one: a populated model menu proves nothing about the key. This
# warning travels with the list because the page has to state it, and because
# whoever reuses this route later finds it attached to the data.
AVVISO_ELENCO_PUBBLICO = (
    "L’elenco dei modelli è pubblico: si carica anche senza chiave e anche con una chiave "
    "scaduta. Che la chiave funzioni lo dice soltanto «Prova la connessione»."
)

# What to tell the user for each status `prova_connessione` can return. The
# tone only controls the warning's color; `ok` in the response is true only for
# `OK`, since that's the only case where key, model id and response format all
# worked together. Saying "key valid" on a 429 would be an inference this page
# deliberately doesn't make.
MESSAGGI_DELLA_PROVA: dict[str, tuple[str, str]] = {
    ai_client.STATO_OK: ("success", "La chiave funziona: il modello ha risposto nel formato previsto."),
    ai_client.STATO_SENZA_CHIAVE: ("danger", "Non c’è nessuna chiave da provare: incollala nel campo qui sopra."),
    "HTTP_401": ("danger", "Chiave non accettata (401): è sbagliata, scaduta o revocata."),
    "HTTP_403": ("danger", "Chiave rifiutata (403): esiste, ma non è abilitata a questo modello."),
    "HTTP_402": ("warning", "Il credito su OpenRouter è esaurito (402): la chiave non è stata respinta, ma finché non lo ricarichi ogni chiamata fallisce."),
    "HTTP_400": ("warning", "Richiesta rifiutata (400): quasi sempre l’identificativo del modello non esiste. La chiave non è stata respinta."),
    "HTTP_404": ("warning", "Questo modello non esiste (404): controlla l’identificativo. La chiave non è stata respinta."),
    "HTTP_429": ("warning", "Il servizio ha chiesto di rallentare (429): riprova fra un minuto. La chiave non è stata respinta."),
    "HTTP_5xx": ("warning", "OpenRouter ha risposto con un guasto suo (5xx): riprova fra un minuto. La chiave non è stata respinta."),
    ai_client.STATO_ERRORE_RETE: ("danger", "OpenRouter non è stato raggiunto: la prova non dice niente sulla chiave. Controlla la connessione."),
    ai_client.STATO_ERRORE_TRASPORTO: ("danger", "La chiamata si è rotta prima di partire: la prova non dice niente sulla chiave."),
    ai_client.STATO_TETTO_SPESA: ("danger", "Il tetto di spesa impostato non lascia partire nemmeno la chiamata di prova: alzalo e riprova."),
    ai_client.STATO_TETTO_CHIAMATE: ("danger", "Il tetto di chiamate impostato non lascia partire nemmeno la chiamata di prova: alzalo e riprova."),
}

# Here the call went through, the model answered, and the answer came back
# malformed. Not a key failure, but still a reason not to use that model: with
# a response like this the AI matching stage couldn't decide anything.
STATI_RISPOSTA_INUTILIZZABILE = frozenset({
    ai_client.STATO_TRONCATA,
    ai_client.STATO_CONTENUTO_VUOTO,
    ai_client.STATO_NON_E_JSON,
    ai_client.STATO_SCHEMA_NON_CONFORME,
    ai_client.STATO_RIGA_FUORI_SHORTLIST,
    ai_client.STATO_JSON_RICHIESTO,
})


def chiave_dal_corpo(payload: Any) -> str:
    """Return the key pasted into the request body, or an empty string.

    Kept in one function because the value returned here is also what the
    handler sets aside to scrub from error messages: if the two call sites read
    the body differently, a key missed by one read would also be missed by the
    scrubbing.
    """

    if not isinstance(payload, dict):
        return ""
    return str(payload.get("chiave") or "").strip()


class ServizioImpostazioni:
    """API key, model and usage caps: everything the Settings page touches.

    Kept outside `ReviewStore` on purpose. `GET /api/review` is the largest and
    most frequently read response in the app, and the AI configuration must
    never be able to end up in it; keeping it in a separate object makes that
    true by construction instead of by convention.

    `leggi_modelli` and `crea_client` are the two points where the network is
    touched, and they're injectable for the same reason `ClientAI(trasporto=…)`
    is: no test in this suite should ever call OpenRouter.
    """

    def __init__(
        self,
        *,
        percorso_secrets: Path | None = None,
        percorso_impostazioni: Path | None = None,
        leggi_modelli: Any = None,
        crea_client: Any = None,
    ) -> None:
        self.percorso_secrets = Path(percorso_secrets) if percorso_secrets else ai_client.PERCORSO_SECRETS
        self.percorso_impostazioni = (
            Path(percorso_impostazioni) if percorso_impostazioni else ai_client.PERCORSO_IMPOSTAZIONI
        )
        self._leggi_modelli = leggi_modelli or ai_client.elenco_modelli
        self._crea_client = crea_client or (
            lambda configurazione, chiave: ai_client.ClientAI(configurazione, chiave=chiave)
        )

    # -- lettura -------------------------------------------------------------

    def configurazione(self) -> dict[str, Any]:
        return ai_client.carica_configurazione(self.percorso_impostazioni)

    def chiave_salvata(self) -> str:
        """Return the key the service would use if the request body carries none.

        Never sent to the browser: it's only used by `_chiave_in_volo`, i.e. to
        know what to scrub if something fails while using it.
        """

        return ai_client.leggi_chiave(self.percorso_secrets) or ""

    def stato(self) -> dict[str, Any]:
        """Return the app's current configuration. Never includes the key itself.

        `stato_chiave` returns only presence, origin, and the last four
        characters — the only thing about a key that's allowed to reach the
        browser.
        """

        configurazione = self.configurazione()
        return {
            "ok": True,
            "chiave": ai_client.stato_chiave(self.percorso_secrets),
            "impostazioni": {nome: configurazione[nome] for nome in VOCI_IMPOSTAZIONI},
            "predefinite": {nome: ai_client.CONFIGURAZIONE_PREDEFINITA[nome] for nome in VOCI_IMPOSTAZIONI},
            "percorsoChiave": str(self.percorso_secrets),
        }

    def modelli(self) -> dict[str, Any]:
        """Return the model list for the dropdown. A failure here isn't a page failure.

        The dropdown is a convenience: the model id can always be typed by hand,
        which is also the only way to use a model released yesterday. So a
        network failure returns `ok: false` with a reason, not a 500 that breaks
        the page.
        """

        configurazione = self.configurazione()
        try:
            elenco = self._leggi_modelli(configurazione["base_url"])
        except ai_client.GUASTI_DI_RETE as errore:
            return {
                "ok": False,
                "modelli": [],
                "avviso": AVVISO_ELENCO_PUBBLICO,
                "messaggio": (
                    f"Elenco dei modelli non disponibile ({type(errore).__name__}): "
                    "scrivi l’identificativo del modello a mano."
                ),
            }
        return {"ok": True, "modelli": elenco, "avviso": AVVISO_ELENCO_PUBBLICO}

    # -- scrittura -----------------------------------------------------------

    def salva_chiave(self, payload: Any) -> dict[str, Any]:
        """Write the key and reply with the status only; the value itself is never echoed back."""

        ai_client.salva_chiave(chiave_dal_corpo(payload), self.percorso_secrets)
        return {
            "ok": True,
            "messaggio": "Chiave salvata sul computer.",
            "chiave": ai_client.stato_chiave(self.percorso_secrets),
        }

    def salva(self, payload: Any) -> dict[str, Any]:
        """Save the exposed settings. A rejected value is an error, not silently coerced.

        `salva_impostazioni` raises `ValueError` with an already user-facing
        message, and that message reaches the user unchanged — rewriting it here
        would mean keeping two texts in sync.
        """

        valori = payload.get("impostazioni") if isinstance(payload, dict) else None
        if not isinstance(valori, dict) or not valori:
            raise ValueError("Non è arrivata nessuna impostazione da salvare.")
        fuori = sorted(set(valori) - set(VOCI_IMPOSTAZIONI))
        if fuori:
            raise ValueError(
                "Queste voci non si cambiano dalla pagina Impostazioni: " + ", ".join(fuori)
                + ". Da qui si cambiano soltanto: " + ", ".join(VOCI_IMPOSTAZIONI) + "."
            )
        configurazione = ai_client.salva_impostazioni(valori, self.percorso_impostazioni)
        return {
            "ok": True,
            "messaggio": "Impostazioni salvate.",
            "impostazioni": {nome: configurazione[nome] for nome in VOCI_IMPOSTAZIONI},
        }

    # -- la prova ------------------------------------------------------------

    def prova(self, payload: Any) -> dict[str, Any]:
        """Make a real call to the model, using the key the user is currently testing.

        The key arrives in the request body and is never saved: this lets the
        user test before saving, without a bad key overwriting a working one. If
        the body carries none, the already-configured key is tested instead.
        """

        chiave = chiave_dal_corpo(payload)
        if not chiave:
            # Read from this service's own path rather than letting the client
            # fall back on its own: `ClientAI(chiave=None)` would look at the
            # module's default path, which matches production but not
            # necessarily every environment. Otherwise the test could report on
            # a different key than the one the page shows above.
            chiave = ai_client.leggi_chiave(self.percorso_secrets) or ""

        configurazione = self.configurazione()
        modello = str((payload or {}).get("model") or "").strip() if isinstance(payload, dict) else ""
        if modello:
            # Test the model id the user just typed, before saving it: a model
            # that doesn't exist responds 400 and is caught here.
            configurazione = {**configurazione, "model": modello}

        # `chiave=""` tells the client the key is absent, producing
        # `SENZA_CHIAVE` instead of an unexpected lookup.
        client = self._crea_client(configurazione, chiave)
        esito = client.prova_connessione()

        if esito.stato in MESSAGGI_DELLA_PROVA:
            tono, messaggio = MESSAGGI_DELLA_PROVA[esito.stato]
        elif esito.stato in STATI_RISPOSTA_INUTILIZZABILE:
            tono = "warning"
            messaggio = (
                "Il modello ha risposto, ma la risposta non è utilizzabile "
                f"({esito.stato}). La chiave non è stata respinta: è questo modello a non "
                "rispettare il formato richiesto."
            )
        else:
            tono = "danger"
            messaggio = f"Prova non riuscita ({esito.stato}). Il dettaglio qui sotto dice che cosa è successo."

        # `ClientAI` already scrubs the key from every field of its result. It's
        # scrubbed again here because this response is what reaches the browser,
        # and the defense that matters is the one at the exit point, not upstream.
        return {
            "ok": esito.stato == ai_client.STATO_OK,
            "tono": tono,
            "stato": senza_la_chiave(esito.stato, chiave),
            "messaggio": messaggio,
            "dettaglio": senza_la_chiave(esito.dettaglio, chiave),
            "modello": senza_la_chiave(esito.modello, chiave),
            "costoUsd": esito.costo_usd,
            "durataS": esito.durata_s,
        }


class AppHandler(BaseHTTPRequestHandler):
    server_version = "ComparaOrdini/0.1"

    # The key currently in flight for this request, and only for its duration.
    # Used by `senza_segreti`: a 500's body echoes `str(exc)` to the browser (and
    # the traceback also goes to `stderr`), and an exception raised while testing
    # a key could carry it along — from the HTTP library, from urllib, from a
    # `KeyError` on a headers dict. It's either the freshly pasted key or the
    # saved one, since the saved key is what's used when the body carries none,
    # which is most of the time. It's an instance attribute, not a module-level
    # one, because each connection gets its own instance and thread.
    _chiave_in_volo: str = ""

    @property
    def store(self) -> ReviewStore:
        return self.server.store  # type: ignore[attr-defined]

    @property
    def impostazioni(self) -> ServizioImpostazioni:
        servizio = getattr(self.server, "impostazioni", None)
        if servizio is None:
            servizio = ServizioImpostazioni()
            self.server.impostazioni = servizio  # type: ignore[attr-defined]
        return servizio

    def senza_segreti(self, testo: Any) -> str:
        """Return text about to go to the browser, with the key scrubbed out."""

        return senza_la_chiave(str(testo), self._chiave_in_volo)

    def log_message(self, format_string: str, *args: Any) -> None:
        # The request line lands here on every call — the reason the key only
        # ever travels in a POST body, never in a query string.
        print(f"{self.address_string()} - {self.senza_segreti(format_string % args)}")

    def end_headers(self) -> None:
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'none'; frame-ancestors 'none'")
        super().end_headers()

    def json_response(self, value: Any, status: int = HTTPStatus.OK) -> None:
        data = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _stampa_il_guasto(self) -> None:
        """Print an unexpected failure's traceback to the app's own console.

        The technical error text also reaches the page (in `DETTAGLIO_TECNICO`),
        but without a full traceback anywhere there was no way to tell which line
        of a large codebase a failure came from — only the exception message,
        with no location. This prints the traceback where whoever is watching the
        app can actually see it.

        Doesn't change the HTTP response and the user never sees it: it goes to
        `stderr`, the launcher's console window.

        Scrubbed through `senza_segreti` like everything else: a key can show up
        inside an `urllib` traceback, which is exactly why `_chiave_in_volo`
        exists.
        """

        tipo, valore, _traccia = sys.exc_info()
        if tipo is None:
            # A 500 raised manually, with no exception actually in flight:
            # `format_exc()` would just print the useless "NoneType: None".
            return
        if isinstance(valore, (BrokenPipeError, ConnectionResetError, ConnectionAbortedError)):
            # Not an app failure: the browser disconnected mid-response — a
            # cancelled download, a tab closed on a large file. A traceback for
            # every one of these would bury the console in noise exactly when it
            # needs to stay readable for a real failure.
            return
        dove = self.senza_segreti(f"{self.command} {self.path}")
        print(
            f"[GUASTO] {dove}\n{self.senza_segreti(traceback.format_exc()).rstrip()}",
            file=sys.stderr,
            # Without this, the text stays buffered: if the process dies right
            # after, the traceback dies with it — exactly when it would matter
            # most.
            flush=True,
        )

    def error_response(self, status: int, message: str, errors: list[dict[str, Any]] | None = None) -> None:
        # Every error message passes through here, including a 500's
        # `str(exc)`: the one place where scrubbing once covers all of them.
        pulito = self.senza_segreti(message)
        voci = list(errors or [])
        if status == HTTPStatus.INTERNAL_SERVER_ERROR:
            self._stampa_il_guasto()
            # An unexpected failure would otherwise reach the page as a raw
            # exception message: English, local file paths, no indication of what
            # to do. The technical message isn't discarded — it's useful for
            # diagnosis — but it goes alongside the response, not in place of it.
            voci = [{"code": "DETTAGLIO_TECNICO", "message": pulito}, *voci]
            pulito = (
                "Ho avuto un problema con questa operazione. Riprova; se succede "
                "ancora, chiudi e riapri il comparatore."
            )
        self.json_response({"ok": False, "message": pulito, "errors": voci}, status)

    def read_json(self) -> Any:
        raw_length = self.headers.get("Content-Length")
        if not raw_length or not raw_length.isdigit():
            raise ValueError("Content-Length mancante o non valido")
        length = int(raw_length)
        if length <= 0 or length > MAX_JSON_BYTES:
            raise ValueError("Dimensione richiesta non valida")
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8"))

    def _origini_della_pagina(self) -> set[str]:
        """Return the only origins the comparator's own page can arrive from.

        The port isn't hardcoded: the launcher falls back to 8766, 8767 and so on
        when the default is already taken, and a check with a fixed port would
        reject the app's own page. `self.server.server_address[1]` is always the
        port this instance is actually listening on.
        """

        porta = self.server.server_address[1]
        return {f"http://{macchina}:{porta}" for macchina in ("127.0.0.1", "localhost", "[::1]")}

    def richiesta_dalla_nostra_pagina(self) -> bool:
        """Return whether this state-changing request actually came from the comparator's page.

        The POST method alone is not a defense: a POST with
        `Content-Type: text/plain` is a browser "simple request" — it's sent
        without a CORS preflight, and its effect happens even if the sender never
        reads the response. Without this check, any page open in the same
        browser could upload a price list, start the pipeline (spending
        OpenRouter credit), compile orders, delete uploads, overwrite the API key
        or shut the app down, while the comparator was running.

        This checks the headers that are present rather than requiring them: the
        launcher calls `/api/spegni` via urllib, which never sends `Origin`, and
        an older browser may not send `Sec-Fetch-Site`. Requests from our own
        page, served by this same instance, always send both and are always
        same-origin.
        """

        sito = str(self.headers.get("Sec-Fetch-Site") or "").strip()
        if sito and sito != "same-origin":
            return False
        origine = str(self.headers.get("Origin") or "").strip()
        if origine and origine not in self._origini_della_pagina():
            return False
        return True

    def rifiuta_l_origine(self) -> None:
        self.error_response(
            HTTPStatus.FORBIDDEN,
            "Questa richiesta non arriva dalla pagina del comparatore: è stata rifiutata.",
        )

    def _host_della_pagina(self) -> set[str]:
        """Return the only `Host` header values the comparator's page can send.

        Both with and without this instance's port: some clients send
        `Host: 127.0.0.1` with no port, and the check has to accept that too. The
        port always comes from `self.server.server_address[1]`, never hardcoded.
        """

        porta = self.server.server_address[1]
        macchine = ("127.0.0.1", "localhost", "[::1]")
        return set(macchine) | {f"{macchina}:{porta}" for macchina in macchine}

    def richiesta_con_host_valido(self) -> bool:
        """Return whether the `Host` header points at this instance, or is absent entirely.

        A malicious site can let its own DNS record expire and re-point its
        domain to 127.0.0.1 (DNS rebinding): from then on the browser sends
        requests to that domain on this app's port, which the browser treats as
        same-origin — so the malicious page can read the response, and on a GET
        `Origin` is often absent anyway, so the CSRF check above isn't enough on
        its own. The `Host` header, however, always carries the attacker's
        domain, never the rebound IP, so checking it closes this gap.

        HTTP/1.0 may omit `Host` entirely, and some clients omit the port: both
        are accepted, or the real page would stop working too.
        """

        host = str(self.headers.get("Host") or "").strip()
        if not host:
            return True
        return host in self._host_della_pagina()

    def rifiuta_l_host(self) -> None:
        self.error_response(
            HTTPStatus.FORBIDDEN,
            "Questo indirizzo non è quello del comparatore: la richiesta è stata rifiutata.",
        )

    def serve_file(self, path: Path, download: bool = False) -> None:
        if not path.is_file():
            self.error_response(HTTPStatus.NOT_FOUND, "Documento non trovato")
            return
        content = path.read_bytes()
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Length", str(len(content)))
        if download:
            # A non-ASCII filename (accented characters, an em dash) can't be
            # encoded in latin-1, which is what `BaseHTTPRequestHandler`'s
            # headers use — a plain `filename=` would raise a
            # `UnicodeEncodeError` after the body has already been promised. The
            # RFC 5987 fallback form is built by `consegna`, in one place shared
            # by every download route.
            self.send_header("Content-Disposition", consegna.intestazione_allegato(path.name))
        self.end_headers()
        self.wfile.write(content)

    def scarica_le_conferme(self) -> None:
        """`GET /api/conferme/esporta`: given confirmations, as a file to save.

        Doesn't go through `json_response`, since this response isn't for the
        page — it's a file the user saves. The header is built by `consegna`,
        like every other download, because the filename carries a date that
        doesn't encode in latin-1.
        """

        conferme, uguaglianze = self.store.esporta_le_conferme()
        momento = datetime.now().astimezone()
        nome = f"Conferme — {consegna.data_leggibile(momento)}.json"
        corpo = json.dumps(
            {
                "esportate_il": momento.isoformat(),
                "quante": len(conferme),
                "conferme": conferme,
                # Equivalences too, not an extra: they live in the same
                # `conferme.db`, are permanent memory just like confirmations, and
                # this file is presented on the page as the exportable copy of
                # that memory. Leaving them out would export only half of it.
                "quante_uguaglianze": len(uguaglianze),
                "uguaglianze": uguaglianze,
            },
            ensure_ascii=False,
            indent=2,
        ).encode("utf-8")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(corpo)))
        self.send_header("Content-Disposition", consegna.intestazione_allegato(nome))
        self.end_headers()
        self.wfile.write(corpo)

    def servi_consegna(self, resto: str) -> None:
        """Serve `/ordini/<cartella>/<nome>` and `/ordini/<cartella>/zip`.

        The path arrives already decoded (`unquote` in `do_GET`), so
        `..%2f..%2fsecrets.json` is `../../secrets.json` by this point: segments
        are counted after decoding, and each one passes through `consegna`'s own
        checks. Error messages never echo back the requested path, so an attacker
        probing this route learns nothing from the response.
        """

        pezzi = resto.split("/")
        if len(pezzi) != 2 or not all(pezzi):
            self.error_response(HTTPStatus.NOT_FOUND, "Percorso non trovato")
            return
        nome_cartella, nome = pezzi
        cartella = consegna.cartella_sicura(self.store.orders_dir, nome_cartella)
        if cartella is None:
            self.error_response(HTTPStatus.NOT_FOUND, "Compilazione non trovata")
            return
        if nome == "zip":
            self.servi_zip(cartella)
            return
        percorso = consegna.file_sicuro(cartella, nome)
        if percorso is None:
            self.error_response(HTTPStatus.NOT_FOUND, "Documento non trovato")
            return
        self.serve_file(percorso, download=True)

    def servi_zip(self, cartella: Path) -> None:
        """Build the delivery zip in memory, never written to disk.

        Contains the compiled price lists and, when present, the list of
        products no supplier carries: whoever downloads the zip is preparing the
        week's order, and that list is as much part of that work as an order is.
        """

        voce = consegna.voce(cartella)
        # The file types to deliver live in `consegna`, in one place: two lists
        # to keep in sync is a list that eventually falls out of sync.
        nomi = [item["nome"] for item in voce["file"]
                if item["tipo"] in consegna.TIPI_DA_CONSEGNARE]
        if not nomi:
            self.error_response(HTTPStatus.NOT_FOUND, "In questa compilazione non ci sono listini da scaricare")
            return
        try:
            contenuto = consegna.zip_in_memoria(cartella, nomi)
        except ValueError:
            # The audit names a price list no longer on disk: the folder was
            # edited by hand. This is a 404, not an app failure.
            self.error_response(HTTPStatus.NOT_FOUND, "In questa compilazione non ci sono listini da scaricare")
            return
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/zip")
        self.send_header("Content-Length", str(len(contenuto)))
        self.send_header("Content-Disposition", consegna.intestazione_allegato(voce["zipNome"]))
        self.end_headers()
        self.wfile.write(contenuto)

    def do_GET(self) -> None:  # noqa: N802
        if not self.richiesta_con_host_valido():
            self.rifiuta_l_host()
            return
        parsed = urlparse(self.path)
        route = unquote(parsed.path)
        try:
            if route == "/api/review":
                self.json_response(self.store.review())
                return
            if route == "/api/products/search":
                parameters = parse_qs(parsed.query)
                query = str((parameters.get("q") or [""])[0]).strip()
                limit = int(number((parameters.get("limit") or [20])[0]) or 20)
                self.json_response(self.store.search_products(query, limit))
                return
            if route == "/api/listino":
                parameters = parse_qs(parsed.query)
                primo = lambda nome: str((parameters.get(nome) or [""])[0]).strip()  # noqa: E731
                self.json_response(self.store.sfoglia_listino(
                    primo("fornitore"),
                    query=primo("q"),
                    da=int(number(primo("da")) or 0),
                    quante=int(number(primo("quante")) or 50),
                    riga=primo("riga") or None,
                ))
                return
            if route == "/api/matches/uguaglianze":
                parameters = parse_qs(parsed.query)
                self.json_response(self.store.elenco_delle_uguaglianze(
                    str((parameters.get("q") or [""])[0])
                ))
                return
            if route == "/api/conferme/esporta":
                self.scarica_le_conferme()
                return
            if route == "/api/schemas/order-column":
                parameters = parse_qs(parsed.query)
                self.json_response(self.store.colonne_d_ordine(
                    str((parameters.get("fornitore") or [""])[0]).strip()
                ))
                return
            if route == "/api/history/pending":
                self.json_response(self.store.history_pending())
                return
            if route == "/api/pipeline/stato":
                self.json_response(self.store.stato_pipeline())
                return
            if route == "/api/schemas/pending":
                self.json_response(self.store.schemi_pendenti())
                return
            if route == "/api/schemas/documento":
                # The document name isn't a path: it's an entry from the list
                # this service just returned to the page, and
                # `colonne_del_documento` re-validates it against the folder.
                parameters = parse_qs(parsed.query)
                nome = str((parameters.get("nome") or [""])[0]).strip()
                self.json_response(self.store.colonne_del_documento(nome))
                return
            if route == "/api/schemas/columns":
                self.json_response(self.store.colonne_dei_documenti())
                return
            if route == "/api/impostazioni":
                self.json_response(self.impostazioni.stato())
                return
            if route == "/api/impostazioni/modelli":
                # This list is public and carries no secret, which is why it's
                # the only settings route allowed to be a GET.
                self.json_response(self.impostazioni.modelli())
                return
            if route == "/api/ordini":
                self.json_response({
                    "ok": True,
                    "compilazioni": consegna.elenco(self.store.orders_dir, etichetta_fornitore=supplier_label),
                })
                return
            if route == "/api/health":
                # `firmaDelCodice` identifies which build is answering, not just
                # that something is. Without it, the launcher could find a
                # healthy server and reuse it even after the source changed
                # underneath, effectively restarting into a stale build.
                self.json_response({
                    "ok": True,
                    "status": "ready",
                    "firmaDelCodice": versione_del_codice.firma(),
                    # The published version's date, currently running: shown at
                    # the top of the page, since a failed sync to the published
                    # source would otherwise go unnoticed.
                    "versionePubblicata": versione_del_codice.pubblicata(),
                })
                return
            if route.startswith(consegna.PREFISSO_URL + "/"):
                self.servi_consegna(route.removeprefix(consegna.PREFISSO_URL + "/"))
                return
            if route.startswith("/outputs/"):
                # Same defense as the sibling `/ordini/<cartella>/<nome>` route:
                # `.name` alone blocks path traversal, but not a symlink placed
                # inside `outputs/` — that's stopped by `file_sicuro`'s own
                # containment check after `resolve()`.
                percorso = consegna.file_sicuro(
                    self.store.output_dir, Path(route.removeprefix("/outputs/")).name
                )
                if percorso is None:
                    self.error_response(HTTPStatus.NOT_FOUND, "Documento non trovato")
                    return
                self.serve_file(percorso, download=True)
                return
            if route in {"/", "/index.html"}:
                self.serve_file(STATIC_DIR / "index.html")
                return
            if route.startswith("/static/"):
                relative = Path(route.removeprefix("/static/"))
                if relative.name != str(relative) or relative.name not in {"index.html", "styles.css", "app.js"}:
                    self.error_response(HTTPStatus.NOT_FOUND, "Risorsa non trovata")
                    return
                self.serve_file(STATIC_DIR / relative.name)
                return
            self.error_response(HTTPStatus.NOT_FOUND, "Percorso non trovato")
        except ValueError as exc:
            self.error_response(HTTPStatus.BAD_REQUEST, str(exc))
        except Exception as exc:
            self.error_response(HTTPStatus.INTERNAL_SERVER_ERROR, str(exc))

    def do_PUT(self) -> None:  # noqa: N802
        # Kept symmetric with `do_POST`. `PUT` isn't a "simple" method, so the
        # browser already runs a CORS preflight and `/api/state` is already out
        # of another site's reach — but a check applied to one state-changing
        # method and not the other invites someone to eventually read it as
        # intentional.
        if not self.richiesta_con_host_valido():
            self.rifiuta_l_host()
            return
        if not self.richiesta_dalla_nostra_pagina():
            self.rifiuta_l_origine()
            return
        if urlparse(self.path).path != "/api/state":
            self.error_response(HTTPStatus.NOT_FOUND, "Percorso non trovato")
            return
        try:
            self.json_response(self.store.save_state(self.read_json()))
        except SnapshotError as exc:
            self.error_response(HTTPStatus.UNPROCESSABLE_ENTITY, str(exc), exc.errors)
        except (ValueError, json.JSONDecodeError) as exc:
            self.error_response(HTTPStatus.BAD_REQUEST, str(exc))
        except Exception as exc:
            self.error_response(HTTPStatus.INTERNAL_SERVER_ERROR, str(exc))

    def _spegni(self) -> dict[str, Any]:
        """Stop the service, but never while the pipeline is running.

        A run killed mid-way leaves an orphaned dated folder and AI work already
        paid for that would have to be redone; whoever asks to shut down needs to
        know something is in progress and decide explicitly. The caller here is
        the launcher, which reuses the running server instead of restarting it in
        that case.

        The actual shutdown starts after the response: `shutdown()` waits for the
        serve loop to stop, and calling it from inside a request handler would
        deadlock the service against itself.
        """

        stato = self.store.stato_pipeline()
        if str(stato.get("stato") or "") in STATI_IN_CORSO:
            return {
                "ok": False,
                "spento": False,
                "motivo": "RUN_IN_CORSO",
                "messaggio": (
                    "C'è un confronto in corso: il programma non si spegne adesso. "
                    "Aspetta che finisca e riprova."
                ),
            }

        # The confirmations database is released before stopping: SQLite holds
        # it open for the connection's lifetime, and on Windows an open file
        # can't be renamed or deleted. Whoever shuts down to restart into new
        # code (`/api/health` reporting a different build signature) needs the
        # folder to be free.
        self.store.chiudi()
        threading.Thread(target=self.server.shutdown, daemon=True).start()
        return {"ok": True, "spento": True}

    def do_POST(self) -> None:  # noqa: N802
        # Before reading the body: none of the routes below may be triggered
        # by a page other than this app's own.
        if not self.richiesta_con_host_valido():
            self.rifiuta_l_host()
            return
        if not self.richiesta_dalla_nostra_pagina():
            self.rifiuta_l_origine()
            return
        route = urlparse(self.path).path
        try:
            payload = self.read_json()
            if route.startswith("/api/impostazioni"):
                # From here on the key may be in the body: it's set aside
                # before it's touched, so any exception between this line and
                # the response is already scrubbed.
                #
                # When the body carries no key, the one about to be used is the
                # saved one — the common case: pressing "test connection"
                # without pasting anything sends a body with no `chiave`, and the
                # service falls back to the saved file. If this only checked the
                # body, `_chiave_in_volo` would stay empty in that case, and a
                # key sourced from an environment variable wouldn't necessarily
                # match the `sk-...` pattern the alternative safeguard relies on.
                self._chiave_in_volo = chiave_dal_corpo(payload) or self.impostazioni.chiave_salvata()
            if route == "/api/impostazioni":
                self.json_response(self.impostazioni.salva(payload))
            elif route == "/api/impostazioni/chiave":
                self.json_response(self.impostazioni.salva_chiave(payload))
            elif route == "/api/impostazioni/prova":
                self.json_response(self.impostazioni.prova(payload))
            elif route == "/api/upload":
                self.json_response(self.store.upload(payload), HTTPStatus.CREATED)
            elif route == "/api/products/add":
                self.json_response(self.store.add_manual_product(payload), HTTPStatus.CREATED)
            elif route == "/api/history/answer":
                self.json_response(self.store.answer_history_order(payload))
            elif route == "/api/matches/answer":
                self.json_response(self.store.answer_rejected_candidate(payload))
            elif route == "/api/matches/abbina":
                self.json_response(self.store.abbina_riga_di_listino(payload))
            elif route == "/api/matches/rifiuta":
                self.json_response(self.store.rifiuta_l_abbinamento(payload))
            elif route == "/api/matches/uguaglianze/togli":
                self.json_response(self.store.togli_uguaglianza(payload))
            elif route == "/api/ordini/elimina":
                self.json_response(self.store.delete_compilation(payload))
            elif route == "/api/uploads/elimina":
                self.json_response(self.store.delete_upload(payload))
            elif route == "/api/comparazione/nuova":
                self.json_response(self.store.nuova_comparazione(payload))
            elif route == "/api/suppliers/discount":
                self.json_response(self.store.set_supplier_discount(payload))
            elif route == "/api/suppliers/move-preview":
                # Read-only: it's a quote, applies nothing.
                self.json_response(self.store.move_preview(payload))
            elif route == "/api/pipeline/avvia":
                # Doesn't wait for the pipeline: returns immediately with the
                # initial status, and the page polls until it finishes.
                self.json_response(self.store.avvia_pipeline(), HTTPStatus.ACCEPTED)
            elif route == "/api/schemas/validate":
                self.json_response(self.store.valida_schemi(payload))
            elif route == "/api/schemas/order-column":
                self.json_response(self.store.cambia_colonna_d_ordine(payload))
            elif route == "/api/schemas/confirm":
                risposta = self.store.conferma_schemi(payload)
                self.json_response(risposta, HTTPStatus.ACCEPTED)
            elif route == "/api/schemas/documento/prova":
                self.json_response(self.store.prova_colonne(payload))
            elif route == "/api/schemas/documento/salva":
                self.json_response(self.store.salva_colonne(payload))
            elif route == "/api/compile":
                self.json_response(self.store.compile(payload))
            elif route == "/api/spegni":
                # Requested by the launcher when the source has changed; it
                # reopens the app right after. POST rather than GET so an
                # `<img src=...>` can't trigger a shutdown — but POST alone isn't
                # the defense: the origin/host guard at the top of `do_POST` is
                # what actually stops requests from other pages.
                self.json_response(self._spegni())
            else:
                self.error_response(HTTPStatus.NOT_FOUND, "Percorso non trovato")
        except SnapshotError as exc:
            self.error_response(HTTPStatus.UNPROCESSABLE_ENTITY, str(exc), exc.errors)
        except LavoroGiaInCorso as exc:
            self.error_response(HTTPStatus.CONFLICT, str(exc))
        except (ValueError, json.JSONDecodeError) as exc:
            self.error_response(HTTPStatus.BAD_REQUEST, str(exc))
        except Exception as exc:
            self.error_response(HTTPStatus.INTERNAL_SERVER_ERROR, str(exc))
        finally:
            # A kept-alive connection serves multiple requests through the same
            # instance: the previous request's key must not stay available to
            # the next one.
            self._chiave_in_volo = ""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review", type=Path, required=True, help="review_data.json generato dalla pipeline")
    parser.add_argument("--state", type=Path, required=True, help="File JSON locale per autosalvataggio")
    parser.add_argument("--uploads", type=Path, required=True, help="Cartella read-only della run per le copie caricate")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--orders-dir", type=Path, help="Radice delle compilazioni datate (predefinito: accanto a output-dir, cartella «ordini»)")
    parser.add_argument("--writer-config", type=Path, help="Configurazione opzionale per creare subito le copie XLSX")
    parser.add_argument("--history", type=Path, help="Archivio degli ordini in attesa di consegna (predefinito: data/history/orders.json)")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.host not in {"127.0.0.1", "localhost", "::1"}:
        raise SystemExit("Per sicurezza l'app locale può ascoltare soltanto sull'interfaccia loopback")
    store = ReviewStore(args.review, args.state, args.uploads, args.output_dir, args.writer_config, args.history, args.orders_dir)
    server = ThreadingHTTPServer((args.host, args.port), AppHandler)
    server.store = store  # type: ignore[attr-defined]
    # AI configuration doesn't go through the store; see ServizioImpostazioni.
    server.impostazioni = ServizioImpostazioni()  # type: ignore[attr-defined]
    print(f"Comparatore locale: http://{args.host}:{args.port}")
    print("Premere Ctrl+C per fermare il server. Gli originali non vengono modificati.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
